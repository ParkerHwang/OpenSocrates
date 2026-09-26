"""Supplemental XLSX unit diagnostic, separate from the frozen 27-group score."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def currency_format(value: str) -> bool:
    # A locale-only Excel prefix has no currency symbol: [$-409] is not dollars.
    visible = re.sub(r"\[\$-[^\]]+\]", "", value)
    return bool(re.search(r"[$₩€£¥원]|\b(?:KRW|USD|EUR|GBP|JPY)\b", visible, re.I))


def check(directory: Path) -> dict:
    path = directory / "operations.xlsx"
    result = {
        "schema": "opensocrates.office-display-check/1",
        "status": "unassessable",
        "counts_not_currency": None,
        "checked_cells": [],
        "violations": [],
        "limitations": ["Number-format metadata only; not a full rendered or semantic review."],
    }
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            schedule = workbook["일정"]
            rows = list(schedule.iter_rows())
            header = [cell.value for cell in rows[0]]
            for field in ("attendees", "capacity"):
                index = header.index(field)
                for row in rows[1:]:
                    if row[index].value is not None:
                        cell = row[index]
                        result["checked_cells"].append(
                            {
                                "sheet": "일정",
                                "cell": cell.coordinate,
                                "field": field,
                                "format": cell.number_format,
                            }
                        )
            for row in workbook["예산"].iter_rows(min_row=2):
                if row[0].value in {"assigned_count", "mandatory_count", "optional_count"}:
                    cell = row[1]
                    result["checked_cells"].append(
                        {
                            "sheet": "예산",
                            "cell": cell.coordinate,
                            "field": row[0].value,
                            "format": cell.number_format,
                        }
                    )
        finally:
            workbook.close()
        result["violations"] = [
            item for item in result["checked_cells"] if currency_format(item["format"])
        ]
        if {item["field"] for item in result["checked_cells"]} != {
            "attendees",
            "capacity",
            "assigned_count",
            "mandatory_count",
            "optional_count",
        }:
            result["reason"] = "count_fields_missing"
            return result
        result.update(status="assessed", counts_not_currency=not result["violations"])
    except (
        OSError,
        ValueError,
        KeyError,
        IndexError,
        BadZipFile,
        ParseError,
        InvalidFileException,
    ) as error:
        result["reason"] = type(error).__name__
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("directory", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    result = check(args.directory)
    with args.report.open("x", encoding="utf-8") as stream:
        json.dump(result, stream, ensure_ascii=False, indent=2)
        stream.write("\n")
    print(json.dumps({k: result[k] for k in ("status", "counts_not_currency")}, ensure_ascii=False))
