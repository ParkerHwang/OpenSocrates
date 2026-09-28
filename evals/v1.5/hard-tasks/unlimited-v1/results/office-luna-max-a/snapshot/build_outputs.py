#!<BUNDLED_PYTHON>
"""Build the regional training planning artifacts from the supplied inputs."""

from __future__ import annotations

import csv
import itertools
import json
import re
from collections import Counter, defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
INPUTS = ROOT / "inputs"
OUTPUT = ROOT / "output"
OUTPUT.mkdir(exist_ok=True)


def read_csv(name: str) -> list[dict[str, str]]:
    with (INPUTS / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_json(name: str) -> dict:
    return json.loads((INPUTS / name).read_text(encoding="utf-8"))


def write_csv(path: Path, headers: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def csv_rows_as_dicts(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def split_list(value: str) -> list[str]:
    return [part.strip() for part in value.split(";") if part.strip()]


def as_bool(value: str) -> bool:
    if value.lower() == "true":
        return True
    if value.lower() == "false":
        return False
    raise ValueError(f"Not a boolean: {value!r}")


def reconcile(rows: list[dict[str, str]], id_column: str, kind: str) -> tuple[list[dict], list[dict]]:
    """Return one disposition row per physical row and selected representative rows."""
    by_id: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        by_id[row[id_column]].append(row)

    selected_ids: set[str] = set()
    dispositions: dict[str, str] = {}
    for business_id, group in by_id.items():
        substantive = [r for r in group if r["status"] != "draft"]
        if not substantive:
            for row in group:
                dispositions[row["row_id"]] = "draft"
            continue
        max_revision = max(int(r["revision"]) for r in substantive)
        latest = [r for r in substantive if int(r["revision"]) == max_revision]
        for row in group:
            if row["status"] == "draft":
                dispositions[row["row_id"]] = "draft"
            elif int(row["revision"]) < max_revision:
                dispositions[row["row_id"]] = "superseded"

        payload_groups: dict[tuple, list[dict[str, str]]] = defaultdict(list)
        for row in latest:
            payload = (
                row["revision"],
                row["status"],
                row["amount_krw"],
                row["invoice_id"],
            )
            payload_groups[payload].append(row)
        for payload_rows in payload_groups.values():
            reps = sorted(payload_rows, key=lambda r: r["row_id"])
            representative = reps[0]
            selected_ids.add(representative["row_id"])
            status = representative["status"]
            dispositions[representative["row_id"]] = {
                "posted": "active",
                "void": "void",
                "reversed": "reversed",
            }[status]
            for duplicate in reps[1:]:
                dispositions[duplicate["row_id"]] = "duplicate"

    records = []
    representatives = []
    for row in rows:
        disposition = dispositions[row["row_id"]]
        recognized = int(row["amount_krw"]) if disposition == "active" else 0
        records.append(
            {
                "row_id": row["row_id"],
                "kind": kind,
                "business_id": row[id_column],
                "disposition": disposition,
                "recognized_krw": recognized,
                "source_id": row["source_id"],
            }
        )
        if row["row_id"] in selected_ids:
            representatives.append({**row, "disposition": disposition, "recognized_krw": recognized})
    return records, representatives


class Edge:
    __slots__ = ("to", "rev", "cap")

    def __init__(self, to: int, rev: int, cap: int):
        self.to = to
        self.rev = rev
        self.cap = cap


class Dinic:
    def __init__(self, node_count: int):
        self.graph: list[list[Edge]] = [[] for _ in range(node_count)]

    def add_edge(self, source: int, target: int, capacity: int) -> Edge:
        forward = Edge(target, len(self.graph[target]), capacity)
        reverse = Edge(source, len(self.graph[source]), 0)
        self.graph[source].append(forward)
        self.graph[target].append(reverse)
        return forward

    def max_flow(self, source: int, sink: int, limit: int | None = None) -> int:
        total = 0
        n = len(self.graph)
        while limit is None or total < limit:
            level = [-1] * n
            queue = [source]
            level[source] = 0
            for node in queue:
                for edge in self.graph[node]:
                    if edge.cap and level[edge.to] < 0:
                        level[edge.to] = level[node] + 1
                        queue.append(edge.to)
            if level[sink] < 0:
                break
            cursor = [0] * n

            def send(node: int, amount: int) -> int:
                if node == sink:
                    return amount
                while cursor[node] < len(self.graph[node]):
                    edge = self.graph[node][cursor[node]]
                    if edge.cap and level[edge.to] == level[node] + 1:
                        sent = send(edge.to, min(amount, edge.cap))
                        if sent:
                            edge.cap -= sent
                            self.graph[edge.to][edge.rev].cap += sent
                            return sent
                    cursor[node] += 1
                return 0

            while limit is None or total < limit:
                amount = (limit - total) if limit is not None else 10**9
                sent = send(source, amount)
                if not sent:
                    break
                total += sent
        return total


def match_people(people: list[dict], sessions: list[dict]) -> tuple[dict[str, str] | None, int]:
    """Cover every mandatory person, then maximize optional assignments."""
    active = sorted(people, key=lambda p: p["person_id"])
    mandatory = [p for p in active if p["tier"] == "mandatory"]
    optional = [p for p in active if p["tier"] == "optional"]
    source = 0
    person_start = 1
    session_start = person_start + len(active)
    sink = session_start + len(sessions)
    flow = Dinic(sink + 1)
    pnodes = {p["person_id"]: person_start + i for i, p in enumerate(active)}
    snodes = {s["session_id"]: session_start + i for i, s in enumerate(sessions)}
    source_edges: dict[str, Edge] = {}
    person_session_edges: dict[tuple[str, str], Edge] = {}

    for p in mandatory:
        source_edges[p["person_id"]] = flow.add_edge(source, pnodes[p["person_id"]], 1)
    for p in active:
        for session in sessions:
            sid = session["session_id"]
            if sid not in p["allowed_sessions"]:
                continue
            if p["access_required"] and not session["accessible"]:
                continue
            person_session_edges[(p["person_id"], sid)] = flow.add_edge(
                pnodes[p["person_id"]], snodes[sid], 1
            )
    for session in sessions:
        flow.add_edge(snodes[session["session_id"]], sink, session["capacity"])

    mandatory_flow = flow.max_flow(source, sink, len(mandatory))
    if mandatory_flow != len(mandatory):
        return None, 0
    for p in optional:
        source_edges[p["person_id"]] = flow.add_edge(source, pnodes[p["person_id"]], 1)
    optional_flow = flow.max_flow(source, sink)

    assigned: dict[str, str] = {}
    for (person_id, session_id), edge in person_session_edges.items():
        if edge.cap == 0:
            if person_id in assigned:
                raise AssertionError(f"Person assigned twice: {person_id}")
            assigned[person_id] = session_id
    if len(assigned) != len(mandatory) + optional_flow:
        raise AssertionError("Flow assignment extraction did not match the flow value")
    if any(pid not in assigned for pid in (p["person_id"] for p in mandatory)):
        raise AssertionError("A mandatory attendee was not assigned")
    return assigned, optional_flow


def is_workday(day: date, holidays: set[date]) -> bool:
    return day.weekday() < 5 and day not in holidays


invoice_rows = read_csv("04_invoices.csv")
credit_rows = read_csv("05_credits.csv")
payment_rows = read_csv("06_payments.csv")
roster_rows = read_csv("01_roster.csv")
availability_rows = read_csv("02_availability.csv")
hr_rows = read_csv("03_hr_changes.csv")
facility_rows = read_csv("07_session_options.csv")
terms = read_json("08_supplier_terms.json")
constraints = read_json("09_approved_constraints.json")

invoice_records, invoice_reps = reconcile(invoice_rows, "invoice_id", "invoice")
credit_records, credit_reps = reconcile(credit_rows, "credit_id", "credit")
payment_records, payment_reps = reconcile(payment_rows, "payment_id", "payment")
records = invoice_records + credit_records + payment_records
records_headers = ["row_id", "kind", "business_id", "disposition", "recognized_krw", "source_id"]
write_csv(OUTPUT / "records.csv", records_headers, records)

balances = []
invoice_ids = [f"INV{i:02d}" for i in range(1, 37)]
for invoice_id in invoice_ids:
    inv_rep = [r for r in invoice_reps if r["invoice_id"] == invoice_id]
    assert len(inv_rep) == 1, f"Expected one selected invoice row for {invoice_id}"
    credit_ids = sorted(
        (r["row_id"] for r in credit_reps if r["invoice_id"] == invoice_id)
    )
    payment_ids = sorted(
        (r["row_id"] for r in payment_reps if r["invoice_id"] == invoice_id)
    )
    gross = inv_rep[0]["recognized_krw"]
    credit_total = sum(
        r["recognized_krw"] for r in credit_reps if r["invoice_id"] == invoice_id
    )
    paid_total = sum(
        r["recognized_krw"] for r in payment_reps if r["invoice_id"] == invoice_id
    )
    balance = gross - credit_total - paid_total
    assert balance >= 0, f"Negative balance for {invoice_id}"
    balances.append(
        {
            "invoice_id": invoice_id,
            "gross_krw": gross,
            "credit_krw": credit_total,
            "paid_krw": paid_total,
            "balance_krw": balance,
            "invoice_row_id": inv_rep[0]["row_id"],
            "credit_row_ids": ";".join(credit_ids),
            "payment_row_ids": ";".join(payment_ids),
        }
    )
balances_headers = [
    "invoice_id", "gross_krw", "credit_krw", "paid_krw", "balance_krw",
    "invoice_row_id", "credit_row_ids", "payment_row_ids",
]
write_csv(OUTPUT / "balances.csv", balances_headers, balances)
historical_outstanding = sum(row["balance_krw"] for row in balances)
gross_total = sum(row["gross_krw"] for row in balances)
credit_total = sum(row["credit_krw"] for row in balances)
paid_total = sum(row["paid_krw"] for row in balances)
assert gross_total - credit_total - paid_total == historical_outstanding
record_disposition_counts = Counter(row["disposition"] for row in records)
total_cash_cap = int(constraints["total_cash_cap_krw"])
launch_budget = total_cash_cap - historical_outstanding
assert launch_budget >= 0

availability_by_person = {r["person_id"]: r for r in availability_rows}
roster_by_person = {r["person_id"]: r for r in roster_rows}
approved_changes: dict[str, list[dict]] = defaultdict(list)
for row in hr_rows:
    if row["approval"] == "approved":
        approved_changes[row["person_id"]].append(row)

people = []
for person_id in sorted(roster_by_person):
    roster = roster_by_person[person_id]
    availability = availability_by_person[person_id]
    attributes = {
        "tier": roster["tier"],
        "active": as_bool(roster["active"]),
        "access_required": as_bool(roster["access_required"]),
        "allowed_sessions": split_list(availability["allowed_sessions"]),
    }
    for change in approved_changes.get(person_id, []):
        field = change["field"]
        value = change["value"]
        if field in ("active", "access_required"):
            attributes[field] = as_bool(value)
        elif field == "allowed_sessions":
            attributes[field] = split_list(value)
        else:
            attributes[field] = value
    source_row_ids = [roster["row_id"], availability["row_id"]]
    source_row_ids.extend(change["row_id"] for change in approved_changes.get(person_id, []))
    if attributes["active"]:
        people.append(
            {
                "person_id": person_id,
                "tier": attributes["tier"],
                "active": attributes["active"],
                "access_required": attributes["access_required"],
                "allowed_sessions": attributes["allowed_sessions"],
                "source_row_ids": source_row_ids,
            }
        )

assert len(people) == 72
mandatory_people = [p for p in people if p["tier"] == "mandatory"]
optional_people = [p for p in people if p["tier"] == "optional"]
assert len(mandatory_people) == 60
assert len(optional_people) == 12

holidays = {date.fromisoformat(d) for d in constraints["company_holidays"]}
window_start = date.fromisoformat(constraints["window_start"])
window_end = date.fromisoformat(constraints["window_end"])
eligible_facilities = []
for row in facility_rows:
    day = date.fromisoformat(row["date"])
    if row["approval"] != "approved" or not (window_start <= day <= window_end):
        continue
    if not is_workday(day, holidays):
        continue
    eligible_facilities.append(
        {
            "session_id": row["session_id"],
            "date": row["date"],
            "facility_capacity": int(row["capacity"]),
            "accessible": as_bool(row["accessible"]),
            "room_cost_krw": int(row["room_cost_krw"]),
            "source_id": row["source_id"],
        }
    )

approved_vendors = [v for v in terms["vendors"] if v["approval"] == "approved"]
candidate_plans = []
for vendor in approved_vendors:
    vendor_sessions = set(vendor["allowed_sessions"])
    vendor_facilities = [s for s in eligible_facilities if s["session_id"] in vendor_sessions]
    for combo in itertools.combinations(vendor_facilities, int(constraints["sessions_required"])):
        dates = [s["date"] for s in combo]
        if len(set(dates)) != len(dates):
            continue
        sessions = []
        for facility in combo:
            sessions.append(
                {
                    "session_id": facility["session_id"],
                    "date": facility["date"],
                    "capacity": min(facility["facility_capacity"], int(vendor["capacity_per_session"])),
                    "accessible": facility["accessible"],
                    "room_cost_krw": facility["room_cost_krw"],
                    "source_id": facility["source_id"],
                }
            )
        sessions.sort(key=lambda s: s["date"])
        assigned_map, optional_max = match_people(people, sessions)
        if assigned_map is None:
            candidate_plans.append(
                {"vendor": vendor, "sessions": sessions, "mandatory_feasible": False}
            )
            continue

        room_cost = sum(s["room_cost_krw"] for s in sessions)
        fixed = int(vendor["fixed_krw"])
        per_person = int(vendor["per_person_krw"])
        mandatory_cost = fixed + room_cost + per_person * len(mandatory_people)
        if mandatory_cost > launch_budget:
            candidate_plans.append(
                {"vendor": vendor, "sessions": sessions, "mandatory_feasible": True,
                 "budget_feasible": False, "optional_max": optional_max}
            )
            continue
        affordable_optional = min(
            optional_max,
            (launch_budget - fixed - room_cost) // per_person - len(mandatory_people),
        )
        affordable_optional = max(0, affordable_optional)
        matched_optional = sorted(
            p["person_id"] for p in optional_people if p["person_id"] in assigned_map
        )
        selected_optional_sets = []
        for subset in itertools.combinations(matched_optional, affordable_optional):
            chosen = set(subset)
            candidate_assignment = {
                pid: sid
                for pid, sid in assigned_map.items()
                if pid not in {p["person_id"] for p in optional_people} or pid in chosen
            }
            counts = defaultdict(int)
            for sid in candidate_assignment.values():
                counts[sid] += 1
            if all(counts[s["session_id"]] > 0 for s in sessions):
                selected_optional_sets.append((subset, candidate_assignment))
                break
        if affordable_optional and not selected_optional_sets:
            candidate_plans.append(
                {"vendor": vendor, "sessions": sessions, "mandatory_feasible": True,
                 "budget_feasible": False, "optional_max": optional_max}
            )
            continue
        if affordable_optional == 0:
            candidate_assignment = {
                pid: sid for pid, sid in assigned_map.items()
                if pid not in {p["person_id"] for p in optional_people}
            }
            counts = defaultdict(int)
            for sid in candidate_assignment.values():
                counts[sid] += 1
            if not all(counts[s["session_id"]] > 0 for s in sessions):
                candidate_plans.append(
                    {"vendor": vendor, "sessions": sessions, "mandatory_feasible": True,
                     "budget_feasible": False, "optional_max": optional_max}
                )
                continue
        else:
            candidate_assignment = selected_optional_sets[0][1]
        assigned_count = len(candidate_assignment)
        launch_cost = fixed + per_person * assigned_count + room_cost
        assert launch_cost <= launch_budget
        candidate_plans.append(
            {
                "vendor": vendor,
                "sessions": sessions,
                "mandatory_feasible": True,
                "budget_feasible": True,
                "optional_max": optional_max,
                "optional_count": affordable_optional,
                "assignment": candidate_assignment,
                "room_cost_krw": room_cost,
                "launch_cost_krw": launch_cost,
            }
        )

feasible_plans = [p for p in candidate_plans if p.get("budget_feasible")]
assert feasible_plans, "No plan covers every mandatory attendee within budget"
best_optional = max(p["optional_count"] for p in feasible_plans)
best_candidates = [p for p in feasible_plans if p["optional_count"] == best_optional]
best = min(
    best_candidates,
    key=lambda p: (
        p["launch_cost_krw"],
        p["vendor"]["vendor_id"],
        tuple(s["session_id"] for s in p["sessions"]),
    ),
)
max_optional_by_vendor = {
    offer["vendor_id"]: max(
        (p["optional_count"] for p in feasible_plans if p["vendor"]["vendor_id"] == offer["vendor_id"]),
        default=-1,
    )
    for offer in approved_vendors
}
max_attendance_by_vendor = {
    vendor_id: len(mandatory_people) + optional_count
    for vendor_id, optional_count in max_optional_by_vendor.items()
}

vendor = best["vendor"]
selected_sessions = best["sessions"]
assignment_by_person = best["assignment"]
assigned_count = len(assignment_by_person)
assigned_mandatory = sum(
    1 for p in mandatory_people if p["person_id"] in assignment_by_person
)
assigned_optional = assigned_count - assigned_mandatory
waitlisted = [p for p in people if p["person_id"] not in assignment_by_person]
assert assigned_mandatory == len(mandatory_people)
assert assigned_optional == best_optional

assignments = []
for person in people:
    pid = person["person_id"]
    session_id = assignment_by_person.get(pid, "")
    assignments.append(
        {
            "person_id": pid,
            "tier": person["tier"],
            "access_required": person["access_required"],
            "allowed_sessions": person["allowed_sessions"],
            "session_id": session_id,
            "state": "assigned" if session_id else "waitlist",
            "source_row_ids": person["source_row_ids"],
        }
    )
assignments_headers = [
    "person_id", "tier", "access_required", "allowed_sessions", "session_id", "state", "source_row_ids"
]

session_rows_out = []
for session in selected_sessions:
    attendees = sum(1 for sid in assignment_by_person.values() if sid == session["session_id"])
    assert attendees > 0
    session_rows_out.append(
        {
            "session_id": session["session_id"],
            "date": session["date"],
            "attendees": attendees,
            "capacity": session["capacity"],
            "accessible": session["accessible"],
            "room_cost_krw": session["room_cost_krw"],
            "source_id": session["source_id"],
        }
    )
sessions_headers = [
    "session_id", "date", "attendees", "capacity", "accessible", "room_cost_krw", "source_id"
]

vendor_fixed = int(vendor["fixed_krw"])
vendor_per_person = int(vendor["per_person_krw"])
room_cost = sum(s["room_cost_krw"] for s in selected_sessions)
launch_cost = vendor_fixed + vendor_per_person * assigned_count + room_cost
cash_remaining = launch_budget - launch_cost
one_more_attendee_cost = launch_cost + vendor_per_person
one_more_over_budget = max(0, one_more_attendee_cost - launch_budget)
pending_extra_budget = int(constraints["pending_extra_budget_krw"])
budget_if_extra_approved = launch_budget + pending_extra_budget
total_selected_capacity = sum(s["capacity"] for s in selected_sessions)
full_capacity_cost = vendor_fixed + vendor_per_person * total_selected_capacity + room_cost
earliest_training = min(date.fromisoformat(s["date"]) for s in selected_sessions)
prep_dates = []
cursor = earliest_training - timedelta(days=1)
while len(prep_dates) < int(constraints["preparation_workdays"]):
    if is_workday(cursor, holidays):
        prep_dates.append(cursor.isoformat())
    cursor -= timedelta(days=1)
prep_dates.reverse()

summary = {
    "total_cash_cap_krw": total_cash_cap,
    "historical_outstanding_krw": historical_outstanding,
    "launch_budget_krw": launch_budget,
    "vendor_id": vendor["vendor_id"],
    "vendor_fixed_krw": vendor_fixed,
    "vendor_per_person_krw": vendor_per_person,
    "room_cost_krw": room_cost,
    "launch_cost_krw": launch_cost,
    "cash_remaining_krw": cash_remaining,
    "assigned_count": assigned_count,
    "mandatory_count": assigned_mandatory,
    "optional_count": assigned_optional,
    "approval_status": "supplier_selection_pending",
    "preparation_dates": prep_dates,
}

evidence_info = [
    ("01_roster.csv", "MULTIPLE", "직원별 기준 활성상태·필수/선택 구분·접근 필요 여부와 HR 변경 전 원본 행을 제공."),
    ("02_availability.csv", "MULTIPLE", "person_id별 원래 가능 세션을 제공하고 승인된 HR 세션 변경의 기준 자료로 사용."),
    ("03_hr_changes.csv", "MULTIPLE", "approved 변경만 반영해 최종 활성상태·등급·접근성·가능 세션을 확정; pending 행은 제외."),
    ("04_invoices.csv", "MULTIPLE", "invoice_id별 최고 비초안 revision과 대표 행을 정해 송장 총액 및 잔액을 계산."),
    ("05_credits.csv", "MULTIPLE", "credit_id별 revision·중복·취소 상태를 반영해 유효 크레딧과 근거 행을 계산."),
    ("06_payments.csv", "MULTIPLE", "payment_id별 재전송·revision·reversed·0원 행을 보존해 유효 지급액을 계산."),
    ("07_session_options.csv", "MULTIPLE", "시설별 승인·날짜·수용인원·접근성·대관료를 후보 일정에 반영."),
    ("08_supplier_terms.json", terms["source_id"], "최종 공급사별 승인상태·기본료·인당료·세션 허용범위·수용인원 기준."),
    ("09_approved_constraints.json", constraints["source_id"], "현금한도·기간·휴일·4회 요건·준비일·목표 순서와 결재 상태의 기준."),
    ("10_prior_maintained_note.md", "OPS-NOTE-V3", "유지본의 과거 가정과 담당 역할을 확인하고 후속 승인 자료로 바뀐 점을 명시."),
    ("11_supplier_correction_email.md", "SUPPLIER-EMAIL-0926", "최종 견적 우선, CEDAR-PILOT 미승인, 수용인원 최소값, 10/9 휴무와 독립 준비 권한을 확인."),
    ("12_domain_rules.md", "RULES-0927", "출처 우선순위·회계 대사·참석 최적화·세션·접근성·준비일·권한 경계를 적용."),
]
evidence = [
    {"file": file, "source_id": source_id, "use": use}
    for file, source_id, use in evidence_info
]
evidence_headers = ["file", "source_id", "use"]

plan = {
    "summary": summary,
    "assignments": assignments,
    "sessions": session_rows_out,
    "evidence": evidence,
}
(OUTPUT / "plan.json").write_text(
    json.dumps(plan, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)


def won(value: int) -> str:
    return f"{value:,}원"


def ids(values: list[str]) -> str:
    return ", ".join(values) if values else "없음"


sessions_md = "\n".join(
    f"| {s['session_id']} | {s['date']} | {s['attendees']}명 / {s['capacity']}명 | "
    f"{'가능' if s['accessible'] else '불가'} | {won(s['room_cost_krw'])} | {s['source_id']} |"
    for s in session_rows_out
)
vendor_comparisons = []
vendor_costs_at_assigned_count = {}
for offer in approved_vendors:
    offer_cost = (
        int(offer["fixed_krw"])
        + int(offer["per_person_krw"]) * assigned_count
        + room_cost
    )
    offer_plan = max(
        (p for p in feasible_plans if p["vendor"]["vendor_id"] == offer["vendor_id"]),
        key=lambda p: (p["optional_count"], -p["launch_cost_krw"]),
    )
    offer_total = max_attendance_by_vendor[offer["vendor_id"]]
    offer_optional = max_optional_by_vendor[offer["vendor_id"]]
    vendor_costs_at_assigned_count[offer["vendor_id"]] = offer_cost
    vendor_comparisons.append(
        f"| {offer['vendor_id']} | {won(int(offer['fixed_krw']))} | "
        f"{won(int(offer['per_person_krw']))}/인 | {won(offer_cost)} | "
        f"{offer_total}명 / 선택 {offer_optional}명 |"
    )
waitlist_reasons = []
for person in waitlisted:
    if person["access_required"] and not any(
        s["session_id"] in person["allowed_sessions"] and s["accessible"]
        for s in selected_sessions
    ):
        waitlist_reasons.append((person["person_id"], "접근 필요 세션 없음"))
budget_waitlist = [pid for pid in (p["person_id"] for p in waitlisted)
                    if pid not in {x[0] for x in waitlist_reasons}]
for pid in budget_waitlist:
    waitlist_reasons.append((pid, "예산 상한"))
waitlist_text = "; ".join(f"{pid} ({reason})" for pid, reason in waitlist_reasons)

facility_by_id = {r["session_id"]: r for r in facility_rows}
assert len(session_rows_out) == int(constraints["sessions_required"])
assert prep_dates == sorted(prep_dates)

memo = f"""# 지역 직원 교육 4회차 개시 권고

수신: 운영책임자  
작성 기준일: 2026-09-27  
결재 상태: 공급사 선택 서명 대기 (`supplier_selection_pending`)

## 권고와 비용

승인된 교육 지출 한도 안에서 **CEDAR와 S2·S3·S5·S6** 조합을 선택해 주십시오. 이 문서는 대사와 배정에 근거한 계획 권고이며 공급사 선택 서명은 미결입니다. 예약·발주·지급·외부 전송은 하지 않았으며, 이 작업에서는 실행하지 않습니다. 서명과 별개로 명단 대사, 접근성 점검, 자료 준비는 진행할 수 있습니다 (`OPS-APPROVAL-0926`; `SUPPLIER-EMAIL-0926`).

승인 총현금 한도는 {won(total_cash_cap)}이고, 36개 논리 송장의 현재 미지급 잔액은 **{won(historical_outstanding)}**입니다. 송장 총액 {won(gross_total)}에서 유효 크레딧 {won(credit_total)}과 유효 지급 {won(paid_total)}을 빼서 대사했습니다. `records.csv`에는 송장 57행·크레딧 17행·지급 27행의 물리 행 101개를 모두 보존하고 대표·중복·초안·취소 상태를 표시했습니다. 예를 들어 `I-R01`은 `I-B01`보다 높은 revision 대표행이고, `C-V09`는 무효 대표행이며, `P-R16`–`P-R18`은 지급 취소 대표행입니다 (`RULES-0927`; `04_invoices.csv`; `05_credits.csv`; `06_payments.csv`). 제안된 추가 70,000원을 제외한 신규 교육 예산은 **{won(launch_budget)}**입니다. 선택안은 CEDAR 기본료 {won(vendor_fixed)} + 배정 1인당 {won(vendor_per_person)} × {assigned_count}명({won(vendor_per_person * assigned_count)}) + 방 비용 {won(room_cost)} = **{won(launch_cost)}**이며, 신규 예산 잔액은 **{won(cash_remaining)}**입니다. 기존 잔액은 예산에서 먼저 유보했고 교육비에 다시 더하지 않았습니다 (`OPS-APPROVAL-0926`; 최종 견적 `SUPPLIER-TERMS-FINAL-0926`; 대사 세부는 `records.csv`, `balances.csv`).

| 항목 | 값 |
|---|---:|
| 승인 총현금 한도 | {won(total_cash_cap)} |
| 과거 송장 미지급 유보액 | {won(historical_outstanding)} |
| 신규 교육 예산 | {won(launch_budget)} |
| 교육 실행비 계획 | {won(launch_cost)} |
| 교육 예산 잔액 | {won(cash_remaining)} |

## 공급사 비교와 세션

최종 승인 견적의 세 공급사를 같은 네 방 조합으로 비교했습니다. 현재 예산에서 CEDAR는 {max_attendance_by_vendor['CEDAR']}명(선택 {max_optional_by_vendor['CEDAR']}명), ORCHID는 {max_attendance_by_vendor['ORCHID']}명(선택 {max_optional_by_vendor['ORCHID']}명), LYRA는 {max_attendance_by_vendor['LYRA']}명(선택 {max_optional_by_vendor['LYRA']}명)까지 가능합니다. 따라서 최대 참석 목표는 CEDAR가 가장 잘 충족합니다. {assigned_count}명을 같은 방 조합으로 비교하면 CEDAR {won(launch_cost)}, ORCHID {won(vendor_costs_at_assigned_count['ORCHID'])}, LYRA {won(vendor_costs_at_assigned_count['LYRA'])}입니다. ORCHID는 인당료가 500원 낮지만 기본료가 45,000원 높아 {assigned_count}명 기준 CEDAR보다 {won(vendor_costs_at_assigned_count['ORCHID'] - launch_cost)} 비쌉니다. LYRA는 기본료가 40,000원 낮지만 인당료가 1,000원 높아 {assigned_count}명 기준 CEDAR보다 {won(vendor_costs_at_assigned_count['LYRA'] - launch_cost)} 비쌉니다. 승인 우선순위는 필수 전원 배정, 선택 인원 최대화, 같은 배정 수에서 신규 비용 최소화입니다 (`OPS-APPROVAL-0926`; `SUPPLIER-TERMS-FINAL-0926`).

| 승인 공급사 | 기본료 | 인당료 | 방 비용 포함 총액({assigned_count}명) | 현재 예산 내 최대 배정 |
|---|---:|---:|---:|---:|
{chr(10).join(vendor_comparisons)}

| 세션 | 날짜 | 참석/유효 정원 | 접근 가능 | 방 비용 | 시설 근거 |
|---|---|---:|---|---:|---|
{sessions_md}

S3의 유효 정원은 시설 16명과 공급사 18명 중 작은 값인 16명입니다. 나머지는 유효 정원 18명입니다 (`S3`, `FACILITY-0926`; `SUPPLIER-TERMS-FINAL-0926`). S2와 S5는 접근 가능하고 S3와 S6는 접근 불가입니다. 선택한 날짜는 승인 기간 안의 서로 다른 회사 근무일이며 하루 한 세션입니다. S4는 제안 상태이고 2026-10-09는 회사 휴무일이라 제외했습니다 (`S4`, `OPS-APPROVAL-0926`; `SUPPLIER-EMAIL-0926`).

세션 조합은 현재 승인 자료의 필수 가용성을 충족합니다. P020은 승인된 가능 세션이 S5뿐이고(H06), P025는 S3뿐이며(H07), P060은 S6뿐입니다(H09; `R020`, `R025`, `R060`). S2를 S1로 바꾸는 조합은 필수 30명(`R001`–`R015`, `R031`–`R045`)을 충족하지 못합니다. 이 30명 중 S1에도 갈 수 있는 사람은 P005·P010·P015·P035·P040 다섯 명뿐입니다 (`A005`, `A010`, `A015`, `A035`, `A040`). P045의 승인된 H08 변경은 S1을 허용하지 않습니다. S2가 없으면 S5 정원 18명과 S1을 실제 이용할 수 있는 다섯 명을 합해도 최대 23명이라 30명을 수용할 수 없습니다. 따라서 S2와 S5가 모두 필요하고, P020·P025·P060 제약과 합치면 S2·S3·S5·S6 네 세션이 남습니다 (`02_availability.csv`; H06–H09, `HR-APPROVED-0926`; `S1`–`S6`, `FACILITY-0926`).

## 참석 범위와 대기자

최종 활성 인원은 {len(people)}명(`R001`–`R072` 및 승인 복귀 H04), 필수 {len(mandatory_people)}명, 선택 {len(optional_people)}명입니다. 필수 {assigned_mandatory}명 전원을 1회 배정하고 선택 인원 중 {assigned_optional}명을 배정합니다. 네 세션의 유효 정원 합은 {total_selected_capacity}명이지만 현재 예산으로는 {assigned_count}명까지 배정할 수 있어 좌석 {total_selected_capacity - assigned_count}개가 남습니다. 이 계획의 현금 잔여는 {won(cash_remaining)}이고 인당 추가비는 {won(vendor_per_person)}이므로 다음 선택 참석자까지 배정하면 신규 비용은 {won(one_more_attendee_cost)}으로 올라 승인 예산을 {won(one_more_over_budget)} 초과합니다. 추가 {won(pending_extra_budget)}은 승인되지 않아 계산에서 제외했습니다 (`OPS-APPROVAL-0926`; `SUPPLIER-TERMS-FINAL-0926`). 추가 예산이 정식 승인되면 신규 한도는 {won(budget_if_extra_approved)}으로 바뀌므로 다시 계산해야 합니다. 현재 세션 구성에서 전체 유효 정원을 채우는 비용은 {won(full_capacity_cost)}이며, 접근성 조건과 물리 정원도 새 배정에 다시 적용해야 합니다 (`pending_extra_budget_krw`, `OPS-APPROVAL-0926`; `SUPPLIER-TERMS-FINAL-0926`).

P064는 접근 필요 변경이 승인됐지만(H05, `R064`) 허용 세션 `A064`의 S3/S6가 모두 접근 불가여서 배정하지 않았습니다. 추가 선택 대기자는 **{waitlist_text}**입니다. 다섯 명은 현재 예산으로 한 명도 더 추가할 수 없어 대기 처리했습니다. 입력에 선택 인원 우선순위가 없어 이 행 순서는 우선순위를 뜻하지 않습니다. 대기자 순서를 따로 정하기 전에는 이 표를 좌석 수립용 계획으로만 취급하십시오 (`01_roster.csv`; 승인 변경 H01–H10, `HR-APPROVED-0926`; H11/H12는 `pending`이므로 미적용).

P010의 등급은 H02에 따라 필수, P061은 H03에 따라 선택이며, P072는 H04에 따라 활성으로 복귀했습니다. P063을 필수로 바꾸자는 H11과 P005 접근성을 false로 바꾸자는 H12는 pending이므로 적용하지 않았습니다. 따라서 P005의 접근 필요는 승인 H01의 true로 유지됩니다 (`03_hr_changes.csv`; `HR-APPROVED-0926`, `HR-PROPOSAL-0927`). 모든 사람의 원본 명단 행·가용성 행과 승인 HR 변경 행은 `plan.json` 및 `operations.xlsx`의 배정 표에 연결했습니다.

## 기존 운영 메모의 정정

유지본 `OPS-NOTE-V3`는 역할 인수인계에만 사용하고 현재 가격·명단·일정 사실은 후속 자료로 갱신했습니다.

- CEDAR의 과거 추정치 210,000원 + 1인당 6,500원은 최종 조건 240,000원 + 1인당 7,000원으로 교체했습니다. CEDAR-PILOT은 제안/미승인이라 기존 승인으로 선택할 수 없습니다 (`SUPPLIER-TERMS-FINAL-0926`; `SUPPLIER-EMAIL-0926`).
- 과거 메모의 P010 선택, P061 필수, P072 철회 상태는 H02/H03/H04 승인에 따라 각각 P010 필수, P061 선택, P072 활성으로 바뀌었습니다. pending인 H11/H12는 반영하지 않았습니다 (`HR-APPROVED-0926`; `HR-PROPOSAL-0927`).
- S4 저가 검토안은 시설 행에서 `proposed`이고 10월 9일은 회사 휴무이므로 일정 후보에서 뺐습니다 (`S4`, `FACILITY-0926`; `OPS-APPROVAL-0926`; `SUPPLIER-EMAIL-0926`).
- 송금 재전송은 합산하지 않았습니다. 예를 들어 PAY01의 `P-B01`/`P-D01` 중 대표행만 인정하고, PAY16–PAY18은 취소 대표행 `P-R16`–`P-R18`로 유효 지급이 0입니다. 금액 0인 게시 크레딧 CR10(`C-B10`)과 PAY19/PAY20(`P-Z19`, `P-Z20`)도 활성 대표 행으로 보존했습니다. 초안 `I-P01`, 무효 송장 `I-V33`–`I-V36`도 빠뜨리지 않고 대사 기록에 남겼습니다 (`04_invoices.csv`; `05_credits.csv`; `06_payments.csv`; `RULES-0927`).
- 과거 메모에는 준비일이 없었으나, 가장 이른 세션 S2(2026-10-06) 전 최근 연속 근무일 3일은 2026-10-01, 2026-10-02, 2026-10-05로 계산했습니다 (`OPS-APPROVAL-0926`; `RULES-0927`).

## 다음 조치

1. **윤구매**는 운영책임자에게 CEDAR와 S2·S3·S5·S6 조합, 신규 비용 {won(launch_cost)}의 선택 서명을 요청하십시오. 첫 준비일인 2026-10-01 시작 때 서명 상태를 확인하되, 일정·자료 준비를 기다리게 하지는 마십시오. 선택 서명은 미결로 기록하고 이 업무에서는 예약·발주·지급·외부 발송을 하지 않습니다 (`OPS-NOTE-V3`; `OPS-APPROVAL-0926`; `SUPPLIER-EMAIL-0926`).
2. **정회계**는 `records.csv`와 `balances.csv`의 36개 잔액 및 {won(historical_outstanding)} 유보액을 검토하고, 결재 요청서에 신규 예산 {won(launch_budget)}과 잔액 {won(cash_remaining)}을 옮기십시오. 중복·취소·0원 행의 출처는 두 표에서 추적할 수 있습니다 (`RULES-0927`; `I-R01`; `I-P01`; `C-B10`; `P-R16`–`P-R18`).
3. **박운영**은 2026-10-01부터 최종 명단과 접근성에 따라 자료를 준비하고, P064의 접근 가능한 허용 세션 변경이 승인될 경우에만 다시 배정 가능성을 계산하십시오. 수용 인원 대기자의 통보 순서는 별도 승인 전까지 확정하지 마십시오 (`H05`, `HR-APPROVED-0926`; `SUPPLIER-EMAIL-0926`).
4. **최시설**은 2026-10-01부터 S2·S3·S5·S6 방과 장비 준비를 진행하고, S2/S5의 접근성 및 S3/S6의 비접근 상태를 운영 계획에 반영하십시오 (`FACILITY-0926`; `RULES-0927`).

현재 독립 진행 가능한 작업은 송장·인사 대사, 60명 필수 좌석 확인, 접근성 검토, 자료 준비입니다. 선택 공급사 서명은 아직 미결이며, 그 확인 전 실행·구매 행위는 보류 상태입니다.
"""
(OUTPUT / "decision_memo.md").write_text(memo, encoding="utf-8")


def excel_rows(headers: list[str], rows: list[dict], json_list_fields: set[str] | None = None) -> list[list]:
    json_list_fields = json_list_fields or set()
    result = [headers]
    for row in rows:
        values = []
        for header in headers:
            value = row[header]
            if header in json_list_fields:
                value = ";".join(value)
            values.append(value)
        result.append(values)
    return result


budget_rows = [
    {"key": key, "value": ";".join(value) if isinstance(value, list) else value}
    for key, value in summary.items()
]

sheet_specs = [
    ("대사기록", records_headers, records, set()),
    ("송장잔액", balances_headers, balances, set()),
    ("배정", assignments_headers, assignments, {"allowed_sessions", "source_row_ids"}),
    ("일정", sessions_headers, session_rows_out, set()),
    ("예산", ["key", "value"], budget_rows, set()),
    ("근거", evidence_headers, evidence, set()),
]

wb = Workbook()
wb.remove(wb.active)
wb.properties.title = "지역 직원 교육 운영 계획"
wb.properties.subject = "회계 대사, 참석 배정, 일정 및 근거"
wb.properties.creator = "운영기획"
header_fill = PatternFill("solid", fgColor="244A69")
header_font = Font(color="FFFFFF", bold=True)
thin_gray = "D9E2F3"
currency_format = '#,##0"원";[Red]-#,##0"원"'

for sheet_name, headers, rows, list_fields in sheet_specs:
    ws = wb.create_sheet(sheet_name)
    for r_index, values in enumerate(excel_rows(headers, rows, list_fields), start=1):
        for c_index, value in enumerate(values, start=1):
            cell = ws.cell(row=r_index, column=c_index)
            if value == "":
                cell.value = None
            else:
                cell.value = value
            if r_index == 1:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="top", wrap_text=(sheet_name in {"근거", "배정"}))
            if r_index > 1 and headers[c_index - 1].endswith("_krw") and isinstance(value, int):
                cell.number_format = currency_format
            if (
                sheet_name == "예산"
                and r_index > 1
                and headers[c_index - 1] == "value"
                and rows[r_index - 2]["key"].endswith("_krw")
                and isinstance(value, int)
            ):
                cell.number_format = currency_format
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.row_dimensions[1].height = 30
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for col_index, header in enumerate(headers, start=1):
        content_lengths = [len(str(header))]
        for row in rows:
            value = row[header]
            if isinstance(value, list):
                value = ";".join(str(v) for v in value)
            content_lengths.append(len(str(value)))
        width = min(max(max(content_lengths) + 2, len(header) + 2, 12), 62)
        if header in {"use"}:
            width = 68
        elif header in {"source_row_ids", "payment_row_ids", "credit_row_ids", "allowed_sessions"}:
            width = min(max(width, 28), 55)
        elif header in {"recognized_krw", "room_cost_krw", "total_cash_cap_krw", "historical_outstanding_krw",
                        "launch_budget_krw", "vendor_fixed_krw", "vendor_per_person_krw", "launch_cost_krw",
                        "cash_remaining_krw", "gross_krw", "credit_krw", "paid_krw", "balance_krw"}:
            width = max(width, 19)
        ws.column_dimensions[get_column_letter(col_index)].width = width

    for r_index in range(2, len(rows) + 2):
        if sheet_name == "근거":
            ws.row_dimensions[r_index].height = 36
        elif sheet_name == "배정":
            ws.row_dimensions[r_index].height = 30

workbook_path = OUTPUT / "operations.xlsx"
wb.save(workbook_path)


# Source-based consistency and constraint checks; these validate the delivered files.
assert len(invoice_rows) + len(credit_rows) + len(payment_rows) == 101
assert len(roster_rows) + len(availability_rows) + len(hr_rows) + len(facility_rows) + len(invoice_rows) + len(credit_rows) + len(payment_rows) == 263
assert len(records) == 101
assert len(balances) == 36
assert len({r["row_id"] for r in records}) == 101
assert len(people) == len(assignments) == 72
assert len(evidence) == 12
assert len(session_rows_out) == int(constraints["sessions_required"]) == 4
assert summary["approval_status"] == constraints["approval_status"] == "supplier_selection_pending"
assert launch_budget == total_cash_cap - historical_outstanding
assert launch_cost == vendor_fixed + vendor_per_person * assigned_count + room_cost
assert cash_remaining == launch_budget - launch_cost >= 0
assert assigned_count == assigned_mandatory + assigned_optional
assert assigned_mandatory == len(mandatory_people)
assert assigned_optional == best_optional
mandatory_feasible_sets = {
    tuple(s["session_id"] for s in p["sessions"])
    for p in candidate_plans if p.get("mandatory_feasible")
}
assert tuple(s["session_id"] for s in selected_sessions) in mandatory_feasible_sets
assert assigned_optional == max(max_optional_by_vendor.values())
assert best["launch_cost_krw"] == min(
    p["launch_cost_krw"] for p in feasible_plans if p["optional_count"] == best_optional
)
assert all(s["attendees"] > 0 and s["attendees"] <= s["capacity"] for s in session_rows_out)
assert sum(s["attendees"] for s in session_rows_out) == assigned_count
assert all(s["accessible"] or not any(
    a["access_required"] and a["session_id"] == s["session_id"] for a in assignments
) for s in session_rows_out)
assert all(
    a["session_id"] in a["allowed_sessions"]
    for a in assignments if a["state"] == "assigned"
)
assert all(
    not a["access_required"] or next(s for s in session_rows_out if s["session_id"] == a["session_id"])["accessible"]
    for a in assignments if a["state"] == "assigned"
)
assert all(a["state"] == "waitlist" and a["session_id"] == "" for a in assignments if a["state"] == "waitlist")
assert all(set(a["source_row_ids"]) >= {roster_by_person[a["person_id"]]["row_id"], availability_by_person[a["person_id"]]["row_id"]} for a in assignments)
for a in assignments:
    approved_ids = {r["row_id"] for r in approved_changes.get(a["person_id"], [])}
    expected_source_ids = {
        roster_by_person[a["person_id"]]["row_id"],
        availability_by_person[a["person_id"]]["row_id"],
        *approved_ids,
    }
    assert set(a["source_row_ids"]) == expected_source_ids
    pending_ids = {r["row_id"] for r in hr_rows if r["person_id"] == a["person_id"] and r["approval"] != "approved"}
    assert not pending_ids.intersection(a["source_row_ids"])
assert all(date.fromisoformat(s["date"]).weekday() < 5 for s in session_rows_out)
assert not any(date.fromisoformat(s["date"]) in holidays for s in session_rows_out)
assert len({s["date"] for s in session_rows_out}) == 4
assert len({s["session_id"] for s in session_rows_out}) == 4
assert not any(facility_by_id[s["session_id"]]["approval"] != "approved" for s in session_rows_out)

for output_name, headers in [("records.csv", records_headers), ("balances.csv", balances_headers)]:
    parsed = csv_rows_as_dicts(OUTPUT / output_name)
    expected = records if output_name == "records.csv" else balances
    assert list(parsed[0]) == headers
    assert len(parsed) == len(expected)
    for actual, expect in zip(parsed, expected):
        assert all(str(actual[key]) == str(expect[key]) for key in headers)
loaded_plan = json.loads((OUTPUT / "plan.json").read_text(encoding="utf-8"))
assert loaded_plan == plan

check_wb = load_workbook(workbook_path, data_only=True, read_only=False)
assert check_wb.sheetnames == [spec[0] for spec in sheet_specs]
for sheet_name, headers, rows, list_fields in sheet_specs:
    ws = check_wb[sheet_name]
    actual_values = list(ws.values)
    expected_values = excel_rows(headers, rows, list_fields)
    assert list(actual_values[0]) == headers
    assert len(actual_values) == len(expected_values)
    for actual_row, expected_row in zip(actual_values[1:], expected_values[1:]):
        for actual, expected in zip(actual_row, expected_row):
            if actual is None:
                actual = ""
            assert actual == expected, f"Workbook mismatch {sheet_name}: {actual!r} != {expected!r}"
    assert ws.freeze_panes == "A2"
    assert ws.max_row == len(rows) + 1
check_wb.close()

print(json.dumps({
    "invoice_rows": len(invoice_rows),
    "credit_rows": len(credit_rows),
    "payment_rows": len(payment_rows),
    "record_rows": len(records),
    "invoice_balances": len(balances),
    "historical_outstanding_krw": historical_outstanding,
    "launch_budget_krw": launch_budget,
    "vendor_id": vendor["vendor_id"],
    "sessions": session_rows_out,
    "assigned_count": assigned_count,
    "mandatory_count": assigned_mandatory,
    "optional_count": assigned_optional,
    "waitlist": [{"person_id": p["person_id"], "access_required": p["access_required"]} for p in waitlisted],
    "launch_cost_krw": launch_cost,
    "cash_remaining_krw": cash_remaining,
    "preparation_dates": prep_dates,
    "candidate_count": len(candidate_plans),
    "feasible_plan_count": len(feasible_plans),
    "outputs": sorted(p.name for p in OUTPUT.iterdir()),
}, ensure_ascii=False, indent=2))
