import csv, json, itertools, datetime, collections, pathlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT=pathlib.Path(__file__).parent
IN=ROOT/'inputs'; OUT=ROOT/'output'; OUT.mkdir(exist_ok=True)
def rows(name):
    with (IN/name).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def writecsv(path,headers,data):
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=headers);w.writeheader();w.writerows(data)
roster=rows('01_roster.csv'); avail=rows('02_availability.csv');changes=rows('03_hr_changes.csv')
invoices=rows('04_invoices.csv');credits=rows('05_credits.csv');payments=rows('06_payments.csv');options=rows('07_session_options.csv')
terms=json.loads((IN/'08_supplier_terms.json').read_text());constraints=json.loads((IN/'09_approved_constraints.json').read_text())

record_headers=['row_id','kind','business_id','disposition','recognized_krw','source_id']
record=[]; selected={}
for kind,data,key in [('invoice',invoices,'invoice_id'),('credit',credits,'credit_id'),('payment',payments,'payment_id')]:
    groups=collections.defaultdict(list)
    for r in data:groups[r[key]].append(r)
    dispositions={}
    for bid,group in groups.items():
        nondraft=[r for r in group if r['status']!='draft']
        if nondraft:
            maxrev=max(int(r['revision']) for r in nondraft)
            leaders=sorted((r for r in nondraft if int(r['revision'])==maxrev),key=lambda r:r['row_id'])
            rep=leaders[0];selected[(kind,bid)]=rep
            for r in group:
                if r['status']=='draft':disp='draft'
                elif int(r['revision'])<maxrev:disp='superseded'
                elif r is not rep:disp='duplicate'
                else:disp={'posted':'active','void':'void','reversed':'reversed'}[r['status']]
                dispositions[r['row_id']]=disp
        else:
            for r in group:dispositions[r['row_id']]='draft'
    for r in data:
        disp=dispositions[r['row_id']]
        record.append(dict(row_id=r['row_id'],kind=kind,business_id=r[key],disposition=disp,recognized_krw=int(r['amount_krw']) if disp=='active' else 0,source_id=r['source_id']))
assert len(record)==len(invoices)+len(credits)+len(payments)
writecsv(OUT/'records.csv',record_headers,record)

rec_by_row={r['row_id']:r for r in record}
balance_headers=['invoice_id','gross_krw','credit_krw','paid_krw','balance_krw','invoice_row_id','credit_row_ids','payment_row_ids']
balances=[]
for i in range(1,37):
    iid=f'INV{i:02d}'
    inv=selected[('invoice',iid)]
    cs=sorted([r for (k,b),r in selected.items() if k=='credit' and r['invoice_id']==iid],key=lambda r:r['row_id'])
    ps=sorted([r for (k,b),r in selected.items() if k=='payment' and r['invoice_id']==iid],key=lambda r:r['row_id'])
    val=lambda r:rec_by_row[r['row_id']]['recognized_krw']
    gross=val(inv);credit=sum(map(val,cs));paid=sum(map(val,ps))
    balances.append(dict(invoice_id=iid,gross_krw=gross,credit_krw=credit,paid_krw=paid,balance_krw=gross-credit-paid,invoice_row_id=inv['row_id'],credit_row_ids=';'.join(r['row_id'] for r in cs),payment_row_ids=';'.join(r['row_id'] for r in ps)))
assert all(r['balance_krw']>=0 for r in balances)
writecsv(OUT/'balances.csv',balance_headers,balances)
outstanding=sum(x['balance_krw'] for x in balances)
budget=constraints['total_cash_cap_krw']-outstanding

avmap={r['person_id']:r for r in avail}; change_map=collections.defaultdict(list)
for c in changes:
    if c['approval']=='approved':change_map[c['person_id']].append(c)
people=[]
for r in roster:
    p=dict(person_id=r['person_id'],tier=r['tier'],access_required=r['access_required']=='true',active=r['active']=='true',allowed_sessions=avmap[r['person_id']]['allowed_sessions'].split(';'),source_row_ids=[r['row_id'],avmap[r['person_id']]['row_id']]+[c['row_id'] for c in change_map[r['person_id']]])
    for c in change_map[r['person_id']]:
        v=c['value']
        p[c['field']]=v=='true' if c['field'] in ('active','access_required') else (v.split(';') if c['field']=='allowed_sessions' else v)
    if p.pop('active'):people.append(p)
mandatory=sum(p['tier']=='mandatory' for p in people);optional=sum(p['tier']=='optional' for p in people)

# Successive shortest augmenting path for the person-to-session bipartite flow.
def allocate(vendor,chosen,max_people):
    n=len(people);m=len(chosen);src=0; pbase=1;sbase=1+n;sink=sbase+m;N=sink+1
    graph=[[] for _ in range(N)]
    def edge(u,v,cap,cost):
        graph[u].append([v,cap,cost,len(graph[v])]);graph[v].append([u,0,-cost,len(graph[u])-1])
    for i,p in enumerate(people):
        edge(src,pbase+i,1,-10000 if p['tier']=='mandatory' else -1)
        for j,s in enumerate(chosen):
            if s['session_id'] in p['allowed_sessions'] and (not p['access_required'] or s['accessible']=='true'):
                edge(pbase+i,sbase+j,1,0)
    for j,s in enumerate(chosen):edge(sbase+j,sink,min(int(s['capacity']),vendor['capacity_per_session']),0)
    flow=0
    while flow<max_people:
        dist=[10**12]*N;prev=[None]*N;dist[src]=0
        for _ in range(N-1):
            changed=False
            for u in range(N):
                if dist[u]==10**12:continue
                for idx,(v,cap,cost,rev) in enumerate(graph[u]):
                    if cap>0 and dist[v]>dist[u]+cost:
                        dist[v]=dist[u]+cost;prev[v]=(u,idx);changed=True
            if not changed:break
        if prev[sink] is None:break
        v=sink
        while v!=src:
            u,idx=prev[v];e=graph[u][idx];e[1]-=1;graph[v][e[3]][1]+=1;v=u
        flow+=1
    assigned={}
    for i,p in enumerate(people):
        for e in graph[pbase+i]:
            if sbase<=e[0]<sink and e[1]==0:assigned[p['person_id']]=chosen[e[0]-sbase]['session_id']
    return assigned

def workday(date):return date.weekday()<5 and date.isoformat() not in constraints['company_holidays']
eligible=[s for s in options if s['approval']=='approved' and constraints['window_start']<=s['date']<=constraints['window_end'] and workday(datetime.date.fromisoformat(s['date']))]
candidates=[]
for v in terms['vendors']:
    if v['approval']!='approved':continue
    for chosen in itertools.combinations(eligible,constraints['sessions_required']):
        if len({s['date'] for s in chosen})!=4 or any(s['session_id'] not in v['allowed_sessions'] for s in chosen):continue
        room=sum(int(s['room_cost_krw']) for s in chosen)
        limit=min(len(people),(budget-v['fixed_krw']-room)//v['per_person_krw'])
        if limit<mandatory:continue
        assigned=allocate(v,chosen,limit)
        mand_count=sum(p['tier']=='mandatory' and p['person_id'] in assigned for p in people)
        if mand_count!=mandatory:continue
        counts=collections.Counter(assigned.values())
        if any(counts[s['session_id']]==0 for s in chosen):continue
        cost=v['fixed_krw']+room+len(assigned)*v['per_person_krw']
        candidates.append(dict(vendor=v,chosen=chosen,assigned=assigned,room=room,cost=cost,optional=len(assigned)-mandatory))
assert candidates,'No feasible approved plan'
candidates.sort(key=lambda c:(-c['optional'],c['cost'],c['vendor']['vendor_id'],tuple(s['session_id'] for s in c['chosen'])))
best=candidates[0];v=best['vendor'];chosen=best['chosen'];assigned=best['assigned']
first=min(datetime.date.fromisoformat(s['date']) for s in chosen);prep=[];d=first
while len(prep)<constraints['preparation_workdays']:
    d-=datetime.timedelta(days=1)
    if workday(d):prep.append(d.isoformat())
prep.sort()
summary=dict(total_cash_cap_krw=constraints['total_cash_cap_krw'],historical_outstanding_krw=outstanding,launch_budget_krw=budget,vendor_id=v['vendor_id'],vendor_fixed_krw=v['fixed_krw'],vendor_per_person_krw=v['per_person_krw'],room_cost_krw=best['room'],launch_cost_krw=best['cost'],cash_remaining_krw=budget-best['cost'],assigned_count=len(assigned),mandatory_count=mandatory,optional_count=best['optional'],approval_status=constraints['approval_status'],preparation_dates=prep)
assignment_headers=['person_id','tier','access_required','allowed_sessions','session_id','state','source_row_ids']
assignments=[]
for p in people:
    assignments.append(dict(person_id=p['person_id'],tier=p['tier'],access_required=p['access_required'],allowed_sessions=p['allowed_sessions'],session_id=assigned.get(p['person_id'],''),state='assigned' if p['person_id'] in assigned else 'waitlist',source_row_ids=p['source_row_ids']))
assert all(a['tier']=='optional' for a in assignments if a['state']=='waitlist')
session_headers=['session_id','date','attendees','capacity','accessible','room_cost_krw','source_id']
sessions=[]
for s in chosen:
    sessions.append(dict(session_id=s['session_id'],date=s['date'],attendees=sum(x==s['session_id'] for x in assigned.values()),capacity=min(int(s['capacity']),v['capacity_per_session']),accessible=s['accessible']=='true',room_cost_krw=int(s['room_cost_krw']),source_id=s['source_id']))
evidence_uses={
'01_roster.csv':'원명단의 인원·등급·활성·접근성 및 원행 ID를 보존',
'02_availability.csv':'인원별 기본 가능 세션과 원행 ID를 결합',
'03_hr_changes.csv':'승인된 인사 변경만 반영하고 보류 행은 제외',
'04_invoices.csv':'송장 업무 ID·revision·상태별 대표 금액 대사',
'05_credits.csv':'크레딧 대표 행과 송장별 차감액 대사',
'06_payments.csv':'지급 대표 행과 취소·중복을 반영한 지급액 대사',
'07_session_options.csv':'승인 시설·일자·접근성·용량·방비 비교',
'08_supplier_terms.json':'승인 공급사별 최종 고정비·인당비·용량 적용',
'09_approved_constraints.json':'현금 한도·기간·휴무·목표·서명 상태 적용',
'10_prior_maintained_note.md':'이전 추정과 역할을 확인하고 최신 근거로 정정',
'11_supplier_correction_email.md':'최종 견적 우선·파일럿 보류·독립 준비 권한 확인',
'12_domain_rules.md':'대사·배정·달력·목표의 우선순위와 산식 적용'}
docids={'08_supplier_terms.json':terms['source_id'],'09_approved_constraints.json':constraints['source_id'],'10_prior_maintained_note.md':'OPS-NOTE-V3','11_supplier_correction_email.md':'SUPPLIER-EMAIL-0926','12_domain_rules.md':'RULES-0927'}
evidence=[dict(file=f.name,source_id='MULTIPLE' if f.suffix=='.csv' else docids[f.name],use=evidence_uses[f.name]) for f in sorted(IN.iterdir()) if f.name in evidence_uses]
plan=dict(summary=summary,assignments=assignments,sessions=sessions,evidence=evidence)
(OUT/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

wb=Workbook();wb.remove(wb.active)
tabs=[('대사기록',record_headers,record),('송장잔액',balance_headers,balances),('배정',assignment_headers,assignments),('일정',session_headers,sessions),('예산',['key','value'],[{'key':k,'value':x} for k,x in summary.items()]),('근거',['file','source_id','use'],evidence)]
for title,headers,data in tabs:
    ws=wb.create_sheet(title);ws.append(headers);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(headers))}{len(data)+1}'
    for r in data:
        ws.append([';'.join(str(x) for x in r[h]) if isinstance(r[h],list) else r[h] for h in headers])
    for c in ws[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='23415B');c.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=28
    for j,h in enumerate(headers,1):
        col=get_column_letter(j);maxlen=max([len(str(h))]+[len(str(ws.cell(i,j).value or '')) for i in range(2,min(ws.max_row+1,100))]);ws.column_dimensions[col].width=min(58,max(13,maxlen+2))
        if h.endswith('_krw') or h=='value' and title=='예산':
            for i in range(2,ws.max_row+1):
                if isinstance(ws.cell(i,j).value,int):ws.cell(i,j).number_format='#,##0" 원"'
    ws.sheet_view.showGridLines=False
wb.save(OUT/'operations.xlsx')

print(json.dumps({'summary':summary,'active_people':len(people),'optional_total':optional,'candidate_count':len(candidates),'comparisons':[{'vendor':c['vendor']['vendor_id'],'sessions':[s['session_id'] for s in c['chosen']],'assigned':len(c['assigned']),'optional':c['optional'],'cost':c['cost']} for c in candidates[:12]],'balance_special':[b for b in balances if b['invoice_id'] in ['INV36']]},ensure_ascii=False,indent=2))
