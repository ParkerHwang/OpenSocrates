import csv
import json
import itertools
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = Path(__file__).parent
IN = BASE / 'inputs'
OUT = BASE / 'output'
OUT.mkdir(exist_ok=True)

def readcsv(name):
    with (IN / name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def writecsv(name, headers, rows):
    with (OUT / name).open('w', encoding='utf-8', newline='') as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader(); w.writerows(rows)

records = []
selected = {}
by_invoice = defaultdict(lambda: {'credit': [], 'payment': []})
for kind, fn, key in [('invoice','04_invoices.csv','invoice_id'),('credit','05_credits.csv','credit_id'),('payment','06_payments.csv','payment_id')]:
    rows = readcsv(fn)
    groups = defaultdict(list)
    for r in rows: groups[r[key]].append(r)
    result = {}
    for business_id, group in groups.items():
        eligible = [r for r in group if r['status'] != 'draft']
        high = max(int(r['revision']) for r in eligible)
        highest = [r for r in eligible if int(r['revision']) == high]
        representative = min(highest, key=lambda r:r['row_id'])
        result[business_id] = representative
        for r in group:
            if r['status'] == 'draft': disp = 'draft'
            elif int(r['revision']) < high: disp = 'superseded'
            elif r is not representative: disp = 'duplicate'
            else: disp = {'posted':'active','void':'void','reversed':'reversed'}[r['status']]
            records.append({'row_id':r['row_id'],'kind':kind,'business_id':business_id,'disposition':disp,
                            'recognized_krw':int(r['amount_krw']) if disp == 'active' else 0,'source_id':r['source_id']})
    selected[kind] = result
    if kind != 'invoice':
        for r in result.values(): by_invoice[r['invoice_id']][kind].append(r)

rec_by_id = {r['row_id']:r for r in records}
balances=[]
for inv_id in sorted(selected['invoice']):
    i=selected['invoice'][inv_id]
    c=sorted(by_invoice[inv_id]['credit'],key=lambda r:r['row_id'])
    p=sorted(by_invoice[inv_id]['payment'],key=lambda r:r['row_id'])
    amt=lambda r: rec_by_id[r['row_id']]['recognized_krw']
    gross=amt(i); credit=sum(map(amt,c)); paid=sum(map(amt,p))
    balances.append({'invoice_id':inv_id,'gross_krw':gross,'credit_krw':credit,'paid_krw':paid,
                     'balance_krw':gross-credit-paid,'invoice_row_id':i['row_id'],
                     'credit_row_ids':';'.join(r['row_id'] for r in c),
                     'payment_row_ids':';'.join(r['row_id'] for r in p)})
reserve=sum(r['balance_krw'] for r in balances)
constraints=json.loads((IN/'09_approved_constraints.json').read_text())
terms=json.loads((IN/'08_supplier_terms.json').read_text())
budget=constraints['total_cash_cap_krw']-reserve

roster={r['person_id']:r for r in readcsv('01_roster.csv')}
avail={r['person_id']:r for r in readcsv('02_availability.csv')}
changes=readcsv('03_hr_changes.csv')
people=[]
for pid,r in sorted(roster.items()):
    d={'person_id':pid,'tier':r['tier'],'access_required':r['access_required']=='true',
       'allowed_sessions':avail[pid]['allowed_sessions'].split(';'),'active':r['active']=='true',
       'source_row_ids':[r['row_id'],avail[pid]['row_id']]}
    for c in changes:
        if c['person_id'] != pid or c['approval'] != 'approved': continue
        field=c['field']; value=c['value']
        d[field]=(value=='true') if field in ('active','access_required') else value.split(';') if field=='allowed_sessions' else value
        d['source_row_ids'].append(c['row_id'])
    if d['active']: people.append(d)

facilities=readcsv('07_session_options.csv')
holidays=set(constraints['company_holidays'])
valid_fac=[r for r in facilities if r['approval']=='approved' and constraints['window_start']<=r['date']<=constraints['window_end']
           and date.fromisoformat(r['date']).weekday()<5 and r['date'] not in holidays]

def solve(people, facilities, vendor, count_cap):
    # Unit-capacity min-cost flow; a mandatory seat has priority over any number of optional seats.
    n=len(people); m=len(facilities); source=n+m; sink=source+1; N=sink+1
    graph=[[] for _ in range(N)]
    def edge(u,v,cap,cost):
        graph[u].append([v,cap,cost,len(graph[v])])
        graph[v].append([u,0,-cost,len(graph[u])-1])
    for j,p in enumerate(people):
        edge(source,j,1,-1000 if p['tier']=='mandatory' else -1)
        for k,s in enumerate(facilities):
            if s['session_id'] in p['allowed_sessions'] and (not p['access_required'] or s['accessible']=='true'):
                edge(j,n+k,1,0)
    for k,s in enumerate(facilities):edge(n+k,sink,min(int(s['capacity']),vendor['capacity_per_session']),0)
    flow=0
    while flow<count_cap:
        dist=[float('inf')]*N; dist[source]=0; prev=[None]*N
        # Bellman-Ford is adequate for the small residual network and handles negative prices.
        for _ in range(N):
            changed=False
            for u in range(N):
                if dist[u]==float('inf'):continue
                for ei,e in enumerate(graph[u]):
                    v,cap,cost,_=e
                    if cap and dist[u]+cost<dist[v]:
                        dist[v]=dist[u]+cost; prev[v]=(u,ei); changed=True
            if not changed:break
        if prev[sink] is None or dist[sink]>=0:break
        v=sink
        while v!=source:
            u,ei=prev[v]; e=graph[u][ei]; e[1]-=1; graph[v][e[3]][1]+=1; v=u
        flow+=1
    assignments={}
    for j,p in enumerate(people):
        for e in graph[j]:
            if n<=e[0]<n+m and e[1]==0:
                assignments[p['person_id']]=facilities[e[0]-n]['session_id']; break
    return assignments

mandatory=sum(p['tier']=='mandatory' for p in people)
options=[]
for vendor in terms['vendors']:
    if vendor['approval']!='approved':continue
    permitted=[f for f in valid_fac if f['session_id'] in vendor['allowed_sessions']]
    for combo in itertools.combinations(permitted,constraints['sessions_required']):
        if len({s['date'] for s in combo})<4:continue
        room=sum(int(s['room_cost_krw']) for s in combo)
        cap=min(len(people),(budget-vendor['fixed_krw']-room)//vendor['per_person_krw'])
        if cap<mandatory:continue
        a=solve(people,combo,vendor,cap)
        if sum(p['person_id'] in a for p in people if p['tier']=='mandatory')<mandatory:continue
        if any(s['session_id'] not in a.values() for s in combo):continue
        cost=vendor['fixed_krw']+room+len(a)*vendor['per_person_krw']
        options.append({'vendor':vendor,'combo':combo,'assignments':a,'cost':cost,'room':room,
                        'optional':len(a)-mandatory})
assert options,'No feasible authorized option'
options.sort(key=lambda o:(-o['optional'],o['cost'],o['vendor']['vendor_id'],[s['session_id'] for s in o['combo']]))
best=options[0]; vendor=best['vendor']; combo=best['combo']; assigned=best['assignments']
first=min(date.fromisoformat(s['date']) for s in combo)
prep=[]; d=first-timedelta(days=1)
while len(prep)<constraints['preparation_workdays']:
    if d.weekday()<5 and d.isoformat() not in holidays:prep.append(d.isoformat())
    d-=timedelta(days=1)
prep.sort()

assignment_rows=[]
for p in people:
    sid=assigned.get(p['person_id'],'')
    assignment_rows.append({'person_id':p['person_id'],'tier':p['tier'],'access_required':p['access_required'],
                            'allowed_sessions':p['allowed_sessions'],'session_id':sid,
                            'state':'assigned' if sid else 'waitlist','source_row_ids':p['source_row_ids']})
session_rows=[]
for s in sorted(combo,key=lambda r:r['date']):
    session_rows.append({'session_id':s['session_id'],'date':s['date'],
                         'attendees':sum(x==s['session_id'] for x in assigned.values()),
                         'capacity':min(int(s['capacity']),vendor['capacity_per_session']),
                         'accessible':s['accessible']=='true','room_cost_krw':int(s['room_cost_krw']),
                         'source_id':s['source_id']})
summary={'total_cash_cap_krw':constraints['total_cash_cap_krw'],'historical_outstanding_krw':reserve,
         'launch_budget_krw':budget,'vendor_id':vendor['vendor_id'],'vendor_fixed_krw':vendor['fixed_krw'],
         'vendor_per_person_krw':vendor['per_person_krw'],'room_cost_krw':best['room'],
         'launch_cost_krw':best['cost'],'cash_remaining_krw':budget-best['cost'],
         'assigned_count':len(assigned),'mandatory_count':mandatory,'optional_count':best['optional'],
         'approval_status':'supplier_selection_pending','preparation_dates':prep}
evidence_uses={
 '01_roster.csv':'기초 인원·등급·활성 상태 및 접근성 원본 행을 보존',
 '02_availability.csv':'인원별 최초 가능 세션과 원본 행을 보존',
 '03_hr_changes.csv':'approved 변경만 적용하고 pending 변경은 제외',
 '04_invoices.csv':'송장 ID별 최종 유효 revision과 물리 행을 대사',
 '05_credits.csv':'크레딧 중복·무효·0원 행을 포함해 대사',
 '06_payments.csv':'지급 중복·취소·0원 행을 포함해 대사',
 '07_session_options.csv':'시설 승인·날짜·접근성·방 비용과 수용 인원 확인',
 '08_supplier_terms.json':'최종 승인 공급사 단가와 세션당 용량 적용',
 '09_approved_constraints.json':'현금 한도·근무일·승인 상태·목표 적용',
 '10_prior_maintained_note.md':'기존 운영 가정과 담당 역할을 갱신할 기준으로 사용',
 '11_supplier_correction_email.md':'최종 견적 우선순위와 서명 전 독립 준비 범위 확인',
 '12_domain_rules.md':'revision·대사·배정·준비일 계산 규칙 적용'}
doc_ids={'08_supplier_terms.json':terms['source_id'],'09_approved_constraints.json':constraints['source_id'],
         '10_prior_maintained_note.md':'OPS-NOTE-V3','11_supplier_correction_email.md':'SUPPLIER-EMAIL-0926',
         '12_domain_rules.md':'RULES-0927'}
evidence=[{'file':f,'source_id':doc_ids.get(f,'MULTIPLE'),'use':evidence_uses[f]} for f in sorted(evidence_uses)]
plan={'summary':summary,'assignments':assignment_rows,'sessions':session_rows,'evidence':evidence}

record_headers=['row_id','kind','business_id','disposition','recognized_krw','source_id']
balance_headers=['invoice_id','gross_krw','credit_krw','paid_krw','balance_krw','invoice_row_id','credit_row_ids','payment_row_ids']
assignment_headers=['person_id','tier','access_required','allowed_sessions','session_id','state','source_row_ids']
session_headers=['session_id','date','attendees','capacity','accessible','room_cost_krw','source_id']
evidence_headers=['file','source_id','use']
writecsv('records.csv',record_headers,records)
writecsv('balances.csv',balance_headers,balances)
(OUT/'plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

wb=Workbook(); wb.remove(wb.active)
tables=[('대사기록',record_headers,records),('송장잔액',balance_headers,balances),
        ('배정',assignment_headers,assignment_rows),('일정',session_headers,session_rows),
        ('예산',['key','value'],[{'key':k,'value':v} for k,v in summary.items()]),
        ('근거',evidence_headers,evidence)]
for title,headers,rows in tables:
    ws=wb.create_sheet(title);ws.append(headers)
    for item in rows:
        vals=[]
        for h in headers:
            v=item[h]
            if isinstance(v,list):v=';'.join(v)
            vals.append(v)
        ws.append(vals)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
        cell.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=24
    for col in ws.columns:
        letter=get_column_letter(col[0].column)
        width=min(65,max(14,max(len(str(c.value or '')) for c in col)+2))
        ws.column_dimensions[letter].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,int) and not isinstance(cell.value,bool):cell.number_format='#,##0;[Red](#,##0)'
            cell.alignment=Alignment(vertical='center')
wb.save(OUT/'operations.xlsx')

print(json.dumps({'summary':summary,'options':[{'vendor':o['vendor']['vendor_id'],'sessions':[s['session_id'] for s in o['combo']],
                 'optional':o['optional'],'cost':o['cost']} for o in options],
                 'waitlist':[p['person_id'] for p in people if p['person_id'] not in assigned],
                 'session_counts':{s['session_id']:s['attendees'] for s in session_rows},
                 'record_dispositions':{d:sum(r['disposition']==d for r in records) for d in set(r['disposition'] for r in records)}} ,ensure_ascii=False,indent=2))
