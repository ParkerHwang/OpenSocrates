#!/usr/bin/env python3
"""Exercise fixture invariants, a complete reference artifact set and bad mutations.
Reference artifacts exist only in a temporary directory; never expose them to candidates.
Run with the supplied Python runtime containing openpyxl. Checker itself is stdlib.
"""
import copy
import csv
import json
import tempfile
from pathlib import Path
from openpyxl import Workbook
from check import SCHEMAS, SHEETS, UnsupportedFormula, Workbook as CheckerWorkbook, check
from oracle_generator import compute, ROOT

def write_csv(path,rows,fields):
    with path.open('w',encoding='utf-8',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
        for r in rows:w.writerow({k:';'.join(v) if isinstance(v,list) else v for k,v in r.items()})

def render(output,oracle,formulas=False):
    output.mkdir(exist_ok=True)
    write_csv(output/'records.csv',oracle['records'],SCHEMAS['records'])
    write_csv(output/'balances.csv',oracle['balances'],SCHEMAS['balances'])
    plan=oracle['example_plan'];(output/'plan.json').write_text(json.dumps(plan,ensure_ascii=False),encoding='utf-8')
    (output/'decision_memo.md').write_text('지역 교육 운영 결정 메모\n'+('필수 인원을 전원 배정하고 승인 예산 범위에서 선택 인원을 배정한다. 공급사 선택 서명은 미결이며 자료 준비는 진행한다. '*12),encoding='utf-8')
    wb=Workbook();wb.remove(wb.active)
    sources=dict(records=oracle['records'],balances=oracle['balances'],assignments=plan['assignments'],sessions=plan['sessions'],evidence=plan['evidence'])
    for title,table in SHEETS.items():
        ws=wb.create_sheet(title);ws.append(SCHEMAS[table])
        for r in sources[table]:ws.append([';'.join(r[k]) if isinstance(r[k],list) else r[k] for k in SCHEMAS[table]])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        if formulas and title=='송장잔액':
            for i in range(2,ws.max_row+1):ws.cell(i,5,f'=B{i}-SUM(C{i}:D{i})')
    ws=wb.create_sheet('예산');ws.append(['key','value'])
    for k,v in plan['summary'].items():ws.append([k,';'.join(v) if isinstance(v,list) else v])
    if formulas:ws['B4']="=SUM('예산'!B2)-'예산'!B3+ROUND(2.5,0)+ROUND(-2.5,0)"
    wb.save(output/'operations.xlsx')

def main():
    oracle=compute();frozen=json.loads((ROOT/'oracle.json').read_text())
    assert oracle==frozen,'Oracle drift; regenerate intentionally before freezing'
    assert len(oracle['records'])==101 and len(oracle['people'])==72
    assert oracle['totals']==dict(gross_krw=881200,credit_krw=6600,paid_krw=87000,balance_krw=787600)
    assert oracle['optimum']['assigned_count']==66 and oracle['optimum']['launch_cost_krw']==962000
    # An independent lower bound: every legal four-room combination costs >=260k;
    # all three approved offers exceed the 965k cap at 67 seats. Reference proves 66 feasible.
    assert min(240000+7000*67,200000+8000*67,285000+6500*67)+260000>965000
    assert any(r['disposition']=='active' and r['recognized_krw']==0 for r in oracle['records'])
    cases=[]
    with tempfile.TemporaryDirectory(prefix='office-selfcheck-') as temp:
        base=Path(temp)
        # Direct evaluator controls isolate Excel semantics from fixture arithmetic.
        control_path=base/'formula-controls.xlsx'
        control_book=Workbook();control_sheet=control_book.active;control_sheet.title='검증'
        for i,(literal,amount) in enumerate([('INV001',2),('inv001',3),('P001',7),('A1"완료"',11)],1):
            control_sheet.cell(i,1,literal);control_sheet.cell(i,2,amount)
        controls=[
            ('=COUNTIF(A1:A4,"INV001")',2),
            ('=SUMIF(A1:A4,"INV001",B1:B4)',5),
            ('=COUNTIF(A1:A4,"P001")',1),
            ('=COUNTIF(A1:A4,"A1""완료""")',1),
            ('="INV001"','INV001'),
            ("=SUM('검증'!B1:B2)",5),
            ('=검증!B3',7),
            ('=ROUND(2.5,0)',3),('=ROUND(-2.5,0)',-3),
            ('=ROUND(1.25,1)',1.3),('=ROUND(-1.25,1)',-1.3),
            ('=ROUND(25,-1)',30),('=ROUND(-25,-1)',-30),
            ('=ROUND(2.675,2)',2.68),('=ROUND(-2.675,2)',-2.68),
        ]
        for i,(formula,_) in enumerate(controls,1):control_sheet.cell(i,3,formula)
        control_book.save(control_path);reader=CheckerWorkbook(control_path)
        for i,(formula,expected) in enumerate(controls,1):
            assert reader.value('검증',f'C{i}')==expected,(formula,reader.value('검증',f'C{i}'),expected)
        for formula in ('COUNTIF(A1:A4,"INV*")','COUNTIF(B1:B4,">1")','SUMIF(B1:B4,"<=3",B1:B4)','SUMIF(A1:A4,"P???",B1:B4)'):
            try:reader.formula(formula,'검증',set())
            except UnsupportedFormula:pass
            else:raise AssertionError(f'Unsupported criterion was silently evaluated: {formula}')
        cases.append(dict(case='formula-lexical-rounding-controls',passed=len(controls),unsupported_criteria_detected=4))
        for formulas in (False,True):
            out=base/('correct-formulas' if formulas else 'correct-static');render(out,oracle,formulas)
            result=check(out);assert result.get('all_objective_checks_pass'),result
            cases.append(dict(case=out.name,status=result['status'],passed=len(result['checks'])))
        # Any seating allocation satisfying the contract is accepted, including swapping optional equals.
        alt=copy.deepcopy(oracle);rows=alt['example_plan']['assignments']
        selected=next(r for r in rows if r['tier']=='optional' and r['state']=='assigned' and not r['access_required'])
        wait=next(r for r in rows if r['state']=='waitlist' and not r['access_required'] and selected['session_id'] in r['allowed_sessions'])
        wait['state'],wait['session_id']='assigned',selected['session_id'];selected['state'],selected['session_id']='waitlist',''
        out=base/'alternative';render(out,alt);assert check(out)['all_objective_checks_pass'];cases.append(dict(case='alternative-allocation',passed=True))
        mutations=[]
        def duplicate(o):next(r for r in o['records'] if r['disposition']=='duplicate')['recognized_krw']=5100
        mutations.append(('duplicate-counted',duplicate,'C04'))
        def zero(o):o['records']=[r for r in o['records'] if r['row_id']!='C-B10']
        mutations.append(('zero-row-deleted',zero,'L01'))
        def approval(o):o['example_plan']['summary']['vendor_id']='CEDAR-PILOT';o['example_plan']['summary']['approval_status']='approved'
        mutations.append(('wrong-approval',approval,'P04'))
        def evidence(o):o['example_plan']['evidence'].pop()
        mutations.append(('missing-evidence',evidence,'L04'))
        def arithmetic(o):o['balances'][0]['balance_krw']+=1
        mutations.append(('arithmetic-error',arithmetic,'C08'))
        def prep(o):o['example_plan']['summary']['preparation_dates']=['2026-10-02','2026-10-05','2026-10-06']
        mutations.append(('training-day-as-preparation',prep,'P09'))
        def access(o):o['example_plan']['assignments'][4]['session_id']='S3'
        mutations.append(('access-exception',access,'P02'))
        def stale(o):o['example_plan']['assignments'][9]['tier']='optional'
        mutations.append(('stale-hr',stale,'P01'))
        def copied(o):o['records'].append(dict(o['records'][0]))
        mutations.append(('duplicate-output-row',copied,'L01'))
        for name,mutate,expected in mutations:
            bad=copy.deepcopy(oracle);mutate(bad);out=base/name;render(out,bad);result=check(out)
            failed={c['id'] for c in result['checks'] if not c['passed']}
            assert result['status']=='assessed' and expected in failed and not result['all_objective_checks_pass'],(name,result)
            cases.append(dict(case=name,expected_failure=expected,failed=sorted(failed)))
        out=base/'workbook-arithmetic';render(out,oracle)
        from openpyxl import load_workbook
        wb=load_workbook(out/'operations.xlsx');wb['송장잔액']['E2']=1;wb.save(out/'operations.xlsx')
        result=check(out);assert not next(c for c in result['checks'] if c['id']=='A02')['passed'];cases.append(dict(case='workbook-only-mismatch',expected_failure='A02'))
        for name,formula in [('unsupported-wildcard',"=COUNTIF('송장잔액'!A2:A37,\"INV*\")"),
                             ('unsupported-comparator',"=SUMIF('송장잔액'!B2:B37,\">0\",'송장잔액'!B2:B37)")]:
            out=base/name;render(out,oracle);wb=load_workbook(out/'operations.xlsx');wb['예산']['B2']=formula;wb.save(out/'operations.xlsx')
            result=check(out)
            assert result['status']=='unassessable' and result['objective_score'] is None and result['checks']==[],(name,result)
            cases.append(dict(case=name,status='unassessable'))
        out=base/'missing';render(out,oracle);(out/'operations.xlsx').unlink();result=check(out)
        assert result['status']=='unassessable' and result['objective_score'] is None and not result['checks'];cases.append(dict(case='missing-artifact',status='unassessable'))
        out=base/'corrupt';render(out,oracle);(out/'operations.xlsx').write_bytes(b'not an xlsx');result=check(out)
        assert result['status']=='unassessable' and result['objective_score'] is None;cases.append(dict(case='corrupt-artifact',status='unassessable'))
    print(json.dumps(dict(status='passed',cases=cases,check_count=27),ensure_ascii=False,indent=2))

if __name__=='__main__':main()
