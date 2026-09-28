"""Independent structural, reconciliation and rendering checks for deliverables."""
from __future__ import annotations

import csv
import hashlib
import json
import zipfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent / "vendor"))
import pymupdf
from PIL import Image, ImageDraw
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "evidence"
OUT = ROOT / "deliverables"
DATA = json.loads((OUT / "metrics.json").read_text())
NUM = ["shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units"]


def D(x):
    return Decimal(str(x))


def rows(path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def check_json():
    required = {"monthly", "countries", "fx_monthly", "market_context", "hub_scenarios", "recommendation", "quality"}
    assert required <= DATA.keys()
    assert (len(DATA["monthly"]), len(DATA["countries"]), len(DATA["fx_monthly"]), len(DATA["market_context"]), len(DATA["hub_scenarios"])) == (72,6,24,18,24)
    assert len(DATA["hub_combinations"]) == 88
    for country in DATA["countries"]:
        monthly = [x for x in DATA["monthly"] if x["country"] == country["country"]]
        assert len(monthly) == 12
        for k in NUM:
            assert sum(D(m[k]) for m in monthly) == D(country[k]), (country["country"], k)
        if D(country["net_sales_eur"]) != 0:
            assert abs(D(country["margin"]) - D(country["contribution_eur"]) / D(country["net_sales_eur"])) < D("1e-12")
    for m in DATA["monthly"]:
        assert D(m["gross_sales_eur"]) - D(m["refunds_eur"]) == D(m["net_sales_eur"])
        assert D(m["net_sales_eur"]) - D(m["net_cogs_eur"]) - D(m["fulfillment_eur"]) == D(m["contribution_eur"])
    assert sum(x["shipped_orders"] for x in DATA["countries"]) == 1254
    assert DATA["recommendation"]["countries"] == ["CZE", "ESP"]
    assert DATA["recommendation"]["base_incremental_contribution_eur"] == 198099.31


def check_ledger():
    ledger = rows(EVIDENCE / "order_ledger.csv")
    assert len(ledger) == 1254
    assert len({x["order_id"] for x in ledger}) == len(ledger)
    totals = defaultdict(lambda: defaultdict(Decimal))
    for x in ledger:
        key = (x["country"], x["month"])
        assert D(x["gross_sales_eur"]) - D(x["refunds_eur"]) == D(x["net_sales_eur"])
        assert D(x["gross_cogs_eur"]) - D(x["recovered_cogs_eur"]) == D(x["net_cogs_eur"])
        assert D(x["net_sales_eur"]) - D(x["net_cogs_eur"]) - D(x["fulfillment_eur"]) == D(x["contribution_eur"])
        totals[key]["shipped_orders"] += 1
        totals[key]["shipped_units"] += D(x["quantity"])
        totals[key]["returned_units"] += D(x["returned_units"])
        for k in NUM[2:-1]:
            totals[key][k] += D(x[k])
    for m in DATA["monthly"]:
        for k in NUM:
            assert totals[(m["country"],m["month"])][k] == D(m[k]), (m["country"],m["month"],k)
    order_audit = rows(EVIDENCE / "order_audit.csv")
    return_audit = rows(EVIDENCE / "return_audit.csv")
    assert len(order_audit) == 1297 and len(return_audit) == 243
    assert sum(x["decision"] == "included" for x in order_audit) == 1254
    assert sum(x["decision"] == "included" for x in return_audit) == 241
    assert sum(x["decision"] == "orphan_quarantined" for x in return_audit) == 1
    assert sum(x["decision"] == "after_inclusive_cutoff" for x in return_audit) == 1


def check_sources():
    register = rows(EVIDENCE / "source_register.csv")
    assert len(register) == 13
    for x in register:
        p = ROOT / x["file"]
        assert p.exists()
        assert hashlib.sha256(p.read_bytes()).hexdigest() == x["sha256"]
        assert int(x["bytes"]) == p.stat().st_size
        assert x["local_downloaded_utc"] and x["units"] and x["period"]
    with zipfile.ZipFile(EVIDENCE / "raw" / "ecb-history.zip") as z:
        names = [x for x in z.namelist() if x.lower().endswith(".csv")]
        assert names
        assert any(z.read(n) == (EVIDENCE / "raw" / "ecb-history.csv").read_bytes() for n in names)
    # Independently recompute all 24 arithmetic means from the archived CSV.
    fx = defaultdict(list)
    for x in rows(EVIDENCE / "raw" / "ecb-history.csv"):
        if x["Date"].startswith("2025-"):
            for c in ("PLN","CZK"):
                if x[c] not in ("", "N/A"):
                    fx[(c,x["Date"][:7])].append(D(x[c]))
    for x in DATA["fx_monthly"]:
        vals = fx[(x["currency"],x["month"])]
        assert vals
        assert abs(sum(vals)/len(vals) - D(x["local_per_eur"])) < D("0.00000000005")


def check_scenarios():
    singles = {(x["country"],x["scenario"]): x for x in DATA["hub_scenarios"]}
    combos = DATA["hub_combinations"]
    assert sum(x["feasible"] for x in combos if x["scenario"] == "base") == 18
    for x in combos:
        assert D(x["incremental_contribution_eur"]) == sum((D(singles[(c,x["scenario"])]["incremental_contribution_eur"]) for c in x["countries"]),D(0))
        assert x["capex_eur"] == sum(singles[(c,x["scenario"])]["capex_eur"] for c in x["countries"])
        assert x["fte"] == sum(singles[(c,x["scenario"])]["fte"] for c in x["countries"])
        assert x["feasible"] == (len(x["countries"]) <= 2 and x["capex_eur"] <= 450000 and x["fte"] <= 7)
        if x["incremental_contribution_eur"] <= 0:
            assert x["payback_years"] is None
    for scenario in ("low","base","high"):
        ranked = sorted((x for x in combos if x["feasible"] and x["scenario"] == scenario),key=lambda x:x["incremental_contribution_eur"],reverse=True)
        assert ranked[0]["option"] == "CZE+ESP"
    ranked_stress = sorted((x for x in combos if x["feasible"] and x["scenario"] == "stress"),key=lambda x:x["incremental_contribution_eur"],reverse=True)
    assert ranked_stress[0]["option"] == "NLD+ESP"


def check_workbook():
    wb = load_workbook(OUT / "meridian_analytical_workbook.xlsx", data_only=False)
    assert len(wb["Monthly"]._cells) >= 12*72
    assert len(wb["Decision dashboard"]._charts) == 1
    assert len(wb["Country 2025"]._charts) == 1
    for i,x in enumerate(DATA["countries"],4):
        ws=wb["Country 2025"]
        assert ws.cell(i,1).value == x["country"]
        assert D(ws.cell(i,9).value) == D(x["contribution_eur"])
        assert ws.cell(i,13).value.startswith("=SUMIF")
    for i,x in enumerate(DATA["monthly"],4):
        ws=wb["Monthly"]
        assert ws.cell(i,1).value == x["country"]
        assert ws.cell(i,2).value == x["month"]
        assert D(ws.cell(i,10).value) == D(x["contribution_eur"])
    wb.close()


def check_pdfs():
    rendered = ROOT / "analysis" / "render"
    rendered.mkdir(exist_ok=True)
    panels=[]
    for prefix,name,min_pages in [("report","meridian_executive_report.pdf",4),("slides","meridian_board_presentation.pdf",7)]:
        doc = pymupdf.open(OUT / name)
        assert len(doc) >= min_pages
        for i,page in enumerate(doc):
            content = page.get_text()
            assert len(content) > 150, (name,i)
            page_rect=page.rect
            # Text touching/beyond page edges indicates a rendering problem.
            for block in page.get_text("blocks"):
                if len(block) > 4 and block[4].strip():
                    x0,y0,x1,y1=block[:4]
                    assert x0 >= -1 and y0 >= -1 and x1 <= page_rect.width+1 and y1 <= page_rect.height+1,(name,i,block[:4])
            pix=page.get_pixmap(matrix=pymupdf.Matrix(1.2,1.2),alpha=False)
            target=rendered/f"{prefix}_{i+1:02d}.png"
            pix.save(target)
            im=Image.open(target).convert("RGB")
            scale=min(360/im.width,270/im.height)
            panels.append((target.name,im.resize((int(im.width*scale),int(im.height*scale)))))
        print(f"{name}: {len(doc)} pages, rendered and text-bounds checked")
        doc.close()
    sheet=Image.new("RGB",(4*390,((len(panels)+3)//4)*310),"white")
    d=ImageDraw.Draw(sheet)
    for i,(name,im) in enumerate(panels):
        x=(i%4)*390;y=(i//4)*310
        sheet.paste(im,(x,y+20))
        d.text((x+5,y+2),name,fill="black")
    sheet.save(rendered/"contact_sheet.png")


def main():
    check_json(); print("JSON schema and monthly/country reconciliation: PASS")
    check_ledger(); print("Order ledger and inclusion audit: PASS")
    check_sources(); print("Source hashes, ZIP contents and 24 FX means: PASS")
    check_scenarios(); print("Scenario addition, feasibility and rankings: PASS")
    check_workbook(); print("XLSX values, formulas and charts: PASS")
    check_pdfs(); print("PDF rendering checks: PASS")


if __name__ == "__main__":
    main()
