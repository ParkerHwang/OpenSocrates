#!/usr/bin/env python3
"""Checks the Meridian Parts analytical outputs and deliverables."""

from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "evidence" / "raw"
PROCESSED = ROOT / "analysis" / "processed"
DELIVERABLES = ROOT / "deliverables"
COUNTRIES = ["DEU", "FRA", "NLD", "POL", "CZE", "ESP"]
REQUIRED_AGG = ["country", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin"]
REQUIRED_MONTHLY = ["month"] + REQUIRED_AGG


def read_csv(path: Path):
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def f(v):
    return float(v) if v not in (None, "") else None


def sha256(path: Path):
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def check(condition, message):
    if not condition:
        raise AssertionError(message)


def main():
    metrics = json.loads((DELIVERABLES / "metrics.json").read_text(encoding="utf-8"))
    required = ["monthly", "countries", "fx_monthly", "market_context", "hub_scenarios", "recommendation", "quality"]
    check(all(k in metrics for k in required), "metrics.json missing required top-level key")
    check(len(metrics["monthly"]) == 72, "monthly must contain 72 country-month rows")
    check(len(metrics["countries"]) == 6, "countries must contain six rows")
    check(len(metrics["fx_monthly"]) == 24, "fx_monthly must contain 12 PLN and 12 CZK rows")
    check(len(metrics["market_context"]) == 18, "market_context must contain 18 rows")
    for row in metrics["monthly"]:
        check(all(k in row for k in REQUIRED_MONTHLY), "monthly schema mismatch")
        check(len(row["month"]) == 7 and row["month"].startswith("2025-"), "monthly month format mismatch")
    for row in metrics["countries"]:
        check(all(k in row for k in REQUIRED_AGG), "country schema mismatch")
    check([r["country"] for r in metrics["countries"]] == COUNTRIES, "country order mismatch")
    check(all(r["currency"] in {"PLN", "CZK"} for r in metrics["fx_monthly"]), "unexpected FX currency")
    check(sorted((r["currency"] for r in metrics["fx_monthly"])).count("PLN") == 12, "PLN months incomplete")
    check(sorted((r["currency"] for r in metrics["fx_monthly"])).count("CZK") == 12, "CZK months incomplete")
    check(all(r["year"] in {2022, 2023, 2024} for r in metrics["market_context"]), "market year mismatch")
    for row in metrics["monthly"] + metrics["countries"] + metrics["fx_monthly"] + metrics["market_context"] + metrics["hub_scenarios"]:
        for value in row.values():
            if isinstance(value, float):
                check(math.isfinite(value), "non-finite numeric metric")

    monthly = metrics["monthly"]
    countries = {r["country"]: r for r in metrics["countries"]}
    money_fields = ["gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur"]
    count_fields = ["shipped_orders", "shipped_units", "returned_units"]
    for country in COUNTRIES:
        rows = [r for r in monthly if r["country"] == country]
        for k in count_fields:
            check(sum(r[k] for r in rows) == countries[country][k], f"monthly/country count mismatch {country}/{k}")
        for k in money_fields:
            check(round(sum(r[k] for r in rows), 2) == round(countries[country][k], 2), f"monthly/country money mismatch {country}/{k}")
        ns = countries[country]["net_sales_eur"]
        expected_margin = countries[country]["contribution_eur"] / ns if ns else None
        if expected_margin is None:
            check(countries[country]["margin"] is None, f"zero-sales margin not null {country}")
        else:
            check(abs(countries[country]["margin"] - expected_margin) < 1e-12, f"margin mismatch {country}")

    order_detail = read_csv(PROCESSED / "order_level.csv")
    check(len(order_detail) == 1254, "order detail must contain 1,254 eligible orders")
    check(sum(int(r["quantity"]) for r in order_detail) == sum(r["shipped_units"] for r in metrics["countries"]), "order detail units mismatch")
    check(all(r["month"].startswith("2025-") for r in order_detail), "non-2025 order entered order detail")
    check(not metrics["quality"]["aggregation"]["reconciliation_issues"], "quality reports reconciliation issues")
    check(len(metrics["quality"]["calculation"]["missing_fx_orders"]) == 0, "missing FX orders")
    check(len(metrics["quality"]["calculation"]["missing_unit_cost_orders"]) == 0, "missing unit cost orders")
    check(len(metrics["quality"]["order_handling"]["return_units_exceed_order_units"]) == 0, "returned units exceed order units")

    # Scenario formula spot checks against saved country figures and options.
    options = {r["country"]: r for r in read_csv(PROCESSED / "hub_options.csv")}
    scenario_rows = read_csv(PROCESSED / "hub_scenarios.csv")
    by_option = defaultdict(dict)
    for r in scenario_rows:
        by_option[r["country"]][r["scenario"]] = r
    check(metrics["recommendation"]["option"] == "NLD+ESP", "recommendation changed unexpectedly")
    check(metrics["recommendation"]["countries"] == ["NLD", "ESP"], "recommendation countries mismatch")
    for option in ["NLD+ESP", "CZE+ESP"]:
        members = option.split("+")
        c = sum(countries[m]["contribution_eur"] for m in members)
        u = sum(countries[m]["shipped_units"] for m in members)
        g = sum(countries[m]["gross_sales_eur"] for m in members)
        fixed = sum(float(options[m]["annual_fixed_eur"]) for m in members)
        saving = sum(float(options[m]["saving_eur_per_unit"]) for m in members)
        capex = sum(float(options[m]["capex_eur"]) for m in members)
        for scenario, uplift in [("low", .10), ("base", .25), ("high", .40)]:
            expected = c * uplift + u * (1 + uplift) * saving - fixed
            actual = float(by_option[option][scenario]["incremental_contribution_eur"])
            check(abs(actual - expected) < 0.02, f"scenario formula mismatch {option}/{scenario}")
            check(abs(float(by_option[option][scenario]["capex_eur"]) - capex) < 0.01, f"scenario capex mismatch {option}")
        fx_shock = sum(.10 * countries[m]["net_sales_eur"] for m in members if m in {"POL", "CZE"})
        c_stress = c - .03 * g - fx_shock
        expected_stress = c_stress * 1.25 - c + u * 1.25 * saving - fixed
        actual_stress = float(by_option[option]["stress"]["incremental_contribution_eur"])
        check(abs(actual_stress - expected_stress) < 0.02, f"stress formula mismatch {option}")
    check(abs(float(metrics["recommendation"]["capex_eur"]) - 270000) < 0.01, "recommendation capex mismatch")
    check(metrics["recommendation"]["fte"] == 6, "recommendation FTE mismatch")

    # Provenance hashes and required raw files.
    register = json.loads((ROOT / "evidence" / "source-register.json").read_text(encoding="utf-8"))["sources"]
    check(len(register) >= 14, "source register incomplete")
    for item in register:
        path = RAW / item["file"]
        check(path.exists(), f"registered raw file missing: {item['file']}")
        check(sha256(path) == item["sha256"], f"source hash mismatch: {item['file']}")

    wb = load_workbook(DELIVERABLES / "meridian_parts_analytical_workbook.xlsx", read_only=False, data_only=False)
    expected_sheets = {"Readme", "Country Results", "Monthly Results", "FX Monthly", "Market Context", "Hub Scenarios", "Scenario Summary", "Order Detail", "Sources", "Quality"}
    check(expected_sheets.issubset(set(wb.sheetnames)), "workbook sheets incomplete")
    check(len(wb["Monthly Results"].tables) == 1, "monthly table missing")
    check(len(wb["Hub Scenarios"].tables) == 1, "scenario table missing")
    check(len(wb["Scenario Summary"]._charts) >= 1, "scenario decision chart missing")
    check(len(wb["FX Monthly"]._charts) >= 1, "FX chart missing")

    pdf = PdfReader(str(DELIVERABLES / "meridian_parts_executive_report.pdf"))
    check(len(pdf.pages) >= 5, "report PDF unexpectedly short")
    report_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    for phrase in ["NLD+ESP", "2025 base diagnosis", "scenario arithmetic", "Source register", "31 January 2026"]:
        check(phrase in report_text, f"report missing phrase: {phrase}")

    prs = Presentation(str(DELIVERABLES / "meridian_parts_board_presentation.pptx"))
    check(len(prs.slides) == 10, "board deck must contain 10 slides")
    deck_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
    for phrase in ["NLD + ESP", "CZE + ESP", "90 days", "€270k"]:
        check(phrase in deck_text, f"deck missing phrase: {phrase}")

    print(json.dumps({
        "status": "PASS",
        "metrics": "schema, counts, monthly/country reconciliation, FX, scenario formula and recommendation checks",
        "sources": f"{len(register)} files with matching SHA-256",
        "workbook": f"{len(wb.sheetnames)} sheets, decision/FX charts loaded",
        "report": f"{len(pdf.pages)} PDF pages extracted",
        "deck": f"{len(prs.slides)} slides loaded",
    }, indent=2))


if __name__ == "__main__":
    main()
