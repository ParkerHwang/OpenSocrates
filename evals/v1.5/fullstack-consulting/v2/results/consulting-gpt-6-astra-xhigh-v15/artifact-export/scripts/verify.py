"""Independent raw-input accounting check, artifact checks and PDF rendering."""
import csv
import hashlib
import json
import math
import platform
import re
import zipfile
from collections import defaultdict
from fractions import Fraction as F
from pathlib import Path
from xml.etree import ElementTree as ET
from datetime import datetime, timezone
import openpyxl
import pypdfium2 as pdfium
from pypdf import PdfReader
from PIL import Image, ImageOps, ImageDraw
from formula_check import verify_formulas

ROOT=Path(__file__).resolve().parents[1]
RAW=ROOT/'raw';OUT=ROOT/'deliverables';VER=ROOT/'verification'
VER.mkdir(exist_ok=True)
M=json.loads((OUT/'metrics.json').read_text())
checks=[]
def check(name,condition,detail=''):
    checks.append(dict(name=name,status='PASS' if condition else 'FAIL',detail=detail))
    if not condition:raise AssertionError(name+': '+str(detail))

def csvrows(path):
    with path.open(newline='',encoding='utf-8-sig') as f:return list(csv.DictReader(f))

def canonical(names,key):
    groups=defaultdict(dict)
    raw_count=0
    fingerprints=set()
    for name in names:
        for row in csvrows(RAW/name):
            raw_count+=1
            fingerprint=tuple(sorted(row.items()))
            if fingerprint in fingerprints:continue
            fingerprints.add(fingerprint)
            revision=int(row['revision'])
            check(f'{name}:{row[key]} revision consistency', revision not in groups[row[key]] or groups[row[key]][revision]==row)
            groups[row[key]][revision]=row
    winners={key:revs[max(revs)] for key,revs in groups.items()}
    return winners,raw_count,len(fingerprints),len(groups)

def cents(x):
    # Rational half-up; separate sign ensures half-away-from-zero for negatives.
    sign=-1 if x<0 else 1;x=abs(F(x))*100
    return sign*((2*x.numerator+x.denominator)//(2*x.denominator))

def approx(a,b,tol=1e-10):return abs(float(a)-float(b))<=tol
moneyfields=['gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur']
unitfields=['shipped_orders','shipped_units','returned_units']
fields=unitfields+moneyfields

def financial_checks():
    registers=json.loads((OUT/'source-register.json').read_text())
    for record in registers:
        check('Hash '+record['file'],hashlib.sha256((ROOT/record['file']).read_bytes()).hexdigest()==record['sha256'])
    supplied=json.loads((RAW/'source-register.json').read_text())
    for item in supplied['public']:
        check('Official archive hash '+item['file'],hashlib.sha256((RAW/item['file']).read_bytes()).hexdigest()==item['sha256'])
    with zipfile.ZipFile(RAW/'ecb-history.zip') as z:
        check('ECB extraction lossless',any(z.read(n)==(RAW/'ecb-history.csv').read_bytes() for n in z.namelist() if n.lower().endswith('.csv')))
    rates=defaultdict(list)
    for row in csvrows(RAW/'ecb-history.csv'):
        if row['Date'].startswith('2025-'):
            for cur in ['PLN','CZK']:
                if row[cur] not in ('','N/A'):rates[(cur,row['Date'][:7])].append(F(row[cur]))
    rates={key:sum(vals,F(0))/len(vals) for key,vals in rates.items()}
    for row in M['fx_monthly']:
        check('FX mean '+row['currency']+row['month'],approx(row['local_per_eur'],rates[row['currency'],row['month']],1e-13))
    o,raw_o,unique_rows_o,ids_o=canonical(['orders-part1.csv','orders-part2.csv','order-corrections.csv'],'order_id')
    r,raw_r,unique_rows_r,ids_r=canonical(['returns.csv'],'return_id')
    check('Raw order dispositions', (raw_o,raw_o-unique_rows_o,unique_rows_o-ids_o,ids_o)==(1328,13,18,1297))
    check('Raw return dispositions',(raw_r,raw_r-unique_rows_r,unique_rows_r-ids_r,ids_r)==(264,20,1,243))
    eligible={k:v for k,v in o.items() if v['status']=='shipped' and v['is_test']=='false' and '2025-01-01'<=v['shipped_at']<='2025-12-31'}
    check('Eligible orders',len(eligible)==1254)
    accepted={k:v for k,v in r.items() if v['order_id'] in eligible and v['received_at']<='2026-01-31'}
    check('Cutoff inclusion / exclusion','R-CUTOFF' in accepted and 'R-LATE' not in accepted)
    check('Orphan quarantine','R-ORPHAN' not in accepted and r['R-ORPHAN']['order_id'] not in o)
    check('Revision replacement',accepted['R-M1-01-003']['refund_local']=='149.00')
    byorder=defaultdict(list)
    for row in accepted.values():byorder[row['order_id']].append(row)
    costs=csvrows(RAW/'unit-costs.csv')
    monthly=defaultdict(lambda:{k:0 for k in fields})
    recomputed=[]
    for oid,row in eligible.items():
        date=row['shipped_at'];mon=date[:7];country=row['country'];U=int(row['quantity'])
        fx=F(1) if row['currency']=='EUR' else rates[row['currency'],mon]
        cost=F(max((x for x in costs if x['sku']==row['sku'] and x['valid_from']<=date),key=lambda x:x['valid_from'])['unit_cost_eur'])
        returns=byorder[oid]
        gross=cents((U*F(row['unit_price_local'])-F(row['discount_local']))/fx)
        refund=cents(sum((F(x['refund_local']) for x in returns),F(0))/fx)
        netcogs=cents(U*cost)-cents(sum(int(x['restocked_quantity']) for x in returns)*cost)
        fulfill=cents(F(row['fulfillment_eur']))
        calc=dict(shipped_orders=1,shipped_units=U,gross_sales_eur=gross,refunds_eur=refund,
                  net_sales_eur=gross-refund,net_cogs_eur=netcogs,fulfillment_eur=fulfill,
                  contribution_eur=gross-refund-netcogs-fulfill,returned_units=sum(int(x['quantity']) for x in returns))
        recomputed.append((oid,calc))
        for field in fields:monthly[country,mon][field]+=calc[field]
    check('Half-up rounding boundary',cents(F('1.005'))==101 and cents(F('-1.005'))==-101)
    check('Zero-price retained',sum(v['gross_sales_eur']==0 for _,v in recomputed)==12)
    check('Monthly shape',len(M['monthly'])==72 and len({(x['country'],x['month']) for x in M['monthly']})==72)
    for row in M['monthly']:
        calc=monthly[row['country'],row['month']]
        for field in fields:
            expected=calc[field]/100 if field in moneyfields else calc[field]
            check('Monthly '+row['country']+row['month']+' '+field,approx(row[field],expected,1e-8))
        expected=None if calc['net_sales_eur']==0 else calc['contribution_eur']/calc['net_sales_eur']
        check('Monthly margin '+row['country']+row['month'],row['margin']==expected if expected is None else approx(row['margin'],expected,1e-14))
    for row in M['countries']:
        sums={key:sum(v[key] for (c,_),v in monthly.items() if c==row['country']) for key in fields}
        for key in fields:check('Country '+row['country']+' '+key,approx(row[key],sums[key]/100 if key in moneyfields else sums[key],1e-8))
        check('Country margin '+row['country'],approx(row['margin'],sums['contribution_eur']/sums['net_sales_eur'],1e-14))
    ledger={r['order_id']:r for r in csvrows(OUT/'order-ledger.csv')}
    for oid,calc in recomputed:
        for k in fields:
            check('Ledger '+oid+' '+k,F(ledger[oid][k])==(F(calc[k],100) if k in moneyfields else calc[k]))
    options={r['country']:r for r in csvrows(RAW/'hub-options.csv')}
    country={r['country']:r for r in M['countries']}
    scmap={}
    for row in M['hub_scenarios']:
        c=row['country'];sc=row['scenario'];base=country[c];option=options[c]
        cv=F(str(base['contribution_eur']));g=F(str(base['gross_sales_eur']));n=F(str(base['net_sales_eur']));u=F({'low':'0.1','base':'0.25','high':'0.4','stress':'0.25'}[sc])
        shock=F('.03')*g+(F('.10')*n if c in ['POL','CZE'] else 0) if sc=='stress' else 0
        expected=cents((cv-shock)*(1+u)-cv+base['shipped_units']*(1+u)*F(option['saving_eur_per_unit'])-F(option['annual_fixed_eur']))/100
        check('Scenario '+c+sc,approx(expected,row['incremental_contribution_eur'],1e-8))
        expectedpayback=float(option['capex_eur'])/expected if expected>0 else None
        check('Payback '+c+sc,row['payback_years'] is None if expectedpayback is None else approx(row['payback_years'],expectedpayback,1e-10))
        scmap[c,sc]=row
    check('Scenario complete grid',len(scmap)==24)
    check('Portfolio complete grid',len(M['portfolio_scenarios'])==22 and len({p['option'] for p in M['portfolio_scenarios']})==22)
    for row in M['portfolio_scenarios']:
        k=sum(int(options[c]['capex_eur']) for c in row['countries']);fte=sum(int(options[c]['fte']) for c in row['countries'])
        check('Portfolio feasibility '+row['option'],row['feasible']==(k<=450000 and fte<=7 and len(row['countries'])<=2))
        for sc in ['low','base','high','stress']:
            expected=sum(scmap[c,sc]['incremental_contribution_eur'] for c in row['countries'])
            check('Portfolio '+row['option']+sc,approx(row[sc+'_incremental_eur'],expected,1e-7))
    for fname,field in [('population.json','population'),('gdp-per-capita.json','gdp_per_capita_usd')]:
        meta,rows=json.loads((RAW/fname).read_text())
        expected={(r['countryiso3code'],int(r['date'])):r['value'] for r in rows}
        check('World Bank vintage '+fname,meta['lastupdated']=='2026-07-13')
        for row in M['market_context']:check('World Bank '+field+row['country']+str(row['year']),row[field]==expected[row['country'],row['year']])
    check('Correct recommendation resource sum',M['recommendation']['countries']==['CZE','ESP'] and M['recommendation']['capex_eur']==225000 and M['recommendation']['fte']==5)

def workbook_checks():
    path=OUT/'Meridian_analytical_workbook.xlsx'
    formulas=openpyxl.load_workbook(path,data_only=False)
    values=openpyxl.load_workbook(path,data_only=True)
    check('Workbook sheets',len(values.sheetnames)==17)
    for tab,key in [('Monthly','monthly'),('Countries','countries')]:
        ws=values[tab];headers=[cell.value for cell in ws[5]]
        for i,row in enumerate(M[key],6):
            for j,field in enumerate(headers,1):
                val=ws.cell(i,j).value;target=row[field]
                check(f'XLSX {tab} {i} {field}',approx(val,target,1e-8) if isinstance(target,(int,float)) else val==target)
    sh=values['Scenarios'];headers=[x.value for x in sh[5]]
    for i,row in enumerate(M['hub_scenarios'],6):
        for j,field in enumerate(headers,1):
            val=sh.cell(i,j).value;target=row[field]
            check(f'XLSX scenario {i} {field}',approx(val,target,1e-8) if isinstance(target,(int,float)) else val in [None,''] if target is None else val==target)
    check('Workbook controls all pass',all(values['Checks'].cell(i,5).value=='PASS' for i in range(6,6+len(fields))))
    formula_count=sum(cell.data_type=='f' for ws in formulas for row in ws for cell in row)
    check('Workbook formulas present',formula_count>1000,str(formula_count))
    evaluated=verify_formulas(formulas,values)
    check('Workbook supported formulas independently evaluated',evaluated==formula_count,str(evaluated)+' formula results match caches; scoped Decimal evaluator, not native Excel')
    for ws in values:
        errors=[cell.coordinate for row in ws for cell in row if cell.data_type=='e']
        check('No cached errors '+ws.title,not errors,str(errors))
    with zipfile.ZipFile(path) as z:
        check('Workbook archive integrity',z.testzip() is None)
        charts=[n for n in z.namelist() if re.fullmatch(r'xl/charts/chart\d+\.xml',n)]
        check('Native decision charts',len(charts)==2)
        ns={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
        for name in charts:
            root=ET.fromstring(z.read(name));points=root.findall('.//c:numCache/c:pt/c:v',ns)
            check('Chart cached values '+name,bool(points))
        check('No broken formula reference',all(b'#REF!' not in z.read(n) for n in z.namelist() if n.endswith('.xml')))
    formulas.close();values.close()

def pdf_checks():
    render=VER/'rendered';render.mkdir(exist_ok=True)
    expected={'Meridian_executive_report.pdf':12,'Meridian_board_presentation.pdf':11}
    for filename,pages in expected.items():
        path=OUT/filename;reader=PdfReader(path)
        check('PDF page count '+filename,len(reader.pages)==pages)
        texts=[page.extract_text() for page in reader.pages]
        check('PDF text exists '+filename,all(len(t)>250 for t in texts))
        check('PDF no placeholder '+filename,not re.search(r'\b(TODO|TBD|Lorem ipsum)\b',' '.join(texts)))
        check('PDF source links '+filename,sum(len(p.get('/Annots',[])) for p in reader.pages)>10)
        (VER/(path.stem+'-text.txt')).write_text('\n\n'.join(f'PAGE {i+1}\n'+t for i,t in enumerate(texts)))
        pdf=pdfium.PdfDocument(path)
        thumbnails=[]
        for i in range(len(pdf)):
            page=pdf[i];bitmap=page.render(scale=1.4);im=bitmap.to_pil().convert('RGB')
            dest=render/f'{path.stem}-{i+1:02}.png';im.save(dest)
            thumb=im.copy();thumb.thumbnail((300,370 if pages==12 else 180))
            tile=Image.new('RGB',(320,400 if pages==12 else 210),'#dce5ec')
            tile.paste(thumb,((320-thumb.width)//2,10));ImageDraw.Draw(tile).text((10,tile.height-20),f'{path.stem} / {i+1}',fill='#162D43')
            thumbnails.append(tile)
            bitmap.close();page.close()
        pdf.close()
        cols=3;rows=math.ceil(len(thumbnails)/cols);th=thumbnails[0].height
        contact=Image.new('RGB',(cols*320,rows*th),'white')
        for i,tile in enumerate(thumbnails):contact.paste(tile,((i%cols)*320,(i//cols)*th))
        contact.save(VER/(path.stem+'-contact.png'))
    for name in ['report-layout.json','deck-layout.json']:
        boxes=json.loads((VER/name).read_text());height=841.89 if name.startswith('report') else 540;width=595.28 if name.startswith('report') else 960
        check('PDF recorded layout inside page '+name,all(b['x']>=0 and b['y']>=49 and b['x']+b['w']<=width+.1 and b['y']+b['h']<=height for b in boxes))

def main():
    financial_checks();workbook_checks();pdf_checks()
    errors=[x for x in checks if x['status']!='PASS']
    output=dict(run_at_utc=datetime.now(timezone.utc).isoformat(),python=platform.python_version(),
                checks_run=len(checks),passed=len(checks)-len(errors),failed=len(errors),checks=checks,
                limitations=['No native Excel/LibreOffice application rendering or recalculation was performed; formula caches, references and XML were inspected.',
                             'PDFs rasterized for visual inspection; contact sheets and full-page PNGs preserved. Human-style visual review is recorded separately.',
                             'No general-ledger, bank, customer, vendor or real-world pilot access; analysis verifies source-room consistency, not real-world causality.'])
    (VER/'checks.json').write_text(json.dumps(output,indent=2)+'\n')
    print(f'PASS: {len(checks):,} computational and structural assertions; PDFs rasterized.')

if __name__=='__main__':main()
