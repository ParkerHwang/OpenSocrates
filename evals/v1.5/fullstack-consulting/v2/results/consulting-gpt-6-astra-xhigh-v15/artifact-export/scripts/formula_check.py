"""Evaluate the workbook's supported formula subset, independent of cached values.

This is a scoped formula check, not a claim of native Excel recalculation. It
uses Decimal semantics for financial arithmetic and lazy IF evaluation.
"""
import re
from decimal import Decimal, ROUND_HALF_UP
from openpyxl.utils.cell import range_boundaries

TOKEN=re.compile(r'''\s*(?:(?P<sheet>'(?:[^']|'')+'!|[A-Za-z_][A-Za-z0-9_]*!)|(?P<string>"(?:[^"]|"")*")|(?P<cell>\$?[A-Z]{1,3}\$?\d+)|(?P<num>\d+(?:\.\d*)?(?:[eE][+-]?\d+)?)|(?P<ident>[A-Za-z_][A-Za-z0-9_]*)|(?P<op><=|>=|<>|[+*/=<>(),:\-]))''')

class Parser:
    def __init__(self,s):
        self.tokens=[];pos=0
        while pos<len(s):
            m=TOKEN.match(s,pos)
            if not m:raise ValueError('Unsupported formula token '+s[pos:])
            self.tokens.append((m.lastgroup,m.group(m.lastgroup)));pos=m.end()
        self.i=0
    def peek(self):return self.tokens[self.i][1] if self.i<len(self.tokens) else None
    def take(self,val=None):
        item=self.tokens[self.i];self.i+=1
        if val is not None and item[1]!=val:raise ValueError(f'Expected {val}, got {item}')
        return item
    def parse(self):
        result=self.expr()
        if self.i!=len(self.tokens):raise ValueError('Unused formula tokens')
        return result
    def expr(self):
        left=self.add()
        while self.peek() in ['=','<>','>','<','>=','<=']:
            op=self.take()[1];left=('op',op,left,self.add())
        return left
    def add(self):
        left=self.mul()
        while self.peek() in ['+','-']:
            op=self.take()[1];left=('op',op,left,self.mul())
        return left
    def mul(self):
        left=self.atom()
        while self.peek() in ['*','/']:
            op=self.take()[1];left=('op',op,left,self.atom())
        return left
    def atom(self):
        kind,val=self.take()
        if val=='-':return ('op','-',('value',Decimal(0)),self.atom())
        if val=='(':
            node=self.expr();self.take(')');return node
        if kind=='num':return ('value',Decimal(val))
        if kind=='string':return ('value',val[1:-1].replace('""','"'))
        if kind=='ident':
            if val in ('TRUE','FALSE'):return ('value',val=='TRUE')
            self.take('(');args=[]
            if self.peek()!=')':
                args.append(self.expr())
                while self.peek()==',':self.take(',');args.append(self.expr())
            self.take(')');return ('call',val,args)
        sheet=None
        if kind=='sheet':
            sheet=val[:-1].strip("'").replace("''", "'");kind,val=self.take()
        if kind=='cell':
            first=val.replace('$','');last=first
            if self.peek()==':':self.take(':');last=self.take()[1].replace('$','')
            return ('ref',sheet,first,last)
        raise ValueError('Unsupported formula atom '+str((kind,val)))

class Evaluator:
    def __init__(self,workbook):self.wb=workbook;self.memo={};self.active=set()
    def cell(self,sheet,coordinate):
        key=(sheet,coordinate)
        if key in self.memo:return self.memo[key]
        if key in self.active:raise ValueError('Circular reference '+str(key))
        self.active.add(key)
        cell=self.wb[sheet][coordinate]
        if cell.data_type=='f':result=self.eval(Parser(cell.value[1:]).parse(),sheet)
        elif isinstance(cell.value,(int,float)) and not isinstance(cell.value,bool):result=Decimal(str(cell.value))
        else:result=cell.value
        self.active.remove(key);self.memo[key]=result;return result
    def eval(self,node,sheet):
        typ=node[0]
        if typ=='value':return node[1]
        if typ=='ref':
            _,source,start,end=node;source=source or sheet
            if start==end:return self.cell(source,start)
            c1,r1,c2,r2=range_boundaries(start+':'+end)
            return [self.cell(source,self.wb[source].cell(r,c).coordinate) for r in range(r1,r2+1) for c in range(c1,c2+1)]
        if typ=='op':
            op=node[1];a=self.eval(node[2],sheet);b=self.eval(node[3],sheet)
            if op=='+':return a+b
            if op=='-':return a-b
            if op=='*':return a*b
            if op=='/':return a/b
            if op=='=':return a==b
            if op=='<>':return a!=b
            if op=='>':return a>b
            if op=='<':return a<b
            if op=='>=':return a>=b
            if op=='<=':return a<=b
        if typ=='call':
            name,args=node[1:]
            if name=='IF':return self.eval(args[1] if self.eval(args[0],sheet) else args[2],sheet)
            vals=[self.eval(x,sheet) for x in args]
            if name=='AND':return all(vals)
            if name=='OR':return any(vals)
            if name=='ROUND':return vals[0].quantize(Decimal(10)**(-int(vals[1])),rounding=ROUND_HALF_UP)
            if name=='SUM':
                flat=[n for v in vals for n in (v if isinstance(v,list) else [v])]
                return sum((x for x in flat if isinstance(x,Decimal)),Decimal(0))
            if name=='SUMIF':
                match,criterion,values=vals
                return sum((v for m,v in zip(match,values) if m==criterion),Decimal(0))
            if name=='SUMIFS':
                values=vals[0];filters=[(vals[i],vals[i+1]) for i in range(1,len(vals),2)]
                return sum((v for i,v in enumerate(values) if all(cells[i]==criterion for cells,criterion in filters)),Decimal(0))
        raise ValueError('Unsupported formula '+str(node))

def verify_formulas(formulas,caches):
    ev=Evaluator(formulas);count=0
    for ws in formulas:
        for row in ws:
            for cell in row:
                if cell.data_type!='f':continue
                actual=ev.cell(ws.title,cell.coordinate);expected=caches[ws.title][cell.coordinate].value
                if isinstance(actual,Decimal):
                    good=abs(actual-Decimal(str(expected)))<Decimal('.00000001')
                elif actual=='':good=expected in (None,'')
                else:good=actual==expected
                if not good:raise AssertionError(f'Formula cache mismatch {ws.title}!{cell.coordinate}: {actual} vs {expected}; {cell.value}')
                count+=1
    return count
