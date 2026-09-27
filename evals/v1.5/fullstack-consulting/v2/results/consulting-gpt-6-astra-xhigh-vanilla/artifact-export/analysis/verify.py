"""Independent exact-rational recalculation plus deliverable consistency checks.

Does not import analyze.py. Uses fractions and integer cents, not Decimal model code.
PDF rendering review is separately documented in VERIFICATION.md.
"""
from pathlib import Path
from fractions import Fraction as F
from collections import defaultdict, Counter
from itertools import combinations
import csv, json, math, hashlib, zipfile, re
import openpyxl
from pypdf import PdfReader
import pypdfium2 as pdfium

ROOT=Path(__file__).resolve().parents[1];RAW=ROOT/'evidence/raw';OUT=ROOT/'deliverables'
M=json.loads((OUT/'metrics.json').read_text());counts=Counter();failures=[]
def test(group,name,condition):
    counts[group]+=1
    if not condition:failures.append({'group':group,'check':name})
def cents(value):
    f=F(value)*100
    return (1 if f>=0 else -1)*((2*abs(f.numerator)+f.denominator)//(2*f.denominator))
def close(a,b,tol=1e-9):return a is None and b is None or a is not None and b is not None and abs(float(a)-float(b))<=tol
def raw(name):
    with (RAW/name).open() as f:return list(csv.DictReader(f))
def final_records(records,key):
    result={}
    for r in sorted(records,key=lambda r:int(r['revision'])):result[r[key]]=r
    return result

test('rounding','positive exact tie',cents(F('1.005'))==101)
test('rounding','negative exact tie',cents(F('-1.005'))==-101)
test('rounding','below tie',cents(F('1.0049999'))==100)

for r in json.loads((ROOT/'evidence/collection_manifest.json').read_text()):
    test('source_integrity',r['file'],hashlib.sha256((RAW/r['file']).read_bytes()).hexdigest()==r['sha256'])
oraw=raw('orders-part1.csv')+raw('orders-part2.csv')+raw('order-corrections.csv');rraw=raw('returns.csv')
orders=final_records(oraw,'order_id');rets=final_records(rraw,'return_id')
eligible={k:r for k,r in orders.items() if r['status']=='shipped' and r['is_test']=='false' and '2025-01-01'<=r['shipped_at']<='2025-12-31'}
test('quality','row and ID counts',(len(oraw),len(orders),len(eligible),len(rraw),len(rets))==(1328,1297,1254,264,243))
test('quality','order duplicate count',len(oraw)-len({tuple(r.items()) for r in oraw})==13)
test('quality','return duplicate count',len(rraw)-len({tuple(r.items()) for r in rraw})==20)
test('quality','18 order corrections selected',sum(int(r['revision'])>1 for r in eligible.values())==18)
test('quality','future shipment absent','FUTURE-2026' not in eligible)
test('quality','zero prices retained',sum(F(r['unit_price_local'])==0 for r in eligible.values())==12)
test('quality','return corrected credit',rets['R-M1-01-003']['refund_local']=='149.00')
included=[r for r in rets.values() if r['order_id'] in eligible and r['received_at']<='2026-01-31']
test('quality','241 return IDs eligible',len(included)==241)
test('quality','cutoff included, late/orphan excluded',{'R-CUTOFF'}.issubset({r['return_id'] for r in included}) and not {'R-LATE','R-ORPHAN'}.intersection({r['return_id'] for r in included}))
linked=defaultdict(list)
for r in included:linked[r['order_id']].append(r)
days=defaultdict(list)
for r in raw('ecb-history.csv'):
    if r['Date'].startswith('2025-'):
        for cc in ['PLN','CZK']:
            if r[cc].strip() not in ['', 'N/A']:days[cc,r['Date'][:7]].append(F(r[cc]))
fx={key:sum(v,F(0))/len(v) for key,v in days.items()}
test('fx','24 currency/month keys',len(fx)==24)
for r in M['fx_monthly']:
    test('fx',str((r['currency'],r['month'])),close(fx[r['currency'],r['month']],r['local_per_eur'],1e-12) and len(days[r['currency'],r['month']])==r['observations'])
uc=raw('unit-costs.csv');agg=defaultdict(Counter);recomputed={}
for oid,r in eligible.items():
    n=int(r['quantity']); month=r['shipped_at'][:7];rate=F(1) if r['currency']=='EUR' else fx[r['currency'],month]
    cost=F(sorted([c for c in uc if c['sku']==r['sku'] and c['valid_from']<=r['shipped_at']],key=lambda c:c['valid_from'])[-1]['unit_cost_eur'])
    rr=linked[oid];returned=sum(int(t['quantity']) for t in rr);restocked=sum(int(t['restocked_quantity']) for t in rr)
    gross=cents((n*F(r['unit_price_local'])-F(r['discount_local']))/rate)
    refund=cents(sum((F(t['refund_local']) for t in rr),F(0))/rate)
    gc=cents(n*cost);recovery=cents(restocked*cost);nc=gc-recovery;fulfill=cents(F(r['fulfillment_eur']))
    fields=dict(shipped_orders=1,shipped_units=n,gross_sales_eur=gross,refunds_eur=refund,net_sales_eur=gross-refund,net_cogs_eur=nc,fulfillment_eur=fulfill,contribution_eur=gross-refund-nc-fulfill,returned_units=returned,gross_cogs_eur=gc,recovered_cogs_eur=recovery,restocked_units=restocked)
    recomputed[oid]=fields
    for key in [(r['country'],month),(r['country'],'annual'),('TOTAL','annual')]:agg[key].update(fields)
with (OUT/'data/order_ledger.csv').open() as f:ledger=list(csv.DictReader(f))
for r in ledger:
    for key,v in recomputed[r['order_id']].items():test('independent_order_recalculation',r['order_id']+' '+key,(cents(F(r[key])) if key.endswith('_eur') else int(r[key]))==v)
for r in M['monthly']+M['countries']+[M['totals']]:
    expected=agg[r['country'],r.get('month','annual')]
    for key,v in expected.items():test('independent_aggregation',str((r['country'],r.get('month'),key)),(cents(str(r[key])) if key.endswith('_eur') else r[key])==v)
    test('independent_aggregation','margin',close(r['margin'],F(expected['contribution_eur'],expected['net_sales_eur']) if expected['net_sales_eur'] else None,1e-12))
test('schema','72 monthly / 6 country results',len(M['monthly'])==72 and len(M['countries'])==6)
test('schema','required top-level fields',all(k in M for k in ['monthly','countries','fx_monthly','market_context','hub_scenarios','recommendation','quality']))
for filename,key in [('population.json','population'),('gdp-per-capita.json','gdp_per_capita_usd')]:
    data=json.loads((RAW/filename).read_text())
    lookup={(o['countryiso3code'],int(o['date'])):o['value'] for o in data[1]}
    for r in M['market_context']:test('world_bank',key+str((r['country'],r['year'])),r[key]==lookup[r['country'],r['year']])
options={r['country']:r for r in raw('hub-options.csv')};income={}
for r in M['hub_scenarios']:
    c=r['country'];o=options[c];a=agg[c,'annual'];u={'low':F(1,10),'base':F(1,4),'high':F(2,5),'stress':F(1,4)}[r['scenario']]
    C=F(a['contribution_eur'],100);G=F(a['gross_sales_eur'],100);N=F(a['net_sales_eur'],100);U=a['shipped_units']
    shock=(F(3,100)*G+(F(1,10)*N if c in ['POL','CZE'] else 0)) if r['scenario']=='stress' else 0
    annual=cents((C-shock)*(1+u)-C+U*(1+u)*F(o['saving_eur_per_unit'])-F(o['annual_fixed_eur']))
    income[c,r['scenario']]=annual
    test('independent_scenarios',str((c,r['scenario'])),cents(str(r['incremental_contribution_eur']))==annual)
    test('independent_scenarios','payback '+str((c,r['scenario'])),close(r['payback_years'],F(o['capex_eur'])/(F(annual,100)) if annual>0 else None))
    test('independent_scenarios','capex/FTE '+c,F(o['capex_eur'])==F(str(r['capex_eur'])) and int(o['fte'])==r['fte'])
test('portfolios','22 combinations',len(M['portfolio_scenarios'])==22)
test('portfolios','11 feasible pairs',sum(len(r['countries'])==2 and r['feasible'] for r in M['portfolio_scenarios'])==11)
for r in M['portfolio_scenarios']:
    capex=sum(F(options[c]['capex_eur']) for c in r['countries']);fte=sum(int(options[c]['fte']) for c in r['countries'])
    test('portfolios','resources '+r['option'],capex==r['capex_eur'] and fte==r['fte'] and r['feasible']==(capex<=450000 and fte<=7 and len(r['countries'])<=2))
    for sc in ['low','base','high','stress']:
        amount=sum(income[c,sc] for c in r['countries'])
        test('portfolios',r['option']+' '+sc,cents(str(r[sc+'_eur']))==amount and close(r[sc+'_payback_years'],capex/F(amount,100) if amount>0 else None))
for sc in ['low','base','high']:
    winner=max([r for r in M['portfolio_scenarios'] if r['feasible']],key=lambda r:r[sc+'_eur'])
    test('recommendation','winner '+sc,winner['countries']==M['recommendation']['countries'])

path=OUT/'Meridian_analytical_workbook.xlsx';wf=openpyxl.load_workbook(path,data_only=False);wv=openpyxl.load_workbook(path,data_only=True)
with zipfile.ZipFile(path) as z:
    test('workbook','ZIP integrity',z.testzip() is None)
    chartfiles=[n for n in z.namelist() if re.match(r'xl/charts/chart\d+.xml$',n)]
    test('workbook','3 native charts',len(chartfiles)==3)
formula_count=0
for sh in wf:
    for row in sh:
        for cell in row:
            if cell.data_type=='f':
                formula_count+=1;v=wv[sh.title][cell.coordinate].value
                test('workbook_formula_cache',sh.title+'!'+cell.coordinate,v is not None or '""' in cell.value)
                test('workbook_formula_cache','no error '+sh.title+'!'+cell.coordinate,'#REF!' not in cell.value and v not in ['#REF!','#VALUE!','#DIV/0!','#NAME?','#N/A','#NUM!'])
    test('workbook','freeze panes '+sh.title,sh.freeze_panes is not None)
for sheet,key in [('Monthly','monthly'),('Countries','countries')]:
    sh=wv[sheet];headers=[c.value for c in sh[5]]
    for i,r in enumerate(M[key],6):
        for j,f in enumerate(headers,1):
            value=sh.cell(i,j).value
            test('workbook_metric_cache',sheet+f'!{i},{j}',close(value,r[f],1e-8) if isinstance(r[f],(int,float)) else value==r[f])
for i,r in enumerate(M['hub_scenarios'],6):
    for col,key in [(13,'incremental_contribution_eur'),(14,'capex_eur'),(15,'fte'),(16,'payback_years')]:test('workbook_metric_cache','hub '+str((i,col)),close(wv['Hub Scenarios'].cell(i,col).value,r[key]))
for i,r in enumerate(M['portfolio_scenarios'],6):
    for col,key in [(2,'capex_eur'),(3,'fte'),(7,'low_eur'),(8,'base_eur'),(9,'high_eur'),(10,'stress_eur'),(11,'low_payback_years'),(12,'base_payback_years'),(13,'high_payback_years'),(14,'stress_payback_years')]:test('workbook_metric_cache','portfolio '+str((i,col)),close(wv['Portfolios'].cell(i,col).value,r[key]))
for row in wv['Reconciliation'].iter_rows(min_row=6,max_row=11,min_col=2,max_col=10):
    for cell in row:test('workbook_metric_cache','zero monthly/annual difference',cell.value==0)

pdf_results=[]
for name,expected in [('Meridian_executive_report.pdf',12),('Meridian_board_presentation.pdf',10)]:
    reader=PdfReader(OUT/name);test('pdf',name+' page count',len(reader.pages)==expected)
    full='\n'.join(p.extract_text() for p in reader.pages)
    test('pdf',name+' no missing glyphs','■' not in full and '\ufffd' not in full)
    test('pdf',name+' recommendation values','225,000' in full and ('198.1' in full or '198,099.31' in full))
    link_count=0
    for i,page in enumerate(reader.pages):
        test('pdf',name+f' page {i+1} nonempty',len(page.extract_text())>400)
        for a in page.get('/Annots',[]):
            obj=a.get_object();action=obj.get('/A',{});uri=action.get('/URI','')
            if uri:
                link_count+=1
                if uri.startswith('../evidence/'):
                    test('pdf_evidence_links',uri,(OUT/uri).is_file())
    pdf=pdfium.PdfDocument(OUT/name);bounds=[]
    for i in range(len(pdf)):
        page=pdf[i];text=page.get_textpage();w,h=page.get_size()
        for j in range(text.count_chars()):
            char=text.get_text_range(j,1)
            if not char.strip():continue
            x0,y0,x1,y1=text.get_charbox(j)
            if x0 < -1 or y0 < -1 or x1>w+1 or y1>h+1:bounds.append((i+1,j,char,(x0,y0,x1,y1)))
        text.close();page.close()
    pdf.close();test('pdf',name+' all glyphs within page',len(bounds)==0)
    pdf_results.append({'file':name,'pages':len(reader.pages),'searchable_text_characters':len(full),'links':link_count,'out_of_page_glyphs':len(bounds)})

result={'status':'PASS' if not failures else 'FAIL','checks':dict(counts),'total_checks':sum(counts.values()),'failed':len(failures),'failures':failures,
    'independent_method':'Raw files recalculated with fractions.Fraction and integer half-up cents, without importing the production Decimal analysis.',
    'workbook':{'sheets':len(wf.sheetnames),'formula_cells':formula_count,'charts':3,'cache_checked':True,'native_spreadsheet_recalculation':False},
    'pdfs':pdf_results,
    'limitations':['No desktop Excel/LibreOffice engine installed; formulas, cached values, source links and native chart structures were checked, but interactive native recalculation was not executed.',
        'PDF page bounds checks do not alone prove absence of overlaps; rendered pages were separately visually inspected.',
        'The analysis verifies implementation against supplied rules, not the commercial validity of synthetic assumptions.']}
(ROOT/'verification/verification_report.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
if failures:raise SystemExit(1)
