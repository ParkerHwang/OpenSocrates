"""Evaluate the workbook's actual formula strings with a restricted local interpreter.

This is deliberately not a native Excel compatibility claim. It verifies all
functions used here and exercises two meaningful changes to editable inputs.
"""
from pathlib import Path
from fractions import Fraction as F
from functools import lru_cache
from fnmatch import fnmatchcase
import json
import openpyxl
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries

ROOT=Path(__file__).resolve().parents[1]
PATH=ROOT/'deliverables/Meridian_analytical_workbook.xlsx'
WB=openpyxl.load_workbook(PATH,data_only=False)
CACHE=openpyxl.load_workbook(PATH,data_only=True)

def halfup(x,digits):
    q=10**int(digits);v=x*q
    return F((1 if v>=0 else -1)*((2*abs(v.numerator)+v.denominator)//(2*v.denominator)),q)

class Parser:
    priority={'=':1,'<>':1,'>':1,'<':1,'>=':1,'<=':1,'&':2,'+':3,'-':3,'*':4,'/':4}
    def __init__(self,formula):self.tokens=[t for t in Tokenizer(formula).items if t.type!='WHITE-SPACE'];self.i=0
    def pop(self):t=self.tokens[self.i];self.i+=1;return t
    def expression(self,minimum=0):
        t=self.pop()
        if t.type=='OPERAND':node=('operand',t.subtype,t.value)
        elif t.type=='OPERATOR-PREFIX':node=('unary',t.value,self.expression(5))
        elif t.type=='PAREN' and t.subtype=='OPEN':
            node=self.expression();assert self.pop().subtype=='CLOSE'
        elif t.type=='FUNC' and t.subtype=='OPEN':
            args=[]
            if self.tokens[self.i].subtype!='CLOSE':
                while True:
                    args.append(self.expression())
                    if self.tokens[self.i].type!='SEP':break
                    self.pop()
            assert self.pop().subtype=='CLOSE';node=('function',t.value[:-1].upper(),args)
        else:raise ValueError((t.value,t.type,t.subtype))
        while self.i<len(self.tokens):
            n=self.tokens[self.i]
            if n.type!='OPERATOR-INFIX' or self.priority[n.value]<minimum:break
            self.pop();node=('binary',n.value,node,self.expression(self.priority[n.value]+1))
        return node
    def parse(self):
        node=self.expression();assert self.i==len(self.tokens);return node

class Evaluator:
    def __init__(self,overrides=None):self.overrides=overrides or {}
    @lru_cache(None)
    def cell(self,sheet,coord):
        coord=coord.replace('$','')
        value=self.overrides.get((sheet,coord),WB[sheet][coord].value)
        if isinstance(value,str) and value.startswith('='):return self.run(Parser(value).parse(),sheet)
        if isinstance(value,bool):return value
        if isinstance(value,(int,float)):return F(str(value))
        return value
    @lru_cache(None)
    def ref(self,sheet,value):
        if '!' in value:sheet,value=value.rsplit('!',1);sheet=sheet.strip("'").replace("''", "'")
        if ':' not in value:return self.cell(sheet,value)
        c0,r0,c1,r1=range_boundaries(value)
        return [self.cell(sheet,WB[sheet].cell(r,c).coordinate) for r in range(r0,r1+1) for c in range(c0,c1+1)]
    def run(self,node,sheet):
        typ=node[0]
        if typ=='operand':
            kind,val=node[1:]
            if kind=='NUMBER':return F(val)
            if kind=='TEXT':return val[1:-1].replace('""','"')
            if kind=='RANGE':return self.ref(sheet,val)
            if kind=='LOGICAL':return val=='TRUE'
            raise ValueError(kind)
        if typ=='unary':return self.run(node[2],sheet)*(-1 if node[1]=='-' else 1)
        if typ=='binary':
            op=node[1];a=self.run(node[2],sheet);b=self.run(node[3],sheet)
            if op=='+':return a+b
            if op=='-':return a-b
            if op=='*':return a*b
            if op=='/':return a/b
            if op=='&':return str(a)+str(b)
            if op=='=':return a==b
            if op=='<>':return a!=b
            if op=='>':return a>b
            if op=='<':return a<b
            if op=='>=':return a>=b
            if op=='<=':return a<=b
            raise ValueError(op)
        name,args=node[1:]
        if name=='IF':return self.run(args[1] if self.run(args[0],sheet) else args[2],sheet)
        values=[self.run(a,sheet) for a in args]
        if name=='AND':return all(values)
        if name=='OR':return any(values)
        if name=='ROUND':return halfup(*values)
        if name in ['SUMIFS','AVERAGEIFS','SUMIF']:
            if name=='SUMIF':values=[values[2],values[0],values[1]]
            numbers=values[0];conditions=list(zip(values[1::2],values[2::2]))
            def match(value,criterion):
                return fnmatchcase(value,criterion) if isinstance(value,str) and isinstance(criterion,str) and ('*' in criterion or '?' in criterion) else value==criterion
            kept=[v for i,v in enumerate(numbers) if all(match(r[i],crit) for r,crit in conditions)]
            return sum(kept,F(0))/len(kept) if name=='AVERAGEIFS' else sum(kept,F(0))
        raise ValueError(name)

ev=Evaluator();checked=0;failures=[]
for sheet in WB:
    for row in sheet:
        for cell in row:
            if cell.data_type!='f':continue
            checked+=1;actual=ev.cell(sheet.title,cell.coordinate);expected=CACHE[sheet.title][cell.coordinate].value
            if actual=='' and expected is None:continue
            passed=abs(float(actual)-float(expected))<1e-8 if isinstance(actual,(F,float,int)) and expected is not None else actual==expected
            if not passed:failures.append({'cell':sheet.title+'!'+cell.coordinate,'evaluated':str(actual),'cached':str(expected)})

changed=Evaluator({('Inputs','B16'):0,('Inputs','B20'):200000})
mutation=[]
for i in range(6,30):
    if WB['Hub Scenarios'].cell(i,2).value=='base':
        row=i;U=changed.cell('Hub Scenarios',f'E{row}');s=changed.cell('Hub Scenarios',f'H{row}');fixed=changed.cell('Hub Scenarios',f'I{row}')
        test=changed.cell('Hub Scenarios',f'M{row}')==halfup(U*s-fixed,2)
        mutation.append({'check':'zero uplift '+str(WB['Hub Scenarios'].cell(i,1).value),'passed':test})
mutation.append({'check':'recommended pair recalculates to -24367.70 with zero uplift','passed':changed.cell('Portfolios','H6')==F('-24367.70')})
mutation.append({'check':'recommended pair becomes infeasible under EUR200000 capex limit','passed':changed.cell('Portfolios','E6') is False})
result={'status':'PASS' if not failures and all(t['passed'] for t in mutation) else 'FAIL','formula_cells_evaluated':checked,'cache_mismatches':failures,'input_change_tests':mutation,
    'method':'Restricted independent formula interpreter reads actual XLSX formulas; exact-rational arithmetic. Supports only functions used in this workbook; not native Excel recalculation.'}
(ROOT/'verification/formula_verification.json').write_text(json.dumps(result,indent=2)+'\n')
print(json.dumps(result,indent=2))
if result['status']!='PASS':raise SystemExit(1)
