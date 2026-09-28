"""Independent-ish structural and reconciliation checks for generated outputs."""

from decimal import Decimal
import hashlib
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "evidence" / "raw"
OUT = ROOT / "analysis_outputs"
DEL = ROOT / "deliverables"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main():
    metrics = json.loads((DEL / "metrics.json").read_text())
    required = {"monthly", "countries", "fx_monthly", "market_context", "hub_scenarios", "recommendation", "quality"}
    assert required.issubset(metrics), required - set(metrics)
    assert len(metrics["monthly"]) == 72
    assert len(metrics["countries"]) == 6
    assert len(metrics["fx_monthly"]) == 24
    assert len(metrics["market_context"]) == 18
    assert metrics["recommendation"]["countries"] == ["CZE", "ESP"]
    monthly_keys = {"country", "month", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin"}
    country_keys = monthly_keys - {"month"}
    assert set(metrics["monthly"][0]) == monthly_keys
    assert set(metrics["countries"][0]) == country_keys
    assert len({(r["country"], r["month"]) for r in metrics["monthly"]}) == 72
    assert all(len(r["month"]) == 7 and r["month"][4] == "-" for r in metrics["monthly"])
    assert all(sum(1 for r in metrics["fx_monthly"] if r["currency"] == c) == 12 for c in ["PLN", "CZK"])
    assert all(sum(1 for r in metrics["market_context"] if r["country"] == c) == 3 for c in ["DEU", "FRA", "NLD", "POL", "CZE", "ESP"])

    # Recompute the archived FX means independently from the raw ECB CSV.
    raw_fx = pd.read_csv(RAW / "ecb-history.csv")
    raw_fx["date"] = pd.to_datetime(raw_fx["Date"])
    raw_fx = raw_fx[raw_fx.date.dt.year == 2025].copy()
    raw_fx["month"] = raw_fx.date.dt.strftime("%Y-%m")
    fx_contract = pd.DataFrame(metrics["fx_monthly"])
    for _, row in fx_contract.iterrows():
        expected = round(float(pd.to_numeric(raw_fx.loc[raw_fx.month == row.month, row.currency], errors="coerce").mean()), 6)
        assert abs(expected - float(row.local_per_eur)) < 0.000001, (row.currency, row.month, expected, row.local_per_eur)

    m = pd.DataFrame(metrics["monthly"])
    c = pd.DataFrame(metrics["countries"])
    for country in c.country:
        cm = m[m.country == country]
        cr = c[c.country == country].iloc[0]
        for field in ["shipped_orders", "shipped_units", "returned_units"]:
            assert int(cm[field].sum()) == int(cr[field]), (country, field)
        for field in ["gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur"]:
            assert abs(float(cm[field].sum()) - float(cr[field])) < 0.011, (country, field, cm[field].sum(), cr[field])
        assert abs((float(cr.contribution_eur) / float(cr.net_sales_eur)) - float(cr.margin)) < 0.000002

    assert metrics["quality"]["order_audit"]["exact_duplicate_order_rows_removed"] == 13
    assert metrics["quality"]["order_audit"]["correction_rows"] == 18
    assert metrics["quality"]["order_audit"]["eligibility_counts"]["eligible_2025_shipped"] == 1254
    assert metrics["quality"]["return_audit"]["exact_duplicate_return_rows_removed"] == 20
    assert metrics["quality"]["return_audit"]["inclusion_counts"]["quarantined_orphan"] == 1
    assert metrics["quality"]["return_audit"]["inclusion_counts"]["excluded_after_cutoff"] == 1
    assert all(0 <= r["margin"] <= 1 for r in metrics["countries"] if r["margin"] is not None)

    # Recompute scenario arithmetic from the saved country totals and client options.
    opts = pd.read_csv(RAW / "hub-options.csv").set_index("country")
    uplift = {"low": 0.10, "base": 0.25, "high": 0.40}
    individual = {(r["country"], r["scenario"]): r for r in metrics["hub_scenarios"] if "+" not in r["country"]}
    for _, cr in c.iterrows():
        op = opts.loc[cr.country]
        for scenario, u in uplift.items():
            expected = float(cr.contribution_eur) * u + float(cr.shipped_units) * (1 + u) * float(op.saving_eur_per_unit) - float(op.annual_fixed_eur)
            actual = individual[(cr.country, scenario)]["incremental_contribution_eur"]
            assert abs(round(expected, 2) - actual) < 0.011, (cr.country, scenario, expected, actual)
        fx_shock = 0.10 * float(cr.net_sales_eur) if cr.country in {"POL", "CZE"} else 0.0
        expected_stress = (float(cr.contribution_eur) - 0.03 * float(cr.gross_sales_eur) - fx_shock) * 1.25 - float(cr.contribution_eur) + float(cr.shipped_units) * 1.25 * float(op.saving_eur_per_unit) - float(op.annual_fixed_eur)
        actual_stress = individual[(cr.country, "stress")]["incremental_contribution_eur"]
        assert abs(round(expected_stress, 2) - actual_stress) < 0.011, (cr.country, "stress", expected_stress, actual_stress)

    pair_contract = [r for r in metrics["pair_scenarios"] if r["country"] in {"CZE+ESP", "POL+ESP", "POL+CZE", "NLD+ESP", "NLD+CZE", "NLD+POL", "FRA+ESP", "FRA+CZE", "FRA+POL", "DEU+CZE", "FRA+NLD"}]
    assert len(pair_contract) == 44
    for row in pair_contract:
        a, b = row["countries"].split(",")
        expected = individual[(a, row["scenario"])]["incremental_contribution_eur"] + individual[(b, row["scenario"])]["incremental_contribution_eur"]
        assert abs(round(expected, 2) - row["incremental_contribution_eur"]) < 0.011, row

    register = json.loads((DEL / "source_register.json").read_text())
    assert len(register) == 13
    for item in register:
        path = ROOT / item["file"]
        assert path.exists(), path
        assert sha256(path) == item["sha256"], item["file"]

    wb = load_workbook(DEL / "meridian_parts_analysis.xlsx", read_only=True, data_only=False)
    for sheet in ["Readme", "Source Register", "Quality", "Orders Audit", "Returns Audit", "Order Metrics", "Monthly", "Countries", "FX Monthly", "Market Context", "Hub Scenarios", "Pair Scenarios", "Option Rankings", "Charts"]:
        assert sheet in wb.sheetnames, sheet
    assert wb["Countries"].max_row == 7
    assert wb["Monthly"].max_row == 73

    report = PdfReader(str(DEL / "meridian_parts_executive_report.pdf"))
    assert len(report.pages) == 10
    report_text = "\n".join(p.extract_text() or "" for p in report.pages)
    for phrase in ["CZE+ESP", "€198,099", "Data quality and reconciliation", "Booked revenue", "90-day", "not causal"]:
        assert phrase in report_text, phrase

    deck = Presentation(str(DEL / "meridian_parts_board_presentation.pptx"))
    assert len(deck.slides) == 10
    deck_text = "\n".join(" ".join(sh.text for sh in slide.shapes if hasattr(sh, "text")) for slide in deck.slides)
    for phrase in ["CZE + ESP", "2.74%", "90-day", "Board ask"]:
        assert phrase in deck_text, phrase

    for chart in ["baseline_contribution.png", "option_comparison.png", "monthly_contribution.png", "fx_monthly.png"]:
        assert (OUT / "charts" / chart).exists()

    print("All metrics, source hashes, reconciliations, workbook, PDF, PPTX and chart checks passed.")


if __name__ == "__main__":
    main()
