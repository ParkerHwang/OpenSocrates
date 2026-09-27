"""Independent SQL/Fraction recalculation, workbook inspection and PDF rendering.

Does not import analyze.py. Monetary recheck uses rational arithmetic and integer
cents, independently of the Decimal implementation in the financial pipeline.
"""
from pathlib import Path
from collections import defaultdict
from fractions import Fraction as Q
from decimal import Decimal
from itertools import combinations
import csv
import hashlib
import json
import sqlite3
import sys
import datetime as dt
import re
import zipfile
import shutil

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.deps'))
import pymupdf
from PIL import Image,ImageDraw
import openpyxl

RAW=ROOT/'evidence/raw';OUT=ROOT/'deliverables'
V=OUT/'verification';V.mkdir(exist_ok=True)
CHECKS=[]
def check(name,condition,detail=''):
    CHECKS.append(dict(check=name,passed=bool(condition),detail=detail))
    if not condition:raise AssertionError(name+': '+str(detail))
def csvrows(name):return list(csv.DictReader((RAW/name).open()))
def cents(x):
    x=Q(x)*100
    if x>=0:return (2*x.numerator+x.denominator)//(2*x.denominator)
    return -cents(-Q(x)/100)
def jcent(x):return int(Decimal(str(x))*100)
def eq(a,b,tol=1e-9):return abs(float(a)-float(b))<=tol

def financial_checks(m):
    con=sqlite3.connect(':memory:');con.row_factory=sqlite3.Row
    order_rows=sum((csvrows(f) for f in ['orders-part1.csv','orders-part2.csv','order-corrections.csv']),[])
    return_rows=csvrows('returns.csv')
    for name,rows in [('orders',order_rows),('returns',return_rows)]:
        keys=list(rows[0]);con.execute('CREATE TABLE '+name+' ('+', '.join(k+' TEXT' for k in keys)+')')
        con.executemany('INSERT INTO '+name+' VALUES ('+','.join('?' for k in keys)+')',[tuple(r[k] for k in keys) for r in rows])
        idcol='order_id' if name=='orders' else 'return_id'
        con.execute(f'CREATE VIEW latest_{name} AS SELECT * FROM (SELECT *, ROW_NUMBER() OVER (PARTITION BY {idcol} ORDER BY CAST(revision AS INT) DESC) AS rn FROM (SELECT DISTINCT * FROM {name})) WHERE rn=1')
    retained=[dict(r) for r in con.execute('SELECT * FROM latest_orders')]
    eligible=[dict(r) for r in con.execute("SELECT * FROM latest_orders WHERE status='shipped' AND is_test='false' AND shipped_at>='2025-01-01' AND shipped_at<='2025-12-31'")]
    ret=[dict(r) for r in con.execute("SELECT r.* FROM latest_returns r JOIN latest_orders o ON r.order_id=o.order_id WHERE r.received_at<='2026-01-31' AND o.status='shipped' AND o.is_test='false' AND o.shipped_at>='2025-01-01' AND o.shipped_at<='2025-12-31'")]
    check('SQL order reconciliation',len(order_rows)==1328 and len(retained)==1297 and len(eligible)==1254)
    check('SQL return reconciliation',len(return_rows)==264 and len(ret)==241)
    orphan=[r['return_id'] for r in con.execute('SELECT r.return_id FROM latest_returns r LEFT JOIN latest_orders o ON o.order_id=r.order_id WHERE o.order_id IS NULL')]
    check('Orphan quarantine',orphan==['R-ORPHAN'])
    check('Inclusive cutoff',any(r['return_id']=='R-CUTOFF' for r in ret) and not any(r['return_id']=='R-LATE' for r in ret))
    check('January 2026 shipment excluded',all(r['order_id']!='FUTURE-2026' for r in eligible))
    check('Return revision replacement',next(r for r in ret if r['return_id']=='R-M1-01-003')['refund_local']=='149.00')
    rawfx=csvrows('ecb-history.csv');fxgroups=defaultdict(list)
    for r in rawfx:
        if r['Date'].startswith('2025-'):
            for c in ['PLN','CZK']:
                if r[c] not in ['','N/A']:fxgroups[c,r['Date'][:7]].append(Q(r[c]))
    fx={k:sum(v,Q(0))/len(v) for k,v in fxgroups.items()}
    for r in m['fx_monthly']:
        key=r['currency'],r['month']
        check('FX mean '+str(key),eq(fx[key],r['local_per_eur'],1e-12) and len(fxgroups[key])==r['observation_count'])
    rr=defaultdict(list)
    for r in ret:rr[r['order_id']].append(r)
    costs=csvrows('unit-costs.csv')
    expected={};groups=defaultdict(lambda:defaultdict(int))
    money_fields=['gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur']
    for r in eligible:
        oid=r['order_id'];month=r['shipped_at'][:7]
        rate=Q(1) if r['currency']=='EUR' else fx[r['currency'],month]
        costrow=sorted([v for v in costs if v['sku']==r['sku'] and v['valid_from']<=r['shipped_at']],key=lambda v:v['valid_from'])[-1]
        cost=Q(costrow['unit_cost_eur']);quantity=int(r['quantity'])
        credit=sum((Q(x['refund_local']) for x in rr[oid]),Q(0))
        returned=sum(int(x['quantity']) for x in rr[oid]);stock=sum(int(x['restocked_quantity']) for x in rr[oid])
        gross=cents((quantity*Q(r['unit_price_local'])-Q(r['discount_local']))/rate)
        refund=cents(credit/rate);nc=cents(quantity*cost)-cents(stock*cost);fulfill=cents(r['fulfillment_eur'])
        e=dict(shipped_orders=1,shipped_units=quantity,gross_sales_eur=gross,refunds_eur=refund,net_sales_eur=gross-refund,net_cogs_eur=nc,fulfillment_eur=fulfill,contribution_eur=gross-refund-nc-fulfill,returned_units=returned)
        expected[oid]=e
        for k,v in e.items():groups[r['country'],month][k]+=v
    for r in m['monthly']:
        g=groups[r['country'],r['month']]
        check('Independent monthly '+r['country']+' '+r['month'],all((jcent(r[k]) if k in money_fields else r[k])==v for k,v in g.items()))
        check('Monthly margin '+r['country']+' '+r['month'],r['margin'] is None if g['net_sales_eur']==0 else eq(r['margin'],g['contribution_eur']/g['net_sales_eur'],1e-12))
    for r in m['countries']:
        for k in list(expected.values())[0]:
            value=sum(g[k] for (c,mo),g in groups.items() if c==r['country'])
            check('Country roll-up '+r['country']+' '+k,(jcent(r[k]) if k in money_fields else r[k])==value)
    ledger=list(csv.DictReader((OUT/'audit/order_ledger.csv').open()))
    check('All ledger rows match independent raw reconstruction',all(all((jcent(row[k]) if k in money_fields else int(row[k]))==v for k,v in expected[row['order_id']].items()) for row in ledger) and len(ledger)==len(expected))
    free=[r for r in eligible if Q(r['unit_price_local'])==0]
    check('Free shipment preservation',len(free)==12 and sum(int(r['quantity']) for r in free)==739 and sum(expected[r['order_id']]['contribution_eur'] for r in free)==-1456590)
    check('Effective July cost',all(Q(r['unit_cost_eur'])==Q(next(v['unit_cost_eur'] for v in costs if v['sku']==r['sku'] and v['valid_from']==('2025-07-01' if r['shipped_at']>='2025-07-01' else '2025-01-01'))) for r in ledger))
    check('Half-up tie / signed rounding',cents('1.005')==101 and cents('-1.005')==-101 and cents('1.0049')==100)
    for file,field in [('population.json','population'),('gdp-per-capita.json','gdp_per_capita_usd')]:
        raw=json.loads((RAW/file).read_text())[1];obs={(r['countryiso3code'],int(r['date'])):r['value'] for r in raw}
        check('World Bank values '+field,len(obs)==18 and all((r[field] is None and obs[r['country'],r['year']] is None) or eq(r[field],obs[r['country'],r['year']],1e-8) for r in m['market_context']))
    cm={r['country']:r for r in m['countries']};opts={r['country']:r for r in csvrows('hub-options.csv')};sx={}
    for r in m['hub_scenarios']:
        c=r['country'];b=cm[c];o=opts[c];sc=r['scenario'];u=Q({'low':'.10','base':'.25','high':'.40','stress':'.25'}[sc])
        base=Q(str(b['contribution_eur']));saving=Q(o['saving_eur_per_unit']);fixed=Q(o['annual_fixed_eur'])
        normal=base*u+b['shipped_units']*(1+u)*saving-fixed
        if sc=='stress':normal-=(1+u)*(Q('.03')*Q(str(b['gross_sales_eur']))+(Q('.10')*Q(str(b['net_sales_eur'])) if c in ['POL','CZE'] else 0))
        v=cents(normal);sx[c,sc]=v
        check('Scenario '+c+' '+sc,jcent(r['incremental_contribution_eur'])==v)
        check('Payback '+c+' '+sc,r['payback_years'] is None if v<=0 else eq(r['payback_years'],Q(o['capex_eur'])*100/v))
    combos={():None,**{(c,):None for c in cm},**{tuple(c):None for c in combinations(cm,2)}}
    check('Complete portfolio coverage',len(m['portfolios'])==88 and set(tuple(r['countries']) for r in m['portfolios'])==set(combos))
    for r in m['portfolios']:
        cs=r['countries'];v=sum(sx[c,r['scenario']] for c in cs);k=sum(int(opts[c]['capex_eur']) for c in cs);fte=sum(int(opts[c]['fte']) for c in cs)
        check('Portfolio '+r['portfolio']+' '+r['scenario'],jcent(r['incremental_contribution_eur'])==v and r['capex_eur']==k and r['fte']==fte and r['feasible']==(k<=450000 and fte<=7))
        check('Portfolio payback '+r['portfolio']+' '+r['scenario'],r['payback_years'] is None if v<=0 else eq(r['payback_years'],k*100/v))
    check('Feasible pair count',sum(len(r['countries'])==2 and r['feasible'] and r['scenario']=='base' for r in m['portfolios'])==11)
    for sc in ['low','base','high']:
        best=max([r for r in m['portfolios'] if r['feasible'] and r['scenario']==sc],key=lambda r:r['incremental_contribution_eur'])
        check('Selected pair wins '+sc,best['countries']==m['recommendation']['countries'])
    worst={p:min(r['incremental_contribution_eur'] for r in m['portfolios'] if r['portfolio']==p) for p in {r['portfolio'] for r in m['portfolios'] if r['feasible']}}
    check('Maximin alternative',max(worst,key=worst.get)=='NLD+ESP')
    r=m['recommendation'];check('Recommendation constraints',r['capex_eur']==225000 and r['fte']==5 and len(r['countries'])==2)
    # Sensitivity results use rational arithmetic, independent of saved results.
    for r in m['sensitivity']:
        b=cm[r['country']];o=opts[r['country']];test=r['test'];u=Q(0) if test=='zero_uplift' else Q('.10') if test=='saving_capped_low' else Q('.25')
        s=Q(o['saving_eur_per_unit']) if test=='zero_uplift' else min(Q(o['saving_eur_per_unit']),Q(str(b['fulfillment_eur']))/b['shipped_units'])
        v=Q(str(b['contribution_eur']))*u+b['shipped_units']*(1+u)*s-Q(o['annual_fixed_eur'])
        if test=='saving_capped_stress':v-=(1+u)*(Q('.03')*Q(str(b['gross_sales_eur']))+(Q('.10')*Q(str(b['net_sales_eur'])) if r['country'] in ['POL','CZE'] else 0))
        check('Supplemental test '+r['country']+' '+test,cents(v)==jcent(r['incremental_contribution_eur']))
    return expected

def workbook_checks(m):
    path=OUT/'meridian_analysis.xlsx'
    f=openpyxl.load_workbook(path,data_only=False);v=openpyxl.load_workbook(path,data_only=True)
    formula_count=sum(c.data_type=='f' for ws in f for row in ws for c in row)
    cache_errors=[]
    for ws in f:
        for row in ws:
            for cell in row:
                if cell.data_type=='f':
                    cv=v[ws.title][cell.coordinate]
                    # XLSX stores an intentionally empty-string formula cache as
                    # an empty <v/>; openpyxl reads it as None. Payback, rank and
                    # issue formulas explicitly return "" in those cases.
                    if cv.data_type=='e' or (cv.value is None and '""' not in cell.value):cache_errors.append(f'{ws.title}!{cell.coordinate}')
    check('Formula caches present / no Excel errors',not cache_errors,cache_errors)
    check('Workbook formulas and charts',formula_count>10000 and len(f['Decision']._charts)==2,{'formulas':formula_count,'charts':len(f['Decision']._charts),'sheets':len(f.sheetnames)})
    check('Workbook no external formula links',not f._external_links)
    for name,records,keys in [('Monthly',m['monthly'],['country','month','shipped_orders','shipped_units','gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur','returned_units','margin']),('Countries',m['countries'],['country','shipped_orders','shipped_units','gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur','returned_units','margin'])]:
        for i,r in enumerate(records,2):
            check(f'Workbook {name} row {i}',all(v[name].cell(i,j+1).value==r[k] or (isinstance(r[k],(float,int)) and eq(v[name].cell(i,j+1).value,r[k])) for j,k in enumerate(keys)))
    for i,r in enumerate(m['hub_scenarios'],2):check(f'Workbook scenario cache {i}',eq(v['Scenarios'].cell(i,15).value,r['incremental_contribution_eur']) and (v['Scenarios'].cell(i,16).value in [None,''] if r['payback_years'] is None else eq(v['Scenarios'].cell(i,16).value,r['payback_years'])))
    for i,r in enumerate(m['portfolios'],2):check(f'Workbook portfolio cache {i}',eq(v['Portfolios'].cell(i,7).value,r['incremental_contribution_eur']) and v['Portfolios'].cell(i,9).value==r['feasible'])
    check('Excel summary formula wiring',f['Monthly']['E2'].value.startswith('=SUMIFS(Orders!$Q$') and f['Countries']['D2'].value.startswith('=SUMIF(Monthly!$A$2:$A$73,A2,Monthly!$E$') and f['Scenarios']['O2'].value=='=ROUND(L2+M2-I2-N2,2)' and f['Scenarios']['N5'].value.startswith('=IF(B5="stress",(1+C5)'))
    check('Workbook chart references',all(s.val.numRef.f for ch in f['Decision']._charts for s in ch.series))
    return dict(sheets=len(f.sheetnames),formula_cells=formula_count,native_charts=2,native_office_recalculation=False,office_engine_available=bool(shutil.which('libreoffice') or shutil.which('soffice')))

def render_checks():
    rendered=V/'rendered';rendered.mkdir(exist_ok=True)
    result={}
    for kind,wanted in [('report',12),('board',10)]:
        doc=pymupdf.open(OUT/f'meridian_{kind}.pdf');check(kind+' page count',len(doc)==wanted)
        bad=[];texts=[];thumbs=[];links=0
        for i,page in enumerate(doc):
            text=page.get_text();texts.append(text);links+=len(page.get_links())
            for block in page.get_text('dict')['blocks']:
                if block['type']!=0:continue
                for line in block['lines']:
                    for span in line['spans']:
                        rect=pymupdf.Rect(span['bbox'])
                        if rect.x0<0 or rect.y0<0 or rect.x1>page.rect.width+.5 or rect.y1>page.rect.height+.5:bad.append(dict(page=i+1,text=span['text'],bbox=list(rect)))
            pix=page.get_pixmap(matrix=pymupdf.Matrix(1.25,1.25),alpha=False);path=rendered/f'{kind}-{i+1:02d}.png';pix.save(path)
            im=Image.open(path).convert('RGB');im.thumbnail((590,835 if kind=='report' else 340));thumbs.append(im)
        check(kind+' no text outside page',not bad,bad)
        check(kind+' no empty pages',all(len(t)>300 for t in texts))
        full='\n'.join(texts)
        check(kind+' key numeric coherence',all(s in full for s in ['225,000','198,099','32,180']))
        per=4 if kind=='report' else 6
        for start in range(0,len(thumbs),per):
            subset=thumbs[start:start+per];ht=860 if kind=='report' else 370
            sheet=Image.new('RGB',(1200,ht*((len(subset)+1)//2)),'#dce4e9');draw=ImageDraw.Draw(sheet)
            for j,im in enumerate(subset):
                x=(j%2)*600+5;y=(j//2)*ht+20;sheet.paste(im,(x,y));draw.text((x+5,y-16),f'{kind} page {start+j+1}',fill='black')
            sheet.save(rendered/f'{kind}-contact-{start//per+1}.png')
        (V/f'{kind}-extracted.txt').write_text(full)
        result[kind]=dict(pages=len(doc),all_pages_rendered=True,text_out_of_bounds=len(bad),link_annotations=links)
    check('Report source link annotations',result['report']['link_annotations']>=30)
    return result

def main():
    m=json.loads((OUT/'metrics.json').read_text())
    for key in ['monthly','countries','fx_monthly','market_context','hub_scenarios']:
        check('JSON array '+key,isinstance(m[key],list))
    check('JSON row dimensions',[len(m[k]) for k in ['monthly','countries','fx_monthly','market_context','hub_scenarios']]==[72,6,24,18,24])
    manifests=json.loads((ROOT/'evidence/collection-manifest.json').read_text())
    for r in manifests:check('Collected hash '+Path(r['file']).name,hashlib.sha256((ROOT/r['file']).read_bytes()).hexdigest()==r['sha256'])
    official=json.loads((RAW/'source-register.json').read_text())['public']
    for r in official:check('Archive hash '+r['file'],hashlib.sha256((RAW/r['file']).read_bytes()).hexdigest()==r['sha256'])
    with zipfile.ZipFile(RAW/'ecb-history.zip') as z:check('ZIP/CSV byte equality',z.read([n for n in z.namelist() if n.endswith('.csv')][0])==(RAW/'ecb-history.csv').read_bytes())
    financial_checks(m);workbook=workbook_checks(m);pdfs=render_checks()
    files=['meridian_report.pdf','meridian_board.pdf','meridian_analysis.xlsx','metrics.json','decision_basis.json']
    result=dict(checked_at_utc=dt.datetime.now(dt.timezone.utc).isoformat(),status='passed',check_count=len(CHECKS),checks=CHECKS,workbook=workbook,pdfs=pdfs,artifact_sha256={f:hashlib.sha256((OUT/f).read_bytes()).hexdigest() for f in files},
        material_limits=['No native Excel/LibreOffice recalculation or native workbook rendering; formula caches and wiring inspected with openpyxl. Workbook browser preview is a separate cached-value representation.',
            'PDF pages rasterized locally; final human/agent visual inspection recorded separately in VERIFICATION.md.',
            'Financial reproducibility does not establish real hub demand, saving attainability, cash flow or implementation feasibility.'])
    (OUT/'verification.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps({k:v for k,v in result.items() if k!='checks'},indent=2))

if __name__=='__main__':main()
