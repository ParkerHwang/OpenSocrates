"""Cached-cell browser preview for visual QA; not a native Excel renderer."""
from pathlib import Path
import base64
import html
import json
import openpyxl

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'deliverables'
def main():
    wb=openpyxl.load_workbook(OUT/'meridian_analysis.xlsx',data_only=True)
    def image(name):return 'data:image/png;base64,'+base64.b64encode((OUT/'charts'/name).read_bytes()).decode()
    css='''body{margin:0;font:14px system-ui,sans-serif;color:#17324d;background:#edf3f6}header{padding:22px 28px;background:#17324d;color:white}h1{margin:0 0 8px;font-size:25px}header p{margin:0;font-size:12px;color:#cad8e2}nav{padding:14px 28px;display:flex;flex-wrap:wrap;gap:7px;background:white;border-bottom:1px solid #c8d7e1}button{border:1px solid #c8d7e1;background:white;color:#17324d;border-radius:4px;padding:7px 10px;cursor:pointer}button.active{background:#087f8c;color:white}main{padding:24px 28px}.panel{display:none}.panel.active{display:block}.cards{display:flex;gap:20px}.card{background:white;padding:20px;flex:1;border-top:4px solid #087f8c}.card strong{display:block;font-size:32px;margin-top:9px;color:#087f8c}.charts{display:grid;grid-template-columns:1fr 1.15fr;gap:20px;margin-top:25px}.charts img{width:100%;background:white}.scroll{max-height:680px;overflow:auto;background:white;border:1px solid #d6e1e9}table{border-collapse:collapse;font-size:12px;min-width:100%}th{position:sticky;top:0;background:#17324d;color:white;white-space:normal;min-width:90px;text-align:left;padding:10px;z-index:1}td{padding:9px 10px;border-bottom:1px solid #dce6ec;vertical-align:top;white-space:nowrap}tr:nth-child(even){background:#f4f8fa}.num{text-align:right}.foot{color:#536879;font-size:12px;margin:12px 0}'''
    names=['Decision','Countries','Monthly','Inputs','Scenarios','Portfolios','Market','FX monthly','Quality','Sources']
    chunks=[f'<!doctype html><html lang="en"><meta charset="utf-8"><title>Meridian workbook preview</title><style>{css}</style><header><h1>Meridian Parts | Analytical workbook preview</h1><p>Cached values read from the delivered XLSX. This browser view checks legibility; it does not emulate Excel calculation or native chart rendering.</p></header><nav>']
    for i,name in enumerate(names):chunks.append(f'<button data-panel="p{i}" class="{"active" if i==0 else ""}">{html.escape(name)}</button>')
    chunks.append('</nav><main>')
    for i,name in enumerate(names):
        chunks.append(f'<section id="p{i}" class="panel {"active" if i==0 else ""}"><h2>{html.escape(name)}</h2>')
        if name=='Decision':
            chunks.append('<p>Reserve Czechia + Spain; Spain first, Czechia subject to demand, saving and FX gates.</p><div class="cards">')
            for label,cell,suffix in [('Base annual increment','A7','EUR'),('Year-zero capex','E7','EUR'),('Hub staff','I7','FTE')]:chunks.append(f'<div class="card">{label}<strong>{wb[name][cell].value:,.0f} {suffix}</strong></div>')
            chunks.append(f'</div><div class="charts"><img alt="2025 contribution by country" src="{image("country_contribution.png")}"><img alt="Portfolio scenarios" src="{image("portfolio_scenarios.png")}"></div>')
            chunks.append('<p class="foot">Charts are independently rendered from metrics.json; the XLSX contains two native charts referencing its formula-linked country and portfolio results.</p>')
        else:
            ws=wb[name];chunks.append('<div class="scroll"><table><thead><tr>')
            def visible(cell):return not any(d.hidden and d.min<=cell.column<=d.max for d in ws.column_dimensions.values())
            for cell in ws[1]:
                if visible(cell):chunks.append('<th>'+html.escape(str(cell.value or ''))+'</th>')
            chunks.append('</tr></thead><tbody>')
            for row in ws.iter_rows(min_row=2):
                chunks.append('<tr>')
                for cell in row:
                    if not visible(cell):continue
                    value=cell.value
                    if value is None:text=''
                    elif isinstance(value,bool):text='YES' if value else 'NO'
                    elif isinstance(value,(int,float)):
                        f=cell.number_format
                        if '%' in f:text=f'{value*100:.1f}%'
                        elif '0.000000' in f:text=f'{value:.6f}'
                        elif '.00' in f or '.0' in f:text=f'{value:,.2f}'
                        else:text=f'{value:,}'
                    else:text=str(value)
                    cl='num' if isinstance(value,(int,float)) else ''
                    chunks.append(f'<td class="{cl}" title="{cell.coordinate}">'+html.escape(text)+'</td>')
                chunks.append('</tr>')
            chunks.append('</tbody></table></div><p class="foot">Full table is scrollable. Actual XLSX provides filters, frozen headers, number formats and formula caches.</p>')
        chunks.append('</section>')
    chunks.append('''</main><script>document.querySelectorAll('button').forEach(b=>b.onclick=()=>{document.querySelectorAll('button,.panel').forEach(x=>x.classList.remove('active'));b.classList.add('active');document.getElementById(b.dataset.panel).classList.add('active')})</script></html>''')
    target=OUT/'verification/workbook_preview.html';target.write_text(''.join(chunks));print(target)

if __name__=='__main__':main()
