#!/usr/bin/env python3
"""Independent, read-only checks for generated Meridian Parts outputs."""
from __future__ import annotations

import json
import math
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"


def close(a, b, tol=0.011):
    return abs(float(a) - float(b)) <= tol


def main():
    m = json.loads((OUT / "metrics.json").read_text(encoding="utf-8"))
    required = {"monthly", "countries", "fx_monthly", "market_context", "hub_scenarios", "recommendation", "quality"}
    assert required <= set(m), "metrics top-level schema incomplete"
    assert len(m["monthly"]) == 72, "expected six countries x twelve months"
    assert len(m["countries"]) == 6
    assert len(m["fx_monthly"]) == 36
    assert len(m["market_context"]) == 18
    assert len(m["hub_scenarios"]) == 68
    fields = ["country", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin"]
    assert all(set(fields) <= set(r) for r in m["countries"])
    assert all(set(fields) <= set(r) for r in m["monthly"])

    # Monthly-to-country reconciliation.
    for c in m["countries"]:
        rows = [r for r in m["monthly"] if r["country"] == c["country"]]
        for f in fields:
            if f in {"country", "margin"}:
                continue
            assert close(sum(r[f] for r in rows), c[f]), f"monthly reconciliation failed: {c['country']} {f}"
        denom = c["net_sales_eur"]
        expected = None if denom == 0 else round(c["contribution_eur"] / denom, 4)
        assert c["margin"] is None or close(c["margin"], expected, 0.00011)

    # Scenario constraints and payback convention.
    rec = m["recommendation"]
    assert rec["capex_eur"] <= 450000 and rec["fte"] <= 7 and len(rec["countries"]) <= 2
    for r in m["hub_scenarios"]:
        assert r["scenario"] in {"low", "base", "high", "stress"}
        if r["incremental_contribution_eur"] <= 0:
            assert r["payback_years"] is None
        else:
            assert r["payback_years"] is not None and r["payback_years"] > 0

    # Workbook structural checks.
    wb = load_workbook(OUT / "meridian_parts_analysis.xlsx", read_only=False, data_only=False)
    expected_sheets = {"Readme", "Monthly", "Countries", "FX_Monthly", "Market_Context", "Hub_Scenarios", "Hub_Options", "Clean_Orders", "Used_Returns", "Excluded_Orders", "Quarantined_Returns", "Quality", "Source_Register", "Charts"}
    assert expected_sheets <= set(wb.sheetnames)
    for name in ["Monthly", "Countries", "FX_Monthly", "Market_Context", "Hub_Scenarios", "Source_Register"]:
        assert len(wb[name].tables) == 1, f"missing Excel table on {name}"

    # PDF and PPTX integrity/content checks.
    pdf = PdfReader(str(OUT / "meridian_parts_board_report.pdf"))
    assert len(pdf.pages) >= 3
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    for needle in ["CZE + ESP", "EUR 198,099", "90-day", "World Bank", "ECB"]:
        assert needle in text, f"PDF missing expected content: {needle}"
    with zipfile.ZipFile(OUT / "meridian_parts_board_presentation.pptx") as z:
        bad = z.testzip()
        assert bad is None, bad
        slide_text = "".join(z.read(n).decode("utf-8", errors="ignore") for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        for needle in ["CZE + ESP", "90-day", "World Bank", "ECB"]:
            assert needle in slide_text, f"PPTX missing expected content: {needle}"

    print("PASS: metrics schema, monthly/country reconciliation, scenario constraints, workbook structure, PDF text, PPTX ZIP integrity/content")


if __name__ == "__main__":
    main()
