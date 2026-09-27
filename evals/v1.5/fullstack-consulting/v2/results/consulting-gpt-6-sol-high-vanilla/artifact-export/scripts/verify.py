#!/usr/bin/env python3
"""Independent delivery checks on saved inputs and generated artifacts."""
from __future__ import annotations

import csv
import hashlib
import json
from decimal import Decimal, ROUND_HALF_UP
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "evidence" / "raw"
OUT = ROOT / "deliverables"
DATA = json.loads((OUT / "metrics.json").read_text())


def D(v):
    return Decimal(str(v))


def cents(v):
    return v.quantize(Decimal(".01"), rounding=ROUND_HALF_UP)


def main():
    checks = []
    def check(name, condition):
        if not condition:
            raise AssertionError(name)
        checks.append(name)

    source = json.loads((RAW / "source-register.json").read_text())
    for r in source["public"]:
        p = RAW / r["file"]
        check(f"official snapshot hash {r['file']}", hashlib.sha256(p.read_bytes()).hexdigest() == r["sha256"])
    register = json.loads((ROOT / "evidence" / "source_register.json").read_text())
    check("all source register hashes", all(hashlib.sha256((ROOT/r["file"]).read_bytes()).hexdigest() == r["sha256"] for r in register))

    quality = DATA["quality"]
    check("order reconciliation", quality["raw_order_rows"] == quality["distinct_order_ids"] + quality["identical_order_duplicates"] + quality["superseded_order_revisions"])
    check("eligible order reconciliation", quality["eligible_2025_orders"] == quality["distinct_order_ids"] - sum(quality["order_exclusions"].values()))
    check("return reconciliation", quality["raw_return_rows"] == quality["distinct_return_ids"] + quality["identical_return_duplicates"] + quality["superseded_return_revisions"])
    check("eligible return reconciliation", quality["eligible_return_ids"] == quality["distinct_return_ids"] - sum(quality["return_exclusions"].values()))
    check("monthly record count", len(DATA["monthly"]) == 72)
    check("country record count", len(DATA["countries"]) == 6)
    check("FX record count", len(DATA["fx_monthly"]) == 24)
    check("market record count", len(DATA["market_context"]) == 18)
    check("only raw cancelled shipment dates missing", quality["missing_values"]["raw_order_shipped_at"] == 24 and all(v == 0 for k,v in quality["missing_values"].items() if k != "raw_order_shipped_at"))
    raw_order_rows=[]
    for name in ("orders-part1.csv","orders-part2.csv","order-corrections.csv"):
        with (RAW/name).open(newline="") as f:
            raw_order_rows.extend(csv.DictReader(f))
    check("blank raw shipment dates are cancelled", all(r["status"] == "cancelled" for r in raw_order_rows if not r["shipped_at"]))

    with (OUT/"order_audit.csv").open(newline="") as f:
        audit = {r["order_id"]:r for r in csv.DictReader(f)}
    check("audit order count", len(audit) == 1254)
    check("latest order revision replaces earlier row", audit["M1-02-006"]["revision"] == "2" and audit["M1-02-006"]["gross_sales_eur"] == "924.00")
    check("latest return revision replaces refund", audit["M1-01-003"]["refunds_eur"] == "149.00")
    check("inclusive return cutoff", audit["M1-12-002"]["return_ids"] == "R-CUTOFF")
    check("post-cutoff return excluded", audit["M1-12-004"]["returned_units"] == "0")
    check("valid zero-price orders retained", sum(r["zero_price"] == "True" for r in audit.values()) == 12)
    check("July cost revision applied", audit["M1-12-004"]["net_cogs_eur"] == "1843.00")

    keys = ["shipped_orders","shipped_units","gross_sales_eur","refunds_eur","net_sales_eur","net_cogs_eur","fulfillment_eur","contribution_eur","returned_units"]
    for country in DATA["countries"]:
        lines = [r for r in DATA["monthly"] if r["country"] == country["country"]]
        for key in keys:
            check(f"monthly to annual {country['country']} {key}", cents(sum((D(r[key]) for r in lines),D(0))) == cents(D(country[key])))
        check(f"accounting identity {country['country']}", cents(D(country["gross_sales_eur"])-D(country["refunds_eur"])-D(country["net_cogs_eur"])-D(country["fulfillment_eur"])) == cents(D(country["contribution_eur"])))

    with (RAW/"ecb-history.csv").open(newline="") as f:
        ecb=list(csv.DictReader(f))
    for fx in DATA["fx_monthly"]:
        values=[D(r[fx["currency"]]) for r in ecb if r["Date"].startswith(fx["month"]) and r[fx["currency"]] not in ("","N/A")]
        check(f"ECB mean {fx['currency']} {fx['month']}", abs(sum(values)/D(len(values))-D(fx["local_per_eur"])) < D("0.00000000001"))
        check(f"ECB observation count {fx['currency']} {fx['month']}", len(values) == fx["observation_days"])

    with (RAW/"hub-options.csv").open(newline="") as f:
        options={r["country"]:r for r in csv.DictReader(f)}
    expected={tuple(sorted(pair)) for pair in combinations(options,2) if sum(int(options[c]["capex_eur"]) for c in pair)<=450000 and sum(int(options[c]["fte"]) for c in pair)<=7}
    actual={tuple(sorted(r["country"].split("+"))) for r in DATA["hub_scenarios"] if "+" in r["country"]}
    check("all and only feasible pairs", actual == expected)
    check("scenario count", len(DATA["hub_scenarios"]) == (6+len(expected)+1)*4)
    country={r["country"]:r for r in DATA["countries"]}
    for c,base in country.items():
        C,U,G,N=(D(base[k]) for k in ("contribution_eur","shipped_units","gross_sales_eur","net_sales_eur"))
        o=options[c]; s=D(o["saving_eur_per_unit"]); F=D(o["annual_fixed_eur"])
        for name,u in (("low",D(".10")),("base",D(".25")),("high",D(".40"))):
            expected_value=cents(C*u+U*(1+u)*s-F)
            got=next(r for r in DATA["hub_scenarios"] if r["country"]==c and r["scenario"]==name)
            check(f"independent {c} {name} formula", cents(D(got["incremental_contribution_eur"]))==expected_value)
        shock=D(".10")*N if c in ("POL","CZE") else D(0)
        expected_stress=cents((C-D(".03")*G-shock)*D("1.25")-C+U*D("1.25")*s-F)
        got=next(r for r in DATA["hub_scenarios"] if r["country"]==c and r["scenario"]=="stress")
        check(f"independent {c} stress formula", cents(D(got["incremental_contribution_eur"]))==expected_stress)

    wb=load_workbook(OUT/"meridian_analysis.xlsx",read_only=False,data_only=False)
    required={"Read me","Annual","Monthly","FX","Market context","Hub inputs","Scenarios","Decision","Quality","Sources"}
    check("workbook tabs", required.issubset(set(wb.sheetnames)))
    check("workbook charts", len(wb["Annual"]._charts)>=1 and len(wb["Decision"]._charts)>=1)
    check("workbook monthly rows", wb["Monthly"].max_row>=77)
    check("workbook source rows", wb["Sources"].max_row>=19)
    wb.close()
    report=PdfReader(str(OUT/"executive_report.pdf")); deck=PdfReader(str(OUT/"board_presentation.pdf"))
    check("report renders to four nonempty pages", len(report.pages)==4 and all(len(p.extract_text())>500 for p in report.pages))
    check("deck renders to six nonempty slides", len(deck.pages)==6 and all(len(p.extract_text())>350 for p in deck.pages))
    check("report includes gate and limitation", "47,500" in "".join(p.extract_text() for p in report.pages) and "not forecasts" in "".join(p.extract_text() for p in report.pages))
    result={"status":"pass","checks":len(checks),"check_names":checks,
            "report_pages":len(report.pages),"deck_slides":len(deck.pages),
            "workbook_sheets":wb.sheetnames,
            "material_limitations":["Volume uplift and unit savings are client scenario assumptions; no causal service uplift estimate.","No observed delivery-time or capacity series, realized FX, full cash-flow or discounted NPV.","Source-room official archive is frozen with historical revisions through its stated vintage."]}
    (OUT/"verification.json").write_text(json.dumps(result,indent=2)+"\n")
    print(f"PASS: {len(checks)} checks; {len(report.pages)} report pages; {len(deck.pages)} deck slides; {len(wb.sheetnames)} workbook tabs")


if __name__=="__main__":
    main()
