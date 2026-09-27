"""Build the executive PDF and board PPTX from analysis/outputs and metrics.json."""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether)
from reportlab.graphics.shapes import Drawing, Rect, String, Line
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
AO = ROOT / "analysis" / "outputs"
metrics = json.loads((OUT / "metrics.json").read_text())
countries = pd.read_csv(AO / "countries.csv")
monthly = pd.read_csv(AO / "monthly.csv")
scenarios = pd.read_csv(AO / "hub_scenarios.csv")
pairs = pd.read_csv(AO / "pair_scenarios.csv")
market = pd.read_csv(AO / "market_context.csv")
audit = pd.read_csv(AO / "quality_audit.csv")

NAVY = colors.HexColor("#16324F")
TEAL = colors.HexColor("#0B6E69")
ORANGE = colors.HexColor("#E28F3D")
PALE = colors.HexColor("#EEF4F7")
INK = colors.HexColor("#20313F")
MUTED = colors.HexColor("#5D6D78")


def eur(v):
    sign = "-" if float(v) < 0 else ""
    return f"{sign}€{abs(float(v)):,.0f}"


def pct(v):
    return "—" if pd.isna(v) else f"{float(v):.1%}"


def table(data, widths=None, header=True, fontsize=7.5, alignments=None):
    t = Table(data, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    ts = [
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("FONTNAME", (0,0), (-1,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,-1), fontsize),
        ("TEXTCOLOR", (0,0), (-1,-1), INK),
        ("GRID", (0,0), (-1,-1), 0.25, colors.HexColor("#D5DEE3")),
        ("ROWBACKGROUNDS", (0,1 if header else 0), (-1,-1), [colors.white, PALE]),
        ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]
    if header:
        ts += [("BACKGROUND", (0,0), (-1,0), NAVY), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold")]
    if alignments:
        for col, align in alignments.items(): ts.append(("ALIGN", (col, 0), (col, -1), align))
    t.setStyle(TableStyle(ts))
    return t


def bar_drawing(items, width=480, height=150, color=TEAL, max_value=None):
    d = Drawing(width, height)
    if max_value is None: max_value = max(float(v) for _, v in items) or 1
    left, top, base, bar_h, gap = 96, height - 18, 20, 16, 7
    d.add(Line(left, base, width-10, base, strokeColor=colors.HexColor("#AAB8C0")))
    for i, (label, value) in enumerate(items):
        y = top - i*(bar_h+gap)
        w = max(0, float(value))/max_value*(width-left-25)
        d.add(String(0, y+3, label, fontName="Helvetica", fontSize=7, fillColor=INK))
        d.add(Rect(left, y, w, bar_h, fillColor=color, strokeColor=None))
        d.add(String(min(width-45, left+w+3), y+3, f"€{float(value):,.0f}", fontName="Helvetica-Bold", fontSize=7, fillColor=INK))
    return d


def footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D5DEE3")); canvas.line(18*mm, 14*mm, 192*mm, 14*mm)
    canvas.setFont("Helvetica", 7); canvas.setFillColor(MUTED)
    canvas.drawString(18*mm, 9*mm, "Meridian Parts | synthetic decision support | 27 Sep 2026")
    canvas.drawRightString(192*mm, 9*mm, f"{doc.page}")
    canvas.restoreState()


def build_report():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=28, leading=32, textColor=NAVY, spaceAfter=10))
    styles.add(ParagraphStyle(name="Sub", parent=styles["Normal"], fontSize=12, leading=17, textColor=MUTED, spaceAfter=16))
    styles.add(ParagraphStyle(name="H1x", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=NAVY, spaceBefore=10, spaceAfter=8))
    styles.add(ParagraphStyle(name="H2x", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=TEAL, spaceBefore=7, spaceAfter=5))
    styles.add(ParagraphStyle(name="Bodyx", parent=styles["BodyText"], fontSize=9.2, leading=13.2, textColor=INK, spaceAfter=6))
    styles.add(ParagraphStyle(name="Smallx", parent=styles["BodyText"], fontSize=7.4, leading=10.2, textColor=MUTED, spaceAfter=3))
    styles.add(ParagraphStyle(name="Callout", parent=styles["BodyText"], fontSize=11, leading=15, textColor=NAVY, backColor=PALE, borderColor=TEAL, borderWidth=1, borderPadding=8, spaceAfter=9))
    doc = BaseDocTemplate(str(OUT / "meridian_parts_board_report.pdf"), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=16*mm, bottomMargin=20*mm)
    doc.addPageTemplates([PageTemplate(id="main", frames=[Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")], onPage=footer)])
    S = []
    S += [Spacer(1, 28*mm), Paragraph("Meridian Parts", styles["CoverTitle"]), Paragraph("European service-hub expansion", styles["Sub"]), Spacer(1, 12*mm)]
    S += [Paragraph("Board decision brief | 2025 operating evidence, option economics and a gated 90-day plan", styles["H2x"]), Spacer(1, 4*mm), Paragraph("Prepared for the synthetic client. Source room collected 27 September 2026; analysis cutoff: 31 January 2026 inclusive.", styles["Smallx"]), PageBreak()]

    rec = metrics["recommendation"]
    S += [Paragraph("Executive answer", styles["H1x"]), Paragraph(f"<b>Fund a staged CZE + ESP service-hub path, subject to operating gates.</b> The option is the highest base-case annual incremental contribution among all feasible singles and pairs under the client’s no-synergy policy: <b>{eur(rec['base_annual_incremental_contribution_eur'])} per year</b> on <b>{eur(rec['capex_eur'])} year-zero capex</b> and <b>{rec['fte']} FTE</b>.", styles["Callout"])]
    exec_data = [["Measure", "Low", "Base", "High", "Defined stress"], ["Annual incremental contribution", eur(rec["low_annual_incremental_contribution_eur"]), eur(rec["base_annual_incremental_contribution_eur"]), eur(float(pairs[(pairs.alternative=="CZE+ESP")&(pairs.scenario=="high")].incremental_contribution_eur.iloc[0])), eur(rec["stress_annual_incremental_contribution_eur"])], ["Simple payback", "3.48 yr", "1.14 yr", "0.68 yr", "6.99 yr"]]
    S += [table(exec_data, widths=[54*mm, 28*mm, 28*mm, 28*mm, 33*mm], alignments={1:"RIGHT",2:"RIGHT",3:"RIGHT",4:"RIGHT"}), Spacer(1, 5), Paragraph("Decision rule and caveat: the recommendation maximizes modeled base annual contribution, not risk-adjusted NPV. The model is scenario arithmetic built from synthetic client inputs; it is not statistically proven, causal or certain. Stage the second site so the board can stop after measured evidence.", styles["Bodyx"])]
    S += [Paragraph("Why this option", styles["H2x"]), Paragraph("CZE combines the strongest modeled hub economics (EUR95k capex, 2 FTE, EUR2.10/unit saving) with a positive low-volume case; ESP adds the highest baseline contribution in the six-market set and remains positive under the defined joint stress. Together they use 5 FTE and EUR225k—within both limits—and beat the strongest base-case alternative POL+ESP by EUR11.8k annual contribution while using one fewer FTE and EUR15k less capex.", styles["Bodyx"])]
    S += [Paragraph("Board approval requested", styles["H2x"]), Paragraph("Approve a two-stage envelope of EUR225k capex and 5 FTE, release CZE first, and authorize ESP only when the gates in the implementation plan are met. If the board requires a one-site-only commitment, CZE is the better first site on modeled base contribution, low-volume contribution and capex/FTE efficiency.", styles["Bodyx"]), PageBreak()]

    S += [Paragraph("1. What the 2025 evidence says", styles["H1x"]), Paragraph("The accounting view is built from shipped orders in 2025, with returns received through 31 January 2026. It uses the source-room dictionary: one order is one row after highest revision selection; identical repeats are collapsed; cancelled/test/future orders are excluded; valid zero-price shipments are retained; distinct return IDs are additive; orphan returns are quarantined.", styles["Bodyx"])]
    total = countries[["shipped_orders","shipped_units","gross_sales_eur","refunds_eur","net_sales_eur","net_cogs_eur","fulfillment_eur","contribution_eur","returned_units"]].sum()
    kpi = [["Shipped orders", f"{int(total.shipped_orders):,}"], ["Shipped units", f"{int(total.shipped_units):,}"], ["Gross sales", eur(total.gross_sales_eur)], ["Refunds", eur(total.refunds_eur)], ["Net sales", eur(total.net_sales_eur)], ["Net COGS", eur(total.net_cogs_eur)], ["Fulfillment", eur(total.fulfillment_eur)], ["Contribution / margin", f"{eur(total.contribution_eur)} / {total.contribution_eur/total.net_sales_eur:.1%}"], ["Returned units", f"{int(total.returned_units):,} ({total.returned_units/total.shipped_units:.1%} of units)"]]
    S += [table([["2025 network measure", "Value"]] + kpi, widths=[70*mm, 55*mm], alignments={1:"RIGHT"}), Spacer(1, 7)]
    cdata = [["Country", "Orders", "Units", "Net sales", "Contribution", "Margin", "Returned units"]]
    for r in countries.itertuples(): cdata.append([r.country, f"{int(r.shipped_orders):,}", f"{int(r.shipped_units):,}", eur(r.net_sales_eur), eur(r.contribution_eur), pct(r.margin), f"{int(r.returned_units):,}"])
    S += [Paragraph("Country diagnosis", styles["H2x"]), table(cdata, widths=[22*mm, 20*mm, 22*mm, 29*mm, 31*mm, 19*mm, 28*mm], alignments={1:"RIGHT",2:"RIGHT",3:"RIGHT",4:"RIGHT",5:"RIGHT",6:"RIGHT"}), Spacer(1, 5)]
    S += [bar_drawing([(r.country, r.contribution_eur) for r in countries.sort_values("contribution_eur", ascending=False).itertuples()], color=TEAL), Paragraph("ESP and CZE lead modeled baseline contribution and margin. This is a starting point for hub economics—not proof that local fulfillment caused the difference. Returned-unit rate is highest in CZE (2.1%) and ESP (1.8%); the implementation KPI plan therefore puts return rate and restock recovery beside volume and savings.", styles["Bodyx"])]
    S += [Paragraph("Booked revenue, cash and contribution", styles["H2x"]), Paragraph("Gross sales are the booked shipment value before explicit refunds. Net sales subtract qualifying refunds known by the cutoff; this is not a cash-receipts measure because customer payment and refund settlement timing are not provided. Contribution subtracts net COGS and actual nonrefundable fulfillment from net sales; it excludes hub capex, hub fixed cost, taxes, financing and corporate overhead. The ECB quote is an analytical translation rate, not proof of the actual cash conversion rate.", styles["Bodyx"])]

    S += [Paragraph("2. Market context: useful context, not demand proof", styles["H1x"]), Paragraph("The archived World Bank responses provide comparable 2022–2024 population and GDP per capita (current US$). Population is a scale indicator and GDP per capita is an economic-context indicator; neither is a direct proxy for Meridian Parts replacement-assembly demand.", styles["Bodyx"])]
    pvt = market.pivot(index="country", columns="year", values="population"); gvt = market.pivot(index="country", columns="year", values="gdp_per_capita_usd")
    mdata = [["Country", "Population 2022", "Population 2024", "Change", "GDP pc 2022", "GDP pc 2024", "Change"]]
    for c in ["DEU","FRA","NLD","POL","CZE","ESP"]:
        mdata.append([c, f"{int(pvt.loc[c,2022]):,}", f"{int(pvt.loc[c,2024]):,}", f"{(pvt.loc[c,2024]/pvt.loc[c,2022]-1):.1%}", eur(gvt.loc[c,2022]).replace("€","US$"), eur(gvt.loc[c,2024]).replace("€","US$"), f"{(gvt.loc[c,2024]/gvt.loc[c,2022]-1):.1%}"])
    S += [table(mdata, widths=[19*mm, 31*mm, 31*mm, 18*mm, 28*mm, 28*mm, 18*mm], alignments={1:"RIGHT",2:"RIGHT",3:"RIGHT",4:"RIGHT",5:"RIGHT",6:"RIGHT"}), Spacer(1, 5), Paragraph("Population grew in CZE and ESP by 2.2% from 2022 to 2024, while POL declined 0.7%. Current-US-dollar GDP per capita rose in all six markets, with the largest percentage change in POL; exchange rates and nominal-price effects are embedded in that observation. Source: archived World Bank API responses in evidence/raw/population.json and gdp-per-capita.json, vintage last updated 13 July 2026.", styles["Smallx"])]
    S += [Paragraph("FX method", styles["H2x"]), Paragraph("PLN and CZK prices and refunds were divided by the original sale-month arithmetic mean of available ECB business-day reference quotes in local currency per EUR. The analysis does not invert the quote or use today’s rate. 2025 mean ranges were PLN 4.1722–4.2658/EUR and CZK 24.2341–25.1633/EUR across the twelve months; EUR transactions use 1.0000. Full monthly rates and observation counts are in the workbook and analysis/outputs/fx_monthly.csv.", styles["Bodyx"]), PageBreak()]

    S += [Paragraph("3. Hub economics and choice", styles["H1x"]), Paragraph("The policy defines annual incremental contribution as C×u + U×(1+u)×s − F, where C is 2025 contribution, U shipped units, u is the low/base/high uplift, s is per-unit saving and F is recurring fixed cost. Capex is kept separate and used only for simple undiscounted payback. Stress applies a 3% gross-sales refund shock plus a 10% net-sales depreciation shock for PLN/CZK markets, then uses 25% volume uplift. Pairs add country figures without synergy.", styles["Bodyx"])]
    sdata = [["Alternative", "Low", "Base", "High", "Stress", "Capex", "FTE", "Base payback"]]
    rank = pairs[pairs.scenario=="base"].sort_values("incremental_contribution_eur", ascending=False)
    for r in rank.head(5).itertuples():
        rr = pairs[pairs.alternative==r.alternative].set_index("scenario")
        sdata.append([r.alternative, eur(rr.loc["low","incremental_contribution_eur"]), eur(rr.loc["base","incremental_contribution_eur"]), eur(rr.loc["high","incremental_contribution_eur"]), eur(rr.loc["stress","incremental_contribution_eur"]), eur(r.capex_eur), int(r.fte), f"{r.payback_years:.2f} yr"])
    S += [Paragraph("Top feasible pairs by base case", styles["H2x"]), table(sdata, widths=[27*mm, 25*mm, 25*mm, 25*mm, 25*mm, 24*mm, 15*mm, 25*mm], alignments={1:"RIGHT",2:"RIGHT",3:"RIGHT",4:"RIGHT",5:"RIGHT",6:"RIGHT",7:"RIGHT"}), Spacer(1, 5)]
    S += [Paragraph("CZE+ESP is the recommended pair. POL+ESP is the strongest alternative on base contribution, but it is EUR11.8k lower, costs EUR15k more capex and requires one additional FTE; its defined stress remains positive but less resilient than CZE+ESP. POL+CZE is cheaper but its joint stress is negative because both markets take the PLN/CZK shock. Full singles and all feasible/infeasible pair combinations are in the workbook’s All options and Pair scenarios tabs.", styles["Bodyx"])]
    S += [Paragraph("What would change the decision", styles["H2x"]), Paragraph("The critical unknown is realized hub uplift and saving—not population or nominal GDP. If a 90-day pilot cannot demonstrate an annualized run-rate near the policy’s 25% volume uplift and contracted/observed saving, defer ESP and retain CZE only. Conversely, verified demand above base with stable returns strengthens the case for the second site. A real decision should add implementation ramp, discounting, taxes, capacity, service-level effects and working-capital cash timing before final long-term commitment.", styles["Bodyx"]), PageBreak()]

    S += [Paragraph("4. 90-day implementation and gates", styles["H1x"]), Paragraph("The plan treats the board’s approval as staged authorization, not an assumption that both sites are live on day one.", styles["Bodyx"])]
    plan = [["Window", "Owner", "Work / dependency", "Decision gate / measure"], ["Days 0–15", "COO + Finance", "Confirm CZE/ESP service promise, lane volumes, SKU mix, baseline SLA, return reason codes and capex release. Finance locks model inputs and cash controls.", "Gate 0: signed scope, data dictionary, cost-center and owner map."], ["Days 16–45", "Ops + 3PL / Procurement", "Run CZE pilot with selected lanes; contract local handling and returns; test inventory positioning and restock inspection. Dependency: approved site/3PL and ERP/WMS changes.", "Gate 1: ≥90% shipped orders within target SLA; saving run-rate ≥ policy saving; return-rate not > baseline +0.5 pp."], ["Days 46–75", "Commercial + CX", "Instrument customer promise, quote-to-ship time, service contacts, conversion and return reasons; calibrate ESP demand and staffing plan.", "Gate 2: CZE annualized contribution run-rate ≥ EUR31.5k low-case equivalent and no material control exceptions."], ["Days 76–90", "Board + COO", "Review pilot evidence, updated scenario with actuals, ramp/capacity/cash impacts; release or defer ESP; document stop-loss.", "Gate 3: release ESP only if volume ≥20% uplift, saving ≥90% of assumption, return rate ≤ baseline +0.5 pp, and capex/FTE remain in envelope."]]
    S += [table(plan, widths=[19*mm, 25*mm, 76*mm, 50*mm], fontsize=7.1), Spacer(1, 6)]
    S += [Paragraph("KPI plan", styles["H2x"]), Paragraph("Weekly by hub: shipped orders, shipped units, gross/net sales, contribution, on-time ship rate, order-to-ship hours, local fulfillment cost per order, saving EUR per unit, return rate by units, restocked recovery rate, open exceptions and FTE utilization. Monthly board dashboard: incremental contribution versus policy low/base, payback progress, capex committed versus budget, SLA, repeat contacts, and cash/refund timing. Each KPI must show numerator, denominator, owner and source system.", styles["Bodyx"]), Paragraph("Key risks and mitigations: demand uplift risk (stage ESP); FX and price translation risk (reprice and refresh scenario monthly); return inflation (reason codes and restock inspection); duplicate/late data (reconciliation controls); staffing/3PL ramp (cross-train and capacity buffer); and service fragmentation (single customer promise and escalation owner).", styles["Bodyx"]), PageBreak()]

    S += [Paragraph("5. Quality, limitations and source register", styles["H1x"]), Paragraph("The analysis did not hide known anomalies. They are visible in the workbook’s Quality audit tab, analysis/outputs/quality_audit.csv and evidence/source-register-collected.json.", styles["Bodyx"])]
    qdata = [["Check", "Value", "Handling"]]
    for r in audit.itertuples(): qdata.append([r.check, str(r.value), r.handling])
    S += [table(qdata, widths=[58*mm, 25*mm, 83*mm], fontsize=7.0), Spacer(1, 6)]
    S += [Paragraph("Material limitations", styles["H2x"]), Paragraph("Client exports are synthetic and generated for this exercise; there are no customer interviews or live customer-account claims. The scenario is not a DCF and has no ramp, tax, financing, synergy, capacity or statistical uncertainty model. World Bank population/GDP are contextual only; GDP is current US$, not PPP or constant-price real income. ECB reference rates are official archived observations used as analytical translation rates. One orphan return (ABSENT) was quarantined and one return received 1 February 2026 was excluded from the cutoff. Returns were assigned to original sale month for accounting, so the contribution view is not a cash-flow timeline.", styles["Bodyx"])]
    S += [Paragraph("Source register / evidence links", styles["H2x"]), Paragraph("[S1] Client source room index: http://127.0.0.1:49683/ (downloaded inputs are preserved under evidence/raw/). [S2] World Bank population archive: https://api.worldbank.org/v2/country/DEU;FRA;NLD;POL;CZE;ESP/indicator/SP.POP.TOTL?date=2022:2024&format=json&per_page=1000. [S3] World Bank GDP per capita archive: https://api.worldbank.org/v2/country/DEU;FRA;NLD;POL;CZE;ESP/indicator/NY.GDP.PCAP.CD?date=2022:2024&format=json&per_page=1000. [S4] ECB historical reference rates: https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip. Original URLs, retrieval vintage and SHA-256 hashes are recorded in evidence/source-register-collected.json; all figures in this report are generated from saved local inputs.", styles["Smallx"])]
    doc.build(S)


# ---- PowerPoint ----
PPT_NAVY = RGBColor(22,50,79); PPT_TEAL = RGBColor(11,110,105); PPT_ORANGE = RGBColor(226,143,61); PPT_INK = RGBColor(32,49,63); PPT_MUTED = RGBColor(93,109,120); PPT_PALE = RGBColor(238,244,247)


def add_text(slide, text, x, y, w, h, size=14, color=PPT_INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text; p.alignment = align
    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = color
    return box


def title(slide, kicker, heading, number):
    add_text(slide, kicker.upper(), 0.55, 0.28, 7.5, 0.25, 9, PPT_TEAL, True)
    add_text(slide, heading, 0.55, 0.58, 12.1, 0.55, 25, PPT_NAVY, True)
    add_text(slide, str(number), 12.45, 0.28, 0.35, 0.3, 10, PPT_MUTED, True, PP_ALIGN.RIGHT)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.22), Inches(12.2), Inches(0.03)).fill.solid(); slide.shapes[-1].fill.fore_color.rgb = PPT_TEAL; slide.shapes[-1].line.fill.background()


def add_footer(slide):
    add_text(slide, "Meridian Parts | synthetic decision support | 27 Sep 2026", 0.55, 7.18, 8.5, 0.2, 7.5, PPT_MUTED)


def add_kpi(slide, label, value, x, y, w=2.75, accent=PPT_TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.95)); shape.fill.solid(); shape.fill.fore_color.rgb = PPT_PALE; shape.line.color.rgb = accent
    add_text(slide, label, x+0.12, y+0.12, w-0.24, 0.22, 8.5, PPT_MUTED, True)
    add_text(slide, value, x+0.12, y+0.38, w-0.24, 0.35, 18, PPT_NAVY, True)


def add_table(slide, data, x, y, widths, row_h=0.34, font=9):
    total = sum(widths)
    for ri, row in enumerate(data):
        xx = x
        for ci, val in enumerate(row):
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(xx), Inches(y+ri*row_h), Inches(widths[ci]), Inches(row_h))
            shape.fill.solid(); shape.fill.fore_color.rgb = PPT_NAVY if ri == 0 else (PPT_PALE if ri % 2 == 0 else RGBColor(255,255,255)); shape.line.color.rgb = RGBColor(210,220,225)
            add_text(slide, str(val), xx+0.04, y+ri*row_h+0.06, widths[ci]-0.08, row_h-0.08, font-0.5 if ri==0 else font, RGBColor(255,255,255) if ri==0 else PPT_INK, ri==0)
            xx += widths[ci]


def add_bar(slide, label, val, maxval, x, y, color=PPT_TEAL, width=4.3):
    add_text(slide, label, x, y, 0.7, 0.22, 8.5, PPT_INK, True)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x+0.75), Inches(y+0.03), Inches(width), Inches(0.18)).fill.solid(); slide.shapes[-1].fill.fore_color.rgb = RGBColor(225,232,235); slide.shapes[-1].line.fill.background()
    bw = max(0, float(val))/maxval*width
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x+0.75), Inches(y+0.03), Inches(bw), Inches(0.18)).fill.solid(); slide.shapes[-1].fill.fore_color.rgb = color; slide.shapes[-1].line.fill.background()
    add_text(slide, eur(val), x+0.75+width+0.08, y-0.02, 1.0, 0.24, 8.5, PPT_INK, True)


def build_deck():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    rec = metrics["recommendation"]
    # 1
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = PPT_PALE
    add_text(s, "MERIDIAN PARTS", 0.65, 0.7, 4, 0.3, 11, PPT_TEAL, True); add_text(s, "European service-hub expansion", 0.65, 1.35, 11, 0.7, 30, PPT_NAVY, True); add_text(s, "Board decision brief | 2025 operating evidence and gated 90-day plan", 0.68, 2.15, 8.5, 0.3, 14, PPT_MUTED)
    add_kpi(s, "RECOMMENDED PATH", "CZE + ESP", 0.68, 3.35, 3.1, PPT_TEAL); add_kpi(s, "YEAR-ZERO CAPEX", "€225k", 4.05, 3.35, 2.4, PPT_ORANGE); add_kpi(s, "FTE", "5", 6.7, 3.35, 1.6, PPT_TEAL); add_kpi(s, "BASE ANNUAL INCREMENT", "€198.1k", 8.55, 3.35, 2.8, PPT_ORANGE)
    add_text(s, "Approve staged funding; release ESP only after measured CZE volume, saving and return-rate gates are met. Scenario arithmetic is not a causal forecast or DCF.", 0.7, 5.2, 10.8, 0.7, 17, PPT_NAVY, True); add_footer(s)
    # 2
    s = prs.slides.add_slide(blank); title(s, "Decision", "The board can fund growth inside both constraints", 2)
    add_text(s, "CZE+ESP is the highest modeled feasible base-case alternative. It stays positive under the defined low and stress cases, but stress payback is 6.99 years—why the second site is gated.", 0.65, 1.5, 11.7, 0.55, 15, PPT_INK)
    add_kpi(s, "LOW CASE", "€64.6k", 0.7, 2.35, 2.7, PPT_TEAL); add_kpi(s, "BASE CASE", "€198.1k", 3.65, 2.35, 2.7, PPT_TEAL); add_kpi(s, "DEFINED STRESS", "€32.2k", 6.6, 2.35, 2.7, PPT_ORANGE); add_kpi(s, "BASE PAYBACK", "1.14 yr", 9.55, 2.35, 2.7, PPT_ORANGE)
    add_text(s, "Decision logic", 0.7, 3.75, 2, 0.3, 12, PPT_TEAL, True)
    add_text(s, "1  Approve the EUR225k / 5 FTE envelope.\n2  Mobilize CZE as the first site.\n3  Release ESP at Gate 3 only if volume ≥20% uplift, saving ≥90% of assumption, returns ≤ baseline +0.5 pp, and the envelope holds.", 0.75, 4.15, 6.0, 1.5, 15, PPT_NAVY, True)
    add_text(s, "Strongest alternative: POL+ESP = €186.3k base, €240k capex, 6 FTE, 1.29-year payback. CZE+ESP adds €11.8k annual contribution with €15k less capex and one fewer FTE.", 7.2, 4.05, 5.3, 1.1, 14, PPT_INK, False); add_footer(s)
    # 3
    s = prs.slides.add_slide(blank); title(s, "Evidence", "2025 network contribution is €2.03m on €4.38m net sales", 3)
    add_kpi(s, "SHIPPED ORDERS", "1,254", 0.7, 1.55, 2.35); add_kpi(s, "SHIPPED UNITS", "77,436", 3.2, 1.55, 2.35); add_kpi(s, "NET SALES", "€4.38m", 5.7, 1.55, 2.35); add_kpi(s, "MARGIN", "46.3%", 8.2, 1.55, 2.35); add_kpi(s, "RETURNED UNITS", "1,081", 10.7, 1.55, 2.0, PPT_ORANGE)
    add_text(s, "Contribution by country", 0.75, 3.0, 3, 0.3, 12, PPT_TEAL, True)
    maxv = countries.contribution_eur.max()
    for i, r in enumerate(countries.sort_values("contribution_eur", ascending=False).itertuples()): add_bar(s, r.country, r.contribution_eur, maxv, 0.8, 3.45+i*0.43, PPT_TEAL if r.country in ["CZE","ESP"] else PPT_ORANGE)
    add_text(s, "Accounting discipline", 7.5, 3.0, 3, 0.3, 12, PPT_TEAL, True)
    add_text(s, "• Highest revision wins; repeated rows collapse\n• 24 cancelled + 18 test + 1 future shipped order excluded\n• 12 valid zero-price shipments retained\n• One orphan return quarantined; one late return excluded\n• Refunds and units attributed to original sale month", 7.55, 3.45, 5.0, 1.9, 14, PPT_INK)
    add_text(s, "Contribution = net sales − net COGS − nonrefundable fulfillment. It is not cash and excludes hub capex/fixed cost.", 7.55, 5.85, 4.9, 0.55, 12, PPT_NAVY, True); add_footer(s)
    # 4
    s = prs.slides.add_slide(blank); title(s, "Context", "Market scale supports screening; it does not prove demand", 4)
    add_text(s, "Archived World Bank context (2022–2024)", 0.7, 1.5, 4, 0.3, 12, PPT_TEAL, True)
    pdata = [["Market", "Pop. change", "GDP pc change"]]
    pvt = market.pivot(index="country", columns="year", values="population"); gvt = market.pivot(index="country", columns="year", values="gdp_per_capita_usd")
    for c in ["DEU","FRA","NLD","POL","CZE","ESP"]: pdata.append([c, f"{(pvt.loc[c,2024]/pvt.loc[c,2022]-1):+.1%}", f"{(gvt.loc[c,2024]/gvt.loc[c,2022]-1):+.1%}"])
    add_table(s, pdata, 0.75, 1.95, [1.3, 1.7, 1.9], 0.38, 11)
    add_text(s, "Implications for this decision", 5.3, 1.5, 4, 0.3, 12, PPT_TEAL, True)
    add_text(s, "• CZE and ESP show +2.2% population growth, providing scale context.\n• POL population declined 0.7% despite nominal GDP pc growth.\n• Current-US-dollar GDP embeds FX and price effects; it is not PPP or real income.\n• Selection is driven by shipped units, contribution, savings and fixed cost—not demographics alone.", 5.35, 1.95, 6.8, 2.0, 15, PPT_INK)
    add_text(s, "FX control: original sale-month ECB mean business-day quote, local currency per EUR; EUR=1. PLN 4.1722–4.2658 and CZK 24.2341–25.1633 across 2025.", 0.8, 5.4, 11.5, 0.5, 13, PPT_NAVY, True); add_footer(s)
    # 5
    s = prs.slides.add_slide(blank); title(s, "Options", "CZE+ESP leads the feasible pair set", 5)
    rank = pairs[pairs.scenario=="base"].sort_values("incremental_contribution_eur", ascending=False).head(5)
    pdata = [["Pair", "Low", "Base", "Stress", "Capex", "FTE"]]
    for r in rank.itertuples():
        rr = pairs[pairs.alternative==r.alternative].set_index("scenario"); pdata.append([r.alternative, eur(rr.loc["low","incremental_contribution_eur"]), eur(r.incremental_contribution_eur), eur(rr.loc["stress","incremental_contribution_eur"]), eur(r.capex_eur), str(int(r.fte))])
    add_table(s, pdata, 0.8, 1.65, [1.6, 1.7, 1.7, 1.7, 1.5, 0.8], 0.43, 11)
    add_text(s, "Why not simply choose the cheapest?", 0.85, 4.55, 3.5, 0.3, 12, PPT_TEAL, True)
    add_text(s, "POL+CZE saves €20k capex versus CZE+ESP but is negative in the defined joint stress because both markets carry the PLN/CZK depreciation shock.\n\nWhy not POL+ESP? It is the strongest alternative on base contribution, but it needs one more FTE and €15k more capex for €11.8k less annual base contribution.", 0.9, 4.95, 5.8, 1.25, 14, PPT_INK)
    add_text(s, "All six singles, all 15 pairs, feasibility constraints and all four scenarios are in the analytical workbook.", 7.3, 4.85, 4.9, 0.6, 14, PPT_NAVY, True); add_footer(s)
    # 6
    s = prs.slides.add_slide(blank); title(s, "Method", "Scenario economics are transparent—and deliberately limited", 6)
    add_text(s, "Annual incremental contribution", 0.8, 1.55, 4, 0.3, 12, PPT_TEAL, True); add_text(s, "C × volume uplift + U × (1 + uplift) × per-unit saving − annual fixed cost", 0.85, 2.0, 5.6, 0.4, 17, PPT_NAVY, True)
    add_text(s, "Stress overlay", 0.8, 2.85, 3, 0.3, 12, PPT_TEAL, True); add_text(s, "3% of gross sales refund shock; plus 10% of net sales depreciation shock in PLN/CZK markets; then 25% uplift.", 0.85, 3.25, 5.6, 0.7, 15, PPT_INK)
    add_text(s, "What the model does not claim", 7.0, 1.55, 4.5, 0.3, 12, PPT_TEAL, True); add_text(s, "• No causal demand estimate\n• No statistical confidence or probability of scenarios\n• No DCF, ramp, tax, financing, synergy or capacity model\n• No proof that population/GDP drives parts demand\n• ECB rates are analytical reference translations", 7.05, 2.0, 5.0, 2.1, 15, PPT_INK)
    add_text(s, "Decision implication: use the model to rank options and set gates, then refresh with observed pilot evidence before full rollout.", 0.85, 5.45, 11.3, 0.5, 16, PPT_NAVY, True); add_footer(s)
    # 7
    s = prs.slides.add_slide(blank); title(s, "Execution", "A 90-day CZE-first plan creates an exit ramp", 7)
    pdata = [["Window", "Owner", "Output / gate"], ["0–15", "COO + Finance", "Scope, data dictionary, SLA baseline, capex/FTE controls; Gate 0 sign-off."], ["16–45", "Ops + 3PL", "CZE pilot, local handling/returns, WMS test; ≥90% SLA and saving run-rate."], ["46–75", "Commercial + CX", "Instrument promise, conversion, contacts, return reasons; Gate 2 CZE contribution run-rate."], ["76–90", "Board + COO", "Update actuals and release ESP only if ≥20% uplift, ≥90% saving, return rate ≤ baseline +0.5 pp."]]
    add_table(s, pdata, 0.75, 1.55, [1.0, 2.1, 8.4], 0.68, 10.5)
    add_text(s, "Weekly KPIs", 0.8, 5.2, 2, 0.3, 12, PPT_TEAL, True); add_text(s, "Units / orders • SLA • order-to-ship hours • saving €/unit • return rate • restocked recovery • contribution run-rate • capex • FTE utilization", 0.85, 5.6, 11.4, 0.45, 14, PPT_INK, True); add_footer(s)
    # 8
    s = prs.slides.add_slide(blank); title(s, "Close", "Fund the path, preserve the option to stop", 8)
    add_text(s, "Board ask", 0.75, 1.55, 2, 0.3, 12, PPT_TEAL, True); add_text(s, "Approve staged CZE + ESP funding within EUR225k capex and 5 FTE; authorize CZE mobilization now; make ESP conditional on Gate 3 evidence.", 0.8, 2.0, 5.8, 1.0, 20, PPT_NAVY, True)
    add_text(s, "Evidence pack", 7.0, 1.55, 2, 0.3, 12, PPT_TEAL, True); add_text(s, "Report PDF\nAnalytical XLSX\nmetrics.json\nanalysis/analysis.py\nanalysis/make_deliverables.py\nevidence/raw + source register", 7.05, 2.0, 4.9, 2.0, 16, PPT_INK)
    add_text(s, "Reproducibility: from project root, run `python3 analysis/analysis.py` then `python3 analysis/make_deliverables.py`.", 0.8, 5.55, 11.2, 0.35, 13, PPT_MUTED, False); add_footer(s)
    prs.save(OUT / "meridian_parts_board_presentation.pptx")


if __name__ == "__main__":
    build_report(); build_deck(); print("built report and presentation")
