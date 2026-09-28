#!/usr/bin/env python3
"""Reconciliation and file-integrity checks for saved Meridian deliverables."""
import hashlib, json, math
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from pptx import Presentation
import pandas as pd

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'deliverables'; D=json.loads((OUT/'metrics.json').read_text())
assert set(['monthly','countries','fx_monthly','market_context','hub_scenarios','recommendation','quality']).issubset(D)
assert len(D['monthly'])==72 and len(D['countries'])==6 and len(D['fx_monthly'])==24
assert len(D['market_context'])==18 and len(D['hub_scenarios'])==24 and len(D['alternatives'])==22
fields=['shipped_orders','shipped_units','gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur','returned_units']
for c in D['countries']:
 rows=[r for r in D['monthly'] if r['country']==c['country']]
 assert len(rows)==12
 for f in fields: assert abs(sum(r[f] for r in rows)-c[f])<0.011,(c['country'],f,sum(r[f] for r in rows),c[f])
 den=c['net_sales_eur']; expected=c['contribution_eur']/den if den else None
 assert expected is None and c['margin'] is None or abs(expected-c['margin'])<1e-10
 assert abs(c['net_sales_eur']-(c['gross_sales_eur']-c['refunds_eur']))<0.011
 assert abs(c['contribution_eur']-(c['net_sales_eur']-c['net_cogs_eur']-c['fulfillment_eur']))<0.011
for a in D['alternatives']:
 assert a['capex_eur']<=450000 and a['fte']<=7 if a['feasible'] else True
assert all(a['feasible'] for a in D['alternatives'] if a['label']=='Defer')
rec=D['recommendation']; candidate=next(a for a in D['alternatives'] if a['countries']==rec['countries'])
assert candidate['capex_eur']==rec['capex_eur'] and candidate['fte']==rec['fte']
best=max((a for a in D['alternatives'] if a['feasible']),key=lambda a:(a['minimum_scenario_eur'],a['base']))
assert best['countries']==rec['countries']
assert {x['scenario'] for x in D['hub_scenarios']}=={'low','base','high','stress'}
base={c['country']:c for c in D['countries']}; opt=pd.read_csv(ROOT/'sources/raw/hub-options.csv').set_index('country')
for x in D['hub_scenarios']:
 c=x['country']; b=base[c]; o=opt.loc[c]; C=b['contribution_eur']; U=b['shipped_units']; G=b['gross_sales_eur']; N=b['net_sales_eur']; s=o.saving_eur_per_unit; F=o.annual_fixed_eur
 if x['scenario']=='low': u=.10; expected=C*u+U*(1+u)*s-F
 elif x['scenario']=='base': u=.25; expected=C*u+U*(1+u)*s-F
 elif x['scenario']=='high': u=.40; expected=C*u+U*(1+u)*s-F
 else:
  depreciation=.10*N if c in ['POL','CZE'] else 0
  expected=(C-.03*G-depreciation)*1.25-C+U*1.25*s-F
 assert abs(expected-x['incremental_contribution_eur'])<0.011,(x,expected)
 assert x['capex_eur']==o.capex_eur and x['fte']==o.fte
for x in D['hub_scenarios']:
 if x['incremental_contribution_eur']<=0: assert x['payback_years'] is None
 else: assert x['payback_years']>0
assert D['quality']['orders']['eligible_2025_shipped_orders']==sum(c['shipped_orders'] for c in D['countries'])
assert D['quality']['returns']['orphan_ids']==['R-ORPHAN']
assert D['quality']['sales_checks']['valid_zero_price_shipments']==12

daily=pd.read_csv(ROOT/'sources/raw/ecb-history.csv',parse_dates=['Date'])
daily=daily[(daily.Date>='2025-01-01')&(daily.Date<'2026-01-01')].copy(); daily['month']=daily.Date.dt.strftime('%Y-%m')
for r in D['fx_monthly']:
 g=daily[daily.month==r['month']][r['currency']].dropna()
 assert len(g)>0 and abs(g.mean()-r['local_per_eur'])<1e-10

# Every registered local source hash must match the retained raw file.
for src in D['source_files']:
 path=ROOT/'sources'/'raw'/src['file']
 assert path.exists() and path.stat().st_size==src['bytes']
 assert hashlib.sha256(path.read_bytes()).hexdigest()==src['sha256']

wb=load_workbook(OUT/'meridian_analysis.xlsx',data_only=False)
assert {'Read me','Monthly','Country totals','FX monthly','Market context','Hub scenarios','Alternatives','Quality & sources'}.issubset(wb.sheetnames)
assert wb['Monthly'].max_row==73 and wb['Country totals'].max_row==7
assert wb['FX monthly'].max_row==25 and wb['Market context'].max_row==19 and wb['Hub scenarios'].max_row==25
assert len(wb['Hub scenarios']._charts)==1
for i,r in enumerate(D['monthly'],start=2):
 ws=wb['Monthly']; assert ws.cell(i,1).value==r['country'] and ws.cell(i,2).value==r['month']
 assert abs(ws.cell(i,10).value-r['contribution_eur'])<0.011
doc=Document(OUT/'meridian_board_brief.docx'); assert len(doc.paragraphs)>20 and len(doc.tables)>=4
prs=Presentation(OUT/'meridian_board_deck.pptx'); assert len(prs.slides)==7
for name in ['meridian_board_brief.pdf','meridian_board_deck.pdf']:
 path=OUT/name; assert path.exists() and path.stat().st_size>10_000
print('PASS: JSON schema shape, 72 monthly-to-country reconciliations, margins, scenario feasibility/payback, recommendation rank, source hashes, workbook crosswalk/chart, DOCX/PPTX structure, and PDFs present.')
