from __future__ import annotations

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def main():
    metrics = json.loads((OUT / "metrics.json").read_text())
    required = {"monthly", "countries", "fx_monthly", "market_context", "hub_scenarios", "recommendation", "quality"}
    assert set(metrics) >= required
    assert len(metrics["monthly"]) == 72
    assert len(metrics["countries"]) == 6
    assert len(metrics["fx_monthly"]) == 24
    assert len(metrics["market_context"]) == 18
    assert len(metrics["hub_scenarios"]) > 0
    expected = {"country", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin"}
    assert expected <= set(metrics["countries"][0])
    assert expected | {"month"} <= set(metrics["monthly"][0])

    monthly = pd.DataFrame(metrics["monthly"])
    countries = pd.DataFrame(metrics["countries"])
    numeric = ["shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units"]
    check = monthly.groupby("country", as_index=False)[numeric].sum().merge(countries[["country"] + numeric], on="country", suffixes=("_monthly", "_country"))
    for col in numeric:
        assert np.allclose(check[f"{col}_monthly"], check[f"{col}_country"], atol=0.005), col
    for _, row in countries.iterrows():
        assert row["margin"] is None or np.isclose(row["margin"], row["contribution_eur"] / row["net_sales_eur"])

    fx = pd.DataFrame(metrics["fx_monthly"])
    assert set(fx["currency"]) == {"PLN", "CZK"}
    assert fx.groupby("currency")["month"].nunique().to_dict() == {"CZK": 12, "PLN": 12}
    assert (fx["local_per_eur"] > 0).all()

    scenarios = pd.DataFrame(metrics["hub_scenarios"])
    base = scenarios[scenarios.scenario == "base"].sort_values("incremental_contribution_eur", ascending=False)
    rec = metrics["recommendation"]
    selected = "+".join(rec["countries"])
    assert selected == base.iloc[0]["country"]
    selected_row = base[base.country == selected].iloc[0]
    assert int(selected_row.capex_eur) == rec["capex_eur"]
    assert int(selected_row.fte) == rec["fte"]
    assert metrics["quality"]["monthly_to_country_reconciliation"] is True
    assert metrics["quality"]["missing_fx_rates"] == 0
    assert metrics["quality"]["missing_unit_costs"] == 0

    register = json.loads((ROOT / "evidence/source_register.json").read_text())
    for row in register["sources"]:
        assert sha256(ROOT / row["local_path"]) == row["sha256"], row["file"]

    wb = load_workbook(OUT / "meridian_parts_analysis.xlsx", read_only=False)
    required_sheets = {"README", "Quality", "SourceRegister", "OrdersReconciled", "ReturnsReconciled", "FXMonthly", "Monthly", "Countries", "MarketContext", "HubScenarios", "DecisionSummary"}
    assert required_sheets <= set(wb.sheetnames)
    assert len(wb["Countries"]._charts) >= 1
    assert len(wb["DecisionSummary"]._charts) >= 1

    pdf = PdfReader(str(OUT / "meridian_parts_board_report.pdf"))
    text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    assert len(pdf.pages) >= 4
    assert "Recommendation" in text and "Czechia + Spain" in text and "€198,099" in text

    prs = Presentation(str(OUT / "meridian_parts_board_presentation.pptx"))
    assert len(prs.slides) == 8
    slide_text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
    assert "CZE+ESP" in slide_text and "198,099" in slide_text and "90-day" in slide_text

    print("PASS: metrics schema, monthly/country reconciliation, FX coverage, scenario ranking, source hashes, workbook sheets/charts, PDF text/pages and PPTX structure verified")


if __name__ == "__main__":
    main()
