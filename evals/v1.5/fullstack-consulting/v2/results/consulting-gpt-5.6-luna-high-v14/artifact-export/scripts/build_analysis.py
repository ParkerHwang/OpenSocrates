from __future__ import annotations

import hashlib
import itertools
import json
import math
import os
import shutil
import subprocess
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from xml.sax.saxutils import escape

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PInches, Pt as PPt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    HRFlowable,
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "evidence" / "source_room"
OUT = ROOT / "deliverables"
ANALYSIS = ROOT / "analysis"
ASSETS = OUT / "assets"
COUNTRIES = ["DEU", "FRA", "NLD", "POL", "CZE", "ESP"]
COUNTRY_NAMES = {
    "DEU": "Germany",
    "FRA": "France",
    "NLD": "Netherlands",
    "POL": "Poland",
    "CZE": "Czechia",
    "ESP": "Spain",
}
SCENARIOS = ["low", "base", "high", "stress"]
UPLIFTS = {"low": 0.10, "base": 0.25, "high": 0.40}
CAPEX_LIMIT = 450000
FTE_LIMIT = 7
CUTOFF = pd.Timestamp("2026-01-31")


def half_up(value: float | int | Decimal | None) -> float:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return 0.0
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def pct(value: float | None, digits: int = 1) -> str:
    return "—" if value is None or (isinstance(value, float) and math.isnan(value)) else f"{value * 100:.{digits}f}%"


def eur(value: float | None, digits: int = 0) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "—"
    return f"€{value:,.{digits}f}"


def num(value: float | None, digits: int = 0) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "—"
    return f"{value:,.{digits}f}"


def save_csv(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def source_register() -> list[dict]:
    provided = json.loads((SRC / "source-register.json").read_text())
    provided_map = {x.get("file"): x for x in provided.get("public", [])}
    period = {
        "orders-part1.csv": "2025-01-01 to 2026-01-05 shipment records",
        "orders-part2.csv": "2025-01-04 to 2026-01-05 shipment records",
        "order-corrections.csv": "2025-02-08 to 2025-11-08 corrections",
        "returns.csv": "2025-01-28 to 2026-02-01 returns",
        "ecb-history.csv": "1999-01-04 to 2026-09-25 observations",
        "ecb-history.zip": "official ECB history archive",
        "population.json": "2022-2024",
        "gdp-per-capita.json": "2022-2024",
        "hub-options.csv": "planning assumptions",
        "unit-costs.csv": "2025 effective-dated costs",
        "scenario-policy.md": "board scenario policy",
        "data-dictionary.md": "client accounting policy",
    }
    units = {
        "population.json": "persons",
        "gdp-per-capita.json": "current US$ per person",
        "ecb-history.csv": "local currency units per EUR",
        "hub-options.csv": "EUR, FTE, EUR per unit",
        "unit-costs.csv": "EUR per unit",
        "orders-part1.csv": "local currency sales; EUR fulfillment",
        "orders-part2.csv": "local currency sales; EUR fulfillment",
        "order-corrections.csv": "local currency sales; EUR fulfillment",
        "returns.csv": "local currency refunds; physical units",
    }
    rows = []
    retrieval = datetime.now(timezone.utc).isoformat()
    for path in sorted(SRC.iterdir()):
        if not path.is_file():
            continue
        name = path.name
        provided_row = provided_map.get(name, {})
        if name == "ecb-history.csv":
            provided_row = provided_map.get("ecb-history.zip", provided_row)
            original = provided_row.get("original_url", f"http://127.0.0.1:49795/{name}")
        else:
            original = provided_row.get("original_url", f"http://127.0.0.1:49795/{name}")
        rows.append(
            {
                "source_id": f"SR-{len(rows)+1:02d}",
                "file": name,
                "local_path": str(path.relative_to(ROOT)),
                "original_url": original,
                "retrieved_utc": provided_row.get("retrieved_utc", retrieval),
                "register_generated_utc": retrieval,
                "sha256": sha256_file(path),
                "bytes": path.stat().st_size,
                "period": period.get(name, "not applicable"),
                "units": units.get(name, "see file / not applicable"),
                "status": "public archived official snapshot" if name in {"population.json", "gdp-per-capita.json", "ecb-history.csv", "ecb-history.zip"} else "synthetic client input",
                "notes": "lossless ZIP extraction" if name == "ecb-history.csv" else "",
            }
        )
    return rows


def load_orders() -> tuple[pd.DataFrame, dict]:
    files = ["orders-part1.csv", "orders-part2.csv", "order-corrections.csv"]
    frames = []
    for name in files:
        frame = pd.read_csv(SRC / name)
        frame["source_file"] = name
        frames.append(frame)
    raw = pd.concat(frames, ignore_index=True)
    raw_rows = len(raw)
    exact_duplicate_rows = int(raw.duplicated().sum())
    dedup = raw.drop_duplicates().copy()
    same_revision_conflicts = int(
        dedup.groupby(["order_id", "revision"]).size().gt(1).sum()
    )
    dedup = dedup.sort_values(["order_id", "revision", "source_file"], ascending=[True, False, True])
    selected = dedup.drop_duplicates("order_id", keep="first").copy()
    selected["revision_selected"] = selected["revision"]
    selected["source_selected"] = selected["source_file"]
    selected["shipped_at"] = pd.to_datetime(selected["shipped_at"])
    selected["ordered_at"] = pd.to_datetime(selected["ordered_at"])
    selected["is_test"] = selected["is_test"].astype(bool)
    eligible = (
        (selected["status"] == "shipped")
        & (~selected["is_test"])
        & (selected["shipped_at"].dt.year == 2025)
    )
    selected["eligible_2025"] = eligible
    selected["exclusion_reason"] = np.where(
        eligible,
        "eligible",
        np.select(
            [selected["status"].ne("shipped"), selected["is_test"], selected["shipped_at"].dt.year.ne(2025)],
            ["status_not_shipped", "test_transaction", "shipment_not_in_2025"],
            default="other_exclusion",
        ),
    )
    stats = {
        "raw_rows": raw_rows,
        "exact_duplicate_rows_removed": exact_duplicate_rows,
        "unique_order_ids_after_dedup": int(len(selected)),
        "same_revision_conflicts": same_revision_conflicts,
        "revision_replacements": int((selected["revision"] > 1).sum()),
        "eligible_orders": int(eligible.sum()),
        "excluded_cancelled": int((selected["exclusion_reason"] == "status_not_shipped").sum()),
        "excluded_test": int((selected["exclusion_reason"] == "test_transaction").sum()),
        "excluded_not_2025": int((selected["exclusion_reason"] == "shipment_not_in_2025").sum()),
        "zero_price_orders": int(((selected["quantity"] * selected["unit_price_local"] - selected["discount_local"]) == 0).sum()),
    }
    return selected, stats


def load_returns(orders: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    raw = pd.read_csv(SRC / "returns.csv")
    raw_rows = len(raw)
    exact_duplicate_rows = int(raw.duplicated().sum())
    dedup = raw.drop_duplicates().copy()
    same_revision_conflicts = int(dedup.groupby(["return_id", "revision"]).size().gt(1).sum())
    dedup = dedup.sort_values(["return_id", "revision"], ascending=[True, False])
    selected = dedup.drop_duplicates("return_id", keep="first").copy()
    selected["received_at"] = pd.to_datetime(selected["received_at"])
    selected["linked_order"] = selected["order_id"].isin(set(orders["order_id"]))
    selected["linked_eligible"] = selected["order_id"].isin(set(orders.loc[orders["eligible_2025"], "order_id"]))
    selected["within_cutoff"] = selected["received_at"] <= CUTOFF
    selected["eligible_return"] = selected["linked_eligible"] & selected["within_cutoff"]
    selected["exclusion_reason"] = np.where(
        selected["eligible_return"],
        "eligible",
        np.select(
            [~selected["linked_order"], ~selected["linked_eligible"], ~selected["within_cutoff"]],
            ["orphan_return", "linked_order_not_eligible", "received_after_cutoff"],
            default="other_exclusion",
        ),
    )
    stats = {
        "raw_rows": raw_rows,
        "exact_duplicate_rows_removed": exact_duplicate_rows,
        "unique_return_ids_after_dedup": int(len(selected)),
        "same_revision_conflicts": same_revision_conflicts,
        "revision_replacements": int((selected["revision"] > 1).sum()),
        "eligible_returns": int(selected["eligible_return"].sum()),
        "eligible_return_units": int(selected.loc[selected["eligible_return"], "quantity"].sum()),
        "orphan_returns": int((selected["exclusion_reason"] == "orphan_return").sum()),
        "returns_on_ineligible_orders": int((selected["exclusion_reason"] == "linked_order_not_eligible").sum()),
        "returns_after_cutoff": int((selected["exclusion_reason"] == "received_after_cutoff").sum()),
    }
    return selected, stats


def load_fx() -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    ecb = pd.read_csv(SRC / "ecb-history.csv")
    ecb["Date"] = pd.to_datetime(ecb["Date"])
    ecb2025 = ecb.loc[ecb["Date"].dt.year == 2025].copy()
    for currency in ["PLN", "CZK"]:
        ecb2025[currency] = pd.to_numeric(ecb2025[currency], errors="coerce")
    ecb2025["month"] = ecb2025["Date"].dt.to_period("M").astype(str)
    fx = (
        ecb2025.groupby("month")[["PLN", "CZK"]]
        .mean()
        .reset_index()
        .melt(id_vars="month", var_name="currency", value_name="local_per_eur")
        .sort_values(["currency", "month"])
        .reset_index(drop=True)
    )
    fx["local_per_eur"] = fx["local_per_eur"].round(10)
    fx_stats = {
        "observations_2025": int(len(ecb2025)),
        "fx_month_rows": int(len(fx)),
        "pln_months": int(fx.loc[fx.currency == "PLN", "month"].nunique()),
        "czk_months": int(fx.loc[fx.currency == "CZK", "month"].nunique()),
        "missing_rate_cells": int(fx["local_per_eur"].isna().sum()),
    }
    return ecb2025, fx, fx_stats


def load_context() -> pd.DataFrame:
    rows = []
    for indicator_file, field in [("population.json", "population"), ("gdp-per-capita.json", "gdp_per_capita_usd")]:
        payload = json.loads((SRC / indicator_file).read_text())
        for row in payload[1]:
            rows.append(
                {
                    "country": row["countryiso3code"],
                    "year": int(row["date"]),
                    field: None if row["value"] is None else float(row["value"]),
                    "indicator": row["indicator"]["id"],
                    "source_file": indicator_file,
                    "source_lastupdated": payload[0].get("lastupdated"),
                }
            )
    context = pd.DataFrame(rows)
    context = context.pivot_table(
        index=["country", "year"], values=["population", "gdp_per_capita_usd"], aggfunc="first"
    ).reset_index()
    for col in ["population", "gdp_per_capita_usd"]:
        if col not in context:
            context[col] = np.nan
    return context.sort_values(["country", "year"]).reset_index(drop=True)


def calculate_orders(orders: pd.DataFrame, returns: pd.DataFrame, fx: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    eligible = orders.loc[orders["eligible_2025"]].copy()
    eligible["month"] = eligible["shipped_at"].dt.to_period("M").astype(str)
    fx_lookup = {(row.currency, row.month): row.local_per_eur for row in fx.itertuples()}

    unit_costs = pd.read_csv(SRC / "unit-costs.csv")
    unit_costs["valid_from"] = pd.to_datetime(unit_costs["valid_from"])

    def cost_for(row):
        choices = unit_costs.loc[(unit_costs.sku == row.sku) & (unit_costs.valid_from <= row.shipped_at)]
        if choices.empty:
            return np.nan
        return float(choices.sort_values("valid_from").iloc[-1].unit_cost_eur)

    eligible["unit_cost_eur"] = eligible.apply(cost_for, axis=1)
    eligible["fx_rate_local_per_eur"] = eligible.apply(
        lambda row: 1.0 if row.currency == "EUR" else fx_lookup.get((row.currency, row.month), np.nan), axis=1
    )
    return_eligible = returns.loc[returns["eligible_return"]].copy()
    return_eligible = return_eligible.merge(
        eligible[["order_id", "month", "currency"]], on="order_id", how="left", validate="many_to_one"
    )
    ret_agg = (
        return_eligible.groupby("order_id", as_index=False)
        .agg(
            refund_local_total=("refund_local", "sum"),
            returned_units=("quantity", "sum"),
            restocked_units=("restocked_quantity", "sum"),
            return_records=("return_id", "count"),
        )
    )
    eligible = eligible.merge(ret_agg, on="order_id", how="left")
    for col in ["refund_local_total", "returned_units", "restocked_units", "return_records"]:
        eligible[col] = eligible[col].fillna(0)
    eligible["gross_sales_local"] = eligible["quantity"] * eligible["unit_price_local"] - eligible["discount_local"]
    eligible["gross_sales_eur"] = eligible.apply(lambda r: half_up(r.gross_sales_local / r.fx_rate_local_per_eur), axis=1)
    eligible["refunds_eur"] = eligible.apply(lambda r: half_up(r.refund_local_total / r.fx_rate_local_per_eur), axis=1)
    eligible["gross_cogs_eur_raw"] = eligible["quantity"] * eligible["unit_cost_eur"]
    eligible["recovered_cogs_eur_raw"] = eligible["restocked_units"] * eligible["unit_cost_eur"]
    eligible["net_cogs_eur"] = eligible.apply(
        lambda r: half_up(r.gross_cogs_eur_raw - r.recovered_cogs_eur_raw), axis=1
    )
    eligible["fulfillment_eur"] = eligible["fulfillment_eur"].map(half_up)
    eligible["net_sales_eur"] = eligible.apply(lambda r: half_up(r.gross_sales_eur - r.refunds_eur), axis=1)
    eligible["contribution_eur"] = eligible.apply(
        lambda r: half_up(r.net_sales_eur - r.net_cogs_eur - r.fulfillment_eur), axis=1
    )
    eligible["returned_units"] = eligible["returned_units"].astype(int)
    eligible["shipped_units"] = eligible["quantity"].astype(int)
    eligible["shipped_orders"] = 1
    eligible["return_rate_units"] = np.where(eligible.shipped_units != 0, eligible.returned_units / eligible.shipped_units, np.nan)
    eligible["month"] = eligible["month"].astype(str)
    checks = {
        "missing_fx_rates": int(eligible["fx_rate_local_per_eur"].isna().sum()),
        "missing_unit_costs": int(eligible["unit_cost_eur"].isna().sum()),
        "returns_exceed_shipped_units": int((eligible["returned_units"] > eligible["shipped_units"]).sum()),
        "negative_gross_sales": int((eligible["gross_sales_local"] < 0).sum()),
        "negative_refunds": int((eligible["refunds_eur"] < 0).sum()),
    }
    return eligible, checks


def aggregate_monthly(orders: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    metrics = [
        "shipped_orders",
        "shipped_units",
        "gross_sales_eur",
        "refunds_eur",
        "net_sales_eur",
        "net_cogs_eur",
        "fulfillment_eur",
        "contribution_eur",
        "returned_units",
    ]
    monthly = (
        orders.groupby(["country", "month"], as_index=False)[metrics]
        .sum()
        .sort_values(["country", "month"])
    )
    all_idx = pd.MultiIndex.from_product([COUNTRIES, sorted(monthly.month.unique())], names=["country", "month"])
    monthly = monthly.set_index(["country", "month"]).reindex(all_idx, fill_value=0).reset_index()
    for col in metrics:
        monthly[col] = monthly[col].fillna(0)
    monthly["margin"] = np.where(monthly.net_sales_eur != 0, monthly.contribution_eur / monthly.net_sales_eur, np.nan)
    countries = monthly.groupby("country", as_index=False)[metrics].sum().sort_values("country")
    countries["margin"] = np.where(countries.net_sales_eur != 0, countries.contribution_eur / countries.net_sales_eur, np.nan)
    countries["country_name"] = countries.country.map(COUNTRY_NAMES)
    monthly["country_name"] = monthly.country.map(COUNTRY_NAMES)
    monthly["month"] = monthly["month"].astype(str)
    monthly_check = (
        monthly.groupby("country")[metrics].sum().reset_index().merge(countries, on="country", suffixes=("_monthly", "_country"))
    )
    recon_ok = True
    for col in metrics:
        recon_ok = recon_ok and np.allclose(monthly_check[f"{col}_monthly"], monthly_check[f"{col}_country"], atol=0.005)
    return monthly, countries, {"monthly_to_country_reconciliation": bool(recon_ok), "monthly_rows": int(len(monthly))}


def build_scenarios(countries: pd.DataFrame, options: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    metrics = countries.set_index("country").to_dict("index")
    opt = options.set_index("country").to_dict("index")
    feasible_pairs = []
    all_pairs = []
    for a, b in itertools.combinations(COUNTRIES, 2):
        label = f"{a}+{b}"
        capex = opt[a]["capex_eur"] + opt[b]["capex_eur"]
        fte = opt[a]["fte"] + opt[b]["fte"]
        row = {"country": label, "members": [a, b], "capex_eur": capex, "fte": fte, "feasible": capex <= CAPEX_LIMIT and fte <= FTE_LIMIT}
        all_pairs.append(row)
        if row["feasible"]:
            feasible_pairs.append(row)
    sets = [{"country": c, "members": [c], "capex_eur": opt[c]["capex_eur"], "fte": opt[c]["fte"], "feasible": True} for c in COUNTRIES] + feasible_pairs
    rows = []
    for choice in sets:
        members = choice["members"]
        C = sum(metrics[m]["contribution_eur"] for m in members)
        U = sum(metrics[m]["shipped_units"] for m in members)
        G = sum(metrics[m]["gross_sales_eur"] for m in members)
        N = sum(metrics[m]["net_sales_eur"] for m in members)
        F = sum(opt[m]["annual_fixed_eur"] for m in members)
        sU = sum(metrics[m]["shipped_units"] * opt[m]["saving_eur_per_unit"] for m in members)
        for scenario in SCENARIOS:
            if scenario in UPLIFTS:
                u = UPLIFTS[scenario]
                annual = C * u + sum(metrics[m]["shipped_units"] * (1 + u) * opt[m]["saving_eur_per_unit"] for m in members) - F
            else:
                fx_shock = sum(0.10 * metrics[m]["net_sales_eur"] for m in members if m in {"POL", "CZE"})
                C_stress = C - 0.03 * G - fx_shock
                annual = C_stress * 1.25 - C + sum(metrics[m]["shipped_units"] * 1.25 * opt[m]["saving_eur_per_unit"] for m in members) - F
            annual = half_up(annual)
            rows.append(
                {
                    "country": choice["country"],
                    "scenario": scenario,
                    "incremental_contribution_eur": annual,
                    "capex_eur": int(choice["capex_eur"]),
                    "fte": int(choice["fte"]),
                    "payback_years": None if annual <= 0 else round(choice["capex_eur"] / annual, 4),
                    "annual_fixed_eur": F,
                    "feasible": bool(choice["feasible"]),
                    "members": ",".join(members),
                    "volume_uplift": None if scenario == "stress" else UPLIFTS[scenario],
                }
            )
    scenarios = pd.DataFrame(rows).sort_values(["scenario", "incremental_contribution_eur"], ascending=[True, False]).reset_index(drop=True)
    base = scenarios.loc[(scenarios.scenario == "base") & scenarios.feasible].sort_values("incremental_contribution_eur", ascending=False)
    selected = base.iloc[0].country if len(base) else "defer"
    selected_members = selected.split("+") if selected != "defer" else []
    selected_base = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "base")].iloc[0] if selected != "defer" else None
    selected_stress = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "stress")].iloc[0] if selected != "defer" else None
    strongest_alt = base.iloc[1].country if len(base) > 1 else "defer"
    alt_base = base.iloc[1] if len(base) > 1 else None
    # Break-even volume uplift for the selected choice under policy arithmetic.
    if selected_members:
        selected_C = sum(metrics[m]["contribution_eur"] for m in selected_members)
        selected_F = sum(opt[m]["annual_fixed_eur"] for m in selected_members)
        selected_sU = sum(metrics[m]["shipped_units"] * opt[m]["saving_eur_per_unit"] for m in selected_members)
        denominator = selected_C + selected_sU
        break_even_uplift = (selected_F - selected_sU) / denominator if denominator else None
    else:
        break_even_uplift = None
    summary = {
        "all_pairs": all_pairs,
        "feasible_pairs": [x["country"] for x in feasible_pairs],
        "selected_base": selected,
        "selected_members": selected_members,
        "strongest_alternative_base": strongest_alt,
        "selected_base_annual_contribution_eur": None if selected_base is None else float(selected_base.incremental_contribution_eur),
        "selected_base_payback_years": None if selected_base is None else float(selected_base.payback_years),
        "selected_stress_annual_contribution_eur": None if selected_stress is None else float(selected_stress.incremental_contribution_eur),
        "strongest_alternative_base_annual_contribution_eur": None if alt_base is None else float(alt_base.incremental_contribution_eur),
        "break_even_volume_uplift": None if break_even_uplift is None else float(break_even_uplift),
    }
    return scenarios, summary


def json_rows(monthly: pd.DataFrame, countries: pd.DataFrame, fx: pd.DataFrame, context: pd.DataFrame, scenarios: pd.DataFrame, summary: dict, quality: dict) -> dict:
    base_cols = [
        "country", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur",
        "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin",
    ]
    def clean(v):
        if pd.isna(v):
            return None
        if isinstance(v, (np.integer, int)):
            return int(v)
        if isinstance(v, (np.floating, float)):
            return round(float(v), 6)
        return v
    monthly_out = []
    for row in monthly.to_dict("records"):
        monthly_out.append({"country": row["country"], "month": row["month"], **{c: clean(row[c]) for c in base_cols[1:]}})
    country_out = [{c: clean(row[c]) for c in base_cols} for row in countries.to_dict("records")]
    context_out = [{"country": row.country, "year": int(row.year), "population": clean(row.population), "gdp_per_capita_usd": clean(row.gdp_per_capita_usd)} for row in context.itertuples()]
    fx_out = [{"currency": row.currency, "month": row.month, "local_per_eur": clean(row.local_per_eur)} for row in fx.itertuples()]
    scenario_cols = ["country", "scenario", "incremental_contribution_eur", "capex_eur", "fte", "payback_years"]
    scenario_out = [{c: clean(row[c]) for c in scenario_cols} for row in scenarios.to_dict("records")]
    recommendation = {
        "countries": summary["selected_members"],
        "capex_eur": int(scenarios.loc[(scenarios.country == summary["selected_base"]) & (scenarios.scenario == "base"), "capex_eur"].iloc[0]) if summary["selected_members"] else 0,
        "fte": int(scenarios.loc[(scenarios.country == summary["selected_base"]) & (scenarios.scenario == "base"), "fte"].iloc[0]) if summary["selected_members"] else 0,
        "rationale": "Fund the highest base-case feasible option under the board policy arithmetic, subject to a 90-day gated launch and downside KPI checks; defer if the base uplift cannot be evidenced before capex commitment.",
        "strongest_alternative": summary["strongest_alternative_base"],
        "break_even_volume_uplift": summary["break_even_volume_uplift"],
        "judgment_status": "supported estimate, not statistically proven or causal",
    }
    return {
        "monthly": monthly_out,
        "countries": country_out,
        "fx_monthly": fx_out,
        "market_context": context_out,
        "hub_scenarios": scenario_out,
        "recommendation": recommendation,
        "quality": quality,
    }


def make_assets(monthly: pd.DataFrame, countries: pd.DataFrame, scenarios: pd.DataFrame, summary: dict) -> dict[str, Path]:
    ASSETS.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "axes.titlesize": 12, "axes.labelsize": 10})
    paths = {}
    fig, ax = plt.subplots(figsize=(8, 4.5))
    cs = countries.sort_values("contribution_eur", ascending=False)
    colors_list = ["#0b6e69" if c in summary["selected_members"] else "#9aa7b2" for c in cs.country]
    ax.bar([COUNTRY_NAMES[c] for c in cs.country], cs.contribution_eur / 1000, color=colors_list)
    ax.set_ylabel("Contribution (€000)")
    ax.set_title("2025 contribution by country")
    ax.tick_params(axis="x", rotation=25)
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    paths["country_contribution"] = ASSETS / "country_contribution.png"
    fig.savefig(paths["country_contribution"], dpi=180)
    plt.close(fig)

    feasible = scenarios.loc[scenarios.feasible].copy()
    pivot = feasible.pivot(index="country", columns="scenario", values="incremental_contribution_eur")
    order = pivot["base"].sort_values(ascending=False).index
    fig, ax = plt.subplots(figsize=(9, 5))
    x = np.arange(len(order))
    width = 0.2
    for i, scenario in enumerate(["low", "base", "high", "stress"]):
        ax.bar(x + (i - 1.5) * width, pivot.loc[order, scenario] / 1000, width, label=scenario.title())
    ax.axhline(0, color="#333333", linewidth=0.7)
    ax.set_xticks(x, order)
    ax.set_ylabel("Annual incremental contribution (€000)")
    ax.set_title("Feasible options: scenario contribution")
    ax.legend(ncol=4, frameon=False)
    ax.tick_params(axis="x", rotation=35)
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    paths["scenario_contribution"] = ASSETS / "scenario_contribution.png"
    fig.savefig(paths["scenario_contribution"], dpi=180)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8, 4.5))
    cm = monthly.groupby("month")["contribution_eur"].sum()
    ax.plot(cm.index, cm.values / 1000, marker="o", color="#0b6e69")
    ax.fill_between(np.arange(len(cm)), cm.values / 1000, alpha=0.12, color="#0b6e69")
    ax.set_ylabel("Contribution (€000)")
    ax.set_title("2025 monthly contribution")
    ax.tick_params(axis="x", rotation=45)
    ax.grid(alpha=0.2)
    fig.tight_layout()
    paths["monthly_contribution"] = ASSETS / "monthly_contribution.png"
    fig.savefig(paths["monthly_contribution"], dpi=180)
    plt.close(fig)
    return paths


def style_excel_sheet(ws, freeze="A2", widths=None):
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="0B6E69")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D5DDE3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    if widths:
        for col, width in widths.items():
            ws.column_dimensions[col].width = width


def write_df_sheet(wb: Workbook, name: str, df: pd.DataFrame, widths=None, number_formats=None):
    ws = wb.create_sheet(name)
    clean = df.copy()
    clean = clean.where(pd.notna(clean), None)
    ws.append(list(clean.columns))
    for row in clean.itertuples(index=False, name=None):
        ws.append(list(row))
    style_excel_sheet(ws, widths=widths)
    if number_formats:
        for col_name, fmt in number_formats.items():
            if col_name in clean.columns:
                idx = list(clean.columns).index(col_name) + 1
                for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
                    for c in cell:
                        c.number_format = fmt
    return ws


def make_workbook(orders, returns, fx, context, monthly, countries, scenarios, source_rows, quality, summary, asset_paths):
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    ws = wb.create_sheet("README")
    readme = [
        ["Meridian Parts — analytical workbook", ""],
        ["Purpose", "2025 shipped-sales and returns reconciliation plus service-hub scenario analysis."],
        ["Scope", "Countries DEU, FRA, NLD, POL, CZE, ESP; shipments in calendar 2025; returns received through 2026-01-31."],
        ["Accounting", "Highest revision per order/return ID; exact repeats removed; original sale-month ECB mean business-day quote; individual order components rounded half-up to cents."],
        ["Scenario policy", "Low/base/high uplifts 10%/25%/40%; stress uses 3% gross-sales refund shock and 10% of net sales FX shock for PLN/CZK markets; capex excluded from annual contribution and used only in payback."],
        ["Interpretation", "Client exports and hub policy are synthetic. World Bank and ECB files are archived official snapshots with provenance below. Scenarios are planning arithmetic, not causal or statistically proven forecasts."],
        ["Reproduce", "Run: python3 scripts/build_analysis.py. See README.md for outputs and verification commands."],
    ]
    for row in readme:
        ws.append(row)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120
    for cell in ws["A"]:
        cell.font = Font(bold=True, color="0B6E69")
    ws.freeze_panes = "A2"

    quality_rows = []
    for key, value in quality.items():
        if isinstance(value, (dict, list)):
            value = json.dumps(value, sort_keys=True)
        quality_rows.append({"check": key, "value": value, "status": "documented"})
    qdf = pd.DataFrame(quality_rows)
    write_df_sheet(wb, "Quality", qdf, widths={"A": 34, "B": 100, "C": 16})
    write_df_sheet(wb, "SourceRegister", pd.DataFrame(source_rows), widths={"A": 12, "B": 24, "C": 34, "D": 78, "E": 34, "F": 20, "G": 70, "H": 14, "I": 30, "J": 34, "K": 20, "L": 24})
    write_df_sheet(wb, "OrdersReconciled", orders, widths={"A": 18, "B": 10, "C": 10, "D": 16, "E": 16, "F": 14, "G": 10, "H": 12, "I": 14, "J": 15, "K": 16, "L": 10, "M": 18, "N": 16, "O": 20, "P": 14, "Q": 18, "R": 15, "S": 14, "T": 15, "U": 15, "V": 15, "W": 15, "X": 15, "Y": 15, "Z": 15, "AA": 15, "AB": 15, "AC": 15, "AD": 15, "AE": 16, "AF": 15, "AG": 15, "AH": 15, "AI": 15, "AJ": 15, "AK": 15, "AL": 15})
    write_df_sheet(wb, "ReturnsReconciled", returns, widths={"A": 20, "B": 10, "C": 18, "D": 16, "E": 12, "F": 18, "G": 16, "H": 18, "I": 18, "J": 18, "K": 18, "L": 20, "M": 24})
    write_df_sheet(wb, "FXMonthly", fx, widths={"A": 14, "B": 14, "C": 20}, number_formats={"local_per_eur": "0.000000"})
    write_df_sheet(wb, "Monthly", monthly, widths={"A": 12, "B": 12, "C": 16, "D": 16, "E": 18, "F": 16, "G": 16, "H": 16, "I": 18, "J": 20, "K": 18, "L": 14, "M": 14, "N": 18}, number_formats={x: "#,##0.00" for x in ["gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur"]} | {"margin": "0.0%"})
    write_df_sheet(wb, "Countries", countries, widths={"A": 12, "B": 16, "C": 16, "D": 18, "E": 16, "F": 16, "G": 16, "H": 16, "I": 18, "J": 20, "K": 18, "L": 14, "M": 18}, number_formats={x: "#,##0.00" for x in ["gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur"]} | {"margin": "0.0%"})
    write_df_sheet(wb, "MarketContext", context, widths={"A": 12, "B": 12, "C": 18, "D": 24}, number_formats={"population": "#,##0", "gdp_per_capita_usd": "$#,##0.00"})
    write_df_sheet(wb, "HubScenarios", scenarios, widths={"A": 14, "B": 12, "C": 24, "D": 14, "E": 10, "F": 16, "G": 20, "H": 12, "I": 14, "J": 14, "K": 14}, number_formats={"incremental_contribution_eur": "€#,##0.00", "capex_eur": "€#,##0", "annual_fixed_eur": "€#,##0", "payback_years": "0.00"})

    decision = pd.DataFrame(
        [
            {"alternative": "DEFER", "scenario": "base", "annual_incremental_contribution_eur": 0, "capex_eur": 0, "fte": 0, "payback_years": None, "feasible": True, "note": "option to defer all"},
        ]
        + scenarios.loc[(scenarios.feasible) & (scenarios.scenario == "base"), ["country", "scenario", "incremental_contribution_eur", "capex_eur", "fte", "payback_years", "feasible"]].rename(columns={"country": "alternative", "incremental_contribution_eur": "annual_incremental_contribution_eur"}).to_dict("records")
    )
    write_df_sheet(wb, "DecisionSummary", decision, widths={"A": 18, "B": 12, "C": 28, "D": 14, "E": 10, "F": 16, "G": 12, "H": 40}, number_formats={"annual_incremental_contribution_eur": "€#,##0.00", "capex_eur": "€#,##0", "payback_years": "0.00"})
    ws = wb["DecisionSummary"]
    base_rows = [i for i, row in enumerate(decision.itertuples(index=False), start=2) if row.scenario == "base"]
    if base_rows:
        chart = BarChart()
        chart.type = "bar"
        chart.style = 10
        chart.title = "Base-case annual incremental contribution"
        chart.y_axis.title = "Alternative"
        chart.x_axis.title = "EUR"
        data = Reference(ws, min_col=3, min_row=1, max_row=ws.max_row)
        cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.height = 8
        chart.width = 16
        chart.legend = None
        ws.add_chart(chart, "J2")
    ws = wb["Countries"]
    chart = BarChart()
    chart.title = "2025 contribution by country"
    chart.y_axis.title = "EUR"
    chart.x_axis.title = "Country"
    data = Reference(ws, min_col=10, min_row=1, max_row=ws.max_row)
    cats = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 8
    chart.width = 14
    ws.add_chart(chart, "O2")
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.row_dimensions[1].height = 30
    wb.save(OUT / "meridian_parts_analysis.xlsx")


def report_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CoverTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=29, textColor=colors.HexColor("#0B6E69"), alignment=TA_LEFT, spaceAfter=12))
    styles.add(ParagraphStyle(name="Subtitle", parent=styles["Normal"], fontSize=12, leading=16, textColor=colors.HexColor("#425466"), spaceAfter=10))
    styles.add(ParagraphStyle(name="H1x", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=colors.HexColor("#0B6E69"), spaceBefore=12, spaceAfter=8))
    styles.add(ParagraphStyle(name="H2x", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, leading=14, textColor=colors.HexColor("#153B50"), spaceBefore=8, spaceAfter=5))
    styles.add(ParagraphStyle(name="Bodyx", parent=styles["BodyText"], fontSize=9.2, leading=13, spaceAfter=6))
    styles.add(ParagraphStyle(name="Smallx", parent=styles["BodyText"], fontSize=7.4, leading=9.5, textColor=colors.HexColor("#425466"), spaceAfter=3))
    styles.add(ParagraphStyle(name="Callout", parent=styles["BodyText"], fontSize=11, leading=15, backColor=colors.HexColor("#E9F4F2"), borderColor=colors.HexColor("#0B6E69"), borderWidth=0.7, borderPadding=8, spaceBefore=5, spaceAfter=8))
    return styles


def pdf_table(data, widths=None, font=7.5, header=True, alignments=None):
    cell_style = ParagraphStyle(name=f"TableCell{font}", fontName="Helvetica", fontSize=font, leading=font + 1.8, textColor=colors.HexColor("#153B50"))
    header_style = ParagraphStyle(name=f"TableHeader{font}", fontName="Helvetica-Bold", fontSize=font, leading=font + 1.8, textColor=colors.white)
    wrapped = []
    for ridx, row in enumerate(data):
        current = []
        for cell in row:
            if isinstance(cell, Paragraph):
                current.append(cell)
            else:
                text = escape(str(cell)).replace("\n", "<br/>")
                current.append(Paragraph(text, header_style if header and ridx == 0 else cell_style))
        wrapped.append(current)
    tbl = Table(wrapped, colWidths=widths, repeatRows=1 if header else 0, hAlign="LEFT")
    style = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D5DDE3")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("FONTSIZE", (0, 0), (-1, -1), font),
    ]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B6E69")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
    if alignments:
        for col, alignment in enumerate(alignments):
            style.append(("ALIGN", (col, 1 if header else 0), (col, -1), alignment))
    tbl.setStyle(TableStyle(style))
    return tbl


def build_report(monthly, countries, context, scenarios, summary, quality, source_rows, asset_paths):
    styles = report_styles()
    path = OUT / "meridian_parts_board_report.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=1.5 * cm, leftMargin=1.5 * cm, topMargin=1.35 * cm, bottomMargin=1.25 * cm, title="Meridian Parts — European service-hub expansion")
    story = []
    selected = summary["selected_base"]
    selected_members = summary["selected_members"]
    selected_label = " + ".join(COUNTRY_NAMES[c] for c in selected_members) if selected_members else "Defer"
    selected_base = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "base")].iloc[0] if selected != "defer" else None
    selected_low = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "low")].iloc[0] if selected != "defer" else None
    selected_high = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "high")].iloc[0] if selected != "defer" else None
    selected_stress = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "stress")].iloc[0] if selected != "defer" else None
    alternative = summary["strongest_alternative_base"]
    alt_base = scenarios.loc[(scenarios.country == alternative) & (scenarios.scenario == "base")].iloc[0] if alternative != "defer" else None
    total = countries[["gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "shipped_orders", "shipped_units", "returned_units"]].sum()
    story += [Paragraph("Meridian Parts", styles["CoverTitle"]), Paragraph("European service-hub expansion — board decision brief", styles["Subtitle"]), HRFlowable(width="100%", color=colors.HexColor("#0B6E69"), thickness=2), Spacer(1, 8)]
    story.append(Paragraph(f"<b>Recommendation:</b> fund <b>{selected_label}</b> as a gated option, with {eur(float(selected_base.capex_eur)) if selected_base is not None else eur(0)} year-zero capex and {int(selected_base.fte) if selected_base is not None else 0} FTE, only after the 90-day launch gates confirm operational readiness. Under the board policy's base arithmetic, the option produces {eur(float(selected_base.incremental_contribution_eur)) if selected_base is not None else eur(0)} annual incremental contribution and pays back in {float(selected_base.payback_years):.2f} years; it is not a statistically proven or causal forecast.", styles["Callout"]))
    story.append(Paragraph("Decision at a glance", styles["H1x"]))
    at_a_glance = [
        ["Metric", "Selected option", "Strongest alternative", "Defer"],
        ["Base annual contribution", eur(float(selected_base.incremental_contribution_eur)) if selected_base is not None else "—", eur(float(alt_base.incremental_contribution_eur)) if alt_base is not None else "—", eur(0)],
        ["Year-zero capex", eur(float(selected_base.capex_eur)) if selected_base is not None else eur(0), eur(float(alt_base.capex_eur)) if alt_base is not None else "—", eur(0)],
        ["FTE", num(float(selected_base.fte)) if selected_base is not None else "0", num(float(alt_base.fte)) if alt_base is not None else "—", "0"],
        ["Payback", f"{float(selected_base.payback_years):.2f} years" if selected_base is not None else "—", f"{float(alt_base.payback_years):.2f} years" if alt_base is not None else "—", "—"],
        ["Stress annual contribution", eur(float(selected_stress.incremental_contribution_eur)) if selected_stress is not None else "—", eur(float(scenarios.loc[(scenarios.country == alternative) & (scenarios.scenario == "stress"), "incremental_contribution_eur"].iloc[0])) if alternative != "defer" else "—", eur(0)],
    ]
    story.append(pdf_table(at_a_glance, widths=[4.0 * cm, 4.0 * cm, 4.0 * cm, 2.5 * cm], alignments=["LEFT", "RIGHT", "RIGHT", "RIGHT"]))
    story.append(Paragraph("Why this is the decision", styles["H2x"]))
    story.append(Paragraph(f"The recommendation is driven by the observed 2025 contribution base and the client-supplied policy arithmetic, not by population or GDP as a demand proxy. The selected option is the highest base-case feasible alternative after applying the hard €{CAPEX_LIMIT:,.0f} capex and {FTE_LIMIT} FTE limits. Its concession is upfront capital and operating complexity; the strongest alternative ({alternative}) gives up {eur(float(selected_base.incremental_contribution_eur - alt_base.incremental_contribution_eur)) if selected_base is not None and alt_base is not None else '—'} of base annual contribution while requiring {eur(float(selected_base.capex_eur - alt_base.capex_eur)) if selected_base is not None and alt_base is not None else '—'} different capex. That trade-off is material enough to require a gate, not a blanket commitment.", styles["Bodyx"]))
    story.append(Paragraph("2025 baseline diagnosis", styles["H1x"]))
    story.append(Paragraph(f"The reconciled 2025 base contains {int(total.shipped_orders):,} shipped orders and {int(total.shipped_units):,} units. Gross sales are {eur(total.gross_sales_eur)}, refunds {eur(total.refunds_eur)}, net sales {eur(total.net_sales_eur)}, net COGS {eur(total.net_cogs_eur)}, fulfillment {eur(total.fulfillment_eur)}, and contribution {eur(total.contribution_eur)} ({pct(total.contribution_eur / total.net_sales_eur)} margin). Returned units are {int(total.returned_units):,}. Booked revenue is gross sales less refunds in the original sale-month translation; cash timing can differ because returns are received through the cutoff, while contribution additionally deducts net COGS and nonrefundable fulfillment. These are not interchangeable measures.", styles["Bodyx"]))
    cdata = [["Country", "Orders", "Units", "Net sales", "Contribution", "Margin", "Returned units"]]
    for r in countries.sort_values("contribution_eur", ascending=False).itertuples():
        cdata.append([COUNTRY_NAMES[r.country], num(r.shipped_orders), num(r.shipped_units), eur(r.net_sales_eur), eur(r.contribution_eur), pct(r.margin), num(r.returned_units)])
    story.append(pdf_table(cdata, widths=[3.0 * cm, 1.6 * cm, 1.8 * cm, 2.8 * cm, 2.8 * cm, 1.6 * cm, 2.0 * cm], alignments=["LEFT", "RIGHT", "RIGHT", "RIGHT", "RIGHT", "RIGHT", "RIGHT"]))
    story.append(Image(str(asset_paths["country_contribution"]), width=16.5 * cm, height=9.3 * cm))
    story.append(Paragraph("Evidence and accounting controls", styles["H1x"]))
    story.append(Paragraph("The two order extracts, correction file, returns file, unit-cost table, scenario policy and data dictionary are synthetic client inputs. The order pipeline removes exact repeated rows, selects the highest revision per order ID, keeps valid zero-price shipments, excludes cancelled/test/future shipments, and counts distinct order IDs. Returns select the highest revision per return ID, remove exact repeats, add distinct return IDs, require an eligible linked order, and use the inclusive 2026-01-31 cutoff. No orphan return is matched by guess. The original sale month's mean ECB business-day quote is used for both sale and refund conversion; PLN and CZK are not inverted.", styles["Bodyx"]))
    quality_bullets = "<br/>".join([f"• {k}: {v}" for k, v in quality.items() if not isinstance(v, list) and not isinstance(v, dict)])
    story.append(Paragraph(quality_bullets, styles["Smallx"]))
    story.append(PageBreak())
    story.append(Paragraph("Market context — context, not demand proof", styles["H1x"]))
    latest = context.loc[context.year == 2024].copy()
    latest["population_change_22_24"] = latest.country.map(dict(context.loc[context.year == 2022, ["country", "population"]].set_index("country").population)).astype(float)
    latest["population_change_22_24"] = latest["population"] / latest["population_change_22_24"] - 1
    mdata = [["Market", "2022–24 population change", "2024 population", "2024 GDP/person (current US$)"]]
    for r in latest.sort_values("population", ascending=False).itertuples():
        mdata.append([COUNTRY_NAMES[r.country], pct(r.population_change_22_24), num(r.population), f"${r.gdp_per_capita_usd:,.0f}"])
    story.append(pdf_table(mdata, widths=[4.0 * cm, 4.0 * cm, 3.3 * cm, 5.0 * cm], alignments=["LEFT", "RIGHT", "RIGHT", "RIGHT"]))
    story.append(Paragraph("The archived World Bank series show differing scale and recent population trajectories, while current-US-dollar GDP per capita supplies economic context. Neither series is a direct measure of Meridian Parts demand, service density, customer concentration, or hub savings. The source room preserves 2022–2024 values, units, missingness and the 2026-07-13 revision vintage; the ECB archive preserves the reference-rate observations used in the calculation.", styles["Bodyx"]))
    story.append(Paragraph("Hub economics and trade-offs", styles["H1x"]))
    story.append(Paragraph("The hub policy defines annual incremental contribution as baseline contribution × volume uplift + projected units × (1 + uplift) × saving per unit − annual fixed cost. Capex is not subtracted from that annual measure; it is year-zero cash and is used in simple undiscounted payback. The stress case jointly applies a 3% gross-sales refund shock and a 10% net-sales FX shock for PLN/CZK markets. Pairs add country figures without synergy.", styles["Bodyx"]))
    sdata = [["Alternative", "Low", "Base", "High", "Stress", "Capex", "FTE", "Base payback"]]
    base_rows = scenarios.loc[scenarios.feasible].pivot(index="country", columns="scenario", values="incremental_contribution_eur")
    for alt in scenarios.loc[scenarios.feasible, "country"].drop_duplicates().tolist():
        rr = scenarios.loc[(scenarios.country == alt) & (scenarios.scenario == "base")].iloc[0]
        sdata.append([alt, eur(float(scenarios.loc[(scenarios.country == alt) & (scenarios.scenario == "low"), "incremental_contribution_eur"].iloc[0])), eur(float(rr.incremental_contribution_eur)), eur(float(scenarios.loc[(scenarios.country == alt) & (scenarios.scenario == "high"), "incremental_contribution_eur"].iloc[0])), eur(float(scenarios.loc[(scenarios.country == alt) & (scenarios.scenario == "stress"), "incremental_contribution_eur"].iloc[0])), eur(float(rr.capex_eur)), num(rr.fte), f"{float(rr.payback_years):.2f}"])
    sdata_sorted = [sdata[0]] + sorted(sdata[1:], key=lambda x: float(x[2].replace("€", "").replace(",", "")), reverse=True)
    story.append(pdf_table(sdata_sorted, widths=[2.6 * cm, 2.25 * cm, 2.25 * cm, 2.25 * cm, 2.25 * cm, 2.1 * cm, 1.1 * cm, 1.9 * cm], font=6.9, alignments=["LEFT", "RIGHT", "RIGHT", "RIGHT", "RIGHT", "RIGHT", "RIGHT", "RIGHT"]))
    story.append(Image(str(asset_paths["scenario_contribution"]), width=17.0 * cm, height=9.4 * cm))
    story.append(Paragraph("Decision conditions and flip points", styles["H2x"]))
    story.append(Paragraph(f"The base-case choice is supported only as a transparent planning estimate. The selected option's policy break-even volume uplift is approximately {pct(summary['break_even_volume_uplift'])} under unchanged unit savings and fixed costs; below that, the base arithmetic does not cover recurring fixed cost. A second flip condition is operational: if the 90-day gate cannot evidence the uplift/savings prerequisites or if the downside KPI indicates stress-like performance, defer capex and retain the central-hub baseline. The strongest alternative remains relevant because it uses different capital and staffing; do not treat its lower base contribution as a universal ranking outside the stated criteria.", styles["Bodyx"]))
    story.append(Paragraph("Staged 90-day implementation", styles["H1x"]))
    idata = [
        ["Stage / owner", "Action and dependency", "Decision gate / KPI"],
        ["Days 0–30 — COO + Finance", "Lock baseline definitions, order/return controls, site shortlist, labor and lease diligence; dependency: finance approves a single contribution bridge and source register.", "Gate 1: no unresolved data or legal red flags. KPI: ≥99.5% order-ID reconciliation; 100% refunds linked or quarantined."],
        ["Days 31–60 — Operations + Procurement", "Validate 3PL/lease quotes, staffing plan, pick-pack process, service-level design and savings measurement; dependency: site and provider quotes.", "Gate 2: all-in capex within budget and staffing within 7 FTE. KPI: quoted cost variance ≤10%; documented saving €/unit by market."],
        ["Days 61–90 — Country lead + IT", "Run a controlled launch readiness test using simulated order flow and returns, monitor service levels and FX exposure; dependency: systems, training and inventory controls.", "Gate 3: release capex only if evidence supports at least the break-even uplift and no critical control fails. KPI: ≥95% on-time dispatch, return processing ≤48h, weekly contribution bridge."],
    ]
    story.append(pdf_table(idata, widths=[3.5 * cm, 8.0 * cm, 5.2 * cm], font=7.0, alignments=["LEFT", "LEFT", "LEFT"]))
    story.append(Paragraph("Key risks and mitigations", styles["H2x"]))
    story.append(Paragraph("Demand uplift and savings are planning assumptions rather than observed hub effects; mitigate with staged release and a measured €/unit bridge. Returns could erode savings; monitor return rate, refund EUR and restocked units separately. FX reference rates are analytical translations, not transaction hedge outcomes; retain monthly PLN/CZK monitoring. Capex and labor quotes may drift; set approval tolerances and a stop-work gate. Hub pairs are modeled without synergy; any claimed network benefit is outside this estimate and requires separate evidence.", styles["Bodyx"]))
    story.append(PageBreak())
    story.append(Paragraph("Limitations, evidence register and reproducibility", styles["H1x"]))
    story.append(Paragraph("Material limits: the client inputs are synthetic; no customer interviews, demand experiment, lease quotes, service-level observations or actual hub pilot are present. The World Bank/ECB material is an archived official snapshot, not a live pull at analysis time. The scenario policy is simple undiscounted arithmetic and does not include tax, working capital, inventory holding, cannibalization, network interaction, ramp timing, or discounting. Returns are known only through the cutoff; later cash timing can differ. Recommendations therefore require gates and should be reopened if the break-even uplift, fixed cost, capex, return rate or FX assumptions change materially.", styles["Bodyx"]))
    rdata = [["ID", "File", "Status", "Period / units", "Original URL"]]
    for r in source_rows:
        rdata.append([r["source_id"], r["file"], r["status"], f"{r['period']} / {r['units']}", r["original_url"]])
    story.append(pdf_table(rdata, widths=[1.3 * cm, 3.0 * cm, 4.0 * cm, 6.0 * cm, 4.0 * cm], font=6.1, alignments=["LEFT", "LEFT", "LEFT", "LEFT", "LEFT"]))
    story.append(Spacer(1, 6))
    story.append(Paragraph("Reproduce from the workspace: <font name='Courier'>python3 scripts/build_analysis.py</font>. Verification outputs include <font name='Courier'>deliverables/metrics.json</font>, <font name='Courier'>deliverables/meridian_parts_analysis.xlsx</font>, the generated report and presentation, reconciled CSVs in <font name='Courier'>analysis/</font>, and hashes in <font name='Courier'>evidence/source_register.json</font>.", styles["Smallx"]))
    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#6B7785"))
        canvas.drawString(1.5 * cm, 0.7 * cm, "Meridian Parts | synthetic client analysis | 2025 base, cutoff 2026-01-31")
        canvas.drawRightString(A4[0] - 1.5 * cm, 0.7 * cm, f"Page {doc.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def ppt_textbox(slide, x, y, w, h, text, size=18, color=(21, 59, 80), bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    return box


def add_slide_title(slide, title, subtitle=None):
    ppt_textbox(slide, 0.55, 0.25, 12.2, 0.5, title, size=25, color=(11, 110, 105), bold=True)
    if subtitle:
        ppt_textbox(slide, 0.58, 0.78, 12.0, 0.3, subtitle, size=10, color=(66, 84, 102))
    line = slide.shapes.add_shape(1, PInches(0.55), PInches(1.08), PInches(12.2), PInches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(11, 110, 105); line.line.fill.background()


def add_bullets(slide, x, y, w, h, bullets, size=17, color=(21, 59, 80)):
    box = slide.shapes.add_textbox(PInches(x), PInches(y), PInches(w), PInches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet; p.level = 0; p.font.size = PPt(size); p.font.color.rgb = RGBColor(*color); p.space_after = PPt(8)
    return box


def build_presentation(countries, monthly, scenarios, context, summary, quality, asset_paths):
    prs = Presentation()
    prs.slide_width = PInches(13.333)
    prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]
    selected = summary["selected_base"]
    selected_members = summary["selected_members"]
    selected_label = "+".join(selected_members) if selected_members else "DEFER"
    sel_base = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "base")].iloc[0] if selected != "defer" else None
    sel_stress = scenarios.loc[(scenarios.country == selected) & (scenarios.scenario == "stress")].iloc[0] if selected != "defer" else None
    alt = summary["strongest_alternative_base"]
    alt_base = scenarios.loc[(scenarios.country == alt) & (scenarios.scenario == "base")].iloc[0] if alt != "defer" else None
    # 1 cover
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(244, 248, 247)
    ppt_textbox(slide, 0.8, 1.2, 11.7, 0.7, "Meridian Parts", size=34, color=(11, 110, 105), bold=True)
    ppt_textbox(slide, 0.8, 2.05, 11.7, 0.7, "European service-hub expansion", size=28, color=(21, 59, 80), bold=True)
    ppt_textbox(slide, 0.8, 3.0, 11.5, 0.5, "Board decision brief | 2025 shipped-sales base | returns cutoff 2026-01-31", size=15, color=(66, 84, 102))
    ppt_textbox(slide, 0.8, 5.9, 11.5, 0.6, f"Recommendation: gated funding for {selected_label}", size=22, color=(11, 110, 105), bold=True)
    # 2 decision
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "Decision: fund the highest feasible base-case option, with gates")
    ppt_textbox(slide, 0.8, 1.45, 5.9, 0.6, f"{selected_label}", size=30, color=(11, 110, 105), bold=True)
    add_bullets(slide, 0.8, 2.2, 5.8, 3.5, [f"Base annual incremental contribution: {eur(float(sel_base.incremental_contribution_eur))}", f"Year-zero capex: {eur(float(sel_base.capex_eur))}; staffing: {int(sel_base.fte)} FTE", f"Simple payback: {float(sel_base.payback_years):.2f} years", f"Stress contribution: {eur(float(sel_stress.incremental_contribution_eur))}", f"Break-even volume uplift: {pct(summary['break_even_volume_uplift'])}"], size=17)
    add_bullets(slide, 7.1, 1.55, 5.4, 3.7, [f"Alternative: {alt} at {eur(float(alt_base.incremental_contribution_eur))} base annual contribution", "Hard constraints: €450k capex, 7 FTE, max two hubs", "Gate release on observed readiness and a measured uplift/savings bridge", "Defer remains the correct action if the gate cannot evidence the premise"], size=17)
    # 3 baseline
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "2025 baseline: profitable, but economics vary by market")
    slide.shapes.add_picture(str(asset_paths["country_contribution"]), PInches(0.7), PInches(1.35), width=PInches(6.1))
    total = countries[["shipped_orders", "shipped_units", "net_sales_eur", "contribution_eur", "returned_units"]].sum()
    add_bullets(slide, 7.1, 1.5, 5.5, 3.8, [f"{int(total.shipped_orders):,} shipped orders / {int(total.shipped_units):,} units", f"{eur(total.net_sales_eur)} net sales / {eur(total.contribution_eur)} contribution", f"{pct(total.contribution_eur / total.net_sales_eur)} contribution margin", f"{int(total.returned_units):,} returned units", "Contribution deducts net COGS and fulfillment; booked revenue and cash timing differ"], size=16)
    # 4 controls
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "Controls: the analysis reconciles revisions, duplicates, returns and historical FX")
    qitems = [f"{quality['orders_raw_rows']} raw order rows → {quality['orders_unique_after_dedup']} unique order IDs", f"{quality['orders_exact_duplicate_rows_removed']} exact repeats removed; {quality['orders_revision_replacements']} correction revisions selected", f"{quality['orders_excluded_cancelled']} cancelled, {quality['orders_excluded_test']} test and {quality['orders_excluded_not_2025']} non-2025 shipment records excluded", f"{quality['returns_raw_rows']} raw return rows → {quality['returns_unique_after_dedup']} return IDs; {quality['returns_orphan_returns']} orphans quarantined", f"{quality['fx_pln_months']} PLN and {quality['fx_czk_months']} CZK monthly ECB means; sale-month quote used for refunds too"]
    add_bullets(slide, 1.0, 1.5, 11.3, 4.8, qitems, size=20)
    ppt_textbox(slide, 1.0, 6.1, 11.2, 0.5, "Archived official snapshots are distinguished from synthetic client inputs; no customer interview or live account access was used.", size=13, color=(66, 84, 102))
    # 5 context
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "Market context: scale and income are context, not demand proof")
    latest = context.loc[context.year == 2024].copy().sort_values("population", ascending=False)
    rows = [["Market", "Population", "2022–24", "GDP/person"]] + [[COUNTRY_NAMES[r.country], f"{r.population:,.0f}", pct(r.population / context.loc[(context.country == r.country) & (context.year == 2022), 'population'].iloc[0] - 1), f"${r.gdp_per_capita_usd:,.0f}"] for r in latest.itertuples()]
    table = slide.shapes.add_table(len(rows), 4, PInches(0.8), PInches(1.45), PInches(6.0), PInches(4.6)).table
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i,j); cell.text = str(val); cell.text_frame.paragraphs[0].font.size = PPt(15 if i else 16); cell.text_frame.paragraphs[0].font.bold = i == 0; cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255,255,255) if i == 0 else RGBColor(21,59,80); cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(11,110,105) if i == 0 else RGBColor(238,244,243)
    add_bullets(slide, 7.3, 1.65, 5.1, 3.7, ["World Bank population and GDP per capita are preserved for 2022–2024", "GDP per capita is current US$, not PPP or constant-price real income", "Neither series measures product demand, service density, customer concentration or realized hub savings", "Use them to frame scale and economic context only"], size=17)
    # 6 economics
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "Scenario economics: the selected pair wins on base contribution within constraints")
    slide.shapes.add_picture(str(asset_paths["scenario_contribution"]), PInches(0.55), PInches(1.25), width=PInches(8.0))
    add_bullets(slide, 8.9, 1.45, 3.8, 4.3, ["Low/base/high volume uplifts: 10% / 25% / 40%", "Stress: 3% gross-sales refund shock + PLN/CZK 10% net-sales FX shock", "Capex is year-zero; recurring fixed cost is in annual contribution", "Pairs have no synergy in the model"], size=15)
    # 7 implementation
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "90-day implementation: staged release with visible gates")
    stages = [("0–30", "COO + Finance", "Lock baseline, site shortlist, labor/lease diligence", "Gate: data and legal readiness"), ("31–60", "Operations + Procurement", "Validate quotes, staffing, process and saving bridge", "Gate: capex/FTE inside limits"), ("61–90", "Country lead + IT", "Readiness test, training, order/return control", "Gate: evidence supports uplift" )]
    y=1.45
    for stage, owner, action, gate in stages:
        ppt_textbox(slide, 0.9, y, 1.0, 0.35, stage, size=18, color=(11,110,105), bold=True)
        ppt_textbox(slide, 2.0, y, 2.7, 0.35, owner, size=16, color=(21,59,80), bold=True)
        ppt_textbox(slide, 4.8, y, 4.6, 0.55, action, size=15, color=(21,59,80))
        ppt_textbox(slide, 9.6, y, 2.7, 0.7, gate, size=14, color=(66,84,102))
        y += 1.35
    ppt_textbox(slide, 0.9, 5.9, 11.6, 0.5, "Release capex only after KPI evidence: ≥99.5% order-ID reconciliation, ≥95% on-time dispatch, returns processed ≤48h, and a weekly contribution bridge.", size=15, color=(11,110,105), bold=True)
    # 8 close
    slide = prs.slides.add_slide(blank); add_slide_title(slide, "Board ask and reopen conditions")
    add_bullets(slide, 0.9, 1.45, 11.5, 4.2, [f"Approve a gated {selected_label} design envelope: {eur(float(sel_base.capex_eur))} capex and {int(sel_base.fte)} FTE", "Authorize 90-day validation, not unconditional deployment", f"Reopen or defer if realized uplift is below {pct(summary['break_even_volume_uplift'])}, savings/return KPIs fail, or capex/fixed-cost assumptions move materially", "Treat the output as a supported estimate from synthetic inputs and archived official context; it is not causal proof"], size=20)
    ppt_textbox(slide, 0.9, 6.2, 11.5, 0.35, "Source detail, calculation workbook and metrics JSON are delivered alongside this deck.", size=13, color=(66,84,102))
    prs.save(OUT / "meridian_parts_board_presentation.pptx")


def main():
    OUT.mkdir(parents=True, exist_ok=True); ANALYSIS.mkdir(parents=True, exist_ok=True); ASSETS.mkdir(parents=True, exist_ok=True)
    source_rows = source_register()
    orders, order_stats = load_orders()
    returns, return_stats = load_returns(orders)
    _, fx, fx_stats = load_fx()
    context = load_context()
    order_lines, calc_checks = calculate_orders(orders, returns, fx)
    monthly, countries, recon = aggregate_monthly(order_lines)
    options = pd.read_csv(SRC / "hub-options.csv")
    scenarios, scenario_summary = build_scenarios(countries, options)
    quality = {
        "scope": "2025 shipped orders and returns received through inclusive 2026-01-31",
        "orders_raw_rows": order_stats["raw_rows"],
        "orders_unique_after_dedup": order_stats["unique_order_ids_after_dedup"],
        "orders_exact_duplicate_rows_removed": order_stats["exact_duplicate_rows_removed"],
        "orders_revision_replacements": order_stats["revision_replacements"],
        "orders_same_revision_conflicts": order_stats["same_revision_conflicts"],
        "orders_eligible": order_stats["eligible_orders"],
        "orders_excluded_cancelled": order_stats["excluded_cancelled"],
        "orders_excluded_test": order_stats["excluded_test"],
        "orders_excluded_not_2025": order_stats["excluded_not_2025"],
        "valid_zero_price_orders_preserved": order_stats["zero_price_orders"],
        "returns_raw_rows": return_stats["raw_rows"],
        "returns_unique_after_dedup": return_stats["unique_return_ids_after_dedup"],
        "returns_exact_duplicate_rows_removed": return_stats["exact_duplicate_rows_removed"],
        "returns_revision_replacements": return_stats["revision_replacements"],
        "returns_eligible": return_stats["eligible_returns"],
        "returns_orphan_returns": return_stats["orphan_returns"],
        "returns_on_ineligible_orders": return_stats["returns_on_ineligible_orders"],
        "returns_after_cutoff": return_stats["returns_after_cutoff"],
        "fx_observations_2025": fx_stats["observations_2025"],
        "fx_pln_months": fx_stats["pln_months"],
        "fx_czk_months": fx_stats["czk_months"],
        "fx_missing_rate_cells": fx_stats["missing_rate_cells"],
        "missing_fx_rates": calc_checks["missing_fx_rates"],
        "missing_unit_costs": calc_checks["missing_unit_costs"],
        "returns_exceed_shipped_units": calc_checks["returns_exceed_shipped_units"],
        "negative_gross_sales": calc_checks["negative_gross_sales"],
        "negative_refunds": calc_checks["negative_refunds"],
        "monthly_to_country_reconciliation": recon["monthly_to_country_reconciliation"],
        "infeasible_pairs": [x for x in scenario_summary["all_pairs"] if not x["feasible"]],
        "limitations": ["synthetic client inputs", "no live demand evidence or customer interviews", "simple undiscounted scenario arithmetic", "no synergy, ramp, tax, working capital or cannibalization model"],
        "handling": "Anomalies are retained in reconciled CSVs and counted here; orphans and ineligible/late returns are excluded per data dictionary, never silently matched.",
    }
    save_csv(order_lines, ANALYSIS / "eligible_order_lines.csv")
    save_csv(orders, ANALYSIS / "orders_reconciled.csv")
    save_csv(returns, ANALYSIS / "returns_reconciled.csv")
    save_csv(monthly, ANALYSIS / "monthly.csv")
    save_csv(countries, ANALYSIS / "countries.csv")
    save_csv(fx, ANALYSIS / "fx_monthly.csv")
    save_csv(context, ANALYSIS / "market_context.csv")
    save_csv(scenarios, ANALYSIS / "hub_scenarios.csv")
    (ROOT / "evidence" / "source_register.json").write_text(json.dumps({"sources": source_rows}, indent=2) + "\n")
    metrics = json_rows(monthly, countries, fx, context, scenarios, scenario_summary, quality)
    (OUT / "metrics.json").write_text(json.dumps(metrics, indent=2, allow_nan=False) + "\n")
    asset_paths = make_assets(monthly, countries, scenarios, scenario_summary)
    make_workbook(orders, returns, fx, context, monthly, countries, scenarios, source_rows, quality, scenario_summary, asset_paths)
    build_report(monthly, countries, context, scenarios, scenario_summary, quality, source_rows, asset_paths)
    build_presentation(countries, monthly, scenarios, context, scenario_summary, quality, asset_paths)
    # Also retain a machine-readable run manifest.
    manifest = {"generated_utc": datetime.now(timezone.utc).isoformat(), "source_hashes": {r["file"]: r["sha256"] for r in source_rows}, "outputs": [str(p.relative_to(ROOT)) for p in [OUT / "metrics.json", OUT / "meridian_parts_analysis.xlsx", OUT / "meridian_parts_board_report.pdf", OUT / "meridian_parts_board_presentation.pptx"]]}
    (ANALYSIS / "run_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps({"selected": scenario_summary["selected_base"], "selected_members": scenario_summary["selected_members"], "base_annual_contribution_eur": scenario_summary["selected_base_annual_contribution_eur"], "strongest_alternative": scenario_summary["strongest_alternative_base"], "quality": quality}, indent=2))


if __name__ == "__main__":
    main()
