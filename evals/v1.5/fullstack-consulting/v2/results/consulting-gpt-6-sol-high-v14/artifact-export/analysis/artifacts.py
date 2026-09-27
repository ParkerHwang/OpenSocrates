"""Create the board-ready PDF report, PDF deck and XLSX from metrics.json."""
from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.platypus import (BaseDocTemplate, Frame, KeepTogether, PageBreak, PageTemplate,
                                Paragraph, Spacer, Table, TableStyle)
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
D = json.loads((OUT / "metrics.json").read_text())
COUNTRIES = ["DEU", "FRA", "NLD", "POL", "CZE", "ESP"]
NAVY = colors.HexColor("#142B40")
TEAL = colors.HexColor("#007D83")
GOLD = colors.HexColor("#DDAA35")
PALE = colors.HexColor("#EAF3F3")
GRAY = colors.HexColor("#566775")
WHITE = colors.white

FONT = "/System/Library/Fonts/Supplemental/Arial.ttf"
FONT_BOLD = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
pdfmetrics.registerFont(TTFont("ArialM", FONT))
pdfmetrics.registerFont(TTFont("ArialMB", FONT_BOLD))


def euro(x, dp=0):
    return f"€{x:,.{dp}f}"


def k(x):
    return f"€{x/1000:,.1f}k"


def pct(x):
    return f"{x*100:.1f}%"


def scen(name, scenario):
    return next(x for x in D["hub_scenarios"] if x["country"] == name and x["scenario"] == scenario)


def country(name):
    return next(x for x in D["countries"] if x["country"] == name)


def para(text, style="Body"):
    return Paragraph(text, STYLES[style])


STYLES = {
    "Title": ParagraphStyle("Title", fontName="ArialMB", fontSize=25, leading=30, textColor=NAVY, spaceAfter=13),
    "H1": ParagraphStyle("H1", fontName="ArialMB", fontSize=15, leading=19, textColor=NAVY, spaceBefore=11, spaceAfter=7),
    "H2": ParagraphStyle("H2", fontName="ArialMB", fontSize=10.5, leading=14, textColor=TEAL, spaceBefore=8, spaceAfter=4),
    "Body": ParagraphStyle("Body", fontName="ArialM", fontSize=8.7, leading=12.7, textColor=NAVY, spaceAfter=6),
    "Small": ParagraphStyle("Small", fontName="ArialM", fontSize=7.4, leading=10.4, textColor=GRAY, spaceAfter=5),
    "Cell": ParagraphStyle("Cell", fontName="ArialM", fontSize=7.1, leading=9.2, textColor=NAVY),
    "CellB": ParagraphStyle("CellB", fontName="ArialMB", fontSize=7.1, leading=9.2, textColor=NAVY),
}


def table(headers, rows, widths=None, font=7.2):
    data = [[para(str(h), "CellB") for h in headers]] + [[para(str(v), "Cell") for v in row] for row in rows]
    t = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND", (0,0),(-1,0), PALE), ("LINEBELOW",(0,0),(-1,0),0.6,TEAL),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[WHITE, colors.HexColor("#F7FAFA")]),
        ("VALIGN",(0,0),(-1,-1),"TOP"), ("TOPPADDING",(0,0),(-1,-1),5),
        ("BOTTOMPADDING",(0,0),(-1,-1),5), ("LEFTPADDING",(0,0),(-1,-1),5),
        ("RIGHTPADDING",(0,0),(-1,-1),5), ("LINEBELOW",(0,-1),(-1,-1),0.4,colors.HexColor("#DCE4E9")),
    ]))
    return t


def footer(c, doc):
    c.saveState()
    c.setStrokeColor(colors.HexColor("#D7E2E7")); c.line(18*mm, 16*mm, 192*mm, 16*mm)
    c.setFont("ArialM", 7); c.setFillColor(GRAY)
    c.drawString(18*mm, 11*mm, "MERIDIAN PARTS  |  Board decision brief  |  27 September 2026")
    c.drawRightString(192*mm, 11*mm, str(doc.page))
    c.restoreState()


def report():
    path = OUT / "executive_report.pdf"
    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
                          topMargin=18*mm, bottomMargin=22*mm, title="Meridian Parts | European service-hub expansion")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates(PageTemplate(id="normal", frames=[frame], onPage=footer))
    s = []
    s += [para("European service-hub expansion", "Title"),
          para("Board decision brief • 2025 shipped base • returns known by 31 January 2026", "H2"),
          Spacer(1, 7*mm), para("Recommendation", "H1"),
          para("Approve a staged <b>Czechia + Spain</b> program with a maximum <b>€225,000 year-zero capex</b> and <b>five FTE</b>. Open Spain first; authorize Czechia procurement only after the 90-day demand, service, returns and FX gate. This pair has the highest modeled base annual increment among feasible choices, <b>€198,099</b> after recurring fixed costs, with simple payback of <b>1.14 years</b> once steady-state benefits are realized. [S1–S5]"),
          para("The modeled joint stress cuts that increment to <b>€32,180</b> and stretches payback to <b>6.99 years</b>. Netherlands + Spain sacrifices €31,168 of base annual increment and requires €45,000 more capex and one extra FTE, but preserves €107,630 under the same stress. If the board requires stress payback below three years, choose Netherlands + Spain instead. The three-year threshold is a proposed risk preference, not a client rule. [S4–S5]"),
          para("Decision frame", "H1"),
          table(["Hard constraint", "Status", "Basis"], [
              ["At most two hubs", "Met: two", "Client scenario policy [S5]"],
              ["Capex ≤ €450,000", "Met: €225,000", "Client scenario policy and option costs [S4–S5]"],
              ["Staff ≤ seven FTE", "Met: five", "Client scenario policy and option staffing [S4–S5]"],
          ], [47*mm, 38*mm, 89*mm]),
          Spacer(1, 4*mm),
          para("Choice criteria are modeled annual contribution, year-zero capital, staffing, payback and resilience under the defined joint stress. The first four are computed from the synthetic client inputs; resilience is the same model under a prescribed, non-probabilistic stress. No weights or causal demand estimates are asserted. [S4–S5]", "Small"),
          para("What the evidence supports", "H1"),
          para("The eligible shipped base totals <b>1,254 orders</b> and <b>77,436 units</b>. Spain and Czechia have the two largest 2025 contributions, €421,900 and €397,335. This is a client-extract accounting observation, not proof that local hubs cause demand growth. Population and GDP per capita supply context only. [S1–S3, S6–S7]"),
          para("Conclusion card: Fund the capped, gated Czechia + Spain program. Grounds: the reconciled client base and scenario arithmetic are computed; hub savings/uplift and fixed costs are synthetic assumptions; ECB/World Bank observations are archived official snapshots. Remaining uncertainty: local hub uplift, implementation timing, actual FX/cash rates and return behavior are unverified. Flip: select Netherlands + Spain if the board adopts a ≤3-year stressed payback rule or if pilot evidence lowers the Czechia base increment by more than the €31,168 pair advantage. Defer the second site if its gate fails. Alternative: Netherlands + Spain is stronger under the joint stress but costs more capital and staff. [S1–S7]", "Small")]
    s.append(PageBreak())
    s += [para("1 | Reconciled 2025 economics", "H1"),
          para("All amounts below are euro. Shipped orders are distinct latest-revision order IDs, not raw extract rows. Returns are assigned to the original shipment month and translated at that month's ECB average quote. [S1–S3, S8]"),
          table(["Market", "Orders", "Units", "Net sales", "Contribution", "Margin", "Returned units"],
                [[c["country"], f'{c["shipped_orders"]:,}', f'{c["shipped_units"]:,}', k(c["net_sales_eur"]), k(c["contribution_eur"]), pct(c["margin"]), f'{c["returned_units"]:,}'] for c in D["countries"]],
                [23*mm,23*mm,23*mm,30*mm,31*mm,23*mm,21*mm]),
          para("The six markets produce €4,464,622 gross sales, €80,831 refunds, €4,383,792 net sales, €2,246,919 net COGS, €106,862 fulfillment and €2,030,010 contribution. Each component is rounded per order half-up before summing. Country totals equal the twelve monthly rows in the workbook and JSON. [S1–S3]"),
          para("Data-quality reconciliation", "H2"),
          table(["Checkpoint", "Treatment / result"], [
              ["Orders: 1,328 raw rows", "1,297 IDs after 13 exact duplicate extras and 18 superseded revisions; no conflicting top revisions."],
              ["Eligibility", "24 cancelled, 18 tests, one January 2026 shipment excluded; 1,254 eligible 2025 shipments. Reasons can overlap: cancelled rows also lack a 2025 shipment date."],
              ["Zero-price", "12 eligible zero-price shipments retained; they carry units, COGS and fulfillment even if sales are zero."],
              ["Returns: 264 raw rows", "243 IDs after 20 duplicate extras and one superseded revision; 241 linked by cutoff, one orphan quarantined, one post-cutoff excluded."],
              ["Missingness", "No missing cost or FX inputs for eligible orders; no missing 2022–2024 World Bank values in the archived response."],
          ], [42*mm,132*mm]),
          para("Accounting interpretation", "H2"),
          para("Booked net revenue is gross shipped sales less credited refunds, excluding taxes. Cash differs because payments, settlements, refund timing and actual bank FX are absent. Contribution then subtracts net product cost (restocked units recover cost at the original SKU rate) and nonrefundable fulfillment; it excludes hub capex, recurring hub fixed costs and wider corporate overhead. A returned unit is not automatically reusable. [S1–S3]"),
          para("Monthly records, an order-level audit, return quarantine and exact formulas are in the accompanying workbook, JSON and analysis scripts.", "Small")]
    s.append(PageBreak())
    s += [para("2 | Market and currency context", "H1"),
          para("The archived World Bank responses cover each market in 2022, 2023 and 2024. The table compares population change and 2024 GDP per capita in current US dollars; the latter is neither PPP nor real-income growth. These figures establish operating context, not product demand. [S6–S7]"),
          table(["Market", "Population 2022", "Population 2024", "Change", "GDP/head 2024 (US$)"],
                [[c, f'{next(x for x in D["market_context"] if x["country"]==c and x["year"]==2022)["population"]:,}',
                  f'{next(x for x in D["market_context"] if x["country"]==c and x["year"]==2024)["population"]:,}',
                  pct(next(x for x in D["market_context"] if x["country"]==c and x["year"]==2024)["population"] / next(x for x in D["market_context"] if x["country"]==c and x["year"]==2022)["population"] - 1),
                  f'{next(x for x in D["market_context"] if x["country"]==c and x["year"]==2024)["gdp_per_capita_usd"]:,.0f}'] for c in COUNTRIES],
                [25*mm,37*mm,37*mm,28*mm,47*mm]),
          para("Spain (+2.22%) and Czechia (+2.18%) show the largest 2022–2024 population increases; Poland declined 0.71%. The Netherlands had the highest 2024 GDP/head in this set. None of these aggregate measures justifies the modeled 25% hub-related volume uplift; that is a client scenario assumption to validate. [S5–S7]"),
          para("2025 ECB monthly reference-rate means", "H2"),
          para("Local units per €1; mean of available published business-day observations. They are used for each order's sale month, including refunds credited later. EUR markets use 1.000000. [S1, S8]"),
          table(["Month", "PLN/€", "CZK/€", "Month", "PLN/€", "CZK/€"],
                [[f"2025-{m:02d}", f'{next(x for x in D["fx_monthly"] if x["currency"]=="PLN" and x["month"]==f"2025-{m:02d}")["local_per_eur"]:.5f}',
                  f'{next(x for x in D["fx_monthly"] if x["currency"]=="CZK" and x["month"]==f"2025-{m:02d}")["local_per_eur"]:.5f}',
                  f"2025-{m+6:02d}", f'{next(x for x in D["fx_monthly"] if x["currency"]=="PLN" and x["month"]==f"2025-{m+6:02d}")["local_per_eur"]:.5f}',
                  f'{next(x for x in D["fx_monthly"] if x["currency"]=="CZK" and x["month"]==f"2025-{m+6:02d}")["local_per_eur"]:.5f}'] for m in range(1,7)],
                [29*mm]*6),
          para("The original source-room archive was retrieved 27 September 2026 and reports World Bank last-updated 13 July 2026. It can reflect revisions to 2022–2024 history; no live replacement series was blended into the calculations. [S6–S8]", "Small")]
    s.append(PageBreak())
    s += [para("3 | Hub choice and sensitivity", "H1"),
          para("Year-zero capex is separate from recurring annual fixed cost. The scenario-policy formula applies the assumed 10% / 25% / 40% volume uplifts to 2025 contribution and per-unit savings, then subtracts annual hub fixed cost. The stress layers extra refunds of 3% of gross sales and, in PLN/CZK markets, depreciation of 10% of net sales into the base-uplift case. These are synthetic assumptions, not probabilities or causal effects. [S4–S5]"),
          table(["Feasible choice", "Capex", "FTE", "Low annual", "Base annual", "High annual", "Stress annual", "Base payback"],
                [[name, k(scen(name,"base")["capex_eur"]), str(scen(name,"base")["fte"]),
                  k(scen(name,"low")["incremental_contribution_eur"]), k(scen(name,"base")["incremental_contribution_eur"]),
                  k(scen(name,"high")["incremental_contribution_eur"]), k(scen(name,"stress")["incremental_contribution_eur"]),
                  (f'{scen(name,"base")["payback_years"]:.2f}y' if scen(name,"base")["payback_years"] is not None else "—")] for name in ["CZE+ESP","POL+ESP","POL+CZE","NLD+ESP","NLD+CZE","ESP","CZE","DEFER"]],
                [29*mm,20*mm,11*mm,25*mm,25*mm,25*mm,25*mm,14*mm]),
          para("The full 18-choice × 4-scenario feasible matrix, including all single hubs and pairs, is in the workbook and JSON. Germany + France breaches both caps; Germany + Spain and Germany + Poland breach the staffing cap. No pair synergy is assumed. [S4–S5]", "Small"),
          para("Why this choice, and when it changes", "H2"),
          para("Czechia + Spain leads the base case by €11,770 versus Poland + Spain and €31,168 versus Netherlands + Spain, while using €45,000 less capex and one fewer FTE than the latter. It also gives up €75,450 of annual contribution in the joint stress relative to Netherlands + Spain. Under the specified full 3% extra refund shock, a Czechia FX depreciation impact around 2.6% of its net sales would erase the base-case advantage over Netherlands + Spain; the policy stress assumes 10%. That switching point is arithmetic, not a forecast. [S4–S5]"),
          para("Defer is financially preferable only if the assumed gains fail enough to make annual increments negative or if the service experiment does not validate a viable operating design. The board has not supplied a stress-risk tolerance or hurdle rate; the comparison therefore remains conditional on that preference. [S5]"),
          para("All paybacks are simple capex / positive annual increment, with no ramp delay, taxes, discounting, working capital or residual value. They are not investment NPVs.", "Small")]
    s.append(PageBreak())
    s += [para("4 | Ninety-day execution and decision gates", "H1"),
          table(["Window", "Accountable owner", "Work and dependency", "Exit gate"], [
              ["Days 0–30", "COO / Finance controller", "Freeze 2025 baseline and return definitions; contract carrier and facility options for Spain; map tax, labor and inventory controls before commitments.", "Signed baseline; facility and carrier quotes within €130k Spain capex and three FTE plan."],
              ["Days 31–60", "Spain operations lead / Customer service", "Pilot Spain order routing and replenishment with limited SKUs; instrument promised versus actual delivery, availability, returns and true savings per unit.", "At least 95% on-time delivery, ≥98% stock availability, and positive measured unit savings after variable handling."],
              ["Days 61–90", "COO / CFO / Czechia lead", "Review Spain pilot, Czechia customer-order evidence and PLN/CZK exposure; test current return mix, supplier quotes, staff readiness and revised payback.", "Release Czechia ≤€95k and two FTE only if revised base increment remains positive and board accepts modeled stress exposure; otherwise retain Spain and reassess Netherlands."],
          ], [23*mm,36*mm,67*mm,48*mm]),
          para("The 95% and 98% pilot thresholds are proposed operational gates, not historical client performance. A two-site opening inside 90 days is not assumed; the day-90 decision releases or withholds second-site capex. [S4–S5]", "Small"),
          para("KPI plan and ownership", "H2"),
          table(["KPI", "Definition / cadence", "Owner", "Decision use"], [
              ["Incremental contribution", "Pilot net sales − net COGS − fulfillment − local recurring cost, versus matched central-routing baseline; weekly and month-end.", "Finance controller", "Recalculate annualized payback and stop loss-making expansion."],
              ["Service level", "On-time delivered orders / promised orders and in-stock order lines / requested lines; weekly.", "Spain operations lead", "Verify hub service advantage and gate operating readiness."],
              ["Returns", "Returned physical units / shipped units; refund EUR / gross sales EUR by original sale cohort; weekly cohort readout.", "Customer service lead", "Detect refund shock and avoid treating returned units as full COGS recovery."],
              ["FX and capacity", "PLN/CZK quote, settled cash rate, unit throughput per FTE and facility utilization; weekly.", "Treasury / COO", "Update Czechia stress and staffing capacity before commitment."],
          ], [31*mm,70*mm,33*mm,40*mm]),
          para("Principal risks", "H2"),
          para("<b>Uplift error:</b> no causal estimate or control group in the extracts; use the Spain routing pilot and cohort comparison. <b>FX/return shock:</b> Czechia's stress annual increment is negative on its own (−€37,827); treasury and customer service must refresh the stress before release. <b>Execution lag:</b> permitting, supplier contracts and staffing can delay savings; CFO should revise the payback clock at each gate. <b>Inventory tied up:</b> working-capital needs are not in the model; finance must price them before final capital release. [S1–S5]"),
          para("Source and calculation notes", "H2"),
          para("[S1] Synthetic data dictionary, orders parts 1–2 and corrections; [S2] synthetic returns; [S3] synthetic effective-dated unit costs; [S4] synthetic hub-options.csv; [S5] synthetic scenario-policy.md; [S6] archived World Bank population response, SP.POP.TOTL; [S7] archived World Bank GDP/head response, NY.GDP.PCAP.CD; [S8] archived ECB eurofxref-hist.zip. Exact URLs, collection timestamps, SHA-256 hashes, periods and units are in sources/source_register.json. Scripts and quality exceptions are supplied with this brief. Source-room collection was local; public observations came from its archived official snapshots.", "Small"),
          para("World Bank: https://api.worldbank.org/v2/country/DEU;FRA;NLD;POL;CZE;ESP/indicator/SP.POP.TOTL?date=2022:2024&amp;format=json&amp;per_page=1000 and /NY.GDP.PCAP.CD with the same parameters. ECB: https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip", "Small")]
    doc.build(s)
    return path


def workbook():
    wb = Workbook()
    ws = wb.active; ws.title = "Read me"
    rows = [
        ["MERIDIAN PARTS | ANALYTICAL WORKBOOK", "2025 shipped base; returns known by 2026-01-31"],
        ["Use", "All figure sheets are generated from saved raw inputs by analysis/build.py."],
        ["Evidence", "Synthetic client exports and archived official ECB/World Bank snapshots; see Sources sheet and sources/source_register.json."],
        ["Accounting", "One latest-revision row/order; sale-month FX; half-up cent rounding per order; returns assigned to shipment month."],
        ["Annual scenario", "C*u + U*(1+u)*s - F, with u=10%/25%/40%; year-zero capex excluded."],
        ["Stress", "(C-3%*G-10%*N for PLN/CZK)*1.25 - C + U*1.25*s - F; no cost recovery for added refund."],
        ["Decision", "Stage CZE+ESP: €225,000 capex, five FTE; gate Czechia release after Spain pilot."],
        ["Limit", "Uplift, savings, fixed cost and stress are synthetic assumptions, not causal forecasts."],
    ]
    for row in rows: ws.append(row)
    ws.column_dimensions["A"].width = 28; ws.column_dimensions["B"].width = 110
    ws.freeze_panes = "B2"
    def data_sheet(name, data, keys, widths=None, money=()):
        sh = wb.create_sheet(name)
        sh.append(keys)
        for r in data: sh.append([r.get(k) for k in keys])
        sh.freeze_panes = "B2"; sh.auto_filter.ref = sh.dimensions
        for j,kname in enumerate(keys,1):
            sh.column_dimensions[get_column_letter(j)].width = (widths[j-1] if widths else max(16, len(kname)+3))
            if kname in money or kname.endswith("_eur"):
                for row in sh.iter_rows(min_row=2, min_col=j, max_col=j): row[0].number_format = '#,##0.00;[Red](#,##0.00)'
            elif kname == "margin":
                for row in sh.iter_rows(min_row=2, min_col=j, max_col=j): row[0].number_format = '0.0%'
        return sh
    monthly_keys = ["country","month","shipped_orders","shipped_units","gross_sales_eur","refunds_eur","net_sales_eur","net_cogs_eur","fulfillment_eur","contribution_eur","returned_units","margin"]
    data_sheet("Monthly", D["monthly"], monthly_keys)
    country_keys = [x for x in monthly_keys if x != "month"]
    csh = data_sheet("Country totals", D["countries"], country_keys)
    context = []
    for x in D["market_context"]:
        y = dict(x); y["population_unit"] = "persons"; y["gdp_unit"] = "current US$ per person"
        y["population_source"] = "World Bank SP.POP.TOTL archive; 2026-07-13 revision"
        y["gdp_source"] = "World Bank NY.GDP.PCAP.CD archive; 2026-07-13 revision"
        context.append(y)
    data_sheet("Market context", context, ["country","year","population","gdp_per_capita_usd","population_unit","gdp_unit","population_source","gdp_source"], [14,12,20,23,19,26,62,68])
    data_sheet("FX monthly", D["fx_monthly"], ["currency","month","local_per_eur"], [16,18,22])
    sc = data_sheet("Hub scenarios", D["hub_scenarios"], ["country","scenario","incremental_contribution_eur","capex_eur","fte","payback_years"], [20,16,36,22,12,19])
    sc["H1"] = "Feasible choices"; sc["H2"] = "18 including defer; each has four scenarios"
    opts = []
    for x in D["hub_scenarios"]:
        if x["scenario"] != "base": continue
        name = x["country"]
        if name == "DEFER": continue
        opts.append({"choice":name, "base_eur": x["incremental_contribution_eur"],
                     "stress_eur": scen(name,"stress")["incremental_contribution_eur"],
                     "capex_eur":x["capex_eur"], "fte":x["fte"]})
    opts.sort(key=lambda x:x["base_eur"], reverse=True)
    osh = data_sheet("Decision chart", opts, ["choice","base_eur","stress_eur","capex_eur","fte"], [24,22,22,22,14])
    chart = BarChart(); chart.type="bar"; chart.style=10; chart.title="Annual increment: base vs joint stress"; chart.y_axis.title="Feasible option"; chart.x_axis.title="EUR/year"
    chart.add_data(Reference(osh,min_col=2,max_col=3,min_row=1,max_row=10),titles_from_data=True)
    chart.set_categories(Reference(osh,min_col=1,min_row=2,max_row=10)); chart.height=10; chart.width=21; osh.add_chart(chart,"G2")
    chart2=BarChart(); chart2.type="col"; chart2.style=11; chart2.title="2025 contribution by market"; chart2.y_axis.title="EUR"
    chart2.add_data(Reference(csh,min_col=9,min_row=1,max_row=7),titles_from_data=True)
    chart2.set_categories(Reference(csh,min_col=1,min_row=2,max_row=7)); chart2.height=9; chart2.width=18; csh.add_chart(chart2,"M2")
    quality = D["quality"]
    qrows = [["Order raw rows",quality["order_reconciliation"]["raw_rows"]],
             ["Order distinct IDs",quality["order_reconciliation"]["distinct_ids"]],
             ["Order exact duplicate extras",quality["order_reconciliation"]["exact_duplicate_extra_rows"]],
             ["Order superseded revisions",quality["order_reconciliation"]["superseded_revision_rows"]],
             ["Eligible 2025 orders",quality["eligible_2025_orders"]],
             ["Excluded latest records: cancelled",quality["excluded_latest_order_records"].get("not_shipped",0)],
             ["Excluded latest records: test",quality["excluded_latest_order_records"].get("test",0)],
             ["Excluded latest records: outside 2025",quality["excluded_latest_order_records"].get("outside_2025_shipment",0)],
             ["Return raw rows",quality["return_reconciliation"]["raw_rows"]],
             ["Return distinct IDs",quality["return_reconciliation"]["distinct_ids"]],
             ["Return exact duplicate extras",quality["return_reconciliation"]["exact_duplicate_extra_rows"]],
             ["Return superseded revisions",quality["return_reconciliation"]["superseded_revision_rows"]],
             ["Included return IDs",quality["included_return_ids"]],
             ["Quarantined orphan",str(quality["quarantined_returns"])],
             ["After-cutoff return",str(quality["after_cutoff_return_ids"])],
             ["Zero-price eligible shipments",quality["zero_price_eligible_shipments"]],
             ["Missing required inputs",str(quality["missing_required_inputs"])]]
    qsh=wb.create_sheet("Quality"); qsh.append(["Checkpoint","Finding / treatment"])
    for row in qrows:qsh.append(row)
    qsh.column_dimensions["A"].width=42;qsh.column_dimensions["B"].width=100;qsh.freeze_panes="B2"
    source_rows=json.loads((ROOT/"sources"/"source_register.json").read_text())
    source_keys=["file","status","collected_from","collected_utc","sha256","units","period","original_public_url","original_retrieved_utc","archive_hash_verified"]
    data_sheet("Sources",source_rows,source_keys,[34,34,65,31,70,32,45,100,31,24])
    # Applied to all sheets for readable decision tables.
    for sh in wb:
        sh.sheet_view.showGridLines=False
        for cell in sh[1]:
            cell.fill=PatternFill("solid",fgColor="142B40"); cell.font=Font(name="Aptos",size=10,bold=True,color="FFFFFF")
            cell.alignment=Alignment(vertical="center",wrap_text=True)
        sh.row_dimensions[1].height=29
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                cell.font=Font(name="Aptos",size=10,color="142B40")
                cell.alignment=Alignment(vertical="top",wrap_text=False)
    path=OUT/"analytical_workbook.xlsx"; wb.save(path); return path


def slide_text(c, text, x, y, width, size=16, leading=None, color=NAVY, bold=False):
    leading = leading or size*1.35
    style=ParagraphStyle("slide",fontName="ArialMB" if bold else "ArialM",fontSize=size,leading=leading,textColor=color)
    p=Paragraph(text,style); w,h=p.wrap(width,1000); p.drawOn(c,x,y-h); return y-h


def deck():
    path=OUT/"board_presentation.pdf"; W,H=landscape(A4)
    c=canvas.Canvas(str(path),pagesize=(W,H)); c.setTitle("Meridian Parts | Board presentation")
    def base(num,title,kicker="BOARD DECISION | SEPTEMBER 2026"):
        c.setFillColor(WHITE);c.rect(0,0,W,H,fill=1,stroke=0)
        c.setFillColor(NAVY);c.rect(0,H-22*mm,W,22*mm,fill=1,stroke=0)
        c.setFont("ArialMB",11);c.setFillColor(WHITE);c.drawString(18*mm,H-14*mm,"MERIDIAN PARTS")
        c.setFont("ArialM",9);c.drawRightString(W-18*mm,H-14*mm,kicker)
        c.setFont("ArialMB",23);c.setFillColor(NAVY);c.drawString(18*mm,H-38*mm,title)
        c.setStrokeColor(colors.HexColor("#D4E0E5"));c.line(18*mm,16*mm,W-18*mm,16*mm)
        c.setFont("ArialM",8);c.setFillColor(GRAY);c.drawString(18*mm,10*mm,"Source: saved synthetic client files + archived ECB / World Bank; details in report and source register")
        c.drawRightString(W-18*mm,10*mm,str(num))
    def box(x,y,w,h,title,value,sub,fill=PALE):
        c.setFillColor(fill);c.roundRect(x,y,w,h,7*mm,fill=1,stroke=0)
        slide_text(c,title,x+7*mm,y+h-7*mm,w-14*mm,11,color=TEAL,bold=True)
        slide_text(c,value,x+7*mm,y+h-23*mm,w-14*mm,22,bold=True)
        slide_text(c,sub,x+7*mm,y+20*mm,w-14*mm,10,color=GRAY)
    base(1,"Fund two sites in stages")
    slide_text(c,"Approve Czechia + Spain within €225k capex / five FTE; open Spain first and gate Czechia at day 90.",18*mm,H-52*mm,W-36*mm,18,bold=True)
    box(18*mm,48*mm,77*mm,75*mm,"BASE ANNUAL INCREMENT","€198.1k","Highest feasible pair")
    box(105*mm,48*mm,77*mm,75*mm,"SIMPLE PAYBACK","1.14 years","At modeled steady state")
    box(192*mm,48*mm,87*mm,75*mm,"JOINT STRESS INCREMENT","€32.2k","Material FX / returns exposure")
    slide_text(c,"Decision condition: if the board requires stressed payback ≤3 years, choose Netherlands + Spain instead.",18*mm,38*mm,W-36*mm,11,color=GRAY)
    c.showPage()
    base(2,"The 2025 base is reconciled")
    c.setFillColor(PALE);c.roundRect(18*mm,45*mm,98*mm,105*mm,6*mm,fill=1,stroke=0)
    y=H-58*mm
    for t,v in [("Raw order rows","1,328"),("Distinct latest IDs","1,297"),("Eligible shipped orders","1,254"),("Shipped units","77,436")]:
        slide_text(c,t,25*mm,y,80*mm,11,color=GRAY);slide_text(c,v,85*mm,y,22*mm,14,bold=True);y-=23*mm
    c.setFillColor(PALE);c.roundRect(126*mm,45*mm,153*mm,105*mm,6*mm,fill=1,stroke=0)
    slide_text(c,"Conservative accounting",134*mm,H-58*mm,136*mm,14,bold=True)
    slide_text(c,"Latest revisions replace whole orders. Cancelled (24), test (18) and one 2026 shipment are excluded. Twelve zero-price shipments stay in the base. One orphan return is quarantined; one late return is excluded.",134*mm,H-78*mm,136*mm,12)
    slide_text(c,"Gross sales − refunds − net COGS − nonrefundable fulfillment = contribution. Cash settlement and actual bank FX are outside this extract.",134*mm,H-121*mm,136*mm,11,color=GRAY)
    c.showPage()
    base(3,"Spain and Czechia lead the baseline")
    vals=[country(x)["contribution_eur"] for x in COUNTRIES]
    maxv=max(vals)
    for i,(name,val) in enumerate(zip(COUNTRIES,vals)):
        y=H-61*mm-i*19*mm
        c.setFillColor(NAVY);c.setFont("ArialMB",11);c.drawString(20*mm,y,name)
        c.setFillColor(PALE);c.roundRect(50*mm,y-3*mm,180*mm,10*mm,3*mm,fill=1,stroke=0)
        c.setFillColor(TEAL if name in ("CZE","ESP") else colors.HexColor("#7FA8AF"));c.roundRect(50*mm,y-3*mm,180*mm*val/maxv,10*mm,3*mm,fill=1,stroke=0)
        c.setFont("ArialMB",11);c.setFillColor(NAVY);c.drawRightString(275*mm,y,k(val))
    slide_text(c,"Observed 2025 contribution is a sizing anchor, not evidence that a local hub causes more demand.",20*mm,30*mm,250*mm,11,color=GRAY)
    c.showPage()
    base(4,"Base upside versus joint-stress resilience")
    names=["CZE+ESP","POL+ESP","NLD+ESP","POL+CZE"]
    c.setFont("ArialMB",11);c.setFillColor(NAVY)
    for xx,lab in [(20,"Choice"),(65,"Base €/yr"),(115,"Stress €/yr"),(170,"Capex"),(218,"FTE"),(244,"Payback")]: c.drawString(xx*mm,H-59*mm,lab)
    for i,name in enumerate(names):
        y=H-(75+i*24)*mm
        if i==0:c.setFillColor(PALE);c.roundRect(18*mm,y-6*mm,261*mm,20*mm,4*mm,fill=1,stroke=0)
        c.setFont("ArialMB" if i==0 else "ArialM",13);c.setFillColor(NAVY)
        vals=[name,k(scen(name,"base")["incremental_contribution_eur"]),k(scen(name,"stress")["incremental_contribution_eur"]),k(scen(name,"base")["capex_eur"]),str(scen(name,"base")["fte"]),f'{scen(name,"base")["payback_years"]:.2f}y']
        for xx,v in zip([20,65,115,170,218,244],vals):c.drawString(xx*mm,y,v)
    slide_text(c,"Low/base/high uplift: 10% / 25% / 40%. Stress: base uplift plus extra refunds and PLN/CZK depreciation. Scenarios are assumptions, not forecast probabilities.",20*mm,35*mm,258*mm,11,color=GRAY)
    c.showPage()
    base(5,"The alternative protects the downside")
    box(18*mm,55*mm,120*mm,90*mm,"CZECHIA + SPAIN","€198.1k base","€32.2k stress • €225k capex • 5 FTE")
    box(158*mm,55*mm,121*mm,90*mm,"NETHERLANDS + SPAIN","€166.9k base","€107.6k stress • €270k capex • 6 FTE")
    slide_text(c,"The chosen pair gains €31.2k/year in the base case, saves €45k capex and one FTE; it concedes €75.5k/year under the defined joint stress.",18*mm,43*mm,260*mm,12)
    c.showPage()
    base(6,"Context is not a demand forecast")
    headers=["Market","Population change 2022–24","GDP/head 2024 (current US$)"]
    for x,h in zip([20,85,188],headers):c.setFont("ArialMB",11);c.drawString(x*mm,H-61*mm,h)
    for i,name in enumerate(COUNTRIES):
        y=H-(76+i*19)*mm
        a=[x for x in D["market_context"] if x["country"]==name]
        change=a[2]["population"]/a[0]["population"]-1
        for x,v in zip([20,85,188],[name,pct(change),f'${a[2]["gdp_per_capita_usd"]:,.0f}']):c.setFont("ArialM",12);c.setFillColor(NAVY);c.drawString(x*mm,y,v)
    slide_text(c,"Archived World Bank series; no missing 2022–2024 cells. Population and current-dollar GDP/head do not establish replacement-assembly demand.",20*mm,26*mm,260*mm,11,color=GRAY)
    c.showPage()
    base(7,"Release capital through measurable gates")
    cards=[("0–30 days","COO + Finance","Freeze baseline; secure Spain facility/carrier quotes; review compliance and inventory controls."),
           ("31–60 days","Spain operations","Pilot limited-SKU routing; track on-time delivery, stock availability, unit savings and refunds."),
           ("61–90 days","COO + CFO","Decide Czechia release ≤€95k / 2 FTE from revised base, return, FX and capacity evidence.")]
    for i,(phase,owner,body) in enumerate(cards):
        x=(18+i*88)*mm;c.setFillColor(PALE);c.roundRect(x,57*mm,82*mm,91*mm,6*mm,fill=1,stroke=0)
        slide_text(c,phase,x+6*mm,139*mm,70*mm,15,bold=True)
        slide_text(c,owner,x+6*mm,117*mm,70*mm,11,color=TEAL,bold=True)
        slide_text(c,body,x+6*mm,100*mm,70*mm,11)
    slide_text(c,"Proposed pilot gates: ≥95% on-time delivery, ≥98% availability, positive unit savings; Czechia release requires board acceptance of revised stress exposure.",18*mm,44*mm,260*mm,11,color=GRAY)
    c.showPage();c.save();return path


if __name__=="__main__":
    print(report());print(deck());print(workbook())
