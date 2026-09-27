"""Independent output checks for the saved-input Meridian deliverables."""
import csv
import hashlib
import json
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
data = json.loads((ROOT / "deliverables/metrics.json").read_text())
by_country = {x["country"]: x for x in data["countries"]}
months = defaultdict(list)
for row in data["monthly"]:
    months[row["country"]].append(row)
assert len(data["monthly"]) == 72 and all(len(v) == 12 for v in months.values())
fields = ["shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur",
          "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units"]
for c, country in by_country.items():
    for key in fields:
        assert abs(sum(x[key] for x in months[c]) - country[key]) < 0.001, (c, key)
    assert abs(country["gross_sales_eur"] - country["refunds_eur"] - country["net_sales_eur"]) < 0.001
    assert abs(country["net_sales_eur"] - country["net_cogs_eur"] - country["fulfillment_eur"] - country["contribution_eur"]) < 0.001
    assert abs(country["contribution_eur"] / country["net_sales_eur"] - country["margin"]) < 1e-7
assert sum(x["shipped_orders"] for x in data["countries"]) == 1254
assert sum(x["shipped_units"] for x in data["countries"]) == 77436
assert len(data["fx_monthly"]) == 24

# Recalculate annual options independently from the published country values and raw option rows.
with (ROOT / "sources/raw/hub-options.csv").open(newline="") as f:
    options = {x["country"]: x for x in csv.DictReader(f)}
scenarios = {(x["country"], x["scenario"]): x for x in data["hub_scenarios"]}
assert len(scenarios) == 72
for c, baseline in by_country.items():
    o = options[c]
    C = Decimal(str(baseline["contribution_eur"]))
    U = Decimal(baseline["shipped_units"])
    s = Decimal(o["saving_eur_per_unit"])
    F = Decimal(o["annual_fixed_eur"])
    expected = C * Decimal("0.25") + U * Decimal("1.25") * s - F
    assert abs(expected - Decimal(str(scenarios[(c, "base")]["incremental_contribution_eur"]))) <= Decimal("0.01")
for (name, scenario), row in scenarios.items():
    if name == "DEFER":
        assert row["incremental_contribution_eur"] == row["capex_eur"] == row["fte"] == 0
        assert row["payback_years"] is None
        continue
    members = name.split("+")
    assert len(members) <= 2 and row["capex_eur"] <= 450000 and row["fte"] <= 7
    assert abs(sum(scenarios[(x, scenario)]["incremental_contribution_eur"] for x in members) - row["incremental_contribution_eur"]) <= 0.02
    if row["incremental_contribution_eur"] > 0:
        assert abs(row["capex_eur"] / row["incremental_contribution_eur"] - row["payback_years"]) < 0.0001
    else:
        assert row["payback_years"] is None

# Verify archive provenance and World Bank panel structure.
register = json.loads((ROOT / "sources/source_register.json").read_text())
archive = json.loads((ROOT / "sources/raw/source-register.json").read_text())
archived = {x["file"]: x for x in archive["public"]}
for item in register:
    path = ROOT / item["file"]
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    assert digest == item["sha256"]
    if path.name in archived:
        assert digest == archived[path.name]["sha256"] and item["archive_hash_verified"] is True
assert len(data["market_context"]) == 18
assert {(x["country"], x["year"]) for x in data["market_context"]} == {(c, y) for c in by_country for y in (2022, 2023, 2024)}

# Deliverable structural checks and text extraction; raster inspection is recorded separately.
wb = load_workbook(ROOT / "deliverables/analytical_workbook.xlsx", read_only=False)
assert {"Monthly", "Country totals", "Market context", "FX monthly", "Hub scenarios", "Decision chart", "Quality", "Sources"}.issubset(wb.sheetnames)
assert wb["Monthly"].max_row == 73 and wb["Market context"].max_row == 19
assert wb["Hub scenarios"].max_row == 73 and wb["FX monthly"].max_row == 25
assert len(wb["Decision chart"]._charts) == 1 and len(wb["Country totals"]._charts) == 1
for name, min_pages, terms in [
    ("executive_report.pdf", 4, ["Czechia + Spain", "Netherlands + Spain", "Data-quality reconciliation", "Ninety-day"]),
    ("board_presentation.pdf", 7, ["Fund two sites in stages", "€198.1k", "Release capital through measurable gates"]),
]:
    pdf = PdfReader(str(ROOT / "deliverables" / name))
    assert len(pdf.pages) >= min_pages
    text = "\n".join(p.extract_text() for p in pdf.pages)
    for term in terms: assert term in text, (name, term)

print("PASS: 72 monthly records, six country reconciliations, 72 feasible scenarios, archive hashes, workbook sheets/charts and PDF text/pages.")
