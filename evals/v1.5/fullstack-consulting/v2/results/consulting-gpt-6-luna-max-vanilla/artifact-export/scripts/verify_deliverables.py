#!/usr/bin/env python3
"""Independent structural and arithmetic checks for generated deliverables."""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
ANALYSIS=ROOT/"analysis"
DEL=ROOT/"deliverables"
SRC=ROOT/"evidence"/"source-room"


def sha(path):
    h=hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""): h.update(b)
    return h.hexdigest()


def close(a,b,tol=0.005): return abs(float(a)-float(b)) <= tol
def money(v): return Decimal(str(v))
def cents(v): return v.quantize(Decimal("0.01"),rounding=ROUND_HALF_UP)


def main():
    m=json.loads((DEL/"metrics.json").read_text(encoding="utf-8"))
    checks=[]
    def record(name,ok,detail):
        checks.append({"check":name,"passed":bool(ok),"detail":detail})
        if not ok: raise AssertionError(f"{name}: {detail}")

    required={"monthly","countries","fx_monthly","market_context","hub_scenarios","recommendation","quality"}
    record("metrics top-level schema",required.issubset(m),f"keys={sorted(m)}")
    monthly_fields={"country","month","shipped_orders","shipped_units","gross_sales_eur","refunds_eur","net_sales_eur","net_cogs_eur","fulfillment_eur","contribution_eur","returned_units","margin"}
    country_fields=monthly_fields-{"month"}
    record("monthly schema",all(set(r)==monthly_fields for r in m["monthly"]),f"rows={len(m['monthly'])}; fields={sorted(m['monthly'][0])}")
    record("metrics copy",sha(DEL/"metrics.json")==sha(ANALYSIS/"metrics.json"),"deliverables/metrics.json matches reproducible analysis output byte for byte")
    record("country schema",all(set(r)==country_fields for r in m["countries"]),f"rows={len(m['countries'])}; fields={sorted(m['countries'][0])}")
    record("country-month grid",len(m["monthly"])==72 and len({(r["country"],r["month"]) for r in m["monthly"]})==72,"6 countries x 12 months, no duplicates")
    record("country rows",len(m["countries"])==6 and len({r["country"] for r in m["countries"]})==6,"six distinct country totals")
    excluded_records=m["quality"]["orders"]["excluded_order_records"]
    record("order exclusions complete",len(excluded_records)==43 and len({r["order_id"] for r in excluded_records})==43 and sum(m["quality"]["orders"]["excluded_by_reason"].values())==len(excluded_records),"all selected latest rows excluded from the 2025 base have an ID and reason")
    fields=sorted(country_fields-{"country","margin"})
    monthly_by={}
    for r in m["monthly"]:
        a=monthly_by.setdefault(r["country"],{f:0 for f in fields})
        for f in fields:a[f]+=r[f]
        expected=None if r["net_sales_eur"]==0 else r["contribution_eur"]/r["net_sales_eur"]
        record(f"monthly margin {r['country']} {r['month']}",r["margin"] is None if expected is None else close(r["margin"],expected,1e-7),"contribution / net sales, null at zero denominator")
    for r in m["countries"]:
        for f in fields:
            record(f"roll-up {r['country']} {f}",close(monthly_by[r["country"]][f],r[f]),"monthly records equal country total")
        expected=None if r["net_sales_eur"]==0 else r["contribution_eur"]/r["net_sales_eur"]
        record(f"country margin {r['country']}",r["margin"] is None if expected is None else close(r["margin"],expected,1e-7),"contribution / net sales, null at zero denominator")
    totals={f:sum(r[f] for r in m["countries"]) for f in fields}
    record("total contribution identity",close(totals["net_sales_eur"]-totals["net_cogs_eur"]-totals["fulfillment_eur"],totals["contribution_eur"]),"net sales minus net COGS and fulfillment")
    record("FX monthly coverage",len(m["fx_monthly"])==24 and {r["currency"] for r in m["fx_monthly"]}=={"PLN","CZK"},"12 archived monthly means for PLN and 12 for CZK")
    record("FX rates positive",all(r["local_per_eur"]>0 for r in m["fx_monthly"]),"local units per EUR; positive rates")
    record("World Bank grid",len(m["market_context"])==18 and {r["year"] for r in m["market_context"]}=={2022,2023,2024},"six countries x three years; raw nulls retained")
    record("World Bank provenance",len(m["metadata"]["market_context_sources"])==2 and all(s["url"] and s["response_lastupdated"] for s in m["metadata"]["market_context_sources"]),"series URLs, units, archive retrieval and vintage in metrics metadata/source register")

    scen=m["hub_scenarios"]
    record("scenario row coverage",len(scen)==68 and {r["scenario"] for r in scen}=={"low","base","high","stress"},"17 feasible hub choices x four cases")
    for r in scen:
        C=money(r["baseline_contribution_eur"]); F=money(r["annual_fixed_eur"]); S=money(r["saving_eur_per_unit_total"]); capex=r["capex_eur"]
        if r["scenario"]=="stress":
            exp=(C-money(r["return_shock_eur"])-money(r["fx_shock_eur"]))*Decimal("1.25")-C+S*Decimal("1.25")-F
        else:
            u=money(r["volume_uplift"]); exp=C*u+S*(Decimal(1)+u)-F
        expected_contribution=cents(exp)
        record(f"scenario formula {r['country']} {r['scenario']}",expected_contribution==money(r["incremental_contribution_eur"]),"policy formula recalculated from saved baseline and assumptions, rounded half-up to cents")
        expected=Decimal(capex)/money(r["incremental_contribution_eur"]) if r["incremental_contribution_eur"]>0 else None
        expected_payback=expected.quantize(Decimal("0.01"),rounding=ROUND_HALF_UP) if expected is not None else None
        actual_payback=money(r["payback_years"]) if r["payback_years"] is not None else None
        record(f"scenario payback {r['country']} {r['scenario']}",expected_payback==actual_payback,"capex / positive annual increment rounded to 0.01 years; otherwise null")
        record(f"scenario feasibility {r['country']}",capex<=450000 and r["fte"]<=7,"all reported hub choices fit stated capex/FTE ceilings")
    selected=m["recommendation"]
    ranked=selected["all_feasible_choices"]
    eligible=[x for x in ranked if x["countries"] and x["stress_eur"]>0]
    expected_choice=max(eligible,key=lambda x:(x["base_eur"],x["stress_eur"]))["choice"]
    record("recommendation selection",selected["choice"]==expected_choice and selected["base_incremental_contribution_eur"]>0 and selected["stress_incremental_contribution_eur"]>0,"highest base case among feasible choices with positive policy stress")
    record("budget limits",selected["capex_eur"]<=450000 and selected["fte"]<=7,"recommendation within board constraints")

    srcrows=list(csv.DictReader((DEL/"source_register.csv").open(encoding="utf-8")))
    for row in srcrows:
        p=SRC/row["file"]
        record(f"source hash {row['file']}",p.exists() and sha(p)==row["sha256"],"saved raw source matches register SHA-256")
    archive_checks=json.loads((ANALYSIS/"source_hash_checks.json").read_text(encoding="utf-8"))
    record("archive source hashes",all(x["match"] for x in archive_checks),"World Bank/ECB archive hashes match frozen register")

    wb=load_workbook(DEL/"Meridian_Analysis.xlsx",read_only=False,data_only=True)
    expected_sheets={"Start here","Sources","Quality & controls","FX 2025","Monthly","Countries","Market context","Scenarios","Decision comparison","Monthly trend","Orders audit","Return exceptions","Excluded orders"}
    record("workbook tabs",expected_sheets.issubset(set(wb.sheetnames)),", ".join(wb.sheetnames))
    record("workbook rows",wb["Monthly"].max_row==73 and wb["Countries"].max_row==7 and wb["Orders audit"].max_row==1255 and wb["Scenarios"].max_row==71 and wb["Excluded orders"].max_row==44,"72 monthly rows, 6 country rows, 1,254 eligible orders, 43 excluded orders, 68 scenario rows")
    record("workbook charts",len(wb["Countries"]._charts)>=1 and len(wb["Decision comparison"]._charts)>=1 and len(wb["Monthly trend"]._charts)>=1,"native Excel charts exist on country, decision and monthly sheets")
    for i,r in enumerate(m["countries"],start=2):
        cell=wb["Countries"].cell(i,9).value
        record(f"workbook country contribution {r['country']}",close(cell,r["contribution_eur"]),"country total matches metrics.json")
    for i,r in enumerate(m["monthly"],start=2):
        cell=wb["Monthly"].cell(i,10).value
        record(f"workbook monthly contribution {r['country']} {r['month']}",close(cell,r["contribution_eur"]),"monthly value matches metrics.json")
    for i,r in enumerate(m["hub_scenarios"],start=4):
        row=wb["Scenarios"].iter_rows(min_row=i,max_row=i,min_col=1,max_col=7,values_only=True).__next__()
        record(f"workbook scenario {r['country']} {r['scenario']}",row[1]==r["country"] and row[2]==r["scenario"] and close(row[3],r["incremental_contribution_eur"]) and row[4]==r["capex_eur"] and row[5]==r["fte"] and ((row[6] is None and r["payback_years"] is None) or close(row[6],r["payback_years"])),"workbook choice/scenario/cash/staff/payback matches metrics.json")

    report=PdfReader(str(DEL/"meridian_executive_report.pdf")); deck=PdfReader(str(DEL/"meridian_board_presentation.pdf"))
    report_text="\n".join(p.extract_text() or "" for p in report.pages)
    deck_text="\n".join(p.extract_text() or "" for p in deck.pages)
    record("report PDF structure",len(report.pages)==8 and "CZE + ESP" in report_text and "World Bank GDP per capita response" in report_text,"8 pages; recommendation and direct source register included")
    record("board PDF structure",len(deck.pages)==6 and "CZE + ESP" in deck_text and "90-day plan" in deck_text,"6 pages; recommendation and implementation slide included")
    record("rendered contact sheets",(ANALYSIS/"render_checks"/"report_contact.png").exists() and (ANALYSIS/"render_checks"/"presentation_contact.png").exists(),"Report and presentation rendered to page contact sheets for visual review")
    out={"generated_utc":datetime.now(timezone.utc).isoformat(),"passed":len(checks),"failed":0,"checks":checks,
         "summary":{"monthly_rows":len(m["monthly"]),"country_rows":len(m["countries"]),"fx_rows":len(m["fx_monthly"]),"market_rows":len(m["market_context"]),"scenario_rows":len(scen),"order_ledger_rows":1254,"excluded_orders":len(excluded_records),"report_pages":len(report.pages),"presentation_pages":len(deck.pages),"workbook_sheets":len(wb.sheetnames)}}
    (DEL/"verification.json").write_text(json.dumps(out,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(json.dumps(out["summary"]|{"passed_checks":out["passed"],"failed_checks":out["failed"]},indent=2))


if __name__=="__main__": main()
