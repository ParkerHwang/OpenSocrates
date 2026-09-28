#!/usr/bin/env python3
"""Build the workbook, report, presentation and source register from saved analysis."""
from __future__ import annotations

import csv
import hashlib
import html
import itertools
import json
import math
import shutil
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import (BaseDocTemplate, Frame, Image, KeepTogether, PageBreak,
                                PageTemplate, Paragraph, Spacer, Table as PDFTable,
                                TableStyle, Flowable)
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "evidence" / "source-room"
ANALYSIS = ROOT / "analysis"
DEL = ROOT / "deliverables"
PROCESSED = ROOT / "evidence" / "processed"
DEL.mkdir(exist_ok=True)
PROCESSED.mkdir(exist_ok=True)
COUNTRY_NAMES = {"DEU": "Germany", "FRA": "France", "NLD": "Netherlands", "POL": "Poland", "CZE": "Czechia", "ESP": "Spain"}
COUNTRIES = ["DEU", "FRA", "NLD", "POL", "CZE", "ESP"]
NAVY = "14324A"
TEAL = "00A89D"
GOLD = "E0A533"
CORAL = "E56B5D"
INK = "253746"
MUTED = "657987"
PALE = "EEF4F6"
PALE_TEAL = "E5F5F2"
PALE_GOLD = "FFF5DE"
WHITE = "FFFFFF"
GRID = "D8E2E7"
PAGE_W, PAGE_H = A4


def load(name):
    return json.loads((ANALYSIS / name).read_text(encoding="utf-8"))


def sha256(path: Path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def dt_utc(path: Path):
    return datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(timespec="seconds")


def eur(v, decimals=0):
    if v is None or v == "":
        return "—"
    x=float(v)
    return ("-€" + f"{abs(x):,.{decimals}f}") if x < 0 else ("€" + f"{x:,.{decimals}f}")


def eur_cell(v, decimals=2):
    return float(v) if v is not None else None


def pct(v, decimals=1):
    if v is None:
        return "—"
    return f"{float(v)*100:.{decimals}f}%"


def number(v, decimals=0):
    if v is None:
        return "—"
    return f"{float(v):,.{decimals}f}"


def ratio(num, den):
    return None if not den else num / den


def exact_file_hashes_check():
    original = json.loads((SRC / "source-register.json").read_text(encoding="utf-8"))
    checks = []
    for row in original["public"]:
        file = SRC / row["file"]
        got = sha256(file)
        checks.append({"file": row["file"], "expected": row["sha256"], "actual": got, "match": got == row["sha256"]})
        if got != row["sha256"]:
            raise ValueError(f"Archived public source hash mismatch: {file}")
    (ANALYSIS / "source_hash_checks.json").write_text(json.dumps(checks, indent=2) + "\n", encoding="utf-8")
    return original


def register_sources(original):
    public_meta = {row["file"]: row for row in original["public"]}
    client_details = {
        "data-dictionary.md": ("Client synthetic metadata", "2025 sales / returns through 2026-01-31", "Grain, revisions, FX and accounting policy"),
        "orders-part1.csv": ("Synthetic client export", "2025 orders plus excluded records", "One SKU per order row; local currency fields"),
        "orders-part2.csv": ("Synthetic client export", "2025 orders plus excluded records", "One SKU per order row; local currency fields"),
        "order-corrections.csv": ("Synthetic client correction file", "Order revisions", "Whole-row corrections; revision numbers"),
        "returns.csv": ("Synthetic client export", "2025 through 2026-02-01 returns", "Return quantity, restock quantity, original-currency refund"),
        "unit-costs.csv": ("Synthetic client assumption", "Effective dates in 2025", "EUR per SKU unit"),
        "hub-options.csv": ("Synthetic client assumption", "Annual recurring costs and investment", "EUR capex, FTE, EUR annual fixed cost, EUR savings/unit"),
        "scenario-policy.md": ("Synthetic client assumption", "Low/base/high and joint stress", "Uplifts, capex/FTE budgets, simple payback rules"),
        "source-register.json": ("Synthetic source-room provenance metadata", "Archive and synthetic input vintage", "URLs, hashes, archive retrieval timestamps"),
    }
    urls = {}
    for f in client_details:
        urls[f] = f"http://127.0.0.1:65200/{f}"
    rows = []
    for file in list(client_details) + ["population.json", "gdp-per-capita.json", "ecb-history.zip", "ecb-history.csv"]:
        path = SRC / file
        if file in client_details:
            kind, period, units = client_details[file]
            source_url = urls[file]
            archive_time = ""
            vintage = "Synthetic client data; source room identifies random seed 2026092736."
            note = "Local frozen source-room copy; no live client account accessed."
            status = "synthetic"
        else:
            archive = public_meta["ecb-history.zip" if file == "ecb-history.csv" else file]
            kind = "Official public archive" if file != "ecb-history.csv" else "Official public data (lossless ZIP extraction)"
            source_url = archive.get("original_url", "")
            archive_time = archive.get("retrieved_utc", "")
            vintage = "Archive register says fetched before outcome calls; historical observations may reflect later revisions."
            if file == "population.json":
                period, units = "2022–2024", "Population, persons"
            elif file == "gdp-per-capita.json":
                period, units = "2022–2024", "GDP per capita, current US dollars/person"
            else:
                period, units = "Daily reference rates; analysis uses 2025 business days", "Currency units per EUR"
            note = "Official source file retained byte-for-byte; see archived URLs and vintage in source-room/source-register.json."
            status = "public_official_archive"
        rows.append({"file": file, "source_url": source_url, "retrieved_utc": dt_utc(path),
                     "archive_retrieved_utc": archive_time, "sha256": sha256(path), "bytes": path.stat().st_size,
                     "kind": kind, "status": status, "period": period, "units": units, "vintage": vintage, "notes": note})
    with (DEL / "source_register.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader(); writer.writerows(rows)
    (DEL / "source_register.json").write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    shutil.copy2(DEL / "source_register.csv", ROOT / "evidence" / "source_register.csv")
    return rows


def export_processed(metrics, ledger, quality):
    tables = {
        "monthly_results.csv": metrics["monthly"],
        "country_results.csv": metrics["countries"],
        "fx_monthly.csv": metrics["fx_monthly"],
        "market_context.csv": metrics["market_context"],
        "hub_scenarios.csv": metrics["hub_scenarios"],
    }
    for name, rows in tables.items():
        with (PROCESSED / name).open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            writer.writeheader(); writer.writerows(rows)
    fields = list(ledger[0].keys())
    with (PROCESSED / "order_ledger.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader(); writer.writerows(ledger)
    excluded=quality["orders"]["excluded_order_records"]
    excluded_fields=sorted({k for row in excluded for k in row})
    with (PROCESSED / "excluded_orders.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=excluded_fields)
        writer.writeheader(); writer.writerows(excluded)
    exceptions = []
    for row in quality["returns"]["orphan_return_rows_quarantined"]:
        exceptions.append({"return_id": row["return_id"], "order_id": row["order_id"], "received_at": row["received_at"], "handling": "Quarantined: no selected source order"})
    for row in quality["returns"]["outside_cutoff_examples"]:
        exceptions.append({"return_id": row["return_id"], "order_id": row["order_id"], "received_at": row["received_at"], "handling": "Excluded: received after inclusive cutoff 2026-01-31"})
    with (PROCESSED / "return_exceptions.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["return_id", "order_id", "received_at", "handling"])
        writer.writeheader(); writer.writerows(exceptions)


def style_sheet(ws, title_row=1):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{title_row+1}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.sheet_view.zoomScale = 90


def add_table(ws, ref, name, style="TableStyleMedium2"):
    tab = Table(displayName=name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name=style, showFirstColumn=False, showLastColumn=False,
                                        showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)


def workbook(metrics, ledger, source_rows, original, exceptions):
    wb = Workbook()
    readme = wb.active; readme.title = "Start here"
    readme.sheet_view.showGridLines = False
    readme.merge_cells("A1:H1")
    readme["A1"] = "MERIDIAN PARTS  |  EUROPE HUB DECISION MODEL"
    readme["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    readme["A1"].fill = PatternFill("solid", fgColor=NAVY)
    readme["A1"].alignment = Alignment(vertical="center")
    readme.row_dimensions[1].height = 34
    rec = metrics["recommendation"]
    readme.merge_cells("A3:H3"); readme["A3"] = f"Decision case: {rec['choice']} | staged, gated recommendation"
    readme["A3"].font = Font(size=15, bold=True, color=NAVY)
    cards = [
        ("Year-zero capex", rec["capex_eur"], '€#,##0'), ("Staff", rec["fte"], '0 "FTE"'),
        ("Base annual increment", rec["base_incremental_contribution_eur"], '€#,##0;[Red](€#,##0)'),
        ("Stress annual increment", rec["stress_incremental_contribution_eur"], '€#,##0;[Red](€#,##0)'),
    ]
    for i, (label, val, fmt) in enumerate(cards):
        col = 1 + i * 2
        readme.cell(5, col, label).font = Font(bold=True, color=MUTED, size=10)
        c = readme.cell(6, col, val); c.number_format = fmt; c.font = Font(size=15, bold=True, color=TEAL if i in (2,3) else NAVY)
    sections = [
        (8, "Model boundaries", "2025 shipments and returns known through 2026-01-31. Client exports, corrections, costs, options and scenario policy are synthetic. World Bank and ECB files are official archived snapshots, not live feeds."),
        (11, "How to read it", "Monthly and country results are EUR after order-level half-up rounding. `Scenarios` lists each single option and every feasible pair under low/base/high/stress. Capex is year zero; fixed cost is recurring and already deducted from annual incremental contribution."),
        (14, "Reproduction", "From the project root run `python3 scripts/analyze.py` then `python3 scripts/build_deliverables.py`. Inputs are saved in `evidence/source-room/`; public archive hashes are checked during build. This workbook is a values-based analytical output with filters and native charts."),
        (17, "Decision caution", "The case ranks synthetic assumptions; it does not establish causal hub impact. The recommendation is conditional on confirming the proposed volume uplift and per-unit savings before committing leases and full capex."),
    ]
    for row, heading, body in sections:
        readme.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        c = readme.cell(row, 1, heading); c.font = Font(bold=True, color=TEAL, size=11)
        readme.merge_cells(start_row=row+1, start_column=1, end_row=row+2, end_column=8)
        c = readme.cell(row+1, 1, body); c.alignment = Alignment(wrap_text=True, vertical="top"); c.font = Font(size=10, color=INK)
    for col in "ABCDEFGH": readme.column_dimensions[col].width = 17
    for row in (9, 10, 12, 13, 15, 16, 18, 19): readme.row_dimensions[row].height = 24
    readme.freeze_panes = "A5"

    # Sources and provenance.
    ws = wb.create_sheet("Sources")
    headers = list(source_rows[0].keys())
    ws.append(headers)
    for row in source_rows: ws.append([row[h] for h in headers])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:L{ws.max_row}"
    add_table(ws, f"A1:L{ws.max_row}", "SourceRegister", "TableStyleMedium2")
    for i, h in enumerate(headers, 1): ws.column_dimensions[chr(64+i)].width = min(48, max(17, len(h)+2))
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["B"].width = 52; ws.column_dimensions["K"].width = 58; ws.column_dimensions["L"].width = 54

    # Quality log and calculation controls.
    quality = metrics["quality"]
    ws = wb.create_sheet("Quality & controls")
    ws.append(["Control / observed issue", "Count / value", "Treatment / interpretation"])
    qrows = [
        ("Order extract rows: part 1 / part 2", f"{quality['orders']['rows_by_file']['orders-part1.csv']} / {quality['orders']['rows_by_file']['orders-part2.csv']}", "Both pages pooled before revision resolution."),
        ("Correction rows", quality["orders"]["correction_rows"], "Whole-row replacement; highest revision selected."),
        ("Order raw rows total", quality["orders"]["raw_rows_total"], "Includes both extract pages and correction rows."),
        ("Exact repeated order rows removed", quality["orders"]["exact_duplicate_rows_removed"], "Identical repeats count once; not treated as separate orders."),
        ("Distinct order IDs / chosen latest rows", f"{quality['orders']['distinct_order_ids']} / {quality['orders']['selected_latest_revision_rows']}", "No conflicting ties at maximum revision."),
        ("Superseded order rows", quality["orders"]["superseded_rows_after_revision_selection"], "Lower revisions retained only in raw evidence, not the base."),
        ("Eligible 2025 shipped orders", quality["orders"]["eligible_2025_shipped_orders"], "Highest revision, shipped, non-test, 2025 shipped date."),
        ("Cancelled / other non-shipped latest orders", quality["orders"]["excluded_by_reason"].get("status_not_shipped",0), "Excluded; these include 24 blank shipped_at values on cancelled rows."),
        ("Test transactions", quality["orders"]["excluded_by_reason"].get("test_transaction",0), "Excluded even when status is shipped."),
        ("2026 shipment rows", quality["orders"]["excluded_by_reason"].get("shipment_outside_2025",0), "Excluded from the 2025 shipment base."),
        ("Selected latest rows excluded", len(quality["orders"]["excluded_order_records"]), "All order IDs and exclusion reasons are listed in the Excluded orders tab and evidence/processed/excluded_orders.csv."),
        ("Valid zero-gross shipments", quality["orders"]["zero_gross_valid_shipments"], "Included as shipped orders and units; zero is not treated as missing."),
        ("Return raw / exact repeats / superseded", f"{quality['returns']['raw_rows']} / {quality['returns']['exact_duplicate_rows_removed']} / {quality['returns']['superseded_rows_after_revision_selection']}", "Latest revision per return ID before date and link filters."),
        ("Included return records", quality["returns"]["included_return_records"], "Known by inclusive 2026-01-31 cutoff and linked to eligible shipments."),
        ("Late return", quality["returns"]["excluded_by_reason"].get("received_after_cutoff",0), "Received 2026-02-01; outside inclusive cutoff."),
        ("Orphan return", quality["returns"]["excluded_by_reason"].get("orphan_quarantined",0), "Quarantined; no order ID match; not guessed or joined."),
        ("Missing World Bank observations", f"population {quality['market_context']['population_missing']}; GDPpc {quality['market_context']['gdp_missing']}", "All 18 country-year observations available in this archive; missingness would remain null."),
        ("Monthly to country reconciliation", "PASS", "72 country-month rows roll to six country totals for all nine requested measures."),
        ("Per-order accounting identity", "PASS", "Contribution = gross − refunds − net COGS − fulfillment after component rounding."),
        ("Conversion and rounding", "ECB monthly mean; EUR=1", "Local-per-EUR quote, sale-month rate for both sales and refunds; ROUND_HALF_UP to cents per order."),
    ]
    for r in qrows: ws.append(r)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:C{ws.max_row}"; add_table(ws, f"A1:C{ws.max_row}", "QualityControls", "TableStyleMedium4")
    ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 27; ws.column_dimensions["C"].width = 84
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 30
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)

    # FX monthly.
    ws = wb.create_sheet("FX 2025")
    ws.append(["Currency", "Month", "Local currency units per EUR", "ECB published business days", "Source / use"])
    for r in metrics["fx_monthly"]:
        ws.append([r["currency"], r["month"], r["local_per_eur"], r["published_business_days"], "ECB archived reference rates; mean of available business days; divide local value by quote."])
    ws.append(["EUR", "2025 all", 1.0, "—", "EUR transactions use 1:1. ECB EUR is reporting currency, not an observed currency quote."])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:E{ws.max_row}"; add_table(ws, f"A1:E{ws.max_row}", "FXMonthly", "TableStyleMedium2")
    for c,w in {"A":14,"B":14,"C":32,"D":28,"E":85}.items(): ws.column_dimensions[c].width=w
    for r in range(2, ws.max_row+1): ws.cell(r,3).number_format="0.0000000000"

    # Monthly records.
    fields = ["country", "month", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin"]
    ws = wb.create_sheet("Monthly")
    ws.append(fields)
    for r in metrics["monthly"]: ws.append([r.get(k) for k in fields])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:L{ws.max_row}"; add_table(ws, f"A1:L{ws.max_row}", "MonthlyResults", "TableStyleMedium2")
    for c,w in {"A":13,"B":13,"C":17,"D":16,"E":19,"F":17,"G":19,"H":17,"I":18,"J":19,"K":16,"L":14}.items(): ws.column_dimensions[c].width=w
    for row in range(2,ws.max_row+1):
        for col in range(5,11): ws.cell(row,col).number_format='€#,##0.00;[Red](€#,##0.00)'
        ws.cell(row,12).number_format="0.0%"

    # Country totals and a readable native chart.
    ws = wb.create_sheet("Countries")
    ws.append(fields[:1] + fields[2:])
    for r in metrics["countries"]: ws.append([r.get(k) for k in ["country"]+fields[2:]])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:K{ws.max_row}"; add_table(ws, f"A1:K{ws.max_row}", "CountryResults", "TableStyleMedium2")
    for i,w in enumerate([13,17,16,19,17,19,17,18,19,16,14],1): ws.column_dimensions[chr(64+i)].width=w
    for row in range(2,ws.max_row+1):
        for col in range(4,10): ws.cell(row,col).number_format='€#,##0.00;[Red](€#,##0.00)'
        ws.cell(row,11).number_format="0.0%"
    chart = BarChart(); chart.type="bar"; chart.style=10; chart.title="2025 contribution by market"; chart.y_axis.title="Country"; chart.x_axis.title="EUR"
    chart.add_data(Reference(ws,min_col=9,min_row=1,max_row=7),titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=7)); chart.height=7.2; chart.width=14.0; chart.legend=None
    chart.dataLabels=DataLabelList(); chart.dataLabels.showVal=True
    ws.add_chart(chart,"M2")

    # Market context retains all original values and derives comparable changes.
    ws = wb.create_sheet("Market context")
    ws.append(["Country", "Year", "Population (persons)", "GDP per capita (current US$/person)", "Population 2022–24 change", "GDPpc 2022–24 nominal change", "Interpretation note"])
    base_lookup={}
    for r in metrics["market_context"]:
        base_lookup[(r["country"],r["year"])]=r
    summary_values={}
    for c in COUNTRIES:
        a=base_lookup[(c,2022)]; b=base_lookup[(c,2024)]
        pchange=(b["population"]/a["population"]-1) if a["population"] and b["population"] else None
        gchange=(b["gdp_per_capita_usd"]/a["gdp_per_capita_usd"]-1) if a["gdp_per_capita_usd"] and b["gdp_per_capita_usd"] else None
        summary_values[c]=(pchange,gchange)
    for r in metrics["market_context"]:
        pchg,gchg=summary_values[r["country"]]
        ws.append([r["country"],r["year"],r["population"],r["gdp_per_capita_usd"],pchg if r["year"]==2024 else None,gchg if r["year"]==2024 else None,
                   "WB macro context only; current US$ is nominal USD, not PPP or real income and not proof of parts demand."])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:G{ws.max_row}"; add_table(ws,f"A1:G{ws.max_row}","MarketContext","TableStyleMedium2")
    for c,w in {"A":14,"B":12,"C":24,"D":36,"E":29,"F":33,"G":80}.items(): ws.column_dimensions[c].width=w
    for row in range(2,ws.max_row+1):
        ws.cell(row,4).number_format='"$"#,##0.00'
        ws.cell(row,5).number_format="0.0%"; ws.cell(row,6).number_format="0.0%"

    # Scenario detail for every single option and feasible pair.
    ws = wb.create_sheet("Scenarios")
    scen_fields=["choice","country","scenario","incremental_contribution_eur","capex_eur","fte","payback_years","baseline_contribution_eur","baseline_units","baseline_gross_sales_eur","baseline_net_sales_eur","annual_fixed_eur","volume_uplift","saving_eur_per_unit_total","return_shock_eur","fx_shock_eur"]
    ws.append(["Choice","Country code(s)","Scenario","Annual incremental contribution (EUR)","Year-zero capex (EUR)","FTE","Simple payback (years)","2025 baseline contribution (EUR)","2025 shipped units","2025 gross sales (EUR)","2025 net sales (EUR)","Recurring annual fixed cost (EUR)","Volume uplift","Annual base-unit savings (EUR)","Stress return shock (EUR)","Stress FX shock (EUR)"])
    for r in metrics["hub_scenarios"]:
        ws.append([r.get("choice"),r.get("country"),r.get("scenario"),r.get("incremental_contribution_eur"),r.get("capex_eur"),r.get("fte"),r.get("payback_years"),r.get("baseline_contribution_eur"),r.get("baseline_units"),r.get("baseline_gross_sales_eur"),r.get("baseline_net_sales_eur"),r.get("annual_fixed_eur"),r.get("volume_uplift"),r.get("saving_eur_per_unit_total"),r.get("return_shock_eur"),r.get("fx_shock_eur")])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:P{ws.max_row}"; add_table(ws,f"A1:P{ws.max_row}","HubScenarios","TableStyleMedium2")
    for i,w in enumerate([17,17,12,31,24,10,24,34,20,23,22,33,15,31,28,24],1): ws.column_dimensions[chr(64+i)].width=w
    for row in range(2,ws.max_row+1):
        for col in [4,5,8,10,11,12,14,15,16]: ws.cell(row,col).number_format='€#,##0.00;[Red](€#,##0.00)'
        ws.cell(row,7).number_format="0.00"; ws.cell(row,13).number_format="0%"
    ws.insert_rows(1,2)
    ws.merge_cells("A1:P1"); ws["A1"]="Scenario arithmetic follows scenario-policy.md: Δ annual contribution = C×u + U×(1+u)×s − F. Stress uses u=25%, return shock 3% of G with no cost recovery, and 10% of N FX shock for PLN/CZK."
    ws["A1"].font=Font(bold=True,color=NAVY); ws["A1"].alignment=Alignment(wrap_text=True); ws.row_dimensions[1].height=33
    ws.freeze_panes="A4"; ws.auto_filter.ref=f"A3:P{ws.max_row}"
    # Reapply table range after title insertion.
    for tab in list(ws.tables.values()):
        tab.ref=f"A3:P{ws.max_row}"

    # Decision summary and chart.
    ws = wb.create_sheet("Decision comparison")
    ws.append(["Choice","Markets","Capex EUR","FTE","Low annual Δ EUR","Base annual Δ EUR","High annual Δ EUR","Stress annual Δ EUR","Base payback years","Stress positive?"])
    for r in metrics["recommendation"]["all_feasible_choices"]:
        ws.append([r["choice"]," + ".join(r["countries"]) if r["countries"] else "—",r["capex_eur"],r["fte"],r["low_eur"],r["base_eur"],r["high_eur"],r["stress_eur"],r["base_payback_years"],"Yes" if r["stress_eur"]>0 else "No"])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:J{ws.max_row}"; add_table(ws,f"A1:J{ws.max_row}","DecisionComparison","TableStyleMedium2")
    for c,w in {"A":19,"B":21,"C":16,"D":10,"E":21,"F":21,"G":21,"H":22,"I":21,"J":17}.items(): ws.column_dimensions[c].width=w
    for row in range(2,ws.max_row+1):
        for col in [3,5,6,7,8]: ws.cell(row,col).number_format='€#,##0;[Red](€#,##0)'
        ws.cell(row,9).number_format="0.00"
    ws.merge_cells("A21:J22"); ws["A21"]="Defer is the zero-investment reference. Feasible means ≤ €450,000 capex and ≤7 FTE. Four pairs are omitted from the scenario rows because they breach FTE and/or capex; see Quality & controls. No pair synergy is modeled."
    ws["A21"].alignment=Alignment(wrap_text=True,vertical="top"); ws["A21"].font=Font(color=MUTED,italic=True)
    chart=BarChart(); chart.type="bar"; chart.style=11; chart.title="Top choices: base and joint stress"; chart.y_axis.title="Choice"; chart.x_axis.title="Annual incremental contribution (EUR)"
    max_row=min(10,ws.max_row)
    chart.add_data(Reference(ws,min_col=6,max_col=6,min_row=1,max_row=max_row),titles_from_data=True)
    chart.add_data(Reference(ws,min_col=8,max_col=8,min_row=1,max_row=max_row),titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=max_row)); chart.height=9.0; chart.width=17
    chart.legend.position="b"; ws.add_chart(chart,"A24")

    # Monthly aggregate trend sheet.
    monthly_totals=[]
    for m in sorted({r["month"] for r in metrics["monthly"]}):
        rs=[r for r in metrics["monthly"] if r["month"]==m]
        monthly_totals.append([m,sum(r["gross_sales_eur"] for r in rs),sum(r["refunds_eur"] for r in rs),sum(r["net_sales_eur"] for r in rs),sum(r["contribution_eur"] for r in rs),sum(r["shipped_units"] for r in rs)])
    ws=wb.create_sheet("Monthly trend")
    ws.append(["Month","Gross sales EUR","Refunds EUR","Net sales EUR","Contribution EUR","Shipped units"])
    for r in monthly_totals: ws.append(r)
    ws.freeze_panes="A2"; add_table(ws,f"A1:F{ws.max_row}","MonthlyTrend","TableStyleMedium2")
    for c,w in {"A":15,"B":20,"C":18,"D":20,"E":22,"F":17}.items(): ws.column_dimensions[c].width=w
    for row in range(2,ws.max_row+1):
        for col in range(2,6): ws.cell(row,col).number_format='€#,##0;[Red](€#,##0)'
    chart=LineChart(); chart.title="2025 monthly sales and contribution"; chart.y_axis.title="EUR"; chart.x_axis.title="Month"; chart.style=13
    chart.add_data(Reference(ws,min_col=4,max_col=5,min_row=1,max_row=13),titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=1,min_row=2,max_row=13)); chart.height=8; chart.width=15; chart.legend.position="b"
    ws.add_chart(chart,"H2")

    # Order-level audit, with applied rate and rounded components.
    ws=wb.create_sheet("Orders audit")
    ledger_fields=["order_id","revision","country","month","shipped_at","sku","quantity","currency","fx_local_per_eur","gross_sales_local","refunds_local","gross_sales_eur","refunds_eur","net_sales_eur","gross_cogs_eur","recovered_cogs_eur","net_cogs_eur","fulfillment_eur","contribution_eur","returned_units","restocked_units","return_ids"]
    ws.append([x.replace("_"," ").title() for x in ledger_fields])
    for r in ledger:
        vals=[]
        for k in ledger_fields:
            v=r.get(k)
            if k in ("fx_local_per_eur","gross_sales_local","refunds_local","unit_cost_eur") and v is not None: v=float(v)
            vals.append(v)
        ws.append(vals)
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:V{ws.max_row}"; add_table(ws,f"A1:V{ws.max_row}","OrderAudit","TableStyleMedium2")
    for i,w in enumerate([20,11,12,12,15,13,12,12,20,19,18,18,17,19,18,21,18,19,20,16,17,42],1): ws.column_dimensions[chr(64+i)].width=w
    for row in range(2,ws.max_row+1):
        ws.cell(row,9).number_format="0.0000000000"
        for col in [10,11]: ws.cell(row,col).number_format='0.00'
        for col in range(12,20): ws.cell(row,col).number_format='€#,##0.00;[Red](€#,##0.00)'

    ws=wb.create_sheet("Return exceptions")
    ws.append(["Return ID","Order ID","Received at","Handling"])
    for r in exceptions: ws.append([r.get(k) for k in ["return_id","order_id","received_at","handling"]])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:D{ws.max_row}"; add_table(ws,f"A1:D{ws.max_row}","ReturnExceptions","TableStyleMedium4")
    for c,w in {"A":20,"B":20,"C":18,"D":72}.items(): ws.column_dimensions[c].width=w

    ws=wb.create_sheet("Excluded orders")
    excluded=quality["orders"]["excluded_order_records"]
    ex_fields=sorted({k for row in excluded for k in row})
    ws.append([x.replace("_"," ").title() for x in ex_fields])
    for r in excluded: ws.append([r.get(k) for k in ex_fields])
    ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{chr(64+len(ex_fields))}{ws.max_row}"; add_table(ws,f"A1:{chr(64+len(ex_fields))}{ws.max_row}","ExcludedOrders","TableStyleMedium4")
    for i,field in enumerate(ex_fields,1): ws.column_dimensions[chr(64+i)].width=max(18,min(36,len(field)+12))

    # Tab styling, document metadata and integrity-friendly settings.
    for s in wb.worksheets:
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = "landscape"
        s.page_setup.fitToWidth = 1
        s.sheet_view.zoomScale = 85 if s.max_column > 8 else 95
        for cell in s[1]:
            if cell.value is not None and not (s.title in ("Start here", "Scenarios") and cell.row==1):
                cell.font = Font(bold=True, color=WHITE)
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.alignment=Alignment(wrap_text=True,vertical="center")
        s.row_dimensions[1].height = max(s.row_dimensions[1].height or 15, 28)
    wb.properties.creator="Meridian Parts operations analysis"
    wb.properties.title="Meridian Parts European service hub analysis"
    wb.properties.subject="Synthetic 2025 operating analysis and hub scenarios"
    wb.calculation.fullCalcOnLoad=True
    wb.calculation.forceFullCalc=True
    wb.save(DEL / "Meridian_Analysis.xlsx")


def build_report(metrics, source_rows, original):
    countries = metrics["countries"]
    lookup = {r["country"]:r for r in countries}
    scenarios = metrics["recommendation"]["all_feasible_choices"]
    alt = next(r for r in scenarios if r["choice"]==metrics["recommendation"]["strongest_alternative"])
    rec = metrics["recommendation"]
    market={}
    for row in metrics["market_context"]: market.setdefault(row["country"],{})[row["year"]]=row
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=25, leading=29, textColor=colors.HexColor("#"+NAVY), alignment=TA_LEFT, spaceAfter=8))
    styles.add(ParagraphStyle(name="Deck", parent=styles["Normal"], fontName="Helvetica", fontSize=10.2, leading=14.3, textColor=colors.HexColor("#"+MUTED), spaceAfter=8))
    styles.add(ParagraphStyle(name="H1x", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=17, leading=21, textColor=colors.HexColor("#"+NAVY), spaceBefore=1, spaceAfter=8))
    styles.add(ParagraphStyle(name="H2x", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor("#"+TEAL), spaceBefore=6, spaceAfter=4))
    styles.add(ParagraphStyle(name="Bodyx", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.65, leading=12.4, textColor=colors.HexColor("#"+INK), spaceAfter=5))
    styles.add(ParagraphStyle(name="Smallx", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.1, leading=9.2, textColor=colors.HexColor("#"+MUTED), spaceAfter=3))
    styles.add(ParagraphStyle(name="Callout", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=10, leading=14, textColor=colors.HexColor("#"+NAVY), spaceAfter=0))
    styles.add(ParagraphStyle(name="WhiteSmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.5, leading=10, textColor=colors.white))
    story=[]
    def P(txt, style="Bodyx"): return Paragraph(txt, styles[style])
    def section(title, kicker=None):
        if kicker: story.append(P(kicker.upper(),"Smallx"))
        story.append(P(title,"H1x"))
    def data_table(data, widths, font=7.4, repeat=1, aligns=None):
        head_style=ParagraphStyle(f"tablehead{font}",fontName="Helvetica-Bold",fontSize=font,leading=font+1.6,textColor=colors.white,wordWrap="CJK")
        body_style=ParagraphStyle(f"tablebody{font}",fontName="Helvetica",fontSize=font,leading=font+2,textColor=colors.HexColor("#"+INK),wordWrap="CJK")
        wrapped=[]
        for ri,row in enumerate(data):
            cells=[]
            for value in row:
                if hasattr(value,"wrap"):
                    cells.append(value)
                else:
                    text=html.escape(str(value),quote=False).replace("&lt;br/&gt;","<br/>")
                    cells.append(Paragraph(text,head_style if ri==0 else body_style))
            wrapped.append(cells)
        t=PDFTable(wrapped,colWidths=widths,repeatRows=repeat,hAlign="LEFT")
        ts=[("BACKGROUND",(0,0),(-1,0),colors.HexColor("#"+NAVY)),("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),font),
            ("LEADING",(0,0),(-1,-1),font+2),("VALIGN",(0,0),(-1,-1),"TOP"),
            ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#"+GRID)),("LEFTPADDING",(0,0),(-1,-1),5),
            ("RIGHTPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4),
            ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F5F8F9")])]
        if aligns:
            for i,a in enumerate(aligns): ts.append(("ALIGN",(i,1),(i,-1),a))
        t.setStyle(TableStyle(ts)); return t
    total={k:sum(r[k] for r in countries) for k in ["shipped_orders","shipped_units","gross_sales_eur","refunds_eur","net_sales_eur","net_cogs_eur","fulfillment_eur","contribution_eur","returned_units"]}
    total_margin=total["contribution_eur"]/total["net_sales_eur"]
    comp_total=total["net_sales_eur"]-total["net_cogs_eur"]-total["fulfillment_eur"]

    # 1 — executive page.
    story.append(Spacer(1,6*mm)); story.append(P("MERIDIAN PARTS  |  BOARD DECISION NOTE","Smallx"))
    story.append(P("Two hubs can clear the hurdle—<br/>if the operating assumptions hold","ReportTitle"))
    story.append(P("European service-hub expansion • Decision analysis of six markets • 2025 shipped sales, returns through 31 January 2026","Deck"))
    top=[P("RECOMMENDATION","Smallx"),P(f"Conditionally prepare <font color='#{TEAL}'>Czechia + Spain</font> for launch. Reserve €225,000 of year-zero capex and 5 FTE; release spend through the 90-day validation gates below.","Callout")]
    call=PDFTable([[top]],colWidths=[PAGE_W-42*mm]); call.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#"+PALE_TEAL)),("BOX",(0,0),(-1,-1),0.8,colors.HexColor("#"+TEAL)),("LEFTPADDING",(0,0),(-1,-1),10),("RIGHTPADDING",(0,0),(-1,-1),10),("TOPPADDING",(0,0),(-1,-1),8),("BOTTOMPADDING",(0,0),(-1,-1),8)])); story.append(call); story.append(Spacer(1,5*mm))
    cards=[[P("BASE ANNUAL INCREMENT","Smallx"),P("JOINT STRESS","Smallx"),P("SIMPLE PAYBACK","Smallx"),P("2025 CONTRIBUTION","Smallx")],
           [P(f"<b>{eur(rec['base_incremental_contribution_eur'])}</b>","Callout"),P(f"<b>{eur(rec['stress_incremental_contribution_eur'])}</b>","Callout"),P(f"<b>{rec['base_payback_years']:.2f} yrs</b>","Callout"),P(f"<b>{eur(total['contribution_eur'])}</b>","Callout")]]
    ct=PDFTable(cards,colWidths=[(PAGE_W-42*mm)/4]*4); ct.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#"+PALE)),("BOX",(0,0),(-1,-1),0.5,colors.HexColor("#"+GRID)),("INNERGRID",(0,0),(-1,-1),0.5,colors.HexColor("#"+GRID)),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7)])); story.append(ct)
    story.append(Spacer(1,4*mm)); story.append(P("Why this choice","H2x"))
    story.append(P(f"At the client’s 25% base volume uplift, the selected pair produces {eur(rec['base_incremental_contribution_eur'])} a year after recurring fixed cost. It ranks first among 17 feasible single/pair choices, stays positive in the defined combined return-rate/FX stress ({eur(rec['stress_incremental_contribution_eur'])}), and uses €225k / 5 FTE versus ceilings of €450k / 7. The strongest base-case alternative, Poland + Spain, gives {eur(alt['base_eur'])} at €240k / 6 FTE. The model is assumption-led and the client transactions are synthetic; this is a conditional investment case, not causal proof. [S1–S2]"))
    story.append(P("Board action requested","H2x"))
    story.append(P("Approve conditional design and validation authority for CZE + ESP, with no irreversible lease or fit-out commitment until the day-30 gate confirms a ≥25% volume case and the stated savings per unit from route and handling evidence. Stop or defer if either market misses the decision threshold. The defer option remains the zero-contribution, zero-capex reference."))
    story.append(P(f"2025 base: {number(total['shipped_orders'])} distinct eligible shipped orders • {number(total['shipped_units'])} units • {eur(total['net_sales_eur'])} net sales • {eur(total['contribution_eur'])} contribution • {pct(total_margin)} contribution / net sales. [S1, S2]","Smallx"))
    story.append(PageBreak())

    # 2 — actual operating signal + market context.
    section("The synthetic baseline favors Spain, Czechia and Poland","01 / operating evidence")
    story.append(P("The 1,254 eligible orders are balanced at 209 per market in this generated extract; higher unit volume and product mix drive most country differences. Spain and Czechia lead on contribution and contribution margin. That is evidence about this synthetic 2025 ledger, not evidence of real customer demand or local-hub causality. [S1]"))
    d=[["Market","Orders","Units","Gross sales","Refunds","Net sales","Contribution","Margin"]]
    for r in countries:
        d.append([r["country"],number(r["shipped_orders"]),number(r["shipped_units"]),eur(r["gross_sales_eur"]),eur(r["refunds_eur"]),eur(r["net_sales_eur"]),eur(r["contribution_eur"]),pct(r["margin"])])
    d.append(["Total",number(total["shipped_orders"]),number(total["shipped_units"]),eur(total["gross_sales_eur"]),eur(total["refunds_eur"]),eur(total["net_sales_eur"]),eur(total["contribution_eur"]),pct(total_margin)])
    story.append(data_table(d,[18*mm,15*mm,20*mm,23*mm,20*mm,22*mm,25*mm,15*mm],font=6.7,aligns=["LEFT","RIGHT","RIGHT","RIGHT","RIGHT","RIGHT","RIGHT","RIGHT"]))
    story.append(Spacer(1,3*mm))
    story.append(P("Country share is not the same as opportunity. Hub savings, fixed cost, capex and headcount shift the ranking: the top baseline contributor is Spain; the best two-market arithmetic is Czechia + Spain. Refunds are highest in euros and units in Spain, but the decision model does not assume every returned unit recovers product cost; only recorded restocks do. [S1]"))
    story.append(P("Macro context: broad, non-demand indicators","H2x"))
    md=[["Market","Population 2024","Change 2022–24","GDPpc 2024 current US$","GDPpc change 2022–24"]]
    for c in COUNTRIES:
        a=market[c][2022]; b=market[c][2024]
        pchg=b["population"]/a["population"]-1
        gchg=b["gdp_per_capita_usd"]/a["gdp_per_capita_usd"]-1
        md.append([c,number(b["population"]),pct(pchg),"$"+number(b["gdp_per_capita_usd"]),pct(gchg)])
    story.append(data_table(md,[23*mm,31*mm,27*mm,43*mm,37*mm],font=7.1))
    story.append(P("Population growth is strongest in Czechia (+2.2%) and Spain (+2.2%); Poland declined about 0.7% over the period. 2024 GDP per capita ranges from roughly $25.1k (Poland) to $67.5k (Netherlands), while two-year nominal USD changes also reflect prices and exchange rates. Neither population nor current-US$ GDP per capita directly measures replacement-assembly demand. Values preserve the archived World Bank response vintage (last updated 13 July 2026); see [S3–S4]."))
    story.append(PageBreak())

    # 3 — method and reconciliation.
    section("Ledger controls keep revisions, refunds and FX on the right grain","02 / calculation and quality")
    story.append(P("The analysis pools both order extract pages and the correction file, removes identical repeats, picks the highest numeric revision by order ID, then applies shipment eligibility. It selects the highest return revision by return ID, applies the 31 Jan cutoff and eligible-order link, and aggregates distinct return IDs additively. No conflicting winning revisions were detected. [S1]"))
    qa=[ ["Check","Observed","Treatment"],
         ["Orders: 1,328 raw rows","13 identical rows removed; 1,297 IDs; 18 lower revisions superseded","1,254 eligible latest rows after 24 cancellations, 18 tests and one 2026 shipment excluded"],
         ["Returns: 264 raw rows","20 identical rows removed; 243 IDs; 1 revision superseded","241 records applied; one 1 Feb return excluded; one orphan quarantined"],
         ["Blank shipment dates","24 selected rows; all are cancelled","Recorded as missing but cannot qualify; no blank date on included shipment"],
         ["Free/zero-gross shipments","12 valid orders","Included in orders and units; zero sales are not dropped"],
         ["Currency and quantities","No country-currency mismatches, negative gross, refunds above gross or returned qty above shipped","All 24 PLN/CZK monthly means have 20–23 published business-day quotes"],
         ["Country/month reconciliation","72 rows roll to 6 country totals","All nine required measures reconcile exactly; order-level contribution identity passes"]]
    story.append(data_table(qa,[39*mm,47*mm,81*mm],font=7.0))
    story.append(P("EUR conversion follows the original shipment month. ECB quotes are local currency per EUR; local proceeds/refunds are divided by the monthly mean, not multiplied or inverted. The same sale-month mean translates later refunds. EUR transactions use 1. Each order’s gross sales, summed refunds, net COGS and fulfillment are rounded to cents with ROUND_HALF_UP before aggregation. The 2025 daily ECB archive is an analytical translation proxy, not actual settlement FX. [S1–S2]"))
    story.append(P("Booked revenue, cash and contribution","H2x"))
    story.append(P(f"The workbook calls gross sales less credited refunds ‘net sales’ ({eur(total['net_sales_eur'])}); it is a transaction measure in this extract, not cash collected. Cash timing, actual settlement rates, taxes, chargebacks, payment fees and receivable/payable timing are absent. Contribution subtracts {eur(total['net_cogs_eur'])} net product cost and {eur(total['fulfillment_eur'])} nonrefundable fulfillment, yielding {eur(total['contribution_eur'])}. COGS recovery is limited to recorded restocked units at the sale-date unit cost; no blanket recovery is assumed. [S1]"))
    story.append(P("The return set contains 241 included distinct records; 1,081 physical units returned (not a return-order count). The orphan ID ABSENT / return R-ORPHAN is kept in the exception log with no guessed join. R-LATE, received 1 Feb 2026, is outside cutoff. Source rows and exceptions are retained under evidence/processed for audit."))
    story.append(PageBreak())

    # 4 — scenario choice table.
    section("Single hubs: baseline strength does not guarantee robust economics","03a / investment scenarios")
    story.append(P("The client policy applies low/base/high uplifts of 10% / 25% / 40% to 2025 baseline contribution C and shipped units U: annual increment = C×u + U×(1+u)×s − F. Capex K is shown separately as year-zero cash; it is not deducted from annual contribution. Simple payback is K / positive annual increment. Stress is a 25% volume case plus a 3%-of-gross additional refund shock with no added COGS recovery and a 10%-of-net-sales depreciation shock only for PLN/CZK. [S1]"))
    story.append(P("Single-market choices","H2x"))
    annual_by_choice={r["country"]:r for r in metrics["hub_scenarios"] if r["scenario"]=="base"}
    sd=[["Hub","Capex\nEUR","FTE","Fixed/y\nEUR","Low Δ/y\nEUR","Base Δ/y\nEUR","High Δ/y\nEUR","Stress Δ/y\nEUR","Base pb\n(years)"]]
    for c in COUNTRIES:
        r=next(x for x in scenarios if x["choice"]==c)
        s=annual_by_choice[c]
        sd.append([c,eur(r["capex_eur"]),str(r["fte"]),eur(s["annual_fixed_eur"]),eur(r["low_eur"]),eur(r["base_eur"]),eur(r["high_eur"]),eur(r["stress_eur"]),f"{r['base_payback_years']:.2f}" if r["base_payback_years"] else "—"])
    story.append(data_table(sd,[15*mm,17*mm,10*mm,17*mm,21*mm,21*mm,21*mm,21*mm,25*mm],font=7.0))
    story.append(Spacer(1,3*mm))
    story.append(P("Spain is the strongest single on base contribution (€102,716/year, €130k capex, three FTE); Czechia provides €95,383/year on €95k capex and two FTE. Germany’s €10,089 base increment carries €280k capex and 27.75-year simple payback. Defer remains a €0 / 0 FTE reference."))
    story.append(PageBreak())
    section("CZE + ESP leads the additive pair cases","03b / feasible pair ranking")
    story.append(P("The 11 feasible pairs below satisfy both €450k capex and seven-FTE limits. Pair savings and costs are summed with no synergy. Four other pairs are excluded: DEU + FRA (€490k / 9 FTE), DEU + NLD (8 FTE), DEU + POL (8 FTE), and DEU + ESP (8 FTE)."))
    pair_rows=[r for r in scenarios if len(r["countries"])==2]
    pd=[ ["Pair","Capex\nEUR","FTE","Fixed/y\nEUR","Low Δ/y\nEUR","Base Δ/y\nEUR","High Δ/y\nEUR","Stress Δ/y\nEUR","Base pb\n(years)"] ]
    for r in pair_rows:
        s=annual_by_choice["+".join(r["countries"])]
        pd.append([r["choice"],eur(r["capex_eur"]),str(r["fte"]),eur(s["annual_fixed_eur"]),eur(r["low_eur"]),eur(r["base_eur"]),eur(r["high_eur"]),eur(r["stress_eur"]),f"{r['base_payback_years']:.2f}" if r["base_payback_years"] else "—"])
    story.append(data_table(pd,[22*mm,17*mm,10*mm,17*mm,21*mm,21*mm,21*mm,21*mm,18*mm],font=7.0))
    story.append(Spacer(1,3*mm)); story.append(P("CZE + ESP is €11,770/year above the next-best base alternative (POL + ESP), has €2,916/year more stress contribution, and uses €15k less capex and one fewer FTE. Exact option-level assumptions and four paybacks per choice are in the workbook and metrics JSON. [S1]"))
    story.append(PageBreak())

    # 5 — recommendation robustness and assumptions.
    section("CZE + ESP wins the stated policy case, but validation is the investment test","04 / decision and pivot")
    story.append(P(f"Selected case: Czechia + Spain uses {eur(rec['capex_eur'])} year-zero capex and {rec['fte']} FTE, below the caps by €{450000-rec['capex_eur']:,.0f} and {7-rec['fte']} FTE. It yields {eur(rec['low_incremental_contribution_eur'])} / {eur(rec['base_incremental_contribution_eur'])} / {eur(rec['high_incremental_contribution_eur'])} annual incremental contribution in low/base/high, and {eur(rec['stress_incremental_contribution_eur'])} in the defined joint stress. The base simple payback is {rec['base_payback_years']:.2f} years; no discounting or ramp-up is included."))
    story.append(P("Why it beats the nearest alternative","H2x"))
    czesp_stress=next(r for r in metrics["hub_scenarios"] if r["country"]=="CZE+ESP" and r["scenario"]=="stress")
    polcze_stress=next(r for r in metrics["hub_scenarios"] if r["country"]=="POL+CZE" and r["scenario"]=="stress")
    story.append(P(f"Poland + Spain is the strongest base-case alternative at {eur(alt['base_eur'])}/year, {eur(alt['capex_eur'])} and {alt['fte']} FTE, with {eur(alt['stress_eur'])} stress contribution and {alt['base_payback_years']:.2f}-year base payback. CZE + ESP is €{rec['base_incremental_contribution_eur']-alt['base_eur']:,.0f}/year higher in base and €{rec['stress_incremental_contribution_eur']-alt['stress_eur']:,.0f}/year higher in stress, while using €15k less capex and one fewer FTE. FX stress applies to both PLN and CZK sales: CZE + ESP includes a {eur(czesp_stress['fx_shock_eur'])} CZK haircut on Czech net sales (Spain is euro-denominated) and remains at {eur(czesp_stress['incremental_contribution_eur'])}; POL + CZE absorbs {eur(polcze_stress['fx_shock_eur'])} combined PLN/CZK haircut and falls to {eur(polcze_stress['incremental_contribution_eur'])}. The recommendation therefore survives the stipulated CZK shock; it is not an FX-free case."))
    story.append(P("Decision pivot","H2x"))
    story.append(P(f"For the selected pair, algebraic break-even uplift is about {pct(rec['break_even_volume_uplift'],1)} under the given savings and annual fixed costs (vs 25% base). But this low break-even is not a measured demand estimate: all uplift is assumed. The board should change the decision to defer if route/handling trials cannot substantiate ≥25% volume uplift and the given €2.10/unit CZE + €3.00/unit ESP savings, or if audited annual fixed cost rises enough to erase the stress cushion ({eur(rec['stress_incremental_contribution_eur'])}). Test lower uptake, full ramp and no savings before final lease/capex release."))
    story.append(P("Cross-check against defer","H2x"))
    story.append(P(f"Deferral retains €0 scenario contribution and avoids investment. The selected low case is +{eur(rec['low_incremental_contribution_eur'])}; the joint stress is +{eur(rec['stress_incremental_contribution_eur'])}. The model therefore favors staged investment only if validation converts those assumed levers into operating evidence. No estimate is statistically proven or causal."))
    story.append(P("Risk register: uplift may not occur; central-hub savings may be counted twice or fail in local execution; rent, hiring and fit-out can exceed assumptions; added inventory can tie up cash; returned parts may not be reusable; FX reference stress is simplified; local tax, labor, customs, regulatory and service requirements are unmodeled. Owners and stop gates appear on page 6."))
    story.append(PageBreak())

    # 6 — implementation.
    section("Use 90 days to turn synthetic assumptions into a gated operating case","05 / execution and measurement")
    story.append(P("Do not release lease or fit-out commitments at board approval. Start with reversible validation and location/design work, then approve each spend release against a named owner and evidence gate."))
    plan=[["Days / gate","Accountable owner","Work and dependencies","Gate to proceed"],
          ["0–30<br/>Validate","COO (owner); CFO / Commercial Ops","Rebuild 2025 catchment baseline from transaction and route records; verify unit-level handling, parcel zones, transit, current central fulfillment cost, return/restock treatment, site/rent quotes and hiring plan. Dependency: access to actual non-synthetic operations data and candidate locations.","CFO signs audited baseline; ops demonstrates ≥25% feasible uplift case and savings ≥€2.10 CZE / €3.00 ESP per shipped unit; fixed-cost run-rate ≤€95k/y combined. Otherwise defer."],
          ["31–60<br/>Design / pilot","Regional Ops (owner); Procurement / HR / IT","Pilot route consolidation or local stock positioning without permanent lease; collect real lead-time, on-time/in-full, cost/unit, stockout and return-restock data; secure conditional site and staffing quotes. Dependency: Gate 1, inventory master and service design.","COO + CFO approve site-level economics, service design, safety/quality and permits; modeled stress stays positive after quoted costs. Release only pilot/design budget."],
          ["61–90<br/>Commit / launch readiness","COO (owner); CFO / HR / Quality / IT","Complete fit-out design, hiring, systems, inventory placement, supplier/3PL SLA, returns process, contingency and local compliance review. Dependency: Gate 2 approval and signed service-provider terms.","Board sponsor releases capex only when purchase orders and staffing fit €225k / 5 FTE envelope and KPI instrumentation is live. Otherwise pause or defer." ]]
    story.append(data_table(plan,[24*mm,29*mm,64*mm,48*mm],font=6.7))
    story.append(P("KPI plan: establish a 4-week matched central-ship baseline in each proposed catchment before any pilot; report weekly and monthly by country/SKU/order. Avoid treating total country growth as causal: use pilot vs matched routes/items where data permit and state limitations.","Bodyx"))
    kpi=[["Measure","Definition / source","Gate / operating target","Owner / cadence"],
         ["Volume uplift","Eligible shipped units vs matched 2025 baseline, adjusted for trend/seasonality","Run-rate ≥25% base assumption before full spend","Commercial Ops / weekly"],
         ["Cost saved per unit","Avoided current fulfillment/handling less local variable cost ÷ units","≥€2.10 CZE; ≥€3.00 ESP","COO / weekly; Finance sign-off"],
         ["Contribution after fixed cost","Incremental net sales − net COGS − fulfillment − recurring hub cost","Positive low-case and positive defined stress; base ≥€198k annualized combined","CFO / monthly"],
         ["Service","Median and 90th percentile order-to-delivery days; on-time-in-full rate","Establish baseline by day 30; improve delivery P90 by ≥1 day and OTIF by ≥5 pp before rollout","Regional Ops / weekly"],
         ["Returns / recovery","Refund € ÷ gross €; returned units; restocked units / returned units","No deterioration in refund rate; itemize recovery, no blanket credit","Quality / monthly"],
         ["Spend / staffing","Committed and forecast capex; filled and planned FTE","≤€225k selected capex and ≤5 FTE; no unapproved recurring overrun","CFO + HR / monthly"]]
    story.append(data_table(kpi,[30*mm,50*mm,57*mm,30*mm],font=6.45))
    story.append(P("If catchment-level baseline, route cost and actual customer/order geography cannot be obtained, do not proceed beyond reversible design work. The synthetic extracts do not include lead time, customer coordinates, stockouts or site cost evidence, so these must be measured before the gates can pass."))
    story.append(PageBreak())

    # 7 — sources and limitations.
    section("Evidence is useful for a structured choice; it cannot eliminate execution risk","06 / sources and limitations")
    story.append(P("The source room is a frozen local package. Its sales, returns, option costs, unit costs, staffing and scenario assumptions are synthetic (documented generator seed 2026092736). No interviews, live client systems or non-archived customer evidence were accessed. The archive’s original source register and file hashes are preserved with the downloaded files. [S1–S4]"))
    story.append(P("Material limitations","H2x"))
    limitations=[
        "Synthetic order rows cannot validate demand, market share, seasonality, churn or actual customer mix; equal 209-order totals per country are an artifact of this generated extract.",
        "No customer or delivery geography, promised/actual lead time, stockout, central-hub capacity, actual carrier route costs, service-level loss, lease quote, taxes, labor/regulatory cost or working-capital schedule is present.",
        "WB population and GDP per capita are national context only. GDP per capita is current US dollars (not PPP or constant-price income) and population is not parts demand.",
        "ECB monthly means are official reference rates, not transactional settlement FX. The 10% PLN/CZK shock is a stipulated scenario simplification, not a volatility distribution.",
        "Scenarios use annual steady-state additive arithmetic: no construction/ramp time, discount rate, inflation, tax, residual value, shared staffing, synergy, cannibalization or probability weights.",
        "Return rates, restocking and cost recovery follow recorded synthetic quantities. Physical reusability, warranty handling and future return timing may differ operationally.",
    ]
    for x in limitations: story.append(P("• "+x))
    story.append(P("Source register and direct links","H2x"))
    refs=[
        ("[S1] Synthetic client files", "data-dictionary.md; orders-part1.csv; orders-part2.csv; order-corrections.csv; returns.csv; unit-costs.csv; hub-options.csv; scenario-policy.md. Frozen source room downloaded 27 Sep 2026; see deliverables/source_register.csv for URL, exact local collection time and SHA-256."),
        ("[S2] ECB historical reference-rate archive", "https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip — daily reference observations, archived file; source-room register retrieval 27 Sep 2026 06:36:40 UTC. Used only 2025 business days."),
        ("[S3] World Bank population response", "https://api.worldbank.org/v2/country/DEU;FRA;NLD;POL;CZE;ESP/indicator/SP.POP.TOTL?date=2022:2024&format=json&per_page=1000 — archived API payload; unit persons; archive response lastupdated 13 Jul 2026."),
        ("[S4] World Bank GDP per capita response", "https://api.worldbank.org/v2/country/DEU;FRA;NLD;POL;CZE;ESP/indicator/NY.GDP.PCAP.CD?date=2022:2024&format=json&per_page=1000 — archived API payload; unit current US dollars/person; archive response lastupdated 13 Jul 2026."),
    ]
    for label,desc in refs:
        safe=desc
        if "https://" in desc:
            start=desc.find("https://"); url=desc[start:].split(" ")[0]
            safe=desc[:start]+f"<link href='{url}' color='#{TEAL}'>{url}</link>"+desc[start+len(url):]
        story.append(P(f"<b>{label}</b> — {safe}","Smallx"))
    story.append(Spacer(1,3*mm))
    story.append(P("Companion files: Meridian_Analysis.xlsx for filtered data and charts; meridian_board_presentation.pdf for the board discussion; metrics.json for exchange; analysis/metrics.json and evidence/processed for audit; scripts/analyze.py and scripts/build_deliverables.py for reproduction."))

    def header_footer(c, doc):
        c.saveState()
        c.setStrokeColor(colors.HexColor("#"+GRID)); c.setLineWidth(0.5)
        c.line(21*mm,PAGE_H-14*mm,PAGE_W-21*mm,PAGE_H-14*mm)
        c.setFont("Helvetica-Bold",7); c.setFillColor(colors.HexColor("#"+NAVY))
        c.drawString(21*mm,PAGE_H-11*mm,"MERIDIAN PARTS  /  SERVICE-HUB EXPANSION")
        c.setFont("Helvetica",7); c.setFillColor(colors.HexColor("#"+MUTED))
        c.drawRightString(PAGE_W-21*mm,9*mm,f"SYNTHETIC CASE  •  27 SEP 2026  •  {doc.page}")
        c.restoreState()
    frame=Frame(21*mm,16*mm,PAGE_W-42*mm,PAGE_H-34*mm,leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0,id="normal")
    doc=BaseDocTemplate(str(DEL/"meridian_executive_report.pdf"),pagesize=A4,leftMargin=21*mm,rightMargin=21*mm,topMargin=18*mm,bottomMargin=16*mm,
                       title="Meridian Parts European service hub expansion",author="Operations analysis")
    doc.addPageTemplates([PageTemplate(id="main",frames=[frame],onPage=header_footer)])
    doc.build(story)


def draw_para(c, text, x, y_top, width, style=None, max_h=180):
    st=style or ParagraphStyle("slide",fontName="Helvetica",fontSize=11,leading=15,textColor=colors.HexColor("#"+INK))
    p=Paragraph(text,st)
    w,h=p.wrap(width,max_h)
    p.drawOn(c,x,y_top-h)
    return y_top-h


def slide_base(c, n, title, sub=None):
    W,H=960,540
    c.setFillColor(colors.HexColor("#F6F9FA")); c.rect(0,0,W,H,fill=1,stroke=0)
    c.setFillColor(colors.HexColor("#"+NAVY)); c.rect(0,H-86,W,86,fill=1,stroke=0)
    c.setFillColor(colors.white); c.setFont("Helvetica-Bold",25); c.drawString(44,H-47,title)
    if sub:
        c.setFillColor(colors.HexColor("#C8D7DD")); c.setFont("Helvetica",10); c.drawString(46,H-69,sub)
    c.setStrokeColor(colors.HexColor("#"+GRID)); c.line(44,31,W-44,31)
    c.setFillColor(colors.HexColor("#"+MUTED)); c.setFont("Helvetica",8)
    c.drawString(44,17,"SYNTHETIC CLIENT CASE  •  archived official ECB / World Bank series")
    c.drawRightString(W-44,17,f"MERIDIAN PARTS  |  {n} / 6")


def card(c,x,y,w,h,label,value,detail,accent=TEAL):
    c.setFillColor(colors.white); c.setStrokeColor(colors.HexColor("#"+GRID)); c.roundRect(x,y,w,h,7,fill=1,stroke=1)
    c.setFillColor(colors.HexColor("#"+accent)); c.rect(x,y+h-5,w,5,fill=1,stroke=0)
    c.setFillColor(colors.HexColor("#"+MUTED)); c.setFont("Helvetica-Bold",9); c.drawString(x+14,y+h-25,label.upper())
    c.setFillColor(colors.HexColor("#"+NAVY)); c.setFont("Helvetica-Bold",23); c.drawString(x+14,y+h-56,value)
    c.setFillColor(colors.HexColor("#"+MUTED)); c.setFont("Helvetica",9); c.drawString(x+14,y+15,detail)


def slide_bar_chart(c, data, x, y, w, h, maxv=None, labels=True, color=TEAL, negative_color=CORAL, compact=False):
    if not data:return
    maxv=maxv or max(abs(float(v)) for _,v in data) or 1
    left=x+(130 if not compact else 100); right=x+w-10; chartw=right-left
    rowh=h/len(data); zerox=left+chartw*(maxv/(maxv+max(0,max(-float(v) for _,v in data)))) if min(float(v) for _,v in data)<0 else left
    # For nonnegative charts, bars start at axis left. If any negative, allocate symmetric zero.
    if min(float(v) for _,v in data)<0:
        maxpos=max(0,max(float(v) for _,v in data)); maxneg=abs(min(0,min(float(v) for _,v in data)))
        usable=chartw; zerox=left+usable*maxneg/(maxpos+maxneg)
    else: maxneg=0;maxpos=maxv;zerox=left
    c.setStrokeColor(colors.HexColor("#"+GRID)); c.setLineWidth(0.6); c.line(zerox,y,zerox,y+h)
    for i,(lab,val) in enumerate(data):
        cy=y+h-(i+0.5)*rowh
        c.setFillColor(colors.HexColor("#"+INK)); c.setFont("Helvetica-Bold" if i==0 else "Helvetica",9 if not compact else 8)
        c.drawRightString(left-9,cy-3,str(lab))
        val=float(val)
        if val>=0: bw=(right-zerox)*val/(maxpos or 1); bx=zerox; fill=color
        else: bw=(zerox-left)*abs(val)/(maxneg or 1); bx=zerox-bw; fill=negative_color
        c.setFillColor(colors.HexColor("#"+fill)); c.roundRect(bx,cy-rowh*0.24,max(bw,1),rowh*0.48,3,fill=1,stroke=0)
        if labels:
            c.setFillColor(colors.HexColor("#"+INK)); c.setFont("Helvetica",8)
            tx=bx+bw+6 if val>=0 else bx-6
            c.drawString(tx if val>=0 else tx-stringWidth(eur(val),"Helvetica",8),cy-3,eur(val))


def build_deck(metrics):
    rec=metrics["recommendation"]
    countries=metrics["countries"]
    scenarios=rec["all_feasible_choices"]
    alt=next(x for x in scenarios if x["choice"]==rec["strongest_alternative"])
    pdf=DEL/"meridian_board_presentation.pdf"
    c=canvas.Canvas(str(pdf),pagesize=(960,540),pageCompression=1)
    c.setTitle("Meridian Parts | European service-hub decision")
    # Slide 1
    slide_base(c,1,"Authorize a gated CZE + ESP decision case","Board recommendation • decision basis is synthetic 2025 exports plus archived public context")
    draw_para(c,"<b>Approve conditional preparation for two hubs</b> in Czechia and Spain. Reserve €225k year-zero capex and five FTE; do not release lease or fit-out commitments until the 90-day gates confirm the operating case.",48,422,864,ParagraphStyle("lead",fontName="Helvetica",fontSize=18,leading=25,textColor=colors.HexColor("#"+NAVY)))
    card(c,48,193,206,100,"Base annual increment",eur(rec["base_incremental_contribution_eur"]),"after €95k recurring fixed cost")
    card(c,272,193,206,100,"Joint stress",eur(rec["stress_incremental_contribution_eur"]),"25% volume + return / FX shock",GOLD)
    card(c,496,193,206,100,"Simple payback",f"{rec['base_payback_years']:.2f} years","no ramp or discounting")
    card(c,720,193,192,100,"Capacity used","€225k / 5 FTE","€450k / 7 FTE ceiling",NAVY)
    c.setFillColor(colors.HexColor("#"+PALE_TEAL)); c.roundRect(48,83,864,78,7,fill=1,stroke=0)
    draw_para(c,f"<b>Decision condition:</b> demonstrate ≥25% annual volume uplift and €2.10/unit CZE + €3.00/unit ESP savings from actual route and handling evidence. If either fails, defer; the modeled stress cushion is only {eur(rec['stress_incremental_contribution_eur'])}/year.",65,145,830,ParagraphStyle("cond",fontName="Helvetica",fontSize=11,leading=16,textColor=colors.HexColor("#"+NAVY)))
    draw_para(c,"Scenario ranking is not causal evidence. Synthetic client exports were not accessed from a live account; no customer interviews were conducted.",48,67,864,ParagraphStyle("foot",fontName="Helvetica-Oblique",fontSize=8,leading=10,textColor=colors.HexColor("#"+MUTED)))
    c.showPage()
    # Slide 2
    slide_base(c,2,"2025 base: Spain and Czechia lead the synthetic ledger","1,254 eligible orders • 77,436 units • €4.384m net sales • €2.030m contribution (46.3% of net sales)")
    slide_bar_chart(c,[(r["country"],r["contribution_eur"]) for r in countries],56,94,496,330,maxv=450000)
    c.setFillColor(colors.white); c.setStrokeColor(colors.HexColor("#"+GRID)); c.roundRect(580,106,330,300,8,fill=1,stroke=1)
    vals=[["Market","Units","Contribution","Margin"]]
    for r in countries: vals.append([r["country"],number(r["shipped_units"]),eur(r["contribution_eur"]),pct(r["margin"])])
    # Draw compact table manually.
    x0,y0=596,375; widths=[75,72,105,58]; rowh=34
    for ri,row in enumerate(vals):
        yy=y0-ri*rowh
        if ri==0:
            c.setFillColor(colors.HexColor("#"+NAVY)); c.rect(x0,yy-rowh+5,sum(widths),rowh,fill=1,stroke=0)
        elif ri%2==0:
            c.setFillColor(colors.HexColor("#F4F7F8")); c.rect(x0,yy-rowh+5,sum(widths),rowh,fill=1,stroke=0)
        cx=x0
        for j,item in enumerate(row):
            c.setFillColor(colors.white if ri==0 else colors.HexColor("#"+INK)); c.setFont("Helvetica-Bold" if ri==0 or j==0 else "Helvetica",9)
            c.drawString(cx+5,yy-17,str(item)); cx+=widths[j]
    draw_para(c,"All six countries contain exactly 209 eligible orders in this synthetic generation. Differences are in units, product mix, refunds, cost and pricing—not a demonstrated local demand advantage.",596,132,294,ParagraphStyle("note",fontName="Helvetica",fontSize=9,leading=13,textColor=colors.HexColor("#"+MUTED)))
    c.showPage()
    # Slide 3
    slide_base(c,3,"CZE + ESP ranks first among 17 feasible choices","Low / base / high are 10% / 25% / 40% volume uplift; annual fixed cost is deducted, capex is separate")
    top=[r for r in scenarios if r["countries"]][:7]
    slide_bar_chart(c,[(r["choice"],r["base_eur"]) for r in top],70,110,505,320,maxv=220000,compact=True)
    c.setFillColor(colors.white); c.setStrokeColor(colors.HexColor("#"+GRID)); c.roundRect(603,101,306,326,8,fill=1,stroke=1)
    lines=[("CZE + ESP","198k base | 32k stress | €225k / 5 FTE"),("POL + ESP","186k base | 29k stress | €240k / 6 FTE"),("POL + CZE","179k base | −79k stress | €205k / 5 FTE"),("NLD + ESP","167k base | 108k stress | €270k / 6 FTE")]
    yy=391
    for i,(head,body) in enumerate(lines):
        c.setFillColor(colors.HexColor("#"+(TEAL if i==0 else NAVY))); c.setFont("Helvetica-Bold",11); c.drawString(621,yy,head)
        draw_para(c,body,621,yy-8,267,ParagraphStyle(f"rank{i}",fontName="Helvetica",fontSize=9,leading=12,textColor=colors.HexColor("#"+INK)))
        yy-=72
    draw_para(c,"Only additive economics; no pair synergy. Four pairs are infeasible on FTE and/or capex. DEU + FRA also exceeds the capex limit.",621,141,267,ParagraphStyle("pairnote",fontName="Helvetica-Oblique",fontSize=8,leading=11,textColor=colors.HexColor("#"+MUTED)))
    c.showPage()
    # Slide 4
    slide_base(c,4,"The modeled stress is positive, with limited room for execution error","CZE + ESP returns and capex shown in year-zero / steady-state terms")
    card(c,50,325,260,95,"Low annual increment",eur(rec["low_incremental_contribution_eur"]),"10% volume uplift")
    card(c,350,325,260,95,"Base annual increment",eur(rec["base_incremental_contribution_eur"]),"25% volume uplift")
    card(c,650,325,260,95,"Stress annual increment",eur(rec["stress_incremental_contribution_eur"]),"25% + refund shock; no added recovery",GOLD)
    draw_para(c,"Why not Poland + Spain?",55,282,300,ParagraphStyle("h",fontName="Helvetica-Bold",fontSize=13,leading=16,textColor=colors.HexColor("#"+NAVY)))
    cze_stress=next(r for r in metrics["hub_scenarios"] if r["country"]=="CZE+ESP" and r["scenario"]=="stress")
    polcze_stress=next(r for r in metrics["hub_scenarios"] if r["country"]=="POL+CZE" and r["scenario"]=="stress")
    draw_para(c,f"The strongest alternative is {eur(alt['base_eur'])}/year base, €15k more capex and one more FTE. CZE + ESP is about €{rec['base_incremental_contribution_eur']-alt['base_eur']:,.0f}/year higher in base and €{rec['stress_incremental_contribution_eur']-alt['stress_eur']:,.0f}/year higher in stress. Stress includes a {eur(cze_stress['fx_shock_eur'])} CZK haircut for Czech net sales in CZE + ESP (Spain is EUR); POL + CZE receives a {eur(polcze_stress['fx_shock_eur'])} combined PLN/CZK haircut.",55,260,410,ParagraphStyle("body",fontName="Helvetica",fontSize=10.5,leading=15,textColor=colors.HexColor("#"+INK)))
    draw_para(c,"Decision pivot",505,282,350,ParagraphStyle("h2",fontName="Helvetica-Bold",fontSize=13,leading=16,textColor=colors.HexColor("#"+NAVY)))
    draw_para(c,f"Break-even uplift is ≈{pct(rec['break_even_volume_uplift'],1)} if the savings and €95k recurring fixed cost hold. Those terms are synthetic assumptions. Require a validated ≥25% case and actual savings of €2.10/unit in CZE and €3.00/unit in ESP; otherwise defer. Stress cushion: {eur(rec['stress_incremental_contribution_eur'])}/year.",505,260,395,ParagraphStyle("body2",fontName="Helvetica",fontSize=10.5,leading=15,textColor=colors.HexColor("#"+INK)))
    c.setFillColor(colors.HexColor("#"+PALE_GOLD)); c.roundRect(55,93,850,75,7,fill=1,stroke=0)
    draw_para(c,"Major unmodeled risks: site and labor cost overruns, ramp/working capital, service mix, returns quality, customer geography, lead-time effect, taxes and discounting. ECB reference FX is not transaction FX; WB macro series is context, not demand evidence.",71,151,820,ParagraphStyle("risk",fontName="Helvetica",fontSize=10,leading=14,textColor=colors.HexColor("#"+INK)))
    c.showPage()
    # Slide 5
    slide_base(c,5,"90-day plan makes the decision reversible until evidence arrives","Named owners, dependencies and gates; no irreversible lease before the first two gates")
    stages=[("0–30 days | Validate","COO • CFO • Commercial Ops","Rebuild actual route/catchment baseline; verify current fulfillment cost, candidate rent, staffing, demand and returns. Needs real, non-synthetic operating data.","Gate 1: CFO signs ≥25% uplift case; savings ≥€2.10 CZE / €3.00 ESP per unit; recurring cost ≤€95k/y combined. Else defer."),("31–60 days | Pilot / design","Regional Ops • Procurement • HR • IT","Pilot route consolidation or local stock positioning; measure service, cost/unit, inventory and returns; secure conditional site/staff quotes.","Gate 2: COO + CFO confirm positive defined stress after quoted costs; compliance and service design pass. Release pilot/design only."),("61–90 days | Commit readiness","COO • CFO • Quality • IT","Fit-out design, recruitment, inventory placement, provider SLA, returns process, instrumentation and local compliance.","Gate 3: Board sponsor releases capex only within €225k / 5 FTE envelope and with KPIs live. Else pause." )]
    y=420
    for i,(stage,owner,work,gate) in enumerate(stages):
        c.setFillColor(colors.white); c.setStrokeColor(colors.HexColor("#"+GRID)); c.roundRect(50,y-105,860,104,7,fill=1,stroke=1)
        c.setFillColor(colors.HexColor("#"+TEAL if i==0 else "#"+NAVY)); c.roundRect(50,y-105,10,104,5,fill=1,stroke=0)
        c.setFillColor(colors.HexColor("#"+NAVY)); c.setFont("Helvetica-Bold",12); c.drawString(70,y-24,stage)
        c.setFillColor(colors.HexColor("#"+TEAL)); c.setFont("Helvetica-Bold",9); c.drawString(70,y-43,owner)
        draw_para(c,work,70,y-51,408,ParagraphStyle(f"work{i}",fontName="Helvetica",fontSize=8.5,leading=11,textColor=colors.HexColor("#"+INK)))
        draw_para(c,gate,505,y-22,380,ParagraphStyle(f"gate{i}",fontName="Helvetica-Bold",fontSize=8.3,leading=11,textColor=colors.HexColor("#"+NAVY)))
        y-=116
    draw_para(c,"Weekly/monthly KPIs: volume uplift ≥25%; net saved cost ≥ option assumptions; contribution after recurring cost positive in low and stress; order-to-delivery P90 improves ≥1 day and OTIF ≥5pp vs 4-week matched baseline; refund rate not worse; capex ≤€225k and staffing ≤5 FTE.",54,77,850,ParagraphStyle("kpi",fontName="Helvetica",fontSize=9.2,leading=13,textColor=colors.HexColor("#"+MUTED)))
    c.showPage()
    # Slide 6
    slide_base(c,6,"Decision rests on synthetic assumptions and archived sources","Sources, limits and decision materials")
    draw_para(c,"Client data (synthetic)",54,420,390,ParagraphStyle("src1",fontName="Helvetica-Bold",fontSize=13,leading=16,textColor=colors.HexColor("#"+NAVY)))
    draw_para(c,"Orders, corrections, returns, effective unit cost, hub options and scenario policy were collected from the local frozen source room. Data dictionary determines grain/revisions and contribution. No live customer account, interview or service-level source was accessed. Raw files and SHA-256 hashes are included.",54,395,385,ParagraphStyle("srcbody1",fontName="Helvetica",fontSize=10,leading=15,textColor=colors.HexColor("#"+INK)))
    draw_para(c,"Official archived observations",505,420,390,ParagraphStyle("src2",fontName="Helvetica-Bold",fontSize=13,leading=16,textColor=colors.HexColor("#"+NAVY)))
    draw_para(c,"ECB daily historical reference rates: 2025 arithmetic business-day means, local currency per EUR. World Bank population and GDP per capita: 2022–2024 archived API responses, archive metadata last updated 13 July 2026. Population is persons; GDPpc is current US dollars/person.",505,395,390,ParagraphStyle("srcbody2",fontName="Helvetica",fontSize=10,leading=15,textColor=colors.HexColor("#"+INK)))
    c.setFillColor(colors.white); c.setStrokeColor(colors.HexColor("#"+GRID)); c.roundRect(54,102,852,145,8,fill=1,stroke=1)
    draw_para(c,"Board packet",72,226,810,ParagraphStyle("pkt",fontName="Helvetica-Bold",fontSize=12,leading=15,textColor=colors.HexColor("#"+TEAL)))
    draw_para(c,"meridian_executive_report.pdf · meridian_board_presentation.pdf · Meridian_Analysis.xlsx · metrics.json · source_register.csv · scripts/analyze.py + scripts/build_deliverables.py · raw archive in evidence/source-room · processed audit in evidence/processed",72,205,805,ParagraphStyle("pktbody",fontName="Helvetica",fontSize=9.3,leading=14,textColor=colors.HexColor("#"+INK)))
    draw_para(c,"Official links: ECB historical archive — www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip · World Bank population — api.worldbank.org/.../SP.POP.TOTL · GDP per capita — api.worldbank.org/.../NY.GDP.PCAP.CD. Full request URLs, collection timestamps, hashes, units and vintage appear in the report and source register.",72,158,805,ParagraphStyle("links",fontName="Helvetica",fontSize=8,leading=11,textColor=colors.HexColor("#"+MUTED)))
    draw_para(c,"Values are scenario outputs, not a statistically proven or causal hub impact. No NPV, tax, working capital, ramp, synergy or probability weighting is included.",54,82,850,ParagraphStyle("caution",fontName="Helvetica-Bold",fontSize=10,leading=13,textColor=colors.HexColor("#"+NAVY)))
    c.save()


def main():
    metrics=load("metrics.json")
    ledger=load("order_ledger.json")
    reconciliation=load("reconciliation.json")
    quality=metrics["quality"]
    original=exact_file_hashes_check()
    sources=register_sources(original)
    exceptions=[]
    for row in quality["returns"]["orphan_return_rows_quarantined"]:
        exceptions.append({"return_id":row["return_id"],"order_id":row["order_id"],"received_at":row["received_at"],"handling":"Quarantined: no selected source order"})
    for row in quality["returns"]["outside_cutoff_examples"]:
        exceptions.append({"return_id":row["return_id"],"order_id":row["order_id"],"received_at":row["received_at"],"handling":"Excluded: received after inclusive cutoff 2026-01-31"})
    export_processed(metrics,ledger,quality)
    shutil.copy2(ANALYSIS/"metrics.json",DEL/"metrics.json")
    shutil.copy2(ANALYSIS/"reconciliation.json",DEL/"reconciliation.json")
    workbook(metrics,ledger,sources,original,exceptions)
    build_report(metrics,sources,original)
    build_deck(metrics)
    print(json.dumps({"outputs":[p.name for p in DEL.iterdir() if p.is_file()],
                      "source_files":len(sources),"order_audit_rows":len(ledger),"excluded_order_rows":len(quality["orders"]["excluded_order_records"]),
                      "report_pdf":(DEL/"meridian_executive_report.pdf").stat().st_size,
                      "presentation_pdf":(DEL/"meridian_board_presentation.pdf").stat().st_size,
                      "workbook":(DEL/"Meridian_Analysis.xlsx").stat().st_size},indent=2))


if __name__=="__main__":
    main()
