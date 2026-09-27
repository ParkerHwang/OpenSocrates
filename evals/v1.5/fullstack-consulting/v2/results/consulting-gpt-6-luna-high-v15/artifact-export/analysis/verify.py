import json,zipfile,sys,csv,hashlib
from pathlib import Path
from decimal import Decimal
from decimal import ROUND_HALF_UP
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation
R=Path(__file__).resolve().parents[1]; M=json.loads((R/'deliverables/metrics.json').read_text())
req=['monthly','countries','fx_monthly','market_context','hub_scenarios','recommendation','quality'];assert all(k in M for k in req)
fields=['shipped_orders','shipped_units','gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur','returned_units']
assert len(M['monthly'])==72 and len(M['countries'])==6 and len(M['fx_monthly'])==24 and len(M['market_context'])==18
for c in M['countries']:
 ms=[x for x in M['monthly'] if x['country']==c['country']]
 for k in fields:
  assert round(sum(x[k] for x in ms),2)==round(c[k],2),(c['country'],k)
 assert c['margin'] is None if c['net_sales_eur']==0 else abs(c['margin']-c['contribution_eur']/c['net_sales_eur'])<1e-7
sc=M['hub_scenarios']; assert len(sc)==68
assert all(x['payback_years'] is None if x['incremental_contribution_eur']<=0 else x['payback_years']>0 for x in sc)
assert M['recommendation']['countries']==['NLD','ESP']
assert M['recommendation']['capex_eur']<=450000 and M['recommendation']['fte']<=7
assert len(M['decision_comparison']['eligible_pairs'])==11
assert M['recommendation']['capex_eur']==270000 and M['recommendation']['fte']==6
# Recompute all scenarios independently from the saved baselines and client option assumptions.
options={}
with (R/'evidence/raw/hub-options.csv').open(newline='') as f:
 for row in csv.DictReader(f):options[row['country']]={'K':Decimal(row['capex_eur']),'F':Decimal(row['annual_fixed_eur']),'s':Decimal(row['saving_eur_per_unit']),'fte':int(row['fte'])}
base={x['country']:x for x in M['countries']}
for x in sc:
 cs=x['country'].split('+'); C0=sum((Decimal(str(base[c]['contribution_eur'])) for c in cs),Decimal(0)); U=sum(base[c]['shipped_units'] for c in cs); G=sum((Decimal(str(base[c]['gross_sales_eur'])) for c in cs),Decimal(0)); N=sum((Decimal(str(base[c]['net_sales_eur'])) for c in cs),Decimal(0))
 K=sum((options[c]['K'] for c in cs),Decimal(0)); F=sum((options[c]['F'] for c in cs),Decimal(0)); S=sum((options[c]['s'] for c in cs),Decimal(0)); fte=sum(options[c]['fte'] for c in cs)
 if x['scenario']=='stress':
  fx=Decimal('.10')*N if any(c in ('POL','CZE') for c in cs) else Decimal(0); increment=(C0-Decimal('.03')*G-fx)*Decimal('1.25')-C0+Decimal(U)*Decimal('1.25')*S-F
 else:
  u={'low':Decimal('.10'),'base':Decimal('.25'),'high':Decimal('.40')}[x['scenario']]; increment=C0*u+Decimal(U)*(1+u)*S-F
 expected=float(increment.quantize(Decimal('.01'),rounding=ROUND_HALF_UP)); assert abs(x['incremental_contribution_eur']-expected)<.005,(x,expected)
 assert x['capex_eur']==int(K) and x['fte']==fte
 expected_pay=round(float(K/increment),4) if increment>0 else None; assert x['payback_years']==expected_pay,(x,expected_pay)
q=M['quality'];assert q['raw_order_rows']==1328 and q['distinct_order_ids']==1297 and q['eligible_orders']==1254
assert q['raw_return_rows']==264 and q['distinct_return_ids']==243 and q['orphan_returns_count']==1
book=load_workbook(R/'deliverables/meridian_analysis.xlsx',read_only=False,data_only=False)
assert {'Monthly','Countries','Market context','FX monthly','Hub scenarios','Quality','Decision chart'}.issubset(book.sheetnames)
assert book['Monthly'].max_row==75 and book['Hub scenarios'].max_row==71
assert len(book['Decision chart']._charts)>=1
assert book['Decision chart'].cell(8,1).value=='Defer'
pdf=PdfReader(str(R/'deliverables/meridian_board_report.pdf'));assert len(pdf.pages)>=3
text=' '.join(p.extract_text() or '' for p in pdf.pages)
for s in ['Netherlands + Spain','Czechia + Spain','90-day','Population','contribution']:assert s.lower() in text.lower(),s
prs=Presentation(str(R/'deliverables/meridian_board_presentation.pptx'));assert len(prs.slides)>=6
deckpdf=PdfReader(str(R/'deliverables/meridian_board_presentation.pdf'));assert len(deckpdf.pages)==7
decktext=' '.join(p.extract_text() or '' for p in deckpdf.pages);assert 'Netherlands + Spain' in decktext and '90-day' in decktext
assert (R/'evidence/source-register.json').exists()
reg=json.loads((R/'evidence/source-register.json').read_text())['register']
for e in reg:
 p=R/e['file'];assert p.exists() and hashlib.sha256(p.read_bytes()).hexdigest()==e['sha256'],e['file']
print(json.dumps({'status':'PASS','monthly_rows':len(M['monthly']),'country_rows':len(M['countries']),'fx_rows':len(M['fx_monthly']),'market_rows':len(M['market_context']),'scenario_rows':len(sc),'feasible_pairs':11,'report_pdf_pages':len(pdf.pages),'presentation_pdf_pages':len(deckpdf.pages),'pptx_slides':len(prs.slides),'xlsx_sheets':book.sheetnames,'reconciliations':'all monthly additive metrics equal country totals','checks':['required JSON schema','monthly-to-country reconciliation','independent low/base/high/stress formula and payback recomputation','capex and staffing limits','quality counts','source SHA-256 hashes','workbook tabs/charts/defer alternative','PDF content extraction','presentation PDF slide count/content','PPTX structure']},indent=2))
