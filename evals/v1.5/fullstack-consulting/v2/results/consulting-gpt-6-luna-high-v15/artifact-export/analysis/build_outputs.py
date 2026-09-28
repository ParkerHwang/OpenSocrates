import json,csv,hashlib
import sys
from pathlib import Path
from datetime import datetime,timezone
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.chart import BarChart,Reference
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Image,KeepTogether
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.deps'))
import matplotlib.pyplot as plt
R=Path(__file__).resolve().parents[1]; RAW=R/'evidence/raw'; O=R/'deliverables'
M=json.loads((O/'metrics.json').read_text()); C=M['countries']; D=M['scenario_details']; names={'DEU':'Germany','FRA':'France','NLD':'Netherlands','POL':'Poland','CZE':'Czechia','ESP':'Spain'}
def get(pair,sc):return next(x for x in D if x['country']==pair and x['scenario']==sc)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
# Provenance register: archive metadata and extraction vintage stay visible.
orig=json.loads((RAW/'source-register.json').read_text()); pub={x['file']:x for x in orig['public']}; reg=[]
descriptions={
 'orders-part1.csv':('synthetic_client_export','2025 shipment/order rows','One order per order_id; local order currency; fulfillment EUR'),
 'orders-part2.csv':('synthetic_client_export','2025 shipment/order rows','One order per order_id; local order currency; fulfillment EUR'),
 'order-corrections.csv':('synthetic_client_corrections','2025 order revisions','Whole-row replacements; numeric revision'),
 'returns.csv':('synthetic_client_export','Returns through 2026-02 source range','Return ID; local refund currency; physical/restocked units'),
 'unit-costs.csv':('synthetic_client_assumption','Effective dates in 2025','EUR per SKU unit'),
 'hub-options.csv':('synthetic_client_assumption','Annual options','Capex EUR; fixed cost EUR/year; saving EUR/unit; FTE'),
 'scenario-policy.md':('synthetic_client_assumption','Low/base/high and stress scenarios','Volume uplifts, refund/FX shock, budget and staffing'),
 'data-dictionary.md':('synthetic_client_metadata','Analysis rules','Grain, revision, FX and accounting definitions'),
 'source-register.json':('source_room_metadata','Archive retrieval vintage','Original URLs, retrieval UTC, hashes, snapshot notes'),
 'source-room-index.html':('source_room_index','Frozen source-room file listing','Index metadata; source-room URL'),
 'ecb-history.csv':('official_public_archive_extraction','Daily rates through archive vintage','Currency units per EUR; extracted from ecb-history.zip')}
for p in sorted(RAW.iterdir()):
 if p.is_file():
  x={'file':'evidence/raw/'+p.name,'sha256':sha(p),'bytes':p.stat().st_size,'workspace_collected_utc':datetime.now(timezone.utc).isoformat()}
  if p.name in pub:
   x.update({k:v for k,v in pub[p.name].items() if k not in ('file','sha256','bytes')})
   x['source_file_name']=p.name
   if p.name in ('population.json','gdp-per-capita.json'):x['world_bank_lastupdated']=json.loads(p.read_text())[0].get('lastupdated')
  elif p.name in ('population.json','gdp-per-capita.json'):
   x.update({k:v for k,v in pub[p.name].items() if k not in ('file','sha256','bytes')});x['source_file_name']=p.name;x['world_bank_lastupdated']=json.loads(p.read_text())[0].get('lastupdated')
  else:
   status,period,units=descriptions.get(p.name,('synthetic_client_file','See file definition','See file definition'))
   x.update({'status':status,'source_url':'http://127.0.0.1:49808/'+p.name,'period':period,'units':units})
  reg.append(x)
(R/'evidence/source-register.json').write_text(json.dumps({'register':reg,'synthetic_seed':'2026092736','vintage_note':orig.get('vintage_note')},indent=2)+'\n')
# Legible charts
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})
fig,ax=plt.subplots(figsize=(8,3.5));ax.bar([names[x['country']] for x in C],[x['contribution_eur']/1000 for x in C],color='#286A6B');ax.set_ylabel('EUR thousands');ax.set_title('2025 contribution by market');ax.grid(axis='y',alpha=.2);fig.tight_layout();cc=R/'analysis/country-contribution.png';fig.savefig(cc,dpi=180);plt.close(fig)
fig,ax=plt.subplots(figsize=(7.5,3.5));ss=['low','base','high','stress'];w=.35
for off,pair,col,label in [(-w/2,'NLD+ESP','#286A6B','Netherlands + Spain'),(w/2,'CZE+ESP','#D89936','Czechia + Spain')]:ax.bar([i+off for i in range(4)],[get(pair,s)['incremental_contribution_eur']/1000 for s in ss],w,color=col,label=label)
ax.set_xticks(range(4),['Low','Base','High','Stress']);ax.axhline(0,color='#333',lw=.8);ax.set_ylabel('EUR thousands/year');ax.set_title('Recommended pair and strongest base alternative');ax.legend(frameon=False);ax.grid(axis='y',alpha=.2);fig.tight_layout();cs=R/'analysis/pair-scenarios.png';fig.savefig(cs,dpi=180);plt.close(fig)
# Workbook
wb=Workbook();s=wb.active;s.title='Start here';s.append(['Meridian Parts | Board decision workbook'])
for line in ['Recommendation: stage Netherlands + Spain; capex after 90-day validation.','Base EUR268,160/year; stress EUR208,859; capex EUR270,000; six FTE.','Czechia + Spain: base EUR284,389; stress EUR12,115; capex EUR225,000; five FTE.','Hard limits: EUR450,000; seven FTE; two hubs; defer allowed.','Client operational data are synthetic; macro/FX are archived official observations.','Source register: evidence/source-register.json. Reproduce with python3 analysis/analyze.py then python3 analysis/build_outputs.py.']:s.append([line])
s.column_dimensions['A'].width=112
for i in range(1,s.max_row+1):s.cell(i,1).alignment=Alignment(wrap_text=True);s.row_dimensions[i].height=25
def tab(name,heads,rows,note):
 q=wb.create_sheet(name);q.append([name]);q.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(heads));q['A1'].font=Font(size=16,bold=True,color='FFFFFF');q['A1'].fill=PatternFill('solid',fgColor='17324D');q.append([note]);q.append(heads)
 for c in q[3]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='286A6B');c.alignment=Alignment(wrap_text=True)
 for r in rows:q.append([r.get(h) for h in heads])
 q.freeze_panes='A4';q.auto_filter.ref=f'A3:{chr(64+len(heads))}{q.max_row}'
 for i,h in enumerate(heads,1):q.column_dimensions[chr(64+i)].width=min(32,max(13,len(h)+2))
 return q
heads=['country','shipped_orders','shipped_units','gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur','returned_units','margin']
tab('Countries',heads,C,'2025 totals; additive metrics reconcile to months; margin=contribution/net sales.')
tab('Monthly',['country','month']+heads[1:],M['monthly'],'EUR; original shipment month FX applies to gross sales and refunds.')
tab('Market context',['country','year','population','gdp_per_capita_usd'],M['market_context'],'World Bank archive; population persons; GDP current US dollars; null preserved.')
tab('FX monthly',['currency','month','local_per_eur'],M['fx_monthly'],'ECB mean of 2025 business days; local units per EUR; EUR=1.')
tab('Hub scenarios',['country','scenario','incremental_contribution_eur','capex_eur','fte','payback_years'],M['hub_scenarios'],'All single options and feasible pairs; capex is year zero; null payback if nonpositive annual increment.')
tab('Quality',['item','value'],[{'item':k,'value':json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v} for k,v in M['quality'].items()],'Detected anomalies and handling; see evidence/source-register.json.')
q=wb.create_sheet('Decision chart');q.append(['Portfolio','Base annual increment','Stress annual increment'])
for p in ['NLD+ESP','CZE+ESP','POL+ESP','NLD+CZE','POL+CZE','ESP']:q.append([p,get(p,'base')['incremental_contribution_eur'],get(p,'stress')['incremental_contribution_eur']])
q.append(['Defer',0,0]);ch=BarChart();ch.type='col';ch.title='Base vs stress annual increment';ch.y_axis.title='EUR';ch.add_data(Reference(q,min_col=2,max_col=3,min_row=1,max_row=8),titles_from_data=True);ch.set_categories(Reference(q,min_col=1,min_row=2,max_row=8));ch.height=9;ch.width=18;q.add_chart(ch,'E2');wb.calculation.fullCalcOnLoad=True;wb.calculation.forceFullCalc=True;wb.save(O/'meridian_analysis.xlsx')
# Report PDF
st=getSampleStyleSheet();st.add(ParagraphStyle(name='MTitle',parent=st['Title'],fontSize=24,leading=28,textColor=colors.HexColor('#17324D'),alignment=0));st.add(ParagraphStyle(name='MH',parent=st['Heading1'],textColor=colors.HexColor('#17324D'),spaceBefore=9,spaceAfter=5))
P=lambda x,y='BodyText':Paragraph(x,st[y])
def table(rows,widths,fs=7):
 t=Table(rows,repeatRows=1,colWidths=widths);t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#286A6B')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('FONTSIZE',(0,0),(-1,-1),fs),('ALIGN',(1,1),(-1,-1),'RIGHT'),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#EEF3F5')]),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#CBD6DC')),('VALIGN',(0,0),(-1,-1),'MIDDLE')]));return t
story=[P('Meridian Parts','MTitle'),P('European service-hub expansion | Board decision report | 27 September 2026','Heading2'),P('<b>Recommendation:</b> stage Netherlands + Spain; release no capex until a 90-day validation confirms demand, service improvement and realizable savings. Supplied assumptions yield EUR268,160 annual increment in base and EUR208,859 in defined stress; capex EUR270,000; six FTE. This is scenario analysis, not a causal estimate.'),P('Decision in brief','MH'),P('Hard limits: EUR450,000 capex, seven FTE, maximum two hubs; defer is allowed. Netherlands + Spain leads on the client-defined stress case. It gives up EUR16,230 base annual contribution against Czechia + Spain but retains EUR196,744 more under stress.'),Image(str(cs),width=6.8*inch,height=3.1*inch),P('Low/base/high uplifts: 10%/25%/40%. Stress: 25% uplift, refund shock 3% gross sales, 10% depreciation shock on PLN/CZK net sales. No probability assigned.'),
P('2025 operating baseline','MH'),Image(str(cc),width=6.8*inch,height=3.0*inch),P('Spain leads contribution at EUR418,333, then Czechia EUR393,792 and Poland EUR353,245. All six markets: EUR4.465m gross, EUR4.036m net sales, EUR2.011m contribution; 1,254 eligible orders, 80,436 units and 1,081 returned units through cutoff. Contribution deducts net COGS and nonrefundable fulfillment; it is not booked revenue or cash. Source files: synthetic orders-part1.csv, orders-part2.csv, order-corrections.csv, returns.csv and unit-costs.csv; archived ECB ecb-history.zip supplies translation rates.'),P('Country totals','MH')]
rows=[['Market','Orders','Units','Gross','Refunds','Net sales','Net COGS','Fulfill','Contribution','Margin','Returns']]
for r in C:rows.append([names[r['country']],str(r['shipped_orders']),f"{r['shipped_units']:,}",f"{r['gross_sales_eur']:,.0f}",f"{r['refunds_eur']:,.0f}",f"{r['net_sales_eur']:,.0f}",f"{r['net_cogs_eur']:,.0f}",f"{r['fulfillment_eur']:,.0f}",f"{r['contribution_eur']:,.0f}",f"{r['margin']:.1%}",str(r['returned_units'])])
story.append(table(rows,[.65*inch,.37*inch,.44*inch,.52*inch,.5*inch,.57*inch,.54*inch,.47*inch,.6*inch,.38*inch,.43*inch],6))
story += [P('Accounting','MH'),P('Gross sales are discounted product sales translated at the original shipment-month ECB mean. Refunds use the same sale-month quote. Net sales=gross less refunds. Net COGS deducts cost only for restocked returns; fulfillment is actual nonrefundable EUR. Contribution=net sales−net COGS−fulfillment. Cash differs with payment timing, fees, settlement FX and receivables not supplied. Hub costs/capex sit outside baseline contribution.'),P('Hub options','MH')]
rows=[['Portfolio','Base EUR/y','Stress EUR/y','Capex','FTE','Stress payback']]
for p in ['NLD+ESP','CZE+ESP','POL+ESP','NLD+CZE','POL+CZE','FRA+ESP','NLD+POL','FRA+NLD','ESP','CZE','POL','NLD','FRA','DEU']:
 a=get(p,'base');b=get(p,'stress');rows.append([p,f"{a['incremental_contribution_eur']:,.0f}",f"{b['incremental_contribution_eur']:,.0f}",f"{a['capex_eur']:,.0f}",str(a['fte']),'—' if b['payback_years'] is None else f"{b['payback_years']:.2f}y"])
story.append(table(rows,[.95*inch,1.02*inch,1.02*inch,.75*inch,.4*inch,.9*inch],7))
story += [P('Recommendation and alternative','MH'),P('Netherlands + Spain is the strongest stress-performing feasible pair: EUR208,859 stress, EUR268,160 base, positive low/base/high, EUR270,000 capex, six FTE. Czechia + Spain has strongest base (EUR284,389), EUR225,000 capex/five FTE, but stress is EUR12,115 and 18.57-year payback vs 1.29 years. The difference is the stipulated 10% CZK shock, not a live forecast. Holding other inputs fixed, Czechia + Spain would lead the stress comparison if the CZK depreciation shock were below about 0.55%. If the board rejects the stress and prioritizes base output, Czechia + Spain is the alternative. If inputs fail validation, defer. Eleven pairs meet limits; workbook includes all singles and feasible pairs across four scenarios; pairs assume no synergy.')]
story.append(P('Criteria provenance: capex and staffing are hard client constraints. Baselines and options are synthetic client inputs; scenario results are computed from them. Ranking on stress is a conservative consultant default because no board weights were supplied; it is not a board-approved preference or blended score. Annual increment = C×u + U×(1+u)×s − F. Capex is year zero; simple payback is capex divided by positive annual increment.'))
pop={(x['country'],x['year']):x['population'] for x in M['market_context']};gdp={(x['country'],x['year']):x['gdp_per_capita_usd'] for x in M['market_context']};rows=[['Market','Population 2022','Population 2024','Change','2024 GDP/capita current USD']]
for c in names:
 a=pop[c,2022];b=pop[c,2024];rows.append([names[c],f'{a:,}' if a else 'missing',f'{b:,}' if b else 'missing',f'{100*(b/a-1):.1f}%' if a and b else '—',f"{gdp[c,2024]:,.0f}" if gdp[c,2024] else 'missing'])
story.append(KeepTogether([P('Market context is not demand proof','MH'),table(rows,[1.05*inch,1.15*inch,1.15*inch,.7*inch,2.05*inch],8)]));story.append(P('Population is persons; GDP current USD/person, not PPP or constant dollars. These are context, not demand proof. All country-years present. World Bank archive collected 27 Sep 2026 reports lastupdated 13 Jul 2026; historical revisions may exist. ECB means are local units/EUR; EUR=1.'))
story += [P('90-day implementation','MH'),table([['Days','Owner','Action / gate'],['0–15','COO + Finance','Reconcile baseline and capture service KPIs; no commitments.'],['16–45','Commercial + Ops','Needs account/customer and carrier data access; verify demand, delivery/returns, SKU savings and quotes.'],['46–70','Finance + HR + Legal','Depends on Gate 1 evidence; refresh fully loaded costs/FX and complete local diligence. Gate: downside increment >0 and payback ≤3 years.'],['71–90','COO + Board sponsor','Depends on Gate 2; commit capped pilot only if it passes, otherwise defer.']],[.7*inch,1.25*inch,4.7*inch]),P('Weekly scorecard vs measured baseline: delivery median/P90 down ≥20% by day 90; on-time-in-full +5pp; returns no more than +0.5pp; realized savings ≥EUR3.10/unit; downside contribution >0; zero unresolved orphan/revision issues. These are proposed gates.'),P('Risks and controls','MH'),P('Demand uplift is unmeasured; validate accounts. Savings may omit labor/inventory; check fully loaded invoices. FX stress is not probability; refresh exposure. Track returns. Local labor, tax, lease and safety need diligence. Cannibalization, working capital, discounting, taxes, ramp and residual value are excluded.'),P('Reconciliation and reproducibility','MH'),P('1,328 raw order rows → 1,297 distinct IDs: 13 identical duplicates removed, 18 superseded revisions replaced, no same-revision conflicts; exclude 24 non-shipped, 18 tests, one 2026 shipment. 264 raw returns →243 IDs: 20 duplicates removed, one superseded revision; exclude one after cutoff; quarantine one orphan; include 241 eligible. No required calculation fields missing. Free/zero-price shipments retained; additive values reconcile month to year.'),P('Synthetic client inputs (seed 2026092736) are separate from archived official ECB/World Bank data. evidence/source-register.json records URL, vintage, hash, period, units/status. Reproduce from root: python3 analysis/analyze.py then python3 analysis/build_outputs.py. No external research/interviews used.')]
SimpleDocTemplate(str(O/'meridian_board_report.pdf'),pagesize=letter,rightMargin=.48*inch,leftMargin=.48*inch,topMargin=.4*inch,bottomMargin=.45*inch).build(story)
# Rendered landscape board presentation PDF; slide order matches the PPTX below.
deck=canvas.Canvas(str(O/'meridian_board_presentation.pdf'),pagesize=(960,540))
def slidebase(title,subtitle,page):
 deck.setFillColor(colors.HexColor('#F5F8FA'));deck.rect(0,0,960,540,fill=1,stroke=0);deck.setFillColor(colors.HexColor('#286A6B'));deck.rect(0,528,960,12,fill=1,stroke=0)
 deck.setFillColor(colors.HexColor('#17324D'));deck.setFont('Helvetica-Bold',26);deck.drawString(46,486,title)
 if subtitle:deck.setFillColor(colors.HexColor('#596C7B'));deck.setFont('Helvetica',11);deck.drawString(48,463,subtitle)
 deck.setFillColor(colors.HexColor('#7B8992'));deck.setFont('Helvetica',8);deck.drawString(46,18,'Meridian Parts | Synthetic client analysis | Decision support');deck.drawRightString(920,18,str(page))
def bullets_pdf(items,x=65,y=400,width=835,size=19,leading=27):
 for item in items:
  deck.setFillColor(colors.HexColor('#286A6B'));deck.circle(x,y+4,3.2,fill=1,stroke=0);deck.setFillColor(colors.HexColor('#263746'));deck.setFont('Helvetica',size)
  lines=simpleSplit(item,'Helvetica',size,width-28)
  for j,line in enumerate(lines):deck.drawString(x+18,y-j*leading,line)
  y-=leading*len(lines)+17
slidebase('Stage Netherlands + Spain','90-day validation before site commitment',1);deck.setFont('Helvetica-Bold',22);deck.setFillColor(colors.HexColor('#17324D'));deck.drawString(65,410,'Recommendation')
bullets_pdf(['EUR 268k base annual increment; EUR 209k under defined stress','EUR 270k year-zero capex and six FTE','Gives up EUR 16k base vs Czechia + Spain, but retains EUR 197k more under stress','Release capex only if verified downside stays positive and payback is at most three years'],65,365,825,18,24);deck.showPage()
slidebase('Observed 2025 contribution','1,254 eligible orders | 80,436 shipped units | returns through 31 January 2026',2);deck.drawImage(str(cc),48,94,width=600,height=330,preserveAspectRatio=True,anchor='c');bullets_pdf(['Spain: EUR 418k','Czechia: EUR 394k','Poland: EUR 353k'],680,350,235,18,24);deck.showPage()
slidebase('Downside changes the pair ranking','Client-defined scenarios; no probability assigned',3);deck.drawImage(str(cs),48,95,width=600,height=325,preserveAspectRatio=True,anchor='c');bullets_pdf(['NL + ES: EUR 268k base / EUR 209k stress','CZ + ES: EUR 284k base / EUR 12k stress','CZ stress includes 10% CZK shock on net sales'],680,370,240,17,23);deck.showPage()
slidebase('Trade-off and switch conditions','Hard limits: EUR 450k capex | seven FTE | at most two hubs',4);bullets_pdf(['NL + ES uses EUR 45k more capex and one more FTE than CZ + ES','CZ + ES leads stress if the CZK shock is below about 0.55%','If the board prioritizes base output, CZ + ES is strongest; if validation fails, defer','Stress is scenario arithmetic, not probability or forecast'],65,395,820,19,27);deck.showPage()
slidebase('90-day gates keep the choice reversible','Owners are roles; assign named accountable people at kickoff',5)
for i,(day,owner,desc) in enumerate([('0–15','COO + Finance','Reconcile baseline and capture service KPIs'),('16–45','Commercial + Ops','Access customer/carrier data; verify demand, service and savings'),('46–70','Finance + HR + Legal','After Gate 1: full costs, local diligence; downside >0; payback ≤3y'),('71–90','COO + Board sponsor','After Gate 2, commit only if evidence passes; else defer')]):
 y=390-i*88;deck.setFillColor(colors.HexColor('#286A6B'));deck.roundRect(55,y-20,125,45,5,fill=1,stroke=0);deck.setFillColor(colors.white);deck.setFont('Helvetica-Bold',15);deck.drawCentredString(117,y-5,'Days '+day);deck.setFillColor(colors.HexColor('#17324D'));deck.setFont('Helvetica-Bold',15);deck.drawString(205,y,owner);deck.setFillColor(colors.HexColor('#263746'));deck.setFont('Helvetica',14)
 for j,line in enumerate(simpleSplit(desc,'Helvetica',14,500)):deck.drawString(430,y-j*18,line)
deck.showPage();slidebase('KPI scorecard','Weekly by country/SKU against measured pre-pilot baseline',6);bullets_pdf(['Delivery median and P90: at least 20% reduction by day 90','On-time-in-full: +5 percentage points; returns: no more than +0.5 points','Realized savings: at least EUR 3.10/unit; downside contribution remains positive','Data quality: zero unresolved orphan/revision exceptions'],65,400,830,19,28);deck.showPage()
slidebase('Evidence boundary','Inputs, public observations and model limits',7);bullets_pdf(['Orders, returns, costs and hub assumptions are synthetic client data','Official archived ECB 2025 rates and World Bank 2022–24 observations','Population/GDP provide context, not proof of parts demand','No causal effect, discounting, ramp, tax or working-capital model'],65,400,830,19,28);deck.save()
# Board presentation
prs=Presentation();prs.slide_width=Inches(13.33);prs.slide_height=Inches(7.5)
def slide(title,sub=''):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor(245,248,250)
 a=s.shapes.add_shape(1,0,0,prs.slide_width,Inches(.15));a.fill.solid();a.fill.fore_color.rgb=RGBColor(40,106,107);a.line.fill.background()
 t=s.shapes.add_textbox(Inches(.6),Inches(.4),Inches(12),Inches(.6));p=t.text_frame.paragraphs[0];p.text=title;p.font.size=Pt(27);p.font.bold=True;p.font.color.rgb=RGBColor(23,50,77)
 if sub:t=s.shapes.add_textbox(Inches(.65),Inches(1.05),Inches(12),Inches(.35));p=t.text_frame.paragraphs[0];p.text=sub;p.font.size=Pt(12)
 return s
def bullets(s,items,size=21):
 t=s.shapes.add_textbox(Inches(.8),Inches(1.55),Inches(11.8),Inches(5.3));f=t.text_frame;f.word_wrap=True;f.clear()
 for i,v in enumerate(items):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text='• '+v;p.font.size=Pt(size);p.font.color.rgb=RGBColor(38,55,68);p.space_after=Pt(15)
s=slide('Stage Netherlands + Spain','90-day validation before site commitment');bullets(s,['EUR268k base annual increment; EUR209k defined stress','EUR270k capex; six FTE','EUR16k less base than CZ + ES; EUR197k more under stress','Release capex only after verified downside remains positive and payback ≤3y'])
s=slide('2025 synthetic contribution','1,254 eligible orders • 80,436 units');s.shapes.add_picture(str(cc),Inches(.7),Inches(1.5),width=Inches(8.4));bullets(s,['Spain EUR418k','Czechia EUR394k','Poland EUR353k'],19)
s=slide('Downside changes the ranking','Client-defined scenarios; no probability assigned');s.shapes.add_picture(str(cs),Inches(.7),Inches(1.55),width=Inches(8));bullets(s,['NL+ES: EUR268k base / EUR209k stress','CZ+ES: EUR284k base / EUR12k stress','CZ stress includes 10% CZK shock'],18)
s=slide('Trade-off and switch conditions','Hard limits: EUR450k • seven FTE • two hubs');bullets(s,['NL + ES costs EUR45k more capex and one more FTE','CZ + ES would lead stress if the modeled CZK shock were below about 0.55%','If demand or savings fail validation: defer','Stress is scenario arithmetic, not forecast probability'],20)
s=slide('90-day gates','Owners are roles; assign people at kickoff')
for i,(day,owner,desc) in enumerate([('0–15','COO + Finance','Reconcile baseline; define KPIs'),('16–45','Commercial + Ops','Access customer/carrier data; validate demand, service and savings'),('46–70','Finance + HR + Legal','After Gate 1: full costs/local diligence; downside >0; payback ≤3y'),('71–90','COO + Board sponsor','After Gate 2, commit only if evidence passes; otherwise defer')]):
 y=1.5+i*1.25;sh=s.shapes.add_shape(1,Inches(.8),Inches(y),Inches(2),Inches(.75));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor(40,106,107);sh.line.fill.background();p=sh.text_frame.paragraphs[0];p.text='Days '+day;p.font.size=Pt(17);p.font.bold=True;p.font.color.rgb=RGBColor(255,255,255)
 for x,w,text,b in [(3.1,2.5,owner,True),(5.8,6.8,desc,False)]:
  t=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(.8));p=t.text_frame.paragraphs[0];p.text=text;p.font.size=Pt(16);p.font.bold=b
s=slide('KPI scorecard','Weekly by country/SKU vs measured baseline');bullets(s,['Delivery median/P90 down ≥20% by day 90','On-time-in-full +5pp; returns no more than +0.5pp','Realized savings ≥EUR3.10/unit; downside contribution >0','Zero unresolved orphan/revision exceptions'],19)
s=slide('Evidence boundary');bullets(s,['Client orders, returns, costs and option assumptions are synthetic','Archived official ECB 2025 rates and WB 2022–24 observations','Macro context is not proof of parts demand','No causal effect, discounting, taxes or ramp modeled'],19)
prs.save(O/'meridian_board_presentation.pptx')
print('Built workbook, report, presentation, charts and source register.')
