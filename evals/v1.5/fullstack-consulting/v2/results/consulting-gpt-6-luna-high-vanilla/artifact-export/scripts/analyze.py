#!/usr/bin/env python3
"""Reproducible Meridian Parts analysis from saved frozen inputs."""
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from itertools import combinations
from datetime import datetime, timezone
import csv, json, hashlib
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]; RAW=ROOT/'evidence/raw'; OUT=ROOT/'deliverables'
OUT.mkdir(exist_ok=True)
COUNTRIES=['DEU','FRA','NLD','POL','CZE','ESP']
NAMES={'DEU':'Germany','FRA':'France','NLD':'Netherlands','POL':'Poland','CZE':'Czechia','ESP':'Spain'}
def D(x): return Decimal(str(x))
def cent(x): return D(x).quantize(Decimal('.01'),rounding=ROUND_HALF_UP)
def safe(x):
    if pd.isna(x): return None
    return round(float(x),6) if isinstance(x,(float,)) else int(x) if isinstance(x,(int,)) else x

# ECB daily observation means, quote local units per EUR.
fxdata=pd.read_csv(RAW/'ecb-history.csv',dtype=str); fxdata['Date']=pd.to_datetime(fxdata.Date)
fxdata=fxdata[fxdata.Date.dt.year==2025]; rates={('EUR',m):1.0 for m in range(1,13)}; fx=[]
for cur in ['PLN','CZK']:
    for m,g in fxdata.groupby(fxdata.Date.dt.month):
        rates[cur,m]=float(pd.to_numeric(g[cur],errors='coerce').mean())
        fx.append({'currency':cur,'month':f'2025-{m:02d}','local_per_eur':round(rates[cur,m],8)})
fx += [{'currency':'EUR','month':f'2025-{m:02d}','local_per_eur':1.0} for m in range(1,13)]

orders=pd.concat([pd.read_csv(RAW/f'orders-part{i}.csv') for i in (1,2)]+[pd.read_csv(RAW/'order-corrections.csv')],ignore_index=True)
raw_order_rows=len(orders); exact_order_duplicates=int(orders.duplicated().sum())
conflicting_order_revision_tie_rows=int(orders.drop_duplicates().duplicated(['order_id','revision'],keep=False).sum())
orders['is_test']=orders.is_test.astype(bool); orders['revision']=orders.revision.astype(int)
orders=orders.sort_values(['order_id','revision']).drop_duplicates(['order_id','revision'],keep='last')
orders=orders[orders.revision==orders.groupby('order_id').revision.transform('max')].drop_duplicates('order_id',keep='last').copy()
orders['shipped_at']=pd.to_datetime(orders.shipped_at,errors='coerce')
orders['eligible']=(orders.status=='shipped')&(~orders.is_test)&orders.shipped_at.between('2025-01-01','2025-12-31')
eligible=orders[orders.eligible].copy(); eligible['month']=eligible.shipped_at.dt.strftime('%Y-%m')
returns=pd.read_csv(RAW/'returns.csv'); raw_return_rows=len(returns); exact_return_duplicates=int(returns.duplicated().sum())
conflicting_return_revision_tie_rows=int(returns.drop_duplicates().duplicated(['return_id','revision'],keep=False).sum())
returns['revision']=returns.revision.astype(int); returns=returns.sort_values(['return_id','revision']).drop_duplicates(['return_id','revision'],keep='last')
returns=returns[returns.revision==returns.groupby('return_id').revision.transform('max')].drop_duplicates('return_id',keep='last')
returns['received_at']=pd.to_datetime(returns.received_at); returns=returns[returns.received_at<=pd.Timestamp('2026-01-31')]
attached0=returns[returns.order_id.isin(set(eligible.order_id))].copy(); orphan=returns[~returns.order_id.isin(set(eligible.order_id))].copy()
attached=attached0.merge(eligible[['order_id','currency','month','country','sku','shipped_at']],on='order_id',how='left',validate='many_to_one')
costs=pd.read_csv(RAW/'unit-costs.csv'); costs.valid_from=pd.to_datetime(costs.valid_from)
calc=[]
for r in eligible.itertuples(index=False):
    rr=attached[attached.order_id==r.order_id]
    refundlocal=sum((D(x) for x in rr.refund_local),D(0)) if len(rr) else D(0)
    returned=int(rr.quantity.sum()) if len(rr) else 0; restocked=int(rr.restocked_quantity.sum()) if len(rr) else 0
    rate=D(1 if r.currency=='EUR' else rates[(r.currency,r.shipped_at.month)])
    grosslocal=D(r.quantity)*D(r.unit_price_local)-D(r.discount_local)
    gross=cent(grosslocal/rate); refund=cent(refundlocal/rate)
    unit=float(costs[(costs.sku==r.sku)&(costs.valid_from<=r.shipped_at)].sort_values('valid_from').iloc[-1].unit_cost_eur)
    cogs=cent(D(r.quantity)*D(unit))-cent(D(restocked)*D(unit))
    fulfillment=cent(r.fulfillment_eur); contribution=gross-refund-cogs-fulfillment
    calc.append({'order_id':r.order_id,'country':r.country,'month':r.month,'units':int(r.quantity),'gross_sales_eur':gross,'refunds_eur':refund,'net_sales_eur':gross-refund,'net_cogs_eur':cogs,'fulfillment_eur':fulfillment,'contribution_eur':contribution,'returned_units':returned})
odf=pd.DataFrame(calc); fields=['shipped_orders','shipped_units','gross_sales_eur','refunds_eur','net_sales_eur','net_cogs_eur','fulfillment_eur','contribution_eur','returned_units']
keys=pd.MultiIndex.from_product([COUNTRIES,pd.period_range('2025-01','2025-12',freq='M').astype(str)],names=['country','month']).to_frame(index=False)
agg=odf.groupby(['country','month']).agg(shipped_orders=('order_id','nunique'),shipped_units=('units','sum'),gross_sales_eur=('gross_sales_eur','sum'),refunds_eur=('refunds_eur','sum'),net_sales_eur=('net_sales_eur','sum'),net_cogs_eur=('net_cogs_eur','sum'),fulfillment_eur=('fulfillment_eur','sum'),contribution_eur=('contribution_eur','sum'),returned_units=('returned_units','sum')).reset_index()
monthly=keys.merge(agg,on=['country','month'],how='left').fillna(0)
for col in fields[2:-1]: monthly[col]=monthly[col].map(lambda x:float(cent(x)))
monthly['margin']=monthly.apply(lambda r:round(r.contribution_eur/r.net_sales_eur,6) if r.net_sales_eur else None,axis=1)
country=monthly.groupby('country')[fields].sum().reset_index(); country['margin']=country.apply(lambda r:round(r.contribution_eur/r.net_sales_eur,6) if r.net_sales_eur else None,axis=1)

# Preserve World Bank series and nulls from the archive.
pop=json.load(open(RAW/'population.json'))[1]; gdp=json.load(open(RAW/'gdp-per-capita.json'))[1]; market=[]
for c in COUNTRIES:
 for y in range(2022,2025):
    market.append({'country':c,'year':y,'population':next((x['value'] for x in pop if x['countryiso3code']==c and int(x['date'])==y),None),'gdp_per_capita_usd':next((x['value'] for x in gdp if x['countryiso3code']==c and int(x['date'])==y),None)})
market_summary=[]
for c in COUNTRIES:
    p22=next((x['population'] for x in market if x['country']==c and x['year']==2022),None); p24=next((x['population'] for x in market if x['country']==c and x['year']==2024),None)
    market_summary.append({'country':c,'population_2022':p22,'population_2024':p24,'population_change_pct_2022_2024':round(100*(p24/p22-1),4) if p22 and p24 else None,'gdp_per_capita_usd_2024':next((x['gdp_per_capita_usd'] for x in market if x['country']==c and x['year']==2024),None)})

opts=pd.read_csv(RAW/'hub-options.csv').set_index('country'); scenarios=[]
for c in COUNTRIES:
    b=country[country.country==c].iloc[0]; o=opts.loc[c]; C=float(b.contribution_eur); U=int(b.shipped_units); G=float(b.gross_sales_eur); N=float(b.net_sales_eur)
    for name,u in [('low',.10),('base',.25),('high',.40)]:
        inc=C*u+U*(1+u)*o.saving_eur_per_unit-o.annual_fixed_eur
        scenarios.append({'country':c,'scenario':name,'incremental_contribution_eur':round(inc,2),'capex_eur':int(o.capex_eur),'fte':int(o.fte),'payback_years':round(o.capex_eur/inc,4) if inc>0 else None})
    fxshock=.10*N if c in ('POL','CZE') else 0; cstress=C-.03*G-fxshock
    inc=cstress*1.25-C+U*1.25*o.saving_eur_per_unit-o.annual_fixed_eur
    scenarios.append({'country':c,'scenario':'stress','incremental_contribution_eur':round(inc,2),'capex_eur':int(o.capex_eur),'fte':int(o.fte),'payback_years':round(o.capex_eur/inc,4) if inc>0 else None})
base={x['country']:x['incremental_contribution_eur'] for x in scenarios if x['scenario']=='base'}
alts=[]
for n in (1,2):
 for group in combinations(COUNTRIES,n):
    cap=sum(int(opts.loc[c].capex_eur) for c in group); fte=sum(int(opts.loc[c].fte) for c in group)
    if cap<=450000 and fte<=7:
        vals={s:sum(next(x['incremental_contribution_eur'] for x in scenarios if x['country']==c and x['scenario']==s) for c in group) for s in ('low','base','high','stress')}
        alts.append({'countries':list(group),'capex_eur':cap,'fte':fte,'low_annual_incremental_eur':round(vals['low'],2),'base_annual_incremental_eur':round(vals['base'],2),'high_annual_incremental_eur':round(vals['high'],2),'stress_annual_incremental_eur':round(vals['stress'],2),'payback_years':round(cap/vals['base'],4) if vals['base']>0 else None})
alts.sort(key=lambda x:x['base_annual_incremental_eur'],reverse=True); best=alts[0]; second=alts[1]
gates=[x for x in alts if x['low_annual_incremental_eur']>0 and x['stress_annual_incremental_eur']>0]
chosen=gates[0] if gates else {'countries':[],'capex_eur':0,'fte':0}
recommendation={'countries':chosen['countries'],'capex_eur':chosen['capex_eur'],'fte':chosen['fte'],'rationale':'Fund the best feasible portfolio that is positive in both low volume and joint stress.' if gates else 'Defer capex: no feasible option is contribution-positive in both the low-volume case and defined joint return/FX stress; reopen after validated evidence clears both gates.'}

noncancel=orders[orders.status!='cancelled']; eligible_status=noncancel[~noncancel.is_test]
recon={'raw_order_rows':raw_order_rows,'exact_duplicate_order_rows':exact_order_duplicates,'conflicting_order_revision_tie_rows':conflicting_order_revision_tie_rows,'distinct_order_ids_after_revisions':int(orders.order_id.nunique()),'eligible_shipped_2025_orders':int(len(eligible)),'excluded_order_ids':int((~orders.eligible).sum()),'excluded_cancelled_ids':int((orders.status=='cancelled').sum()),'excluded_test_ids':int((noncancel.is_test).sum()),'excluded_out_of_year_or_missing_shipment_ids':int((eligible_status.shipped_at.isna()|~eligible_status.shipped_at.between('2025-01-01','2025-12-31')).sum()),'missing_shipment_date_cancelled_ids':int(((orders.status=='cancelled')&orders.shipped_at.isna()).sum()),'eligible_zero_price_orders':int(((eligible.quantity*eligible.unit_price_local-eligible.discount_local)==0).sum()),'eligible_missing_required_values':int(eligible[['country','shipped_at','sku','quantity','unit_price_local','discount_local','currency','fulfillment_eur']].isna().any(axis=1).sum()),'raw_return_rows':raw_return_rows,'exact_duplicate_return_rows':exact_return_duplicates,'conflicting_return_revision_tie_rows':conflicting_return_revision_tie_rows,'return_ids_after_revisions_and_cutoff':int(len(returns)),'post_cutoff_return_ids':int((pd.to_datetime(pd.read_csv(RAW/'returns.csv').received_at)>pd.Timestamp('2026-01-31')).sum()),'attached_return_ids':int(attached.return_id.nunique()),'quarantined_or_unlinked_return_ids':int(orphan.return_id.nunique()),'monthly_rows':len(monthly),'monthly_country_totals_reconcile':bool(all(abs(float(monthly.groupby('country')[col].sum().loc[c])-float(country.set_index('country').loc[c,col]))<.011 for c in COUNTRIES for col in fields))}
quality={'reconciliation':recon,'handling':['Highest numeric revision replaces whole row; exact repeated records collapse.','Include shipped, non-test orders shipped in 2025; retain valid zero-price shipments.','Returns received through 2026-01-31 inclusive attach only to eligible order IDs; orphans quarantined.','Use original sale month ECB quote, local units per EUR; round per-order components half-up.','Use effective-dated EUR unit cost and recover only restocked quantity.','Empty country-months have zero activity; margin null at zero net sales.'],'missing_values':['World Bank nulls preserved without imputation.','ECB 2025 PLN/CZK means use available published business days.'],'limitations':['Client order, return, cost and hub options are synthetic, not observed performance.','World Bank snapshot may include revisions after historical years.','ECB reference rates are translation assumptions, not realized FX.','Scenario arithmetic is simple, undiscounted and non-causal; no ramp or working capital.']}
def cleanrows(df):
    out=df.to_dict('records')
    return [{k:(v if isinstance(v,list) else safe(v)) for k,v in r.items()} for r in out]
metrics={'monthly':cleanrows(monthly),'countries':cleanrows(country),'fx_monthly':fx,'market_context':market,'market_summary':market_summary,'hub_scenarios':scenarios,'alternatives':alts,'recommendation':recommendation,'quality':quality}
(OUT/'metrics.json').write_text(json.dumps(metrics,indent=2,allow_nan=False))

# Charts
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})
fig,ax=plt.subplots(figsize=(9,4.5)); csorted=country.sort_values('contribution_eur',ascending=False)
ax.bar([c for c in csorted.country],csorted.contribution_eur/1000,color='#237d83'); ax.set_ylabel('2025 contribution (€000)'); ax.set_title('Synthetic 2025 contribution by country'); fig.tight_layout(); fig.savefig(OUT/'country_contribution.png',dpi=180); plt.close(fig)
fig,ax=plt.subplots(figsize=(9,4.5)); top=alts[:10]; x=range(len(top)); ax.bar(x,[a['base_annual_incremental_eur']/1000 for a in top],color='#237d83'); ax.axhline(0,color='#333',lw=.8); ax.set_xticks(list(x),['+'.join(a['countries']) for a in top],rotation=24,ha='right'); ax.set_ylabel('Annual incremental contribution (€000)'); ax.set_title('Feasible hub portfolios: base case'); fig.tight_layout(); fig.savefig(OUT/'alternatives.png',dpi=180); plt.close(fig)

# XLSX analytical workbook
wb=Workbook(); readme=wb.active; readme.title='Read me'
notes=[['Meridian Parts | analysis and notes'],['Cutoff','2025 shipments; returns through 2026-01-31 inclusive'],['Client inputs','All transaction, unit-cost and hub-option data are synthetic.'],['Archives','ECB historic reference rate snapshot and World Bank API snapshots in evidence/raw. WB lastupdated 2026-07-13; captured by room 2026-09-27.'],['Accounting','Highest revision wins; sale month local units per EUR; order-level cent half-up rounding. See data-dictionary.md.'],['Scenario','Low/base/high uplift 10/25/40%; stress as specified in scenario-policy.md. Capex year zero; annual increments after recurring fixed costs.'],['Limitations','Population/GDP context is not direct product demand. Scenarios are assumption-driven, non-causal, undiscounted.'],['Quality counts',json.dumps(recon,sort_keys=True)]]
for row in notes: readme.append(row)
def sheet(name,df):
    sh=wb.create_sheet(name); sh.append(list(df.columns))
    for row in df.itertuples(index=False,name=None): sh.append([None if pd.isna(v) else (list(v) if isinstance(v,tuple) else v) for v in row])
    return sh
sheet('Monthly',monthly); sheet('Countries',country); sheet('FX monthly',pd.DataFrame(fx)); sheet('Market context',pd.DataFrame(market)); sheet('Market summary',pd.DataFrame(market_summary)); sheet('Hub scenarios',pd.DataFrame(scenarios));
alt_df=pd.DataFrame([{**a,'countries':'+'.join(a['countries'])} for a in alts]); ps=sheet('Portfolios',alt_df)
for sh in wb.worksheets:
    sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
    for cell in sh[1]: cell.fill=PatternFill('solid',fgColor='123843'); cell.font=Font(color='FFFFFF',bold=True)
    for col in sh.columns:
        width=min(40,max(12,max(len(str(v.value or '')) for v in col)+2)); sh.column_dimensions[get_column_letter(col[0].column)].width=width
    for row in sh.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,float): cell.number_format='#,##0.00;[Red](#,##0.00);-'
ch=BarChart(); ch.type='bar'; ch.title='Top feasible portfolios: annual base increment'; ch.x_axis.title='EUR'; ch.y_axis.title='Portfolio'; ch.add_data(Reference(ps,min_col=5,min_row=1,max_row=min(11,ps.max_row)),titles_from_data=True); ch.set_categories(Reference(ps,min_col=1,min_row=2,max_row=min(11,ps.max_row))); ch.width=18; ch.height=9; ps.add_chart(ch,'J2')
country_sheet=wb['Countries']; ch2=BarChart(); ch2.type='bar'; ch2.title='2025 contribution by country'; ch2.x_axis.title='EUR'; ch2.add_data(Reference(country_sheet,min_col=9,min_row=1,max_row=7),titles_from_data=True); ch2.set_categories(Reference(country_sheet,min_col=1,min_row=2,max_row=7)); ch2.width=15; ch2.height=8; country_sheet.add_chart(ch2,'N2')
wb.save(OUT/'meridian_analysis.xlsx')

# Source register with hashes and provenance. Room register preserves official retrieval vintage.
orig=json.load(open(RAW/'source-register.json')); official={x['file']:x for x in orig['public']}; entries=[]; capture=datetime.now(timezone.utc).isoformat()
for path in sorted(RAW.iterdir()):
    if not path.is_file(): continue
    info=official.get(path.name,{}); data=path.read_bytes()
    entries.append({'file':path.name,'url':info.get('original_url','https://www.ecb.europa.eu/stats/eurofxref/eurofxref-hist.zip' if path.name=='ecb-history.csv' else 'http://127.0.0.1:53350/'+path.name),'retrieved_utc':info.get('retrieved_utc',capture),'sha256':hashlib.sha256(data).hexdigest(),'bytes':len(data),'units':{'population.json':'persons','gdp-per-capita.json':'current US dollars/person','ecb-history.csv':'local units per EUR','ecb-history.zip':'local units per EUR','unit-costs.csv':'EUR/unit','hub-options.csv':'EUR, FTE, EUR/unit savings','orders-part1.csv':'local currency, EUR cost, units','orders-part2.csv':'local currency, EUR cost, units','order-corrections.csv':'local currency, EUR cost, units','returns.csv':'local currency, units'}.get(path.name,'definitions/metadata'),'period':'2022–2024' if path.name in ('population.json','gdp-per-capita.json') else '2025 calculations; archive 1999–2026' if path.name.startswith('ecb-') else '2025 shipments; returns cutoff 2026-01-31' if path.name in ('orders-part1.csv','orders-part2.csv','order-corrections.csv','returns.csv') else 'as described in file','status':info.get('kind','synthetic_client_input' if path.name not in ('data-dictionary.md','scenario-policy.md','source-register.json') else 'synthetic_metadata'),'transformation':'Lossless extraction of ecb-history.zip' if path.name=='ecb-history.csv' else ''})
with open(ROOT/'evidence/source-register.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(entries[0])); w.writeheader(); w.writerows(entries)
print(json.dumps({'reconciliation':recon,'country_totals':cleanrows(country),'best_portfolio':best,'recommendation':recommendation},indent=2,default=str))
