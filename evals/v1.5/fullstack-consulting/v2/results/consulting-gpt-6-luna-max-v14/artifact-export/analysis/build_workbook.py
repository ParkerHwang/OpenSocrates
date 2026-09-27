"""Create the reviewable Meridian analytical workbook from saved outputs."""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "analysis" / "output"
DEL = ROOT / "deliverables"
SRC = ROOT / "evidence" / "source_room"
BLUE = "17324D"
TEAL = "087E8B"
LIGHT = "E9F2F5"
GOLD = "E4A23A"
GRAY = "5E6B75"
WHITE = "FFFFFF"
GREEN = "E1F2EA"
RED = "F9E4E2"


def read_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def read_json(path):
    return json.loads(path.read_text())


def as_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    try:
        return float(value)
    except (ValueError, TypeError):
        return value


def clean_name(name):
    return "".join(ch if ch.isalnum() else "_" for ch in str(name))[:200]


def style_table(ws, header_row=1, filter_table=True, tablename=None):
    max_row, max_col = ws.max_row, ws.max_column
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    for c in ws[header_row]:
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.font = Font(color=WHITE, bold=True, size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 34
    if filter_table and max_row > header_row:
        ref = f"A{header_row}:{ws.cell(max_row, max_col).coordinate}"
        table = Table(displayName=tablename or clean_name(ws.title), ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)
    for col_cells in ws.columns:
        col = col_cells[0].column_letter
        header = str(ws.cell(header_row, col_cells[0].column).value or "")
        width = max([len(str(c.value)) if c.value is not None else 0 for c in col_cells[: min(max_row, 80)] ] + [len(header)])
        ws.column_dimensions[col].width = min(max(width + 2, 12), 52)
        if "url" in header.lower() or "description" in header.lower() or "handling" in header.lower() or "rationale" in header.lower() or "limitation" in header.lower():
            ws.column_dimensions[col].width = min(max(width + 2, 28), 58)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=(isinstance(cell.value, str) and len(cell.value) > 50))
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def add_rows(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        if isinstance(row, dict):
            vals = [row.get(h) for h in headers]
        else:
            vals = row
        ws.append([as_number(v) for v in vals])


def add_formats(ws, header_row=1):
    for col in range(1, ws.max_column + 1):
        h = str(ws.cell(header_row, col).value or "").lower()
        for row in range(header_row + 1, ws.max_row + 1):
            cell = ws.cell(row, col)
            if any(x in h for x in ["eur", "capex", "fixed", "sales", "refund", "cogs", "contribution", "fulfillment", "saving"]):
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '€#,##0.00;[Red](€#,##0.00);–'
            if "payback_years" in h:
                cell.number_format = '0.00;[Red](0.00);–'
            if "margin" in h or "change_2022_2024_pct" in h or "percent" in h:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.0%;[Red](0.0%);–'
            if h in {"population", "shipped_orders", "shipped_units", "returned_units", "fte", "year", "quantity", "restocked_units", "observation_days"}:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0;[Red](#,##0);–'
    ws.auto_filter.ref = ws.dimensions


def main():
    metrics = read_json(DEL / "metrics.json")
    source_register = read_json(DEL / "source_register.json")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "Meridian Parts — European service-hub analysis"
    wb.properties.subject = "Reconciled 2025 shipments, historical market context and hub scenarios"
    wb.properties.creator = "Meridian Parts analytical workstream"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    # Start page: source status, scope, key decision and navigation.
    ws = wb.create_sheet("Start Here")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "MERIDIAN PARTS | EUROPEAN SERVICE HUBS"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    summary_rows = [
        ("Decision", "Stage a Czechia + Spain two-hub program; release Spain first, hold Czechia funding to the day-90 evidence gate."),
        ("Base-case option", "CZE + ESP: €198,099 annual incremental contribution after recurring fixed cost; €225,000 year-zero capex; 5 FTE; 1.14-year simple undiscounted payback."),
        ("Defined joint stress", "CZE + ESP: €32,180 annual incremental contribution; positive, but €75,450 below the €107,630 NLD + ESP stress outcome."),
        ("Scope", "1,254 distinct eligible orders shipped in 2025; returns received through 2026-01-31 inclusive; all amounts EUR unless named otherwise."),
        ("Evidence state", "Client exports and assumptions are synthetic. World Bank and ECB values are archived official public snapshots in the source room; this analysis did not make a live API pull."),
        ("Flip condition", metrics["recommendation"]["flip_condition"]),
        ("Rebuild", "Run analysis/build_analysis.py then analysis/build_workbook.py and analysis/build_documents.py with the Python environment in TOOLING.md. The frozen inputs are under evidence/source_room/.")
    ]
    for i, (label, value) in enumerate(summary_rows, 3):
        ws.cell(i, 1, label).font = Font(bold=True, color=BLUE)
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(i, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(i, 2, value).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = 48 if i in (3, 4, 5, 7, 8, 9) else 38
    ws["A11"] = "WORKSHEETS"
    ws["A11"].font = Font(bold=True, color=WHITE)
    ws["A11"].fill = PatternFill("solid", fgColor=TEAL)
    tabs = [
        "Countries: annual orders, units, gross/refunds/net sales, COGS, fulfillment, contribution and margin.",
        "Monthly: 72 country-month rows; sums reconcile to Countries.",
        "Order Ledger: per-order rounded EUR components and FX rate for audit.",
        "Hub Scenarios / Alternatives: six options, all 11 feasible pairs, defer, low/base/high/stress and payback.",
        "Sensitivity: volume assumptions, defined joint stress, board-floor flip and limits on untested cost ranges.",
        "Market Context / FX Monthly: archived World Bank values and 2025 ECB monthly means.",
        "Quality / Exclusions / Source Register: data decisions, anomalies, source URLs, retrieval times and hashes.",
        "Decision Charts: annual and monthly visual comparisons; Chart Data is hidden support data."
    ]
    for i, line in enumerate(tabs, 12):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
        ws.cell(i, 1, "• " + line).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 26
    ws.column_dimensions["A"].width = 24
    for c in "BCDEF":
        ws.column_dimensions[c].width = 20
    ws.freeze_panes = "A3"

    # Core country and monthly results.
    country_headers = ["country", "shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units", "margin"]
    ws = wb.create_sheet("Countries")
    country_rows = [{k: r.get(k) for k in country_headers} for r in metrics["countries"]]
    add_rows(ws, country_headers, country_rows); style_table(ws, tablename="CountryTotals"); add_formats(ws)
    ws.column_dimensions["A"].width = 12
    monthly_headers = ["country", "month"] + country_headers[1:]
    ws = wb.create_sheet("Monthly")
    add_rows(ws, monthly_headers, [{k: r.get(k) for k in monthly_headers} for r in metrics["monthly"]]); style_table(ws, tablename="MonthlyCountryResults"); add_formats(ws)
    ws.column_dimensions["A"].width = 12; ws.column_dimensions["B"].width = 12

    # Saved order-level computation ledger and exclusion ledger.
    comp = read_csv(OUT / "order_components.csv")
    ws = wb.create_sheet("Order Ledger")
    ledger_headers = ["order_id", "order_revision", "country", "shipped_at", "month", "sku", "currency", "fx_local_per_eur", "quantity", "unit_price_local", "discount_local", "gross_local", "refund_local", "gross_sales_eur", "refunds_eur", "unit_cost_eur", "gross_cogs_eur", "restocked_units", "recovered_cogs_eur", "net_cogs_eur", "fulfillment_eur_raw", "fulfillment_eur", "net_sales_eur", "contribution_eur", "returned_units"]
    add_rows(ws, ledger_headers, [{k: r.get(k) for k in ledger_headers} for r in comp]); style_table(ws, tablename="OrderLedger"); add_formats(ws)
    ws.column_dimensions["A"].width = 18
    exclusions = read_csv(OUT / "order_exclusions.csv")
    ex_headers = list(exclusions[0]) if exclusions else ["order_id", "reason"]
    ws = wb.create_sheet("Exclusions")
    add_rows(ws, ex_headers, exclusions); style_table(ws, tablename="OrderExclusions"); add_formats(ws)

    # Exchange rates with count of available daily observations.
    observation_counts = metrics["fx_metadata"]["observation_counts"]
    fx_rows = [{"currency": r["currency"], "month": r["month"], "local_per_eur": r["local_per_eur"], "observation_days": observation_counts[f"{r['currency']}:{r['month']}"]} for r in metrics["fx_monthly"]]
    ws = wb.create_sheet("FX Monthly")
    fx_headers = ["currency", "month", "local_per_eur", "observation_days"]
    add_rows(ws, fx_headers, fx_rows); style_table(ws, tablename="FXMonthlyMeans"); add_formats(ws)
    ws["F1"] = "Method"
    ws["F1"].font = Font(bold=True, color=WHITE); ws["F1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["F2"] = "Arithmetic mean of available daily ECB reference quotes in the calendar month. Quote is local currency units per EUR; never inverted. Used for both shipment and refunds in the original shipment month."
    ws["F2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["F"].width = 48
    ws.row_dimensions[2].height = 65

    # Market context includes full year panels, current USD unit, and two-year population change.
    pop_change = metrics["country_population_change_2022_2024_pct"]
    market_rows = [{**r, "population_change_2022_2024_pct": pop_change[r["country"]] / 100 if pop_change[r["country"]] is not None else None} for r in metrics["market_context"]]
    market_headers = ["country", "year", "population", "gdp_per_capita_usd", "population_change_2022_2024_pct"]
    ws = wb.create_sheet("Market Context")
    add_rows(ws, market_headers, market_rows); style_table(ws, tablename="WorldBankMarketContext"); add_formats(ws)
    ws["G1"] = "Interpretation"; ws["G1"].font = Font(bold=True, color=WHITE); ws["G1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["G2"] = "Population and GDP per capita describe macro context; neither is a direct measure of replacement-assembly demand. GDP per capita is current US$, not PPP or constant-price income. World Bank archive last updated 2026-07-13."
    ws["G2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["G"].width = 48; ws.row_dimensions[2].height = 66

    # Scenario inputs/results and feasible portfolio comparisons.
    hub_rows = metrics["hub_scenarios"]
    hub_headers = ["country", "scenario", "incremental_contribution_eur", "capex_eur", "annual_fixed_eur", "saving_eur_per_unit", "fte", "payback_years"]
    ws = wb.create_sheet("Hub Scenarios")
    add_rows(ws, hub_headers, [{k: r.get(k) for k in hub_headers} for r in hub_rows]); style_table(ws, tablename="IndividualHubScenarios"); add_formats(ws)
    ws.conditional_formatting.add(f"C2:C{ws.max_row}", ColorScaleRule(start_type="min", start_color="F4B7B2", mid_type="percentile", mid_value=50, mid_color="FFF0BF", end_type="max", end_color="9AD1B7"))

    alternatives = metrics["option_comparison"]
    alt_headers = ["label", "countries", "capex_eur", "annual_fixed_eur", "fte", "low", "base", "high", "stress", "base_payback_years"]
    ws = wb.create_sheet("Alternatives")
    add_rows(ws, alt_headers, [{k: r.get(k) for k in alt_headers} for r in alternatives]); style_table(ws, tablename="FeasibleAlternativeComparison"); add_formats(ws)
    ws.conditional_formatting.add(f"F2:I{ws.max_row}", ColorScaleRule(start_type="min", start_color="F4B7B2", mid_type="percentile", mid_value=50, mid_color="FFF0BF", end_type="max", end_color="9AD1B7"))
    ws["L1"] = "Scenario formula"
    ws["L1"].font = Font(bold=True, color=WHITE); ws["L1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["L2"] = "Low/base/high: C×u + U×(1+u)×s − F; u=10%/25%/40%. Stress: C_stress×1.25 − C + U×1.25×s − F; C_stress=C−3%×G−10%×N for PLN/CZK, otherwise C−3%×G. Payback=capex/positive contribution."
    ws["L2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["L"].width = 54; ws.row_dimensions[2].height = 92

    raw_options = read_csv(SRC / "hub-options.csv")
    ws = wb.create_sheet("Hub Inputs")
    input_rows = [{**r, "budget_capex_eur": 450000, "budget_fte": 7, "max_hubs": 2} for r in raw_options]
    input_headers = ["country", "capex_eur", "fte", "annual_fixed_eur", "saving_eur_per_unit", "budget_capex_eur", "budget_fte", "max_hubs"]
    add_rows(ws, input_headers, [{k: r.get(k) for k in input_headers} for r in input_rows]); style_table(ws, tablename="HubAssumptions"); add_formats(ws)
    ws["J1"] = "Status"; ws["J1"].font = Font(bold=True, color=WHITE); ws["J1"].fill = PatternFill("solid", fgColor=BLUE)
    ws["J2"] = "Client-defined synthetic option assumptions; annual fixed cost excluded from year-zero capex. No demand-uplift observations or measured service savings were supplied."
    ws["J2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["J"].width = 52; ws.row_dimensions[2].height = 62

    # Quality notes include all dedupe/exclusion actions and declared limitations.
    q = metrics["quality"]
    quality_rows = [
        {"topic": "Order rows", "result": "1,328 raw rows; 13 exact repeats removed; 18 lower revisions superseded; 1,297 selected order IDs.", "handling": "One maximum numeric revision per order_id; whole row replaces prior revision."},
        {"topic": "Eligible shipment base", "result": "1,254 eligible unique orders across six countries; no calculation-field missing values among eligible candidates.", "handling": "Status shipped, is_test=false, shipped_at in 2025; latest effective unit cost applied."},
        {"topic": "Excluded orders", "result": "24 cancelled/non-shipped, 18 tests, and 1 shipment outside 2025 (January 2026) excluded.", "handling": "No revenue or units from these orders included."},
        {"topic": "Free shipments", "result": "12 valid zero-price shipments.", "handling": "Included as orders and units with zero gross sales."},
        {"topic": "Returns revisions", "result": "264 raw rows; 20 exact repeats removed; 1 superseded revision; 243 selected return IDs.", "handling": "Highest numeric revision per return_id."},
        {"topic": "Returns included", "result": "241 return IDs and 1,081 returned units applied to eligible linked 2025 orders.", "handling": "Cutoff inclusive through 2026-01-31; refunds aggregate by original order and use original sale-month FX."},
        {"topic": "Late and orphan returns", "result": "1 return after cutoff; 1 return linked to an unknown order ID.", "handling": "Late return excluded; orphan quarantined and not guessed into an order."},
        {"topic": "Reconciliation", "result": "Passed for orders, units, returned units and all EUR monetary fields by country.", "handling": "Country totals equal the sum of 12 monthly rows per country."},
        {"topic": "Rounding", "result": "Per-order monetary components rounded half-up to EUR cents.", "handling": "Values then summed to country/month; margin = contribution/net sales or blank for zero denominator."},
        {"topic": "ECB", "result": "2025 monthly PLN and CZK arithmetic means from published business-day reference-rate quotes.", "handling": "Quote stays local units per EUR; archived reference rates are translation assumptions, not settlement FX."},
        {"topic": "Synthetic inputs", "result": "All order, return, unit cost and hub-option data are synthetic.", "handling": "No actual customer interviews, customer records or measured hub outcomes are claimed."},
        {"topic": "Model limitation", "result": "Scenario model has no causal validation, synergy, discounting, tax, working capital or terminal value.", "handling": "Results are scenario arithmetic only; no statistical or causal certainty claimed."},
    ]
    quality_rows += [{"topic": "Limitation", "result": limitation, "handling": "Read alongside scenario assumptions and report."} for limitation in q["limitations"]]
    ws = wb.create_sheet("Quality")
    add_rows(ws, ["topic", "result", "handling"], quality_rows); style_table(ws, tablename="QualityNotes")
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 65; ws.column_dimensions["C"].width = 68
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row[0].row].height = 45

    # Register retains official public URLs, room collection timestamps, hashes and units.
    source_rows = [{k: r.get(k) for k in ["file", "description", "source_room_url", "original_source_url", "retrieved_from_room_utc", "original_archive_retrieved_utc", "sha256", "bytes", "units", "period", "status", "revision_vintage"]} for r in source_register]
    ws = wb.create_sheet("Source Register")
    source_headers = list(source_rows[0])
    add_rows(ws, source_headers, source_rows); style_table(ws, tablename="EvidenceSourceRegister")
    for col, width in {"A":22,"B":46,"C":42,"D":60,"E":31,"F":31,"G":66,"H":12,"I":34,"J":36,"K":34,"L":62}.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if isinstance(c.value, str) and c.value.startswith("http"):
                c.hyperlink = c.value; c.style = "Hyperlink"
        ws.row_dimensions[row[0].row].height = 65

    # Sensitivity ledger makes scenario provenance and decision flips easy to review.
    sens = metrics["sensitivity_analysis"]
    sensitivity_rows = []
    baseline = sens["baseline"]
    sensitivity_rows.extend([
        {"section":"Baseline", "item":"Primary metric / decision rule", "uplift_pct":baseline["volume_uplift_pct"], "base_contribution_eur":baseline["incremental_contribution_eur"], "base_or_stress_leader":baseline["winner"], "provenance_status":baseline["decision_rule"]},
        {"section":"Baseline", "item":"Year-zero capex / annual fixed cost / FTE", "base_contribution_eur":baseline["capex_eur"], "base_or_stress_leader":f"€{baseline['annual_fixed_eur']:,.0f} fixed/year; {baseline['fte']} FTE", "provenance_status":"Capex is separate from recurring cost; synthetic client option assumptions."},
    ])
    for r in sens["tested_ranges"]:
        summary = r.get("status", "")
        if r.get("provenance"):
            summary += " | " + r["provenance"]
        sensitivity_rows.append({"section":"Tested range", "item":r["input"], "uplift_pct":r.get("base_pct",r.get("volume_uplift_pct")), "base_contribution_eur":r.get("illustrative_threshold_eur"), "base_or_stress_leader":str({k:v for k,v in r.items() if k not in {"input","status","provenance"}}), "provenance_status":summary})
    for r in sens["perturbation_results"]:
        sensitivity_rows.append({"section":"Perturbation", "item":r["variation"], "uplift_pct":r.get("tested_uplift_pct"), "base_contribution_eur":r.get("incremental_contribution_eur"), "base_or_stress_leader":r.get("winner",r.get("winner_by_stress_contribution")), "stress_contribution_eur":r.get("stress_incremental_contribution_eur",r.get("CZE_ESP_stress_eur")), "provenance_status":r["status"]})
    for r in sens["switching_values"]:
        sensitivity_rows.append({"section":"Switching value", "item":r.get("input", "volume uplift"), "uplift_pct":r.get("uplift_pct"), "base_contribution_eur":r.get("switch_above_eur"), "base_or_stress_leader":r.get("switch_to",r.get("to")), "stress_contribution_eur":r.get("highest_qualifying_stress_eur"), "provenance_status":r.get("status",r.get("provenance",""))})
    sensitivity_rows.append({"section":"Conclusion", "item":"Robust or fragile", "provenance_status":sens["robust_or_fragile_conclusion"]})
    sensitivity_rows.extend({"section":"Research priority", "item":f"{i+1}. {item}", "provenance_status":"Required before the corresponding capital gate."} for i,item in enumerate(sens["research_priority"]))
    ws = wb.create_sheet("Sensitivity")
    sensitivity_headers = ["section", "item", "uplift_pct", "base_contribution_eur", "base_or_stress_leader", "stress_contribution_eur", "provenance_status"]
    add_rows(ws, sensitivity_headers, sensitivity_rows); style_table(ws, tablename="SensitivityAndSwitchingValues"); add_formats(ws)
    for col,width in {"A":20,"B":52,"C":16,"D":25,"E":45,"F":25,"G":90}.items(): ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
        ws.row_dimensions[row[0].row].height=50

    # Hidden chart table for native Excel charts.
    ws = wb.create_sheet("Chart Data")
    ws.append(["month"] + [r["country"] for r in metrics["countries"]])
    for month in [f"2025-{m:02d}" for m in range(1, 13)]:
        ws.append([month] + [next(r["contribution_eur"] for r in metrics["monthly"] if r["country"] == c["country"] and r["month"] == month) for c in metrics["countries"]])
    ws.append([])
    ws.append(["Country", "2025 contribution EUR", "Contribution margin"])
    for c in metrics["countries"]:
        ws.append([c["country"], c["contribution_eur"], c["margin"]])
    ws.append([])
    ws.append(["Country", "Base incremental EUR", "Stress incremental EUR"])
    for country in [c["country"] for c in metrics["countries"]]:
        base = next(r["incremental_contribution_eur"] for r in hub_rows if r["country"] == country and r["scenario"] == "base")
        stress = next(r["incremental_contribution_eur"] for r in hub_rows if r["country"] == country and r["scenario"] == "stress")
        ws.append([country, base, stress])
    ws.append([])
    ws.append(["Portfolio", "Base incremental EUR", "Stress incremental EUR"])
    pair_recs = [r for r in alternatives if "+" in r["label"]]
    pair_recs = sorted(pair_recs, key=lambda r: r["base"], reverse=True)[:6]
    for r in pair_recs:
        ws.append([r["label"], r["base"], r["stress"]])

    # Charts sheet: annual outcomes, scenario exposure and monthly contribution trend.
    ws = wb.create_sheet("Decision Charts")
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:Q1")
    ws["A1"] = "2025 BASELINE & HUB SCENARIOS"
    ws["A1"].font = Font(size=16, bold=True, color=WHITE); ws["A1"].fill = PatternFill("solid", fgColor=BLUE)
    ws.row_dimensions[1].height = 30
    source = wb["Chart Data"]
    annual = BarChart(); annual.type = "bar"; annual.style = 10; annual.title = "2025 country contribution"; annual.y_axis.title = "Country"; annual.x_axis.title = "EUR"
    annual.add_data(Reference(source, min_col=2, max_col=2, min_row=15, max_row=21), titles_from_data=True)
    annual.set_categories(Reference(source, min_col=1, min_row=16, max_row=21)); annual.height = 7.4; annual.width = 15.2; annual.legend = None
    ws.add_chart(annual, "A3")
    scenarios = BarChart(); scenarios.type = "col"; scenarios.style = 12; scenarios.title = "Single hubs: base vs defined joint stress"; scenarios.y_axis.title = "Annual incremental contribution (EUR)"; scenarios.x_axis.title = "Country"
    scenarios.add_data(Reference(source, min_col=2, max_col=3, min_row=23, max_row=29), titles_from_data=True)
    scenarios.set_categories(Reference(source, min_col=1, min_row=24, max_row=29)); scenarios.height = 7.4; scenarios.width = 17.5
    ws.add_chart(scenarios, "I3")
    monthly = LineChart(); monthly.style = 13; monthly.title = "Monthly contribution by country, 2025"; monthly.y_axis.title = "EUR"; monthly.x_axis.title = "Month"
    monthly.add_data(Reference(source, min_col=2, max_col=7, min_row=1, max_row=13), titles_from_data=True)
    monthly.set_categories(Reference(source, min_col=1, min_row=2, max_row=13)); monthly.height = 8.7; monthly.width = 30
    monthly.legend.position = "b"
    ws.add_chart(monthly, "A19")
    pairs = BarChart(); pairs.type = "bar"; pairs.style = 11; pairs.title = "Feasible two-hub pairs: base vs stress"; pairs.y_axis.title = "Portfolio"; pairs.x_axis.title = "Annual incremental contribution (EUR)"
    pairs.add_data(Reference(source, min_col=2, max_col=3, min_row=31, max_row=37), titles_from_data=True)
    pairs.set_categories(Reference(source, min_col=1, min_row=32, max_row=37)); pairs.height = 8.3; pairs.width = 30
    ws.add_chart(pairs, "A36")
    ws.merge_cells("A53:Q54")
    ws["A53"] = "Stress is the client-defined joint case: extra refund shock equal to 3% of gross sales, plus 10% of net sales for PLN/CZK; 25% base volume uplift; no extra cost recovery. All charts are scenario comparisons, not forecasts."
    ws["A53"].alignment = Alignment(wrap_text=True, vertical="top"); ws["A53"].font = Font(italic=True, color=GRAY, size=10)
    for c in range(1, 18): ws.column_dimensions[chr(64+c)].width = 12
    ws.freeze_panes = "A2"
    source.sheet_state = "hidden"

    # Workbook-level baseline checks are recorded as numeric values, not unexplained formulas.
    check = wb.create_sheet("Reconciliation")
    recon = []
    for country in [r["country"] for r in metrics["countries"]]:
        year_row = next(r for r in metrics["countries"] if r["country"] == country)
        month_rows = [r for r in metrics["monthly"] if r["country"] == country]
        for field in ["shipped_orders", "shipped_units", "gross_sales_eur", "refunds_eur", "net_sales_eur", "net_cogs_eur", "fulfillment_eur", "contribution_eur", "returned_units"]:
            monthly_sum = sum((r[field] for r in month_rows), 0)
            recon.append({"country": country, "metric": field, "monthly_sum": monthly_sum, "country_total": year_row[field], "difference": monthly_sum - year_row[field], "status": "PASS" if abs(monthly_sum - year_row[field]) < 0.000001 else "FAIL"})
    recon_headers = ["country", "metric", "monthly_sum", "country_total", "difference", "status"]
    add_rows(check, recon_headers, recon); style_table(check, tablename="MonthlyToAnnualReconciliation"); add_formats(check)
    for row in range(2, check.max_row + 1):
        cell = check.cell(row, 6)
        cell.fill = PatternFill("solid", fgColor=GREEN if cell.value == "PASS" else RED)

    # Sensible tab ordering.
    order = ["Start Here", "Countries", "Monthly", "Decision Charts", "Alternatives", "Sensitivity", "Hub Scenarios", "Hub Inputs", "FX Monthly", "Market Context", "Order Ledger", "Exclusions", "Quality", "Reconciliation", "Source Register", "Chart Data"]
    wb._sheets = [wb[s] for s in order]
    target = DEL / "meridian_analysis.xlsx"
    wb.save(target)
    # Reopen and check structure/expected row counts and stored values.
    check_wb = load_workbook(target, read_only=False, data_only=False)
    assert check_wb["Monthly"].max_row == 73
    assert check_wb["Countries"].max_row == 7
    assert check_wb["Order Ledger"].max_row == 1255
    assert check_wb["Reconciliation"].max_row == 55
    assert len(check_wb["Decision Charts"]._charts) == 4
    assert "Sensitivity" in check_wb.sheetnames
    assert all(check_wb["Reconciliation"].cell(r, 6).value == "PASS" for r in range(2, 56))
    print(f"Wrote {target} ({target.stat().st_size:,} bytes); {len(check_wb.sheetnames)} sheets; 4 charts; 54 reconciliation checks passed.")


if __name__ == "__main__":
    main()
