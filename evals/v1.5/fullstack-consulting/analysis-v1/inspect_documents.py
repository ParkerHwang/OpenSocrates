"""Read-only artifact inventory, text extraction and representative render evidence."""
from pathlib import Path
import hashlib,json,subprocess,tempfile
from concurrent.futures import ThreadPoolExecutor
from pypdf import PdfReader
import openpyxl
HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
STORAGE=Path('/private/var/folders/k6/8pv_0f_17kq34j9v5swgq67c0000gp/T/opensocrates-matrix-outcomes-v2-arxilet1')
PDFTOPPM='/Users/parkerhwang/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/pdftoppm'
SOFFICE='/Users/parkerhwang/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/override/soffice'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()

def inspect(cell):
    source=STORAGE/cell/'locked-artifacts';target=HERE/'documents'/cell;target.mkdir(parents=True)
    snapshot=json.loads((ROOT/'v2/results'/cell/'snapshot.json').read_text())['files']
    files=[n for n in snapshot if n.startswith('deliverables/') and n.count('/')==1 and Path(n).suffix.lower() in ('.pdf','.pptx','.xlsx','.docx')]
    receipt={'cell':cell,'read_only':True,'artifacts':[],'render_scope':'Representative first, densest-text and middle pages; not a complete visual audit. Text extracted from every PDF page. LibreOffice derivatives are separate and may differ from PowerPoint.'}
    for name in files:
        p=source/name;assert sha(p)==snapshot[name]['sha256']
        item={'source':name,'sha256':sha(p),'bytes':p.stat().st_size}
        pdf=p
        if p.suffix=='.pptx':
            out=target/'pptx-render'/p.stem;out.mkdir(parents=True)
            profile=Path(tempfile.mkdtemp(prefix='opensocrates-lo-review-')).resolve()
            cp=subprocess.run([SOFFICE,'-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(out),str(p)],capture_output=True,text=True)
            item['render_conversion']={'exit_code':cp.returncode,'stdout':cp.stdout.replace(str(source),'<LOCKED_ORIGINAL>').replace(str(HERE),'<ANALYSIS>'),'stderr':cp.stderr.replace(str(source),'<LOCKED_ORIGINAL>').replace(str(HERE),'<ANALYSIS>')}
            pdf=out/(p.stem+'.pdf')
            if not pdf.exists():item['render_unavailable']=True
        if p.suffix=='.xlsx':
            wb=openpyxl.load_workbook(p,data_only=False,read_only=False)
            sheets=[]
            for ws in wb:
                formulas=[];errors=[]
                for row in ws.iter_rows():
                    for c in row:
                        if c.data_type=='f':formulas.append({'cell':c.coordinate,'formula':c.value})
                        if c.data_type=='e':errors.append({'cell':c.coordinate,'error':c.value})
                sheets.append({'name':ws.title,'dimensions':ws.dimensions,'freeze_panes':str(ws.freeze_panes) if ws.freeze_panes else None,'autofilter':ws.auto_filter.ref,'formula_count':len(formulas),'formula_examples':formulas[:5],'stored_errors':errors,'preview':[[c.value for c in row] for row in ws.iter_rows(min_row=1,max_row=min(6,ws.max_row),max_col=min(10,ws.max_column))]})
            item['sheets']=sheets
            chosen=next((s['name'] for s in sheets if any(k in s['name'].lower() for k in ('dashboard','decision','summary','country','countries'))),sheets[0]['name'])
            item['render_sheet']=chosen
            (target/'workbook-source.json').write_text(json.dumps({'path':str(p),'sheet':chosen,'output':str(target/'workbook.png')},indent=2,default=str)+'\n')
        elif pdf.suffix=='.pdf' and pdf.exists():
            reader=PdfReader(pdf);texts=[page.extract_text() or '' for page in reader.pages]
            item['pages']=len(texts);item['empty_text_pages']=[i+1 for i,t in enumerate(texts) if not t.strip()]
            textpath=target/(p.stem+'.txt');textpath.write_text('\n\n'.join(f'PAGE {i+1}\n{t}' for i,t in enumerate(texts)))
            selected=sorted(set([0,len(texts)//2,max(range(len(texts)),key=lambda i:len(texts[i]))]))
            item['rendered_pages']=[i+1 for i in selected]
            item['render_images']=[]
            for index in selected:
                prefix=target/(p.stem+f'-page-{index+1:02d}')
                cmd=[PDFTOPPM,'-f',str(index+1),'-l',str(index+1),'-scale-to','1250','-singlefile','-png',str(pdf),str(prefix)]
                cp=subprocess.run(cmd,capture_output=True,text=True)
                item['render_images'].append({'page':index+1,'file':prefix.name+'.png','exit_code':cp.returncode,'stderr':cp.stderr[:500]})
        assert sha(p)==item['sha256']
        receipt['artifacts'].append(item)
    (target/'inventory.json').write_text(json.dumps(receipt,indent=2,default=str)+'\n')
    print(json.dumps({'cell':cell,'files':len(files),'pdf_pages':sum(a.get('pages',0) for a in receipt['artifacts'])}),flush=True)

if __name__=='__main__':
    cells=[p.name for p in sorted((ROOT/'v2/results').glob('consulting-*'))]
    with ThreadPoolExecutor(max_workers=4) as pool:list(pool.map(inspect,cells))
